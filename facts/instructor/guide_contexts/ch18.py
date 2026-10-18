"""Chapter 18's instructor guides (the capstone case "Make, Buy, or Reprice"): the values the guides state, and the
claims their wording makes.

The case (chapters/18-make-buy-or-reprice/chapter.qmd) is an independent final-project option with ten requirements in
four phases (four milestones) and no tutorials.  Its model answer is three files, one for each tool the case names:

- MakeBuy.sql, the SQL part (outputs/companion/solutions/Chapter18-capstone-sql/), 31 queries headed
  "Requirement N, QM",
- Make Buy Reprice.xlsx, the Excel part (outputs/companion/solutions/Chapter18-capstone-excel/),
- Product Costs, the Power BI project of Requirement 5 (outputs/companion/solutions/Chapter18-capstone-pbi/).

The registered notes ch18.r1 to ch18.r10 and ch18.m1 to ch18.m4 (facts/notes/ch18.py) already supply the expected values
of the requirements and milestones.  These contexts add what the model paths need beyond them:

- the text of each model file the guides quote (Excel formulas from the workbook, DAX and M from the Power BI project),
  read from the file itself, so a quote cannot drift from the file the instructor opens (the SQL is quoted into the
  templates verbatim and `guides.py check` verifies it against MakeBuy.sql),
- the symptom of every wrong version a guide names (a join chain that fans out, a margin divided by the units sold, a
  payroll grouped by the month of the period, setup counted once for every completion, a rate that counts the Factory
  Overhead entries twice, an explorer with no year filter, ...), computed read-only on the dataset, on the model
  script's own text with one change where the wrong version is a change to a query,
- the deck's slide numbers and claims on every table and column the guides name.

What the Power BI project proves and what only Desktop shows: the project encodes the model and the structure of every
page, and its Checks tab (DAX query view) recomputes the values of Requirement 5 and of Milestone 2, so a result is
"proved" when a Checks row names it.  How a visual renders (the matrix's expanded rows, the slicer state) is visible only
in Desktop.

Keys:
- ch18.g0: the case note (slides, the model files, counts);
- ch18.g1 to ch18.g5: Requirements 1 to 5, each used next to the note ch18.r1 to ch18.r5 that supplies the values of the
  requirement's instructor note.  ch18.g6 to ch18.g10 (Requirements 6 to 10) follow the same pattern and reuse its helpers:
  script(), quote(), sub(), temp_view(), run(), formula(), clean(), dax(), step(), check_expr(), title(), has_visual(),
  test_block(), deck_titles(), at(), span().
"""

from __future__ import annotations

import json
import re
from contextlib import contextmanager
from functools import lru_cache
from pathlib import Path

from guide_contexts import context
from guide_contexts.ch01 import schema
from guide_contexts.ch13 import _block
from notes import ch18 as n18

REPO = Path(__file__).resolve().parents[3]
CHAPTER = REPO / "chapters" / "18-make-buy-or-reprice" / "chapter.qmd"
DECK = REPO / "slides" / "chapter-18" / "index.qmd"
DOWNLOADS = REPO / "front-matter" / "downloads.qmd"
SOLUTIONS = REPO / "outputs" / "companion" / "solutions"
SQL = SOLUTIONS / "Chapter18-capstone-sql" / "MakeBuy.sql"
XLSX = SOLUTIONS / "Chapter18-capstone-excel" / "Make Buy Reprice.xlsx"
PROJECT = SOLUTIONS / "Chapter18-capstone-pbi" / "Product Costs"
MODEL = PROJECT / "Product Costs.SemanticModel"
REPORT = PROJECT / "Product Costs.Report"
GUIDES = REPO / "facts" / "instructor" / "guides" / "chapter-18-make-buy-or-reprice"

SCHEMA = {
    "Item": ["ItemID", "ItemCode", "ItemGroup", "ItemType", "SupplyMode", "StandardCost", "ListPrice",
             "StandardConversionCost", "StandardLaborHoursPerUnit", "RoutingID"],
    "SalesInvoice": ["SalesInvoiceID", "InvoiceDate"],
    "SalesInvoiceLine": ["SalesInvoiceID", "ItemID", "Quantity", "LineTotal", "BaseListPrice"],
    "WorkOrder": ["WorkOrderID", "ItemID", "RoutingID"],
    "WorkOrderClose": ["WorkOrderID", "CloseDate", "MaterialVarianceAmount", "DirectLaborVarianceAmount",
                       "OverheadVarianceAmount", "TotalVarianceAmount"],
    "ProductionCompletion": ["ProductionCompletionID", "WorkOrderID", "CompletionDate"],
    "ProductionCompletionLine": ["ProductionCompletionID", "ItemID", "QuantityCompleted", "ExtendedStandardMaterialCost",
                                 "ExtendedStandardDirectLaborCost", "ExtendedStandardVariableOverheadCost",
                                 "ExtendedStandardFixedOverheadCost", "ExtendedStandardConversionCost",
                                 "ExtendedStandardTotalCost"],
    "PurchaseOrder": ["PurchaseOrderID", "OrderDate"],
    "PurchaseOrderLine": ["PurchaseOrderID", "ItemID", "Quantity", "UnitCost"],
    "ShipmentLine": ["ShipmentID", "ItemID", "QuantityShipped", "ExtendedStandardCost"],
    "MaterialIssue": ["MaterialIssueID", "WorkOrderID", "IssueDate"],
    "MaterialIssueLine": ["MaterialIssueID", "ItemID", "ExtendedStandardCost"],
    "GoodsReceipt": ["GoodsReceiptID", "ReceiptDate"],
    "GoodsReceiptLine": ["GoodsReceiptID", "ItemID", "ExtendedStandardCost"],
    "SalesOrder": ["SalesOrderID", "OrderDate", "CustomerID"],
    "SalesOrderLine": ["SalesOrderID", "ItemID", "Quantity", "BaseListPrice", "PromotionID", "PriceListLineID"],
    "Customer": ["CustomerID", "CustomerSegment"],
    "PromotionProgram": ["PromotionID", "PromotionCode", "ScopeType", "CustomerSegment", "ItemGroup", "CollectionName",
                         "EffectiveStartDate", "EffectiveEndDate"],
    "PriceList": ["PriceListID", "PriceListName", "ScopeType", "Status", "EffectiveStartDate", "EffectiveEndDate"],
    "PriceListLine": ["PriceListLineID", "PriceListID", "ItemID", "UnitPrice"],
    "PriceOverrideApproval": ["Status", "ApprovedByEmployeeID"],
    "FixedAsset": ["OriginalCost", "UsefulLifeMonths", "InServiceDate", "DisposalDate", "DepreciationDebitAccountID"],
    "Shipment": ["ShipmentID", "ShipmentDate"],
    "SalesReturn": ["SalesReturnID", "ReturnDate"],
    "SalesReturnLine": ["SalesReturnID", "ItemID", "QuantityReturned"],
    "GLEntry": ["PostingDate", "AccountID", "Debit", "Credit", "VoucherNumber", "SourceDocumentType", "SourceDocumentID",
                "FiscalYear"],
    "Account": ["AccountID", "AccountNumber"],
    "JournalEntry": ["JournalEntryID", "EntryNumber", "EntryType"],
    "PayrollRegister": ["PayrollRegisterID", "PayrollPeriodID", "CostCenterID", "EmployeeID", "GrossPay",
                        "EmployerPayrollTax", "EmployerBenefits"],
    "PayrollRegisterLine": ["PayrollRegisterID", "LineType", "Hours", "Amount", "LaborTimeEntryID"],
    "PayrollPeriod": ["PayrollPeriodID", "PeriodEndDate", "PayDate", "FiscalYear"],
    "CostCenter": ["CostCenterID", "CostCenterName"],
    "Employee": ["EmployeeID", "PayClass", "CostCenterID", "JobTitle"],
    "LaborTimeEntry": ["LaborTimeEntryID", "LaborType", "WorkDate", "WorkOrderOperationID", "RegularHours", "OvertimeHours"],
    "WorkOrderOperation": ["WorkOrderOperationID", "WorkCenterID"],
    "RoutingOperation": ["RoutingID", "WorkCenterID", "StandardRunHoursPerUnit", "StandardSetupHours"],
    "WorkCenter": ["WorkCenterID", "WorkCenterName"],
    "WorkCenterCalendar": ["WorkCenterID", "CalendarDate", "AvailableHours", "IsWorkingDay"],
}
# The student time of each requirement in minutes, and its weight in the case rubric (the weights add up to 100).  The
# requirement guides' front matter must agree; g0 claims it for every guide written so far.
PLAN = {1: 90, 2: 150, 3: 90, 4: 180, 5: 150, 6: 150, 7: 120, 8: 120, 9: 150, 10: 90}
WEIGHTS = {1: 7, 2: 10, 3: 7, 4: 12, 5: 9, 6: 12, 7: 8, 8: 9, 9: 14, 10: 12}
SOLVER_MINUTES = 55     # the optional Solver test inside Requirement 7's 120 minutes (its guide's "How to run it")


def need(d, claim, *tables: str) -> None:
    schema(d, claim, {t: SCHEMA[t] for t in (tables or SCHEMA)})


# --- the SQL model script ---------------------------------------------------------------------------------------------

def sub(sql: str, old: str, new: str) -> str:
    """The statement with `old` replaced by `new`, ignoring differences in whitespace; the text must be there."""
    pattern = r"\s+".join(re.escape(w) for w in old.split())
    if not re.search(pattern, sql):
        raise ValueError(f"the text to change is not in the statement: {old[:70]}")
    return re.sub(pattern, lambda m: new, sql, count=1)


@lru_cache(maxsize=None)
def script() -> tuple[tuple[str, str, str], ...]:
    """(first comment line, second comment line, statement) of every query of MakeBuy.sql, in order; the trailing result
    comment is dropped and the closing notes block is left out."""
    body = SQL.read_text(encoding="utf-8").split("\n/* Model answer", 1)[0]
    out = []
    for chunk in re.split(r"(?m)^(?=-- Requirement \d+, Q\d+:)", body):
        if not re.match(r"-- Requirement \d+, Q\d+:", chunk):
            continue
        lines = chunk.rstrip().split("\n")
        stmt = re.sub(r";[ \t]*-- Result:[^\n]*$", ";", "\n".join(lines[2:]), flags=re.M)
        out.append((lines[0], lines[1], stmt.strip()))
    return tuple(out)


def quote(key: str) -> str:
    """The statement of the one query whose first comment line starts with `key` ('Requirement 1, Q2')."""
    hits = [s for first, _, s in script() if first.removeprefix("-- ").startswith(key + ":")]
    if len(hits) != 1:
        raise KeyError(f"{len(hits)} queries of MakeBuy.sql start with {key!r}")
    return hits[0]


def run(d, sql: str) -> list:
    return d.q(sql.strip().rstrip(";"))


@contextmanager
def temp_view(d, name: str, new_name: str | None = None, *changes: tuple[str, str]):
    """A saved view of the model script created as a TEMP view (for the length of the block), optionally under another
    name and with changes to its text; the Data object is shared, so the view is dropped again."""
    new_name = new_name or name
    stmt = next((s for first, _, s in script() if re.match(rf"(?s)DROP VIEW IF EXISTS {name};", s)), None)
    if stmt is None:
        raise KeyError(f"MakeBuy.sql does not save the view {name}")
    text = re.match(rf"(?s)DROP VIEW IF EXISTS {name};\s*(CREATE VIEW {name} AS.*)", stmt).group(1).rstrip().rstrip(";")
    for old, new in changes:
        text = sub(text, old, new)
    text = text.replace(f"CREATE VIEW {name} AS", f"CREATE TEMP VIEW {new_name} AS", 1)
    d.con.execute(f"DROP VIEW IF EXISTS temp.{new_name}")
    d.con.execute(text)
    try:
        yield
    finally:
        d.con.execute(f"DROP VIEW IF EXISTS temp.{new_name}")


# --- the Excel model workbook -----------------------------------------------------------------------------------------------

@lru_cache(maxsize=None)
def workbook():
    import openpyxl
    return openpyxl.load_workbook(XLSX, read_only=False, data_only=False)


def sheets() -> list[str]:
    return list(workbook().sheetnames)


def clean(text) -> str:
    """A formula as the formula bar reads it: the file's prefixes written the way a person types them, and a Table's
    this-row references as @Column."""
    t = str(getattr(text, "text", text))
    t = re.sub(r"_xl(?:fn|pm|ws)\.", "", t)
    t = re.sub(r"(\w+)\[\[#This Row\],\[([^\]]+)\]\]", r"[@\2]", t)
    return t


def formula(claim, sheet: str, cell: str) -> str:
    value = workbook()[sheet][cell].value
    ok = isinstance(value, str) and value.startswith("=") or hasattr(value, "text")
    claim(ok, f"the cell {sheet}!{cell} of the model workbook holds a formula")
    return clean(value) if ok else ""


def item_formula(claim, column: str) -> str:
    """The formula of one calculated column of the ItemCosts Table (the first data row), as the formula bar reads it."""
    ws = workbook()["ItemCosts"]
    names = [c.value for c in ws[1]]
    ok = column in names
    claim(ok, f"the ItemCosts Table has the column {column}")
    return clean(ws.cell(2, names.index(column) + 1).value) if ok else ""


def defined(claim, name: str) -> str:
    wb = workbook()
    ok = name in wb.defined_names
    claim(ok, f"the model workbook has the defined name {name}")
    return wb.defined_names[name].attr_text if ok else ""


# --- the Power BI model answer -------------------------------------------------------------------------------------------

def _model_text(relative: str) -> str:
    return (MODEL / "definition" / relative).read_text(encoding="utf-8")


def table_m(name: str) -> list[str]:
    return _block(_model_text(f"tables/{name}.tmdl"), "source =")


def step(claim, lines: list[str], *names: str) -> str:
    """The named steps of an M query as the formula bar reads them."""
    out = []
    for n in names:
        found = [l for l in lines if l.startswith(f'#"{n}" =') or l.startswith(f"{n} =")]
        claim(len(found) == 1, f"the solution project's query has the step {n}")
        out += ["    " + l.rstrip(",") for l in found]
    return "\n".join(out)


@lru_cache(maxsize=None)
def measure_table() -> dict[str, str]:
    """Every measure of every table of the project as the formula bar reads it: `Name = expression`, the body dedented."""
    out: dict[str, str] = {}
    for path in sorted((MODEL / "definition" / "tables").glob("*.tmdl")):
        lines = path.read_text(encoding="utf-8").split("\n")
        for i, line in enumerate(lines):
            m = re.match(r"^\tmeasure (?:'([^']+)'|(\S+)) =(.*)$", line)
            if not m:
                continue
            name, first = m.group(1) or m.group(2), m.group(3).strip()
            body, j = [], i + 1
            while j < len(lines) and lines[j].startswith("\t\t\t"):
                body.append(lines[j][3:].rstrip())
                j += 1
            out[name] = f"{name} = {first}" if first and not body else f"{name} =\n" + "\n".join(body)
    return out


def dax(claim, *names: str) -> str:
    table = measure_table()
    for n in names:
        claim(n in table, f"the solution project has the measure {n}")
    return "\n\n".join(table.get(n, "") for n in names)


def check_expr(claim, label: str) -> str:
    """The DAX expression of one row of the project's Checks tab (DAX query view), found by its label."""
    text = (MODEL / "DAXQueries" / "Checks.dax").read_text(encoding="utf-8")
    m = re.search(r'"Check", "' + re.escape(label) + r'", "Expected", .*?, "Actual", (.*?), "Agrees"', text)
    claim(m is not None, f"the Checks tab has the row {label!r}")
    return m.group(1) if m else ""


def check_sections() -> dict[str, int]:
    text = (MODEL / "DAXQueries" / "Checks.dax").read_text(encoding="utf-8")
    counts: dict[str, int] = {}
    for sec in re.findall(r'ROW \( "Section", "([^"]*)"', text):
        counts[sec] = counts.get(sec, 0) + 1
    return counts


def title(claim, page: str, name: str) -> str:
    """The title of a visual of the solution project's report page, as Desktop shows it."""
    try:
        v = json.loads((REPORT / "definition" / "pages" / page / "visuals" / name / "visual.json").read_text(encoding="utf-8"))
        text = v["visual"]["visualContainerObjects"]["title"][0]["properties"]["text"]["expr"]["Literal"]["Value"].strip("'")
        text = text.replace("''", "'")
    except (OSError, KeyError, IndexError, ValueError):
        text = ""
    claim(text != "", f"the visual {name} of page {page} in the solution project has a title")
    return text


def has_visual(claim, page: str, *names: str) -> None:
    for n in names:
        claim((REPORT / "definition" / "pages" / page / "visuals" / n / "visual.json").is_file(),
              f"the page {page} of the solution project has the visual {n}")


def test_block(claim, n: int) -> str:
    """Test n of the Tests tab, comment line first."""
    text = (MODEL / "DAXQueries" / "Tests.dax").read_text(encoding="utf-8").strip()
    found = [p.strip() for p in re.split(r"(?m)^(?=// Test \d)", text) if p.startswith(f"// Test {n}:")]
    claim(len(found) == 1, f"the Tests tab holds Test {n}")
    return found[0] if found else ""


# --- the deck ------------------------------------------------------------------------------------------------------------------

def deck_titles() -> list[str]:
    titles = ["(title slide)"]
    for line in DECK.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^#{1,2} (.+?)\s*(\{[^}]*\})?\s*$", line)
        if m:
            titles.append(m.group(1).strip())
    return titles


def at(titles: list[str], first: str, nth: int = 1) -> int:
    hits = [n for n, t in enumerate(titles, 1) if t.startswith(first)]
    return hits[nth - 1]


def span(titles: list[str], first: str, after: str, nth: int = 1, after_nth: int = 1) -> str:
    start, end = at(titles, first, nth), at(titles, after, after_nth) - 1
    return str(start) if end <= start else f"{start} to {end}"


def dates_of(d, year) -> dict:
    """The dates through which the base year's records run (the first query of Requirement 1)."""
    row = run(d, quote("Requirement 1, Q1"))[0]
    return dict(invoice=row[0], completion=row[1], close=row[2], time=row[3])


# --- the case note -----------------------------------------------------------------------------------------------------------

@context("ch18.g0")
def g0(d, claim):
    need(d, claim)
    t = deck_titles()
    claim(len(t) == 66, "the case's deck has 66 slides")
    text = CHAPTER.read_text(encoding="utf-8")
    n_req = len(re.findall(r"^\*\*Requirement \d+:", text, re.M))
    claim(n_req == 10, "the case has ten requirements")
    n_ms = len(re.findall(r"^\d+\. \*\*[^*]+\*\* \(Requirements? ", text, re.M))
    claim(n_ms == 4, "the case has four milestones, one for each phase")
    claim("## Tutorial" not in text and "Guided Tutorial" not in text, "the case has no tutorials")
    downloads = DOWNLOADS.read_text(encoding="utf-8")
    claim(all(z in downloads for z in ("Chapter18-capstone-excel.zip", "Chapter18-capstone-pbi.zip", "Chapter18-capstone-sql.zip")),
          "the Downloads page lists the three solution files of the case")
    claim(SQL.is_file() and XLSX.is_file() and MODEL.is_dir() and REPORT.is_dir(),
          "the three model files of the case are in outputs/companion/solutions")
    queries = script()
    claim(len(queries) == 31, "the model script holds 31 queries")
    per: dict[int, int] = {}
    for first, _, _ in queries:
        k = int(re.match(r"-- Requirement (\d+),", first).group(1))
        per[k] = per.get(k, 0) + 1
    wb = sheets()
    claim(all(s in wb for s in ("Documentation", "Assumptions", "Families", "Cost Behavior", "Activities", "Costs",
                                "Make or Buy", "Capacity", "Prices", "Standards", "Memo", "ItemCosts", "Solution Notes")),
          "the model workbook has the worksheets the guides name")
    pages = json.loads((REPORT / "definition" / "pages" / "pages.json").read_text(encoding="utf-8"))["pageOrder"]
    claim(sorted(pages) == sorted(["about", "families", "concentration", "capacity", "validation", "notes"]),
          "the solution project has the six pages the guides name")
    sections = check_sections()
    claim(set(sections) == {"Requirement 5", "Milestone 2"}, "the Checks tab holds the rows of Requirement 5 and Milestone 2")
    m = n18.model(d)
    claim(sum(WEIGHTS.values()) == 100, "the rubric weights of the ten requirements add up to 100")
    r7_guide = (GUIDES / "req-7-value-the-capacity-that-buying-would-free.md").read_text(encoding="utf-8")
    claim(f"{SOLVER_MINUTES} for the optional Solver test" in r7_guide,
          "Requirement 7's guide gives the optional Solver test the minutes the case note leaves out of the required time")
    claim(set(PLAN) == set(WEIGHTS) == set(range(1, 11)), "the plan and the weights cover the ten requirements")
    for path in sorted(GUIDES.glob("req-*.md")):
        k = int(re.match(r"req-(\d+)-", path.name).group(1))
        mm = re.search(r"(?m)^minutes: (\d+)", path.read_text(encoding="utf-8"))
        claim(not mm or int(mm.group(1)) in (0, PLAN[k]), f"requirement {k}'s guide gives its student time as {PLAN[k]} minutes")
    # the data conditions of the at-a-glance table
    dt = dates_of(d, d.C)
    claim(dt["time"] < dt["close"] <= dt["completion"], "the time records end before the closes and the completions")
    claim(d.one("SELECT COUNT(*) FROM sqlite_master WHERE type = 'table' AND (name LIKE '%History%' OR name LIKE '%ChangeLog%')") == 0,
          "the dataset keeps no history of master-data changes")
    s = dict(
        kick=span(t, "Is the company selling below cost?", "The Situation"), situation=span(t, "The Situation", "The Data You Will Use"),
        data=span(t, "The Data You Will Use", "Getting Started"), setup=span(t, "Getting Started", "Phase 1: What the Records Show"),
        tool_check=span(t, "Check your understanding", "Phase 1: What the Records Show"), handin=at(t, "What you will hand in"),
        criteria=span(t, "How your work will be judged", "The Situation"), watch=span(t, "Watch out: units", "Getting Started"),
        milestones=at(t, "Four milestones"), tools=at(t, "Choose the tool for each step"),
        p1=span(t, "Phase 1: What the Records Show", "Phase 2: What a Product Costs"), r1=at(t, "Test the CFO's numbers"),
        r2=at(t, "Take the overhead apart"), r3=at(t, "What each family uses"), m1=at(t, "Milestone 1"),
        p2=span(t, "Phase 2: What a Product Costs", "Phase 3: Make, Buy, or Reprice"), views=span(t, "Three costing views", "Compute three costs"),
        r4=at(t, "Compute three costs"), r5=at(t, "Build the explorer"), check2=span(t, "Check your understanding", "Milestone 2", nth=2),
        m2=at(t, "Milestone 2"), p3=span(t, "Phase 3: Make, Buy, or Reprice", "Phase 4: Standards and the Memo"),
        flow=span(t, "Compare the quote with the costs that stop", "Decide make or buy"), r6=at(t, "Decide make or buy"),
        spiral=at(t, "Watch out: absorption cost"), r7=at(t, "Value the capacity"), r8=at(t, "Reprice the families"),
        check3=span(t, "Check your understanding", "Milestone 3", nth=3), m3=at(t, "Milestone 3"),
        p4=span(t, "Phase 4: Standards and the Memo", "Deliverables and Checklist"), r9=span(t, "Revise the standards", "Write the decision memo"),
        p4g=at(t, "Phase 4 at a glance"), r10=at(t, "Write the decision memo"), m4=at(t, "Milestone 4"), deliver=at(t, "Deliverables and Checklist"),
        before=at(t, "Before you submit"), ahead=span(t, "Looking Ahead", "About these slides"))
    required = sum(PLAN.values()) - SOLVER_MINUTES
    return dict(plan=PLAN, weights=WEIGHTS, student_minutes=sum(PLAN.values()), student_hours=sum(PLAN.values()) // 60,
                solver_minutes=SOLVER_MINUTES, required=f"{required // 60} hours {required % 60} minutes",
                n_req=n_req, n_ms=n_ms, slides=len(t), s=s, n_queries=len(queries), per=per, n_sheets=len(wb), n_pages=len(pages),
                checks=sum(sections.values()), per_section=sections, n_measures=len(measure_table()), dt=dt,
                n_families=len(m["fam"]), n_quoted=len(m["quoted"]), n_items=m["explorer"]["n_items"])


# --- Requirement 1: the family results and the premises --------------------------------------------------------------------

@context("ch18.g1")
def g1(d, claim):
    need(d, claim)
    m = n18.model(d)
    C = d.C
    first, last = f"{C}-01-01", f"{C}-12-31"
    claim(sheets()[:3] == ["Documentation", "Assumptions", "Families"], "the model workbook starts with Documentation, Assumptions and Families")
    q2 = quote("Requirement 1, Q2")
    claim("Sold AS" in q2 and "Closes AS" in q2 and "Made AS" in q2,
          "the model's family results keep the sales, the closes and the completions in a CTE each")
    with temp_view(d, "FamilyResults"):
        rows = run(d, quote("Requirement 1, Q3"))
        n_rows = len(rows)
        q4 = run(d, quote("Requirement 1, Q4"))
        q5 = run(d, quote("Requirement 1, Q5"))
        claim(n_rows == len(m["fam"]) == len(q4), "the view has one row for each manufactured family sold in the year")
        # two margins that collapse: the variance per unit applied to the units sold gives the whole year's variance
        per_unit_sold = all(abs((f["variance"] / f["units"]) * f["units"] - f["variance"]) < 1e-6 for f in m["fam"].values())
        claim(per_unit_sold, "a variance divided by the units sold and charged on the units sold is the year's variance")
        claim(any(abs(f["m_unit"] - f["m_year"]) > 0.001 for f in m["fam"].values()),
              "the per-unit-completed margin differs from the whole-year margin for some family")
        q7 = run(d, quote("Requirement 1, Q7"))[0]
    claim(abs(q7[0] - q7[1]) < 0.005 and abs(q7[0] - m["totals"]["total"]) < 0.005, "the families' variances equal account 5080")
    # wrong version 1: one join chain from the invoice lines to the closes by item
    lines_by_item = dict(d.q("SELECT sil.ItemID, COUNT(*) FROM SalesInvoiceLine sil JOIN SalesInvoice si ON si.SalesInvoiceID = "
                             "sil.SalesInvoiceID WHERE si.InvoiceDate BETWEEN ? AND ? GROUP BY 1", first, last))
    var_by_item = dict(d.q("SELECT wo.ItemID, SUM(c.TotalVarianceAmount) FROM WorkOrderClose c JOIN WorkOrder wo ON wo.WorkOrderID = "
                           "c.WorkOrderID WHERE c.CloseDate BETWEEN ? AND ? GROUP BY 1", first, last))
    fan = sum(lines_by_item.get(i, 0) * v for i, v in var_by_item.items())
    claim(fan > 10 * m["totals"]["total"], "joining the sales lines to the closes by item repeats each close once for every line, many times the ledger's variance")
    # wrong version 2: families by the three-letter product type
    by_type = d.q("SELECT SUBSTR(i.ItemCode, 5, 3), COUNT(DISTINCT SUBSTR(i.ItemCode, 1, 7)), COUNT(DISTINCT i.ItemGroup) FROM Item i "
                  "WHERE i.SupplyMode = 'Manufactured' AND i.ItemType = 'Finished Good' AND SUBSTR(i.ItemCode, 1, 7) IN "
                  "(SELECT SUBSTR(i2.ItemCode, 1, 7) FROM Item i2 JOIN SalesInvoiceLine l ON l.ItemID = i2.ItemID JOIN SalesInvoice si "
                  "ON si.SalesInvoiceID = l.SalesInvoiceID WHERE si.InvoiceDate BETWEEN ? AND ? AND i2.SupplyMode = 'Manufactured') GROUP BY 1",
                  first, last)
    merged = [r for r in by_type if r[1] > 1]
    claim(len(merged) == 1 and merged[0][2] == 2, "exactly one three-letter product type merges two families, and they sit in two item groups")
    # wrong version 3: starting from every manufactured family keeps the family that was not sold
    every = d.one("SELECT COUNT(DISTINCT SUBSTR(ItemCode, 1, 7)) FROM Item WHERE SupplyMode = 'Manufactured' AND ItemType = 'Finished Good'")
    claim(every == len(m["fam"]) + len(m["unsold"]), "the manufactured families are those sold and the one that was not")
    # wrong version 4: rows as units
    lines_mfg = d.one("SELECT COUNT(*) FROM SalesInvoiceLine l JOIN SalesInvoice si ON si.SalesInvoiceID = l.SalesInvoiceID JOIN Item i "
                      "ON i.ItemID = l.ItemID WHERE si.InvoiceDate BETWEEN ? AND ? AND i.SupplyMode = 'Manufactured' "
                      "AND i.ItemType = 'Finished Good'", first, last)
    units_mfg = sum(f["units"] for f in m["fam"].values())
    claim(lines_mfg < units_mfg, "the manufactured invoice lines are fewer than the units sold (quantities are fractional and often above one)")
    # closes dated in the year for work completed before it
    prior = d.one("SELECT COUNT(*) FROM WorkOrderClose WHERE CloseDate BETWEEN ? AND ? AND WorkOrderID NOT IN "
                  "(SELECT WorkOrderID FROM ProductionCompletion WHERE CompletionDate BETWEEN ? AND ?)", first, last, first, last)
    claim(prior > 0, "some closes of the year are for work completed in an earlier year")
    dt = dates_of(d, C)
    g_f = n18.story(m)
    pts = (g_f["m_std"] - g_f["m_year"]) * 100
    return dict(
        dt=dt, n_rows=n_rows, q7=dict(family=q7[0], ledger=q7[1]), fan=round(fan, 2), fan_times=fan / m["totals"]["total"],
        n_types=len(by_type), merged=merged[0][0], every=every, lines_mfg=lines_mfg, units_mfg=units_mfg, prior=prior,
        pts=pts, n_fam=len(m["fam"]), unsold=m["unsold"],
        f_vpu=formula(claim, "Families", "R5"), f_m1=formula(claim, "Families", "S5"), f_m2=formula(claim, "Families", "T5"),
        f_m3=formula(claim, "Families", "U5"), f_group=formula(claim, "Families", "G26"),
        recon=formula(claim, "Families", "B34"), n_q1=sum(1 for f, _, _ in script() if f.startswith("-- Requirement 1,")))


# --- Requirement 2: the inputs to 1090 and their behavior ------------------------------------------------------------------

def _series(d, where: str) -> dict:
    from collections import defaultdict
    return defaultdict(float, d.q("SELECT substr(g.PostingDate, 1, 7), SUM(g.Debit) FROM GLEntry g LEFT JOIN JournalEntry je "
                                  "ON je.JournalEntryID = g.SourceDocumentID AND g.SourceDocumentType = 'JournalEntry' "
                                  f"WHERE g.AccountID = ? AND {where} GROUP BY 1", d.account("1090")))


def _r(xs: list, ys: list) -> float:
    import statistics as st
    return st.correlation(xs, ys)


@context("ch18.g2")
def g2(d, claim):
    need(d, claim)
    from collections import defaultdict
    m = n18.model(d)
    C = d.C
    first, last = f"{C}-01-01", f"{C}-12-31"
    a1090 = d.account("1090")
    std_m, months, with_start = m["std_m"], m["months"], m["with_start"]
    pay = _series(d, "g.SourceDocumentType = 'PayrollSummary'")
    fo = _series(d, "je.EntryType = 'Factory Overhead'")
    xs = lambda ms: [std_m[k] for k in ms]                                                     # noqa: E731
    # the whole payroll without the six months that carry three pay dates
    three = m["three"]
    no3 = [k for k in months if k not in three]
    r_no3 = _r(xs(no3), [pay[k] for k in no3])
    claim(abs(r_no3) < 0.3, "leaving out the months with three pay dates leaves payroll unrelated to output")
    r_fo_start = _r(xs(with_start), [fo[k] for k in with_start])
    # the crew's regular pay, direct and indirect together, by pay date
    comp = defaultdict(lambda: defaultdict(float))
    alt = defaultdict(lambda: defaultdict(float))
    sql = ("SELECT substr(pp.{col}, 1, 7), l.LineType, lte.LaborType, SUM(l.Amount) FROM PayrollRegisterLine l JOIN PayrollRegister pr "
           "ON pr.PayrollRegisterID = l.PayrollRegisterID JOIN PayrollPeriod pp ON pp.PayrollPeriodID = pr.PayrollPeriodID JOIN CostCenter cc "
           "ON cc.CostCenterID = pr.CostCenterID LEFT JOIN LaborTimeEntry lte ON lte.LaborTimeEntryID = l.LaborTimeEntryID "
           "WHERE cc.CostCenterName = 'Manufacturing' AND l.LineType NOT IN ('Benefits Deduction', 'Employee Tax Withholding') GROUP BY 1, 2, 3")
    for mo, t, k, a in d.q(sql.format(col="PayDate")):
        comp[mo][(t, k)] += a
    for mo, t, k, a in d.q(sql.format(col="PeriodEndDate")):
        alt[mo][(t, k)] += a
    dk, ik = ("Regular Earnings", "Direct Manufacturing"), ("Regular Earnings", "Indirect Manufacturing")
    r_reg = _r(xs(months), [comp[k][dk] + comp[k][ik] for k in months])
    r_reg_alt = _r(xs(months), [alt[k][dk] + alt[k][ik] for k in months])
    r_dreg_alt = _r(xs(months), [alt[k][dk] for k in months])
    gap_alt = max(abs(sum(alt[k].values()) - pay[k]) for k in months)
    claim(abs(r_reg) < 0.2 and r_dreg_alt > 0.8 and gap_alt > 1000,
          "the crew's regular pay does not move with output, while grouping the direct part by the period's month shows a strong "
          "relation and no longer ties to the ledger")
    # the debits to 1090 with the closes' debits left in
    with_closes = d.one("SELECT SUM(Debit) FROM GLEntry WHERE AccountID = ? AND PostingDate BETWEEN ? AND ? AND Debit > 0", a1090, first, last)
    claim(with_closes > m["inputs"], "the favorable closes post debits to 1090, so leaving them in overstates the inputs")
    # the queries
    q10 = run(d, quote("Requirement 2, Q10"))
    q11 = run(d, quote("Requirement 2, Q11"))
    q14 = run(d, quote("Requirement 2, Q14"))
    q12 = run(d, quote("Requirement 2, Q12"))[0]
    claim(len(q10) == 7 and len(q11) == 3 and len(q14) == 36, "the extracts return seven payroll lines, three sources and thirty-six months")
    claim(abs(q12[5] - m["conversion"]) < 0.005 and abs(q12[6] - m["released_gl"]) < 0.005,
          "the completion lines and the ledger's credits are the note's")
    pay_months = [r for r in q14 if r[6] == 3]
    claim([r[0] for r in pay_months] == three, "the monthly extract shows three pay dates in the same months as the note")
    f_by = {k: formula(claim, "Cost Behavior", c) for k, c in (("sumifs", "C20"), ("per_hour", "D20"), ("slope", "R60"),
                                                                ("correl", "T60"), ("rsq", "U60"), ("test_from", "R57"),
                                                                ("share", "B50"), ("times", "B49"))}
    claim(defined(claim, "TestFrom").endswith("$R$57") and defined(claim, "StdHours").endswith("$B$35"),
          "the workbook names TestFrom and StdHours")
    sm = m["comp_fits"]
    return dict(
        r_no3=r_no3, n_no3=len(no3), three=three, r_fo_start=r_fo_start, r_reg=r_reg, r_reg_alt=r_reg_alt, r_dreg_alt=r_dreg_alt,
        gap_alt=round(gap_alt, 2), with_closes=round(with_closes, 2), closes_extra=round(with_closes - m["inputs"], 2),
        n_q10=len(q10), n_q11=len(q11), n_q14=len(q14), f=f_by, n_start=len(with_start), n_test=len(months),
        r_dreg=sm[dk]["r"], gap=m["comp_gap"])


# --- Requirement 3: activities and capacity --------------------------------------------------------------------------------

@context("ch18.g3")
def g3(d, claim):
    need(d, claim)
    m = n18.model(d)
    C = d.C
    first, last = f"{C}-01-01", f"{C}-12-31"
    with temp_view(d, "FamilyActivity"):
        q16 = run(d, quote("Requirement 3, Q16"))
        q17 = run(d, quote("Requirement 3, Q17"))
        q18 = run(d, quote("Requirement 3, Q18"))
        claim(len(q16) == len(m["fam"]) and len(q17) == len(q16) and len(q18) == 3, "the activity queries return a row for each family and group")
        # the report's batch: units over work orders (a work order counts once however many completions)
        furn = next(r for r in q18 if r[0] == n18.STORY_GROUP)
    q19 = run(d, quote("Requirement 3, Q19"))
    q20 = run(d, quote("Requirement 3, Q20"))
    claim(len(q19) == 5 and len(q20) == 6, "the calendar has five work centers and the capacity table adds the plant")
    # wrong version 1: setup hours counted once for every completion instead of once for every work order
    setup_wrong = d.one("SELECT SUM(ro.StandardSetupHours) FROM ProductionCompletionLine l JOIN ProductionCompletion pc ON "
                        "pc.ProductionCompletionID = l.ProductionCompletionID JOIN WorkOrder wo ON wo.WorkOrderID = pc.WorkOrderID "
                        "JOIN RoutingOperation ro ON ro.RoutingID = wo.RoutingID WHERE pc.CompletionDate BETWEEN ? AND ?", first, last)
    claim(setup_wrong > m["setup_total"] * 1.05, "counting setup once for every completion line overstates it")
    # wrong version 2: issue lines joined through the completions
    issues_wrong = d.one("SELECT COUNT(*) FROM MaterialIssueLine mil JOIN MaterialIssue mi ON mi.MaterialIssueID = mil.MaterialIssueID "
                         "JOIN ProductionCompletion pc ON pc.WorkOrderID = mi.WorkOrderID WHERE pc.CompletionDate BETWEEN ? AND ?", first, last)
    issues_true = d.one("SELECT COUNT(*) FROM MaterialIssueLine mil JOIN MaterialIssue mi ON mi.MaterialIssueID = mil.MaterialIssueID "
                        "WHERE mi.WorkOrderID IN (SELECT WorkOrderID FROM ProductionCompletion WHERE CompletionDate BETWEEN ? AND ?)",
                        first, last)
    claim(issues_wrong > issues_true, "joining the issue lines through the completions repeats them for a work order completed in parts")
    # wrong version 3: a completion counted as a work order
    n_completions = d.one("SELECT COUNT(*) FROM ProductionCompletion WHERE CompletionDate BETWEEN ? AND ?", first, last)
    n_orders = d.one("SELECT COUNT(DISTINCT WorkOrderID) FROM ProductionCompletion WHERE CompletionDate BETWEEN ? AND ?", first, last)
    claim(n_completions > n_orders, "some work orders are completed in more than one part")
    furn_units = furn[2]
    furn_completions = d.one("SELECT COUNT(*) FROM ProductionCompletion pc JOIN WorkOrder wo ON wo.WorkOrderID = pc.WorkOrderID JOIN Item i "
                             "ON i.ItemID = wo.ItemID WHERE pc.CompletionDate BETWEEN ? AND ? AND i.ItemGroup = ?", first, last, n18.STORY_GROUP)
    # wrong version 4: the recorded direct time as the measure of use
    recorded = {r[0]: r[1] for r in d.q(
        "SELECT woo.WorkCenterID, SUM(lte.RegularHours + lte.OvertimeHours) FROM LaborTimeEntry lte JOIN WorkOrderOperation woo ON woo.WorkOrderOperationID = "
        "lte.WorkOrderOperationID WHERE lte.LaborType = 'Direct Manufacturing' AND lte.WorkDate BETWEEN ? AND ? GROUP BY 1", first, last)}
    centers = {c["id"]: c for c in m["centers"]}
    over = {k: recorded.get(k, 0.0) / centers[k]["pract"] for k in centers}
    top = max(over, key=over.get)
    claim(over[top] > 1 and centers[top]["short"] == "Packing", "the direct hours recorded at Packing exceed its practical capacity")
    claim(recorded[top] > centers[top]["avail"], "the direct hours recorded at Packing exceed even its available hours")
    low = min(over, key=over.get)
    return dict(
        n_q16=len(q16), n_q17=len(q17), n_q18=len(q18), n_q19=len(q19), n_q20=len(q20), setup_wrong=round(setup_wrong, 2),
        setup_right=round(m["setup_total"], 2), issues_wrong=issues_wrong, issues_true=issues_true, n_completions=n_completions,
        n_orders=n_orders, furn_batch_wrong=furn_units / furn_completions, furn_batch=furn[3],
        rec_top=dict(name=centers[top]["short"], hours=recorded[top], pract=centers[top]["pract"], avail=centers[top]["avail"],
                     times=over[top]), rec_low=dict(name=centers[low]["short"], times=over[low]),
        n_plant=len(centers))


# --- Requirement 4: three costs for each family ----------------------------------------------------------------------------

@context("ch18.g4")
def g4(d, claim):
    need(d, claim)
    from collections import defaultdict
    m = n18.model(d)
    C = d.C
    first, last = f"{C}-01-01", f"{C}-12-31"
    q21 = run(d, quote("Requirement 4, Q21"))
    claim(len(q21) == 2 and {r[0] for r in q21} == {"Hourly", "Salary"}, "the pay class query returns the hourly crew and the salaried staff")
    hourly = next(r for r in q21 if r[0] == "Hourly")
    claim(hourly[1] == hourly[2] == m["pclass"]["Hourly"]["n"], "the crew is every hourly manufacturing employee on record")
    # wrong version 1: the rate on the available hours
    rate_avail = m["committed"] / m["avail"]
    # wrong version 2: the Factory Overhead entries counted twice
    rate_twice = (m["committed"] + m["foh"]) / m["pract"]
    twice_gap = m["pract"] * rate_twice + m["foh"] - m["inputs"]
    claim(abs(twice_gap - m["foh"]) < 0.01, "a rate that includes the Factory Overhead entries, plus those entries per standard hour, overstates the total by exactly them")
    # wrong version 3: setup left out of the routing time
    setup = dict(d.q("SELECT RoutingID, SUM(StandardSetupHours) FROM RoutingOperation GROUP BY 1"))
    ex = m["fam"][n18.EXAMPLE]
    items, sold, batch = m["items"], m["sold"], m["batch"]
    su = sum(sold[i][0] * setup[items[i]["routing"]] / batch[n18.EXAMPLE] for i in ex["ids"]) / ex["units"]
    td_no_setup = ex["unit"]["td"] - su * m["td_rate"]
    claim(td_no_setup < ex["unit"]["td"], "leaving setup out lowers the family's time-driven cost")
    # wrong version 4: Power Query without the year filter
    rev_all = d.one("SELECT SUM(l.LineTotal) FROM SalesInvoiceLine l JOIN Item i ON i.ItemID = l.ItemID WHERE i.ItemType = 'Finished Good'")
    units_all = d.one("SELECT SUM(l.Quantity) FROM SalesInvoiceLine l JOIN Item i ON i.ItemID = l.ItemID WHERE i.ItemType = 'Finished Good' "
                      "AND i.SupplyMode = 'Manufactured'")
    units_year = sum(f["units"] for f in m["fam"].values())
    claim(rev_all > 2.5 * m["explorer"]["rev"], "three years of invoice lines carry about three times the year's revenue")
    # wrong version 5: the plain mean of the family margins
    fur = [f for f in m["fam"].values() if f["group"] == n18.STORY_GROUP]
    mean_td = sum(f["unit"]["m_td"] for f in fur) / len(fur)
    group = next(g for g in m["groups"] if g["name"] == n18.STORY_GROUP)
    claim(abs(mean_td - group["m_td"]) > 0.001, "the plain mean of the families' margins differs from the group's margin")
    # the workbook
    cols = {c: item_formula(claim, c) for c in ("Material", "MaterialAtActual", "CostKept", "RoutingTime", "AbsorptionCost", "TimeDrivenCost",
                                                "CostReduced", "Batch", "VariancePerUnit")}
    cells = {k: formula(claim, "Costs", c) for k, c in (("committed", "B12"), ("rate", "B16"), ("practical", "B15"), ("products", "B24"),
                                                        ("unused", "B26"), ("setup_diff", "B34"), ("fam_td", "G48"), ("top", "B67"),
                                                        ("group", "E60"), ("recon", "B30"))}
    dt = workbook()["Costs"]["B77"].value
    claim(getattr(dt, "ref", "") == "B77:D80" and getattr(dt, "r1", "") == "B14", "the Data Table of Costs reads its column input from B14")
    claim(defined(claim, "PracticalUsed") == "Costs!$B$14" and defined(claim, "PracticalShare") == "Assumptions!$B$22",
          "PracticalUsed is the cell on the Costs sheet and PracticalShare the Assumptions value")
    claim(len(m["quoted"]) == 6 and ex["name"] == n18.EXAMPLE, "the example family is one the quote names")
    return dict(
        n_q21=len(q21), rate_avail=rate_avail, rate_twice=rate_twice, twice_gap=round(twice_gap, 2), su=su, td_no_setup=td_no_setup,
        rev_all=round(rev_all, 2), units_all=units_all, units_year=units_year, mean_td=mean_td, group_td=group["m_td"],
        cols=cols, cells=cells, dt_ref=dt.ref, n_ic=219, n_fg=d.one("SELECT COUNT(*) FROM Item WHERE ItemType = 'Finished Good'"),
        n_input=m["explorer"]["n_items"], n_hours=len(m["pclass"]), example=n18.EXAMPLE,
        ex=dict(price=ex["unit"]["price"], std=ex["unit"]["std"], absorp=ex["unit"]["absorp"], td=ex["unit"]["td"]))


# --- Requirement 5: the explorer -------------------------------------------------------------------------------------------

@context("ch18.g5")
def g5(d, claim):
    need(d, claim)
    m = n18.model(d)
    C = d.C
    e = m["explorer"]
    first, last = f"{C}-01-01", f"{C}-12-31"
    sections = check_sections()
    claim(sections.get("Requirement 5", 0) > 150, "the Checks tab holds a check for every view, family and total of Requirement 5")
    for page, names in (("families", ("familyMatrix", "itemsChart", "viewSlicer", "viewTotals", "viewCards")),
                        ("capacity", ("capacityCards", "capacityChart", "capacityTable", "familyHours")),
                        ("validation", ("byGroup", "byYear", "inputs1090", "ledgerCards", "viewsWorkbook", "valAnswer")),
                        ("about", ("aboutText", "lastInvoice", "pageNavigator")),
                        ("concentration", ("curve", "concAnswer"))):
        has_visual(claim, page, *names)
    tests = (MODEL / "DAXQueries" / "Tests.dax").read_text(encoding="utf-8")
    n_tests = len(re.findall(r"(?m)^// Test \d", tests))
    # wrong version 1: no year filter (the explorer then holds three years of revenue against one year's item costs)
    r_all, c_all = d.q("SELECT SUM(l.LineTotal), SUM(l.Quantity * i.StandardCost) FROM SalesInvoiceLine l JOIN Item i ON i.ItemID = l.ItemID "
                       "WHERE i.ItemType = 'Finished Good'")[0]
    margin_all = 1 - c_all / r_all
    margin_year = e["std"] / e["rev"]
    claim(abs(margin_all - margin_year) < 0.02 and r_all > 2.5 * e["rev"],
          "without the year filter the standard margin still looks plausible while revenue is three years' worth")
    # wrong version 2: the design services left in the margin
    margin_with = 1 - e["all_std"] / e["all_rev"]
    claim(margin_with > margin_year + 0.01, "the design services, with no standard cost, raise the margin when they are left in")
    blank_lines = d.one("SELECT COUNT(*) FROM SalesInvoiceLine WHERE ItemID IN (SELECT ItemID FROM Item WHERE ItemType <> 'Finished Good')")
    claim(blank_lines > 0 and abs(e["services"] - (e["all_rev"] - e["rev"])) < 0.005, "the lines with no item cost row are the design services")
    ic = table_m("ItemCosts")
    si = table_m("SalesInvoice")
    cv = _model_text("tables/Costing View.tmdl")
    claim("Standard" in cv or "Table.FromRows" in cv, "the Costing View table is typed with Enter data")
    km = _model_text("tables/Key Measures.tmdl")
    share_read = "SELECTEDVALUE ( 'Model Parameters'[PracticalShare] )" in km
    claim(share_read or "'Practical Share' = 0.8" in km,
          "the Practical Share measure is typed (the first build) or read from the workbook's name PracticalUsed through the Model Parameters query")
    # the practical share changed to 85%: the probe that tells a drifting share from a synchronized one (MOD18-03)
    p85 = m["avail"] * 0.85
    rate85 = m["committed"] / p85
    at85 = dict(share=0.85, pract=p85, rate=rate85, unused_cost=(p85 - m["used"]) * rate85)
    claim(abs(at85["rate"] - m["td_rate"]) > 1 and abs(at85["pract"] - m["pract"]) > 1000,
          "a different practical share moves the capacity, the rate and the unused capacity's cost")
    rel = _model_text("relationships.tmdl")
    claim("SalesInvoiceLine.ItemID" in rel and "toColumn: ItemCosts.ItemID" in rel, "the lines relate to the item costs on ItemID")
    return dict(
        n_checks=sections.get("Requirement 5", 0), n_milestone=sections.get("Milestone 2", 0), n_tests=n_tests,
        n_ic=d.one("SELECT COUNT(*) FROM Item WHERE ItemType = 'Finished Good'"), n_lines=d.one("SELECT COUNT(*) FROM SalesInvoiceLine"),
        n_invoices=d.one("SELECT COUNT(*) FROM SalesInvoice"), margin_all=margin_all, margin_year=margin_year,
        margin_with=margin_with, rev_all=round(r_all, 2), blank_lines=blank_lines, blank_rev=e["services"],
        m_year_step=step(claim, si, "Added Custom"), m_cols=step(claim, ic, "Removed Other Columns"),
        d_views=dax(claim, "Revenue", "Units", "Standard Cost", "Absorption Cost", "Time-Driven Cost"),
        d_switch=dax(claim, "Product Cost", "Margin", "Margin %", "Cost per Unit"),
        d_share=dax(claim, "Practical Share"), share_read=share_read, at85=at85, base_share=n18.PRACTICAL, base_pract=m["pract"], base_unused_cost=m["unused_h"] * m["td_rate"],
        c_nodefault=check_expr(claim, "2026 margin with no view selected (falls back to Standard)"),
        c_diff=check_expr(claim, "2026 margin less the workbook's, Time-driven"),
        c_blank=check_expr(claim, "lines with no item cost row (the design services, every year)"),
        t_matrix=title(claim, "families", "familyMatrix"), t_chart=title(claim, "families", "itemsChart"),
        t_slicer=title(claim, "families", "viewSlicer"), t_totals=title(claim, "families", "viewTotals"),
        t_cap_chart=title(claim, "capacity", "capacityChart"), t_cap_table=title(claim, "capacity", "capacityTable"),
        t_family_hours=title(claim, "capacity", "familyHours"), t_by_group=title(claim, "validation", "byGroup"),
        t_by_year=title(claim, "validation", "byYear"), t_inputs=title(claim, "validation", "inputs1090"),
        t_views_wb=title(claim, "validation", "viewsWorkbook"))


# --- helpers for Requirements 6 to 10 ---------------------------------------------------------------------------------------

def row_of(claim, sheet: str, text: str, col: int = 1) -> int:
    """The row of a worksheet of the model workbook whose cell in `col` starts with `text`."""
    ws = workbook()[sheet]
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and v.startswith(text):
            return r
    claim(False, f"the worksheet {sheet} has a row that starts {text!r}")
    return 1


def table_formula(claim, sheet: str, table: str, column: str) -> str:
    """The formula of one calculated column of an Excel Table of the model workbook (the first data row)."""
    from openpyxl.utils import range_boundaries
    ws = workbook()[sheet]
    t = ws.tables.get(table)
    claim(t is not None, f"the worksheet {sheet} has the Table {table}")
    if t is None:
        return ""
    c0, r0, _, _ = range_boundaries(t.ref)
    names = [c.name for c in t.tableColumns]
    claim(column in names, f"the Table {table} has the column {column}")
    return clean(ws.cell(r0 + 1, c0 + names.index(column)).value) if column in names else ""


def cell_value(sheet: str, cell: str):
    return workbook()[sheet][cell].value


def two_way(m: dict, share: float, adj: float) -> float:
    """The annual effect of buying the families to buy, at a share of the freed crew time removed and a quote adjustment."""
    return sum(((x["kept"] + share * m["crew_rate"] * x["time"]) - x["price"] * (1 + adj) * (1 + n18.FREIGHT)) * x["units"]
               for x in m["buy"])


# --- Requirement 6: make or buy ---------------------------------------------------------------------------------------------

@context("ch18.g6")
def g6(d, claim):
    need(d, claim)
    m = n18.model(d)
    n18.make_buy_claims(m, claim)
    quoted, buy, lose = m["quoted"], m["buy"], m["lose"]
    names = [x["name"] for x in buy]
    ws = workbook()["Make or Buy"]
    claim([ws.cell(r, 1).value for r in range(10, 10 + len(quoted))] == [x["name"] for x in quoted],
          "the Make or Buy sheet lists the quoted families in the quote's order from row 10")
    claim(all(x["a_kept"] <= 0 for x in quoted), "with the crew kept, no quoted family is worth buying")
    # the SQL: the materials of the families to buy (Q22)
    q22 = quote("Requirement 6, Q22")
    claim(all(q22.count(f"'{f}'") == 2 for f in names) and len(re.findall(r"'[A-Z]{3}-[A-Z]{3}'", q22)) == 2 * len(names),
          "Query 22 names exactly the families to buy, in its two lists")
    row = run(d, q22)
    claim(len(row) == 1 and all(abs(a - b) < 0.006 for a, b in zip(row[0], (m["buy_material"], m["fam_issued"], m["six_issued"], m["issues"]))),
          "Query 22 returns the note's materials row")
    # wrong version 1: the families by the three-letter product type (FUR-TBL and LGT-TBL both end in TBL)
    by_type = run(d, sub(q22, "SUBSTR(i.ItemCode, 1, 7) IN ('FUR-BNH', 'FUR-NGT', 'FUR-TBL')", "SUBSTR(i.ItemCode, 5, 3) IN ('BNH', 'NGT', 'TBL')"))[0][1]
    claim(by_type > m["fam_issued"] * 1.05, "grouping the families by their three-letter type adds the table lamps to the dining tables' issues")
    # wrong version 2: the CFO's absorption view
    absorp_sum = sum(x["absorp"] for x in quoted if x["absorp"] > 0)
    # wrong version 3: the quote without its inbound freight (crew reduced)
    nf = [x for x in quoted if (x["reduced"] - x["price"]) * x["units"] > 0]
    nf_flip = [x["short"] for x in nf if x not in buy]
    claim(len(nf_flip) > 0, "leaving the freight out makes a family that loses money look worth buying")
    # wrong version 4: the whole time-driven rate as the crew's cost
    tdw = [x for x in quoted if (x["kept"] + m["td_rate"] * x["time"] - x["delivered"]) * x["units"] > 0]
    tdw_flip = [x["short"] for x in tdw if x not in buy]
    claim(len(tdw_flip) > 0, "charging the whole time-driven rate as the crew's cost makes a family that loses money look worth buying")
    # wrong version 5: positions on the available hours of a position instead of its practical hours
    pos_wrong = m["hours"] / (m["pos_hours"] / n18.PRACTICAL)
    sev_wrong = pos_wrong * m["severance"]
    # wrong version 6: one year of savings against the one-time costs
    net_one = m["gain"] - m["positions"] * m["severance"] - m["tooling"]
    # wrong version 7: the Data Table's input cells swapped
    corner = two_way(m, 1.0, 0.0)
    claim(abs(corner - m["gain"]) < 0.01, "the Data Table's corner (all the crew time removed, the quote as it is) is the annual effect")
    claim(all(abs(two_way(m, s, a) - m["two_way"][i][j]) < 0.01 for i, s in enumerate((0.0, 0.5, 1.0)) for j, a in enumerate((-0.10, 0.0, 0.10))),
          "the note's two-way table is the effect at each share of the crew time and each quote adjustment")
    swapped = two_way(m, 0.0, 1.0)
    dt = ws["B42"].value
    claim(getattr(dt, "ref", "") == "B42:D44" and getattr(dt, "r1", "") == "B5" and getattr(dt, "r2", "") == "B6",
          "the Data Table reads its row input from B5 (the quote) and its column input from B6 (the crew share)")
    be_share = -sum(x["a_kept"] for x in buy) / (m["gain"] - sum(x["a_kept"] for x in buy))
    claim(0 < be_share < 1, "the three families break even at a share of the freed crew time between none and all of it")
    payback = 12 * (m["positions"] * m["severance"] + m["tooling"]) / m["gain"]
    gs = [abs(ws[f"B{19 + i}"].value - x["be_kept"]) < 0.005 and abs(ws[f"D{19 + i}"].value - x["be_reduced"]) < 0.005
          for i, x in enumerate(quoted)]
    claim(all(gs), "the workbook's Goal Seek cells hold the break-even quotes")
    f = {k: formula(claim, "Make or Buy", c) for k, c in (
        ("deliv", "G10"), ("kept", "I10"), ("reduced", "J10"), ("abs_eff", "L10"), ("kept_eff", "M10"), ("red_eff", "N10"),
        ("decision", "O10"), ("flag", "P10"), ("be_kept", "Q10"), ("be_red", "R10"), ("hours", "T10"), ("pos", "U10"),
        ("sens", "V10"), ("gs_kept", "C19"), ("gs_red", "E19"), ("families", "B27"), ("gain", "B28"), ("severance", "B33"),
        ("tooling", "B34"), ("two_years", "B35"), ("corner", "A41"))}
    f["freight"], f["crew_removed"], f["quote_adj"] = (defined(claim, n) for n in ("Freight", "CrewRemoved", "QuoteAdjustment"))
    # the two-year gate on each family (columns X to Z of the sheet, MOD18-04): present once the workbook is built with it
    gate = str(ws["Y9"].value or "").startswith("Effect over the two years")
    if gate:
        f["once"], f["two"], f["be_two"] = (formula(claim, "Make or Buy", c) for c in ("X10", "Y10", "Z10"))
        claim(all(q["two_year"] > 0 for q in buy) and all(q["two_year"] < 0 for q in lose),
              "each family to buy covers its own one-time costs in the two years, and each family to make does not")
    return dict(
        gate=gate, f=f, example=buy[0]["short"], by_type=round(by_type, 2), absorp_sum=round(absorp_sum, 2), nf_flip=nf_flip,
        nf_buy=word_of(len(nf)), tdw_flip=tdw_flip, tdw_buy=word_of(len(tdw)), pos_wrong=pos_wrong, sev_wrong=round(sev_wrong, 2),
        net_one=round(net_one, 2), net_free=round(2 * m["gain"], 2), swapped=round(swapped, 2), n_q22=len(row),
        be_share=be_share, payback=payback)


def word_of(n: int) -> str:
    return n18.word(n)


# --- Requirement 7: freed capacity and the growth test -------------------------------------------------------------------------

@lru_cache(maxsize=None)
def _lp_inputs(d):
    """The routing load of each family's base-year sales at each work center, and the calendar's available hours."""
    from collections import defaultdict
    m = n18.model(d)
    items, fam, sold, batch = m["items"], m["fam"], m["sold"], m["batch"]
    run_wc, setup_wc = defaultdict(dict), defaultdict(dict)
    for rid, wc, sh, rh in d.q("SELECT RoutingID, WorkCenterID, SUM(StandardSetupHours), SUM(StandardRunHoursPerUnit) FROM RoutingOperation GROUP BY 1, 2"):
        run_wc[rid][wc], setup_wc[rid][wc] = rh, sh
    avail = dict(d.q("SELECT WorkCenterID, SUM(AvailableHours) FROM WorkCenterCalendar WHERE CalendarDate BETWEEN ? AND ? GROUP BY 1",
                     f"{d.C}-01-01", f"{d.C}-12-31"))
    load = {f: defaultdict(float) for f in fam}
    for f, v in fam.items():
        for i in v["ids"]:
            rid = items[i]["routing"]
            for wc, rh in run_wc[rid].items():
                load[f][wc] += sold[i][0] * (rh + setup_wc[rid][wc] / batch[f])
    return load, avail


def _solve(d, growth: float, share: float = n18.PRACTICAL, sense: int = 1, cap: bool = True):
    """Solver's model of Requirement 7, one variation at a time: minimize (sense 1) or maximize (-1) the added cost, with the
    work centers' limit at `share` of the available hours, and the units bought capped at the demand at growth or not."""
    from scipy.optimize import linprog
    m = n18.model(d)
    load, avail = _lp_inputs(d)
    qf, wcs = m["quoted"], m["wcs"]
    a_ub = [[-load[q["name"]][wc] / q["units"] for q in qf] for wc in wcs]
    b_ub = [avail[wc] * share - sum(load[f][wc] * (1 + growth) for f in m["fam"]) for wc in wcs]
    bounds = [(0, q["units"] * (1 + growth) if cap else None) for q in qf]
    return linprog([sense * (q["delivered"] - q["kept"]) for q in qf], A_ub=a_ub, b_ub=b_ub, bounds=bounds, method="highs")


@context("ch18.g7")
def g7(d, claim):
    need(d, claim)
    m = n18.model(d)
    first, last = f"{d.C}-01-01", f"{d.C}-12-31"
    names = [x["name"] for x in m["buy"]]
    quoted = m["quoted"]
    total = sum(m["freed"].values())
    with temp_view(d, "FamilyActivity"):
        q23 = run(d, quote("Requirement 7, Q23"))
    claim(len(q23) == 1 and abs(q23[0][5] - total) < 0.06 and abs(q23[0][6] - m["used"] / m["pract"]) < 0.0006
          and abs(q23[0][8] - m["unused_h"]) < 0.06, "Query 23 returns the note's freed hours, use and unused hours")
    claim(all(f"'{f}'" in quote("Requirement 7, Q23") for f in names), "Query 23 names the families to buy")
    # the two bases of the hours freed: the year's completions (the plant) and the year's sales (the make-or-buy table)
    claim(abs(total - m["hours"]) > 1, "the hours freed on the year's completions differ from the routing time of the year's sales")
    # wrong version 1: run hours only (setup left out)
    marks = ",".join("?" * len(names))
    run_only = d.one(f"SELECT SUM(l.QuantityCompleted * i.StandardLaborHoursPerUnit) FROM ProductionCompletionLine l JOIN ProductionCompletion pc "
                     f"ON pc.ProductionCompletionID = l.ProductionCompletionID JOIN Item i ON i.ItemID = l.ItemID WHERE pc.CompletionDate "
                     f"BETWEEN ? AND ? AND SUBSTR(i.ItemCode, 1, 7) IN ({marks})", first, last, *names)
    claim(run_only < total, "leaving setup out frees fewer hours")
    # wrong version 2: use measured on the available hours instead of the practical hours
    avail_before, avail_after = m["used"] / m["avail"], (m["used"] - total) / m["avail"]
    # wrong versions 3 to 5: Solver set to maximize, the demand limits left out, the work centers limited at the available hours
    wrong = []
    for g in n18.GROWTH:
        base = _solve(d, g)
        mx = _solve(d, g, sense=-1)
        nocap = _solve(d, g, cap=False)
        av = _solve(d, g, share=1.0)
        claim(base.status == 0 and mx.status == 0 and nocap.status == 0 and av.status == 0, "each variation of the Solver model solves")
        demand = {q["name"]: q["units"] * (1 + g) for q in quoted}
        over = [dict(short=q["short"], units=float(x), demand=demand[q["name"]]) for q, x in zip(quoted, nocap.x) if x > demand[q["name"]] + 0.5]
        wrong.append(dict(growth=g, max_cost=float(-mx.fun), max_n=sum(1 for x in mx.x if x > 0.5), over=over,
                          avail_n=sum(1 for x in av.x if x > 0.5), avail_cost=float(av.fun)))
    claim(all(w["max_cost"] > 10 * m["lp"][k]["cost"] and w["max_n"] == len(quoted) for k, w in enumerate(wrong)),
          "maximizing instead of minimizing buys every family and costs more than ten times the optimum")
    claim(any(w["over"] for w in wrong), "without the demand limits, Solver buys more of a family than it would sell at the higher growth")
    claim(all(w["avail_n"] == 0 for w in wrong), "limiting the work centers at their available hours, the model buys nothing at either growth rate")
    ws40 = workbook()[n18_growth_name(0)]
    claim(all(isinstance(ws40[f"L{r}"].value, (int, float)) for r in range(14, 20)) and ws40["B21"].value is not None,
          "the growth sheet holds Solver's variable cells (L14 to L19) and its objective (B21)")
    sens = workbook()[n18_growth_name(0) + " Sensitivity"]
    shadow = m["lp"][0]["binding"][0]["price"]
    claim(any(isinstance(c.value, (int, float)) and abs(c.value + shadow) < 1e-3 for row_ in sens.iter_rows() for c in row_),
          "the Sensitivity sheet reports the binding work center's shadow price as a negative number (the model minimizes)")
    f = {k: formula(claim, "Capacity", c) for k, c in (("freed", "C6"), ("after", "D6"), ("use", "F6"), ("use_after", "G6"),
                                                         ("unused_pos", "B15"), ("crew_cost", "B21"), ("support", "B22"),
                                                         ("remain", "B24"), ("spread", "B25"), ("td_after", "B26"),
                                                         ("head", "D30"), ("first", "B35"))}
    gname = n18_growth_name(0)
    f.update({k: formula(claim, gname, c) for k, c in (("load_growth", "D6"), ("used", "E6"), ("status", "G6"), ("extra", "F14"),
                                                        ("cost", "B21"), ("minimum", "O14"), ("plan_added", "Q14"), ("plan_all", "B30"))})
    remain_h = sum(v["unit"]["time"] * v["units"] for v in m["fam"].values()) - m["hours"]
    return dict(
        f=f, n_q23=len(q23), positions=m["positions"], remain_h=remain_h, spread=m["hours"] * m["td_rate"] / remain_h,
        unused_after=(m["unused_h"] + total) * m["td_rate"], mixed_after=(m["unused_h"] + m["hours"]) * m["td_rate"],
        run_only=round(run_only, 2), avail_before=avail_before, avail_after=avail_after, wrong=wrong,
        sales_hours=m["hours"], sheet=gname, sheets=[n18_growth_name(i) for i in range(len(n18.GROWTH))],
        growths=list(n18.GROWTH), n_wcs=len(m["wcs"]))


def n18_growth_name(k: int) -> str:
    return f"Growth {n18.GROWTH[k]:.0%}"


# --- Requirement 8: the floor prices and the evidence on customers --------------------------------------------------------------

@context("ch18.g8")
def g8(d, claim):
    need(d, claim)
    from notes.case2 import volume
    m = n18.model(d)
    fam = m["fam"]
    q24, q25, q26 = (run(d, quote(f"Requirement 8, Q{k}")) for k in (24, 25, 26))
    vol = volume(d)
    claim(len(q24) == len(vol) and len(q25) == d.one("SELECT COUNT(*) FROM PriceList") and len(q26) == 1 and q26[0][0] == "Sales Manager",
          "the three queries return a row for each promotion, each price list and the one approver")
    expired = sorted(r[0] for r in q25 if r[3] == "Expired")
    claim(all(r[8] > 0 for r in q25 if r[3] == "Expired") and all(r[8] == 0 for r in q25 if r[3] != "Expired"),
          "only the expired price lists have order lines after their end date")
    lines_after = sum(r[8] for r in q25 if r[3] == "Expired")
    active = [r for r in q25 if r[3] == "Active" and r[6] is not None]
    # wrong version 1: a markup of the floor on the cost instead of a margin on the price
    mk, loss_td = [], []
    for r in m["reprice"]:
        w = fam[r["name"]]["unit"]
        mk.append(dict(short=r["short"], incr=w["td"] * (1 + n18.FLOOR) / w["price"] - 1))
        cm = w["price"] - w["td"]
        loss_td.append(dict(short=r["short"], loss=1 - cm / (cm + r["increase"] * w["price"]), right=r["loss"]))
    claim(all(x["incr"] < 0 for x in mk), "a 40% markup on time-driven cost is below today's price for every family below the floor")
    claim(all(x["loss"] > x["right"] for x in loss_td),
          "measuring contribution over time-driven cost overstates the volume an increase can lose")
    # wrong version 2: the increases of the families that would be bought, counted as well
    buy_names = {q["name"] for q in m["buy"]}
    all_rev = sum(r["revenue"] for r in m["reprice"])
    kept_rev = sum(r["revenue"] for r in m["reprice"] if r["name"] not in buy_names)
    claim(all_rev > kept_rev, "the families bought are among those below the floor")
    # wrong version 3: the months of a promotion taken from its effective dates
    ends = d.q("SELECT PromotionID, EffectiveStartDate, EffectiveEndDate FROM PromotionProgram WHERE EffectiveEndDate <= EffectiveStartDate")
    n_before = sum(1 for _, s, e in ends if e < s)
    n_same = sum(1 for _, s, e in ends if e == s)
    claim(len(ends) == 3 and n_before == 2, "three promotions end on or before their start date, two of them before it")
    ex = m["buy"][0]["name"]
    claim(ex in [r["name"] for r in m["reprice"]], "the first family to buy is also below the floor")
    b = row_of(claim, "Prices", ex)
    f = {k: formula(claim, "Prices", f"{c}{b}") for k, c in (("margin", "E"), ("below", "F"), ("check", "H"), ("direct", "I"), ("incr", "J"),
                                                              ("contrib", "L"), ("loss", "M"), ("revenue", "N"))}
    claim(abs(cell_value("Prices", f"G{b}") - fam[ex]["unit"]["td"] / (1 - n18.FLOOR)) < 0.005,
          "the workbook's Goal Seek result for the first family is the floor price")
    # the increase, the volume lost and the revenue read the direct price, with a column that says Goal Seek still agrees (MOD18-05)
    agrees = str(workbook()["Prices"]["O4"].value or "").startswith("Goal Seek agrees")
    f.update({k: table_formula(claim, "Prices", "PromotionVolume", c) for k, c in (("months", "Months"), ("value", "PromotionValue"),
                                                                                    ("expected", "Expected"), ("lift", "Lift"), ("z", "Z"))})
    f["total_rev"], f["still"], f["still_rev"] = (formula(claim, "Prices", f"B{row_of(claim, 'Prices', t)}")
                                                   for t in ("Revenue a year if volume holds", "If the families of Requirement 6 are bought",
                                                             "  their revenue"))
    f["bought_margin"] = formula(claim, "Prices", f"B{row_of(claim, 'Prices', 'Delivered margins') + 2}")
    w_ex = fam[ex]["unit"]
    end_date = sorted({r[5] for r in q25 if r[3] == "Active"})
    claim(len(end_date) == 1, "every active price list ends on one date")
    return dict(
        agrees=agrees, f=f, example=n18.short(ex), n_below=n18.word(len(m["reprice"])), n_rest=n18.word(len(m["reprice"]) - len(buy_names & {r["name"] for r in m["reprice"]})),
        ex=dict(margin=w_ex["m_td"], price=w_ex["price"], td=w_ex["td"], kept=w_ex["kept"], new=w_ex["td"] / (1 - n18.FLOOR)),
        end_date=end_date[0], n_q24=len(q24), n_q25=len(q25), n_q26=len(q26), expired=expired, lines_after=lines_after,
        low=min(r[6] for r in active), high=max(r[7] for r in active), mk=mk, loss_td=loss_td, all_rev=round(all_rev, 2),
        kept_rev=round(kept_rev, 2), n_before=n_before, n_same=n_same, n_inverted=len(ends), n_lists=len(q25),
        ratio_bought=[dict(short=q["short"], margin=q["margin"]) for q in m["buy"]], approver=q26[0][0], approved=q26[0][1])


# --- Requirement 9: the standard and the statements -----------------------------------------------------------------------------

@context("ch18.g9")
def g9(d, claim):
    need(d, claim)
    m = n18.model(d)
    P, Cc, Fy = d.P, d.C, d.F
    # the rates on one basis: per standard hour completed (N18-01)
    stdh = m["stdh"]
    used_conv = m["used"] * m["td_rate"] + m["foh"]
    unused_cost = m["unused_h"] * m["td_rate"]
    claim(abs(used_conv / stdh - m["td_std"]) < 1e-9,
          "the time-driven rate is the Factory Overhead entries plus the routing time used at the committed rate, per standard hour")
    claim(abs((used_conv + unused_cost) / stdh - m["actual_rate"]) < 1e-9,
          "the time-driven conversion plus the unused capacity's cost is the actual rate, so the time-driven rate leaves the unused capacity out")
    claim(abs((m["conversion"] + unused_cost) / stdh - m["td_std"]) > 1,
          "the released conversion plus the unused capacity's cost is not the time-driven rate")
    rates = dict(td_rate=m["td_rate"], unused_per_hour=unused_cost / stdh, stdh=stdh)
    q27, q28, q29, q30, q31 = (run(d, quote(f"Requirement 9, Q{k}")) for k in (27, 28, 29, 30, 31))
    claim(len(q27) == 1 and abs(q27[0][0] - m["plan_dep"]) < 0.005, "Query 27 returns the next year's depreciation, the plan's")
    claim(len(q29) == len(d.years) and all(abs(r[3] - s["stdh"]) < 0.06 and abs(r[4] - s["conv"]) < 0.006 for r, s in zip(q29, m["stock"])),
          "Query 29 returns the note's standard hours and conversion at each year-end")
    claim(len(q30) == len(d.years) and len(q31) == 1, "Query 30 returns an income for each year and Query 31 the opening finished goods")
    claim(abs(q31[0][1] - d.one("SELECT SUM(Debit) - SUM(Credit) FROM GLEntry WHERE AccountID = ? AND VoucherNumber = ?", d.account("1040"), q31[0][0])) < 0.006,
          "Query 31 is the opening entry's line to finished goods")
    neg = [(s["year"], c) for s in m["stock"] for c in s["negative"]]
    claim(len(neg) == 2, "two items go negative in the rebuild, in two different years")
    s = {x["year"]: x for x in m["stock"]}
    p = m["parts"]
    # wrong versions of the rate
    no_burden = m["plan_salary"] / m["plan_stdh"]
    rate_nb = m["new_rate"] - p["supervision"] + no_burden
    dep_base = d.one("SELECT SUM(g.Debit) FROM GLEntry g JOIN JournalEntry je ON je.JournalEntryID = g.SourceDocumentID AND g.SourceDocumentType = "
                     "'JournalEntry' WHERE g.AccountID = ? AND je.EntryType = 'Depreciation' AND g.FiscalYear = ?", d.account("1090"), Cc)
    rate_dep = m["new_rate"] - p["depreciation"] + dep_base / m["plan_stdh"]
    claim(abs(dep_base - m["plan_dep"]) > 1, "the base year's depreciation differs from the next year's plan")
    full_value = s[Cc]["stdh"] * m["new_rate"]
    mat = {y: n18.materiality(d, y) for y in d.years}
    with_closes = d.one("SELECT SUM(g.Credit) - SUM(g.Debit) FROM GLEntry g JOIN Account a ON a.AccountID = g.AccountID WHERE a.AccountType IN "
                        "('Revenue', 'Expense') AND g.FiscalYear = ?", Cc)
    claim(abs(with_closes) < 0.01, "with the year-end closes left in, revenue and expenses net to zero")
    income = {y: n18.net_income(d, y) for y in d.years}
    claim(s[Cc]["plan"] > mat[Cc] > s[Cc]["plan"] - s[P]["plan"], "the year's balance effect is above the screen and its income effect below")
    ws = workbook()["Standards"]
    rv = row_of(claim, "Standards", "Revaluation at the proposed rate")
    f = {k: formula(claim, "Standards", c) for k, c in (("direct", "B5"), ("variable", "B6"), ("indirect", "B7"), ("supervision", "B8"),
                                                         ("depreciation", "B9"), ("rate", "B10"), ("one", "B11"), ("current", "B12"),
                                                         ("td", "B13"), ("actual", "B14"), ("low", "B15"), ("high", "B16"),
                                                         ("revised", "B25"), ("opening", "B49"), ("income_prior", "B50"),
                                                         ("mat_prior", "B54"), ("both", "B56"))}
    f["units"], f["hours"], f["conv"] = (formula(claim, "Standards", f"B{row_of(claim, 'Standards', t)}")
                                         for t in ("Units on hand", "Standard hours", "Conversion at the current"))
    f["reval"] = formula(claim, "Standards", f"B{rv}")
    f["on_hand"] = table_formula(claim, "Standards", "OnHand", "StdHoursPerUnit")
    claim(ws.cell(rv, 1).value.startswith("Revaluation at the proposed rate"), "the revaluation row is found by its label")
    return dict(
        rates=rates, f=f, n_q28=len(q28), n_items=[r[1] for r in q29], neg=neg, rate_nb=rate_nb, dep_base=round(dep_base, 2), rate_dep=rate_dep,
        full_value=round(full_value, 2), with_closes=round(with_closes, 2), income_prior=income[P], mat=mat,
        bal_cur=s[Cc]["plan"], inc_cur=s[Cc]["plan"] - s[P]["plan"], no_burden=no_burden, years=list(d.years), p=p, hours_cur=s[Cc]["stdh"],
        flags=(s[P]["plan"] > mat[P], s[Cc]["plan"] > mat[Cc], s[P]["one"] > mat[P], s[Cc]["one"] > mat[Cc]))


# --- Requirement 10: the decision table and the memo -----------------------------------------------------------------------------

@context("ch18.g10")
def g10(d, claim):
    need(d, claim)
    m = n18.model(d)
    n18.make_buy_claims(m, claim)
    buy, lose = m["buy"], m["lose"]
    below = [r["name"] for r in m["reprice"]]
    order = [x["name"] for x in buy] + [x["name"] for x in lose] + [f for f in below if f not in n18.QUOTE]
    ws = workbook()["Memo"]
    claim([ws.cell(5 + i, 1).value for i in range(len(order))] == order, "the Memo sheet's decision table lists the families bought, those made and the others below the floor")
    paths = n18.paths_compared(m)
    tot = row_of(claim, "Memo", "Total annual effect")
    claim(tot == 5 + len(order), "the decision table ends in its totals")
    f = {k: formula(claim, "Memo", c) for k, c in (("what", "B5"), ("effect", "C5"), ("once", "D5"), ("total", f"B{tot}"), ("total_once", f"B{tot + 1}"),
                                                    ("depends", "E8"))}
    claim(len(order) == len(buy) + len(lose) + len([f for f in below if f not in n18.QUOTE]), "every family of the table is counted once")
    # wrong versions of the table
    both = m["gain"] + sum(r["revenue"] for r in m["reprice"])
    claim(both > paths["make"] > paths["buy"], "counting both the purchase's saving and the increases of the families bought beats both honest paths")
    unused = m["unused_h"] * m["crew_rate"]
    # the decision table's rows, as the Memo sheet computes them
    inc = {r["name"]: r for r in m["reprice"]}
    q = {x["name"]: x for x in m["quoted"]}
    rows = []
    for name in order:
        if name in {x["name"] for x in buy}:
            x = q[name]
            rows.append(dict(short=n18.short(name), action="Buy only with the crew reduced", effect=x["a_reduced"],
                             once=x["tooling"] + x["time"] * x["units"] / m["pos_hours"] * m["severance"],
                             increase=inc[name]["increase"] if name in inc else None))
        else:
            r = inc.get(name)
            rows.append(dict(short=n18.short(name), action=("Make; raise the price" if r else "Make") if name in q else "Raise the price",
                             effect=r["revenue"] if r else 0.0, once=0.0, increase=r["increase"] if r else None))
    claim(abs(sum(r["effect"] for r in rows) - paths["buy"]) < 0.01 and abs(sum(r["once"] for r in rows) - paths["once"]) < 0.01,
          "the decision table's rows add up to the path that buys and reprices the rest")
    return dict(
        f=f, rows=rows, kept_cost=-sum(x["a_kept"] for x in buy), n_rows=len(order), order=order, n_buy_rows=len(buy), n_make_rows=len(lose), n_only=len([x for x in order if x not in n18.QUOTE]),
        both=round(both, 2), once_tooling=m["tooling"], once=paths["once"], unused_all=m["unused_h"] * m["td_rate"], unused_crew=unused,
        sheets=[ws.title for ws in workbook().worksheets if ws.title in ("Costs", "Make or Buy", "Capacity", "Prices", "Standards", "Memo")])

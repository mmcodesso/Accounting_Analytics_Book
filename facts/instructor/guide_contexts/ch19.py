"""Chapter 19's instructor guides (the capstone case "Auditing the Customer Credits Cycle"): the values the guides
state, and the claims their wording makes.

The case (chapters/19-customer-credits-audit/chapter.qmd) is an independent final-project option with ten requirements
in three phases (three milestones) and no tutorials.  Its model answer is four files, one or two for each tool the case
names:

- Credits.sql, the SQL part (outputs/companion/solutions/Chapter19-capstone-sql/), 40 queries headed
  "Requirement N, test X" (M1 to M8, G1 to G9, P1 and P2, T1 to T12, K1 and K2, S1 to S6, E1),
- Credits Audit.xlsx, the Excel part (outputs/companion/solutions/Chapter19-capstone-excel/), a planning workbook
  whose worksheets (Cycle Map, Reconciliation, Materiality, Risks, Tests, Sample, Timeline, Deficiencies, Report) are
  the appendix of the report and whose Tables are Power Query queries on CharlesRiver.xlsx,
- Credits Analytics, the Power BI project of Requirement 5 (outputs/companion/solutions/Chapter19-capstone-pbi/),
- Audit Monitoring, updated for Requirement 10 (outputs/companion/solutions/Chapter19-capstone-pbi-monitoring/).

The registered notes ch19.r1 to ch19.r10 and ch19.m1 to ch19.m3 (facts/notes/ch19.py) already supply the expected values
of the requirements and milestones.  These contexts add what the model paths need beyond them:

- the text of each model file the guides quote (Excel formulas from the workbook, M and DAX from the Power BI project),
  read from the file itself, so a quote cannot drift from the file the instructor opens (the SQL is quoted into the
  templates verbatim and `guides.py check` verifies it against Credits.sql),
- the symptom of every wrong version a guide names (an anti-join written as an inner join, a join on the document ID
  alone, a year read from the wrong characters of a number, a materiality base with the closes left in, a balance of
  today's status instead of the year-end's, a past-due test on today's balances, a refund test with the wrong
  comparison, a rate read by the wrong date, ...), computed read-only on the dataset, on the model script's own text
  where the wrong version is a change to a query,
- the deck's slide numbers and claims on every table and column the guides name.

What the Power BI project proves and what only Desktop shows: the project encodes the model and the structure of every
page, and its Checks tab (DAX query view) recomputes the values of Requirement 5 and of Milestone 2, so a result is
"proved" when a Checks row names it.  How a visual renders (the matrix's expanded rows, the drill-through menu, the
slicer state) is visible only in Desktop.

Keys:
- ch19.g0: the case note (slides, the model files, counts, the plan and the weights);
- ch19.g1 to ch19.g5: Requirements 1 to 5, each used next to the note ch19.r1 to ch19.r5 that supplies the values of the
  requirement's instructor note.  ch19.g6 to ch19.g10 (Requirements 6 to 10) follow the same pattern and reuse its
  helpers: script(), quote(), sub(), run(), workbook(), values(), formula(), clean(), cell(), dax(), table_m(), step(),
  check_expr(), title(), has_visual(), test_block(), deck_titles(), at(), span(), need().
"""

from __future__ import annotations

import base64
import json
import re
import zlib
from collections import Counter
from functools import lru_cache
from pathlib import Path

from scipy.stats import beta, hypergeom

from guide_contexts import context
from guide_contexts.ch01 import schema
from guide_contexts.ch13 import _block
from notes import ch19 as n19

REPO = Path(__file__).resolve().parents[3]
CHAPTER = REPO / "chapters" / "19-customer-credits-audit" / "chapter.qmd"
DECK = REPO / "slides" / "chapter-19" / "index.qmd"
DOWNLOADS = REPO / "front-matter" / "downloads.qmd"
SOLUTIONS = REPO / "outputs" / "companion" / "solutions"
SQL = SOLUTIONS / "Chapter19-capstone-sql" / "Credits.sql"
XLSX = SOLUTIONS / "Chapter19-capstone-excel" / "Credits Audit.xlsx"
PROJECT = SOLUTIONS / "Chapter19-capstone-pbi" / "Credits Analytics"
MODEL = PROJECT / "Credits Analytics.SemanticModel"
REPORT = PROJECT / "Credits Analytics.Report"
FLAGS_CSV = SOLUTIONS / "Chapter19-capstone-pbi" / "RefundFlags.csv"
MONITOR = SOLUTIONS / "Chapter19-capstone-pbi-monitoring" / "Audit Monitoring" / "Audit Monitoring.SemanticModel"
MON_REPORT = MONITOR.parent / "Audit Monitoring.Report"
XL_BUILDER = REPO / "scripts" / "companion" / "xlbuild" / "cap19.py"
GUIDES = REPO / "facts" / "instructor" / "guides" / "chapter-19-customer-credits-audit"

SCHEMA = {
    "SalesReturn": ["SalesReturnID", "ReturnNumber", "ReturnDate", "CustomerID", "ReceivedByEmployeeID", "ReasonCode", "Status"],
    "SalesReturnLine": ["SalesReturnLineID", "SalesReturnID", "ShipmentLineID", "ItemID", "QuantityReturned",
                        "ExtendedStandardCost"],
    "CreditMemo": ["CreditMemoID", "CreditMemoNumber", "CreditMemoDate", "SalesReturnID", "CustomerID",
                   "OriginalSalesInvoiceID", "SubTotal", "FreightCreditAmount", "TaxAmount", "GrandTotal", "Status",
                   "ApprovedByEmployeeID", "ApprovedDate"],
    "CreditMemoLine": ["CreditMemoLineID", "CreditMemoID", "SalesReturnLineID", "ItemID", "Quantity", "UnitPrice", "Discount",
                       "LineTotal"],
    "CustomerRefund": ["CustomerRefundID", "RefundNumber", "RefundDate", "CustomerID", "CreditMemoID", "Amount",
                       "PaymentMethod", "ApprovedByEmployeeID", "ClearedDate"],
    "SalesCommissionAdjustment": ["SalesCommissionAdjustmentID", "AdjustmentNumber", "AdjustmentDate",
                                  "SalesCommissionAccrualID", "CreditMemoID", "CreditMemoLineID",
                                  "CommissionBaseReductionAmount", "CommissionRatePct", "CommissionAdjustmentAmount",
                                  "ApprovedByEmployeeID", "Status", "CustomerID"],
    "SalesCommissionAccrual": ["SalesCommissionAccrualID", "AccrualNumber", "AccrualDate", "CommissionAmount",
                               "CommissionRatePct", "SalesInvoiceLineID", "CreatedByEmployeeID", "CustomerID"],
    "SalesCommissionPayment": ["SalesCommissionPaymentID", "PaymentNumber", "PaymentDate", "NetPaymentAmount",
                               "ApprovedByEmployeeID"],
    "SalesCommissionPaymentLine": ["SourceDocumentType", "SourceDocumentID"],
    "SalesCommissionRate": ["SalesCommissionRateID", "EffectiveStartDate", "EffectiveEndDate", "ApprovedByEmployeeID"],
    "CashReceipt": ["CashReceiptID", "ReceiptNumber", "ReceiptDate", "CustomerID", "SalesInvoiceID", "Amount",
                    "PaymentMethod", "DepositDate", "RecordedByEmployeeID"],
    "CashReceiptApplication": ["CashReceiptApplicationID", "CashReceiptID", "SalesInvoiceID", "ApplicationDate",
                               "AppliedAmount", "AppliedByEmployeeID"],
    "SalesInvoice": ["SalesInvoiceID", "InvoiceNumber", "CustomerID", "InvoiceDate", "DueDate", "GrandTotal", "PaymentDate"],
    "SalesInvoiceLine": ["SalesInvoiceLineID", "SalesInvoiceID", "ShipmentLineID", "ItemID", "UnitPrice", "Discount", "LineTotal"],
    "Shipment": ["ShipmentID", "ShipmentNumber", "ShipmentDate", "ShippedBy", "WarehouseID", "DeliveryDate"],
    "ShipmentLine": ["ShipmentLineID", "ShipmentID", "QuantityShipped"],
    "GLEntry": ["AccountID", "Debit", "Credit", "VoucherNumber", "SourceDocumentType", "SourceDocumentID",
                "CreatedByEmployeeID", "FiscalYear", "PostingDate"],
    "Account": ["AccountID", "AccountNumber", "AccountName", "AccountType"],
    "JournalEntry": ["JournalEntryID", "EntryNumber", "EntryType", "PostingDate"],
    "Employee": ["EmployeeID", "JobTitle", "MaxApprovalAmount", "HireDate", "TerminationDate", "CostCenterID"],
    "CostCenter": ["CostCenterID", "CostCenterName"],
    "PayrollRegister": ["PayrollRegisterID", "ApprovedByEmployeeID"],
    "Item": ["ItemID", "ItemCode", "ItemGroup", "StandardCost"],
    "Customer": ["CustomerID", "CustomerName", "CustomerSegment", "PaymentTerms"],
}
# The student time of each requirement in minutes, and its weight in the case rubric (the weights add up to 100).  The
# requirement guides' front matter must agree; g0 claims it for every guide written so far.
PLAN = {1: 120, 2: 120, 3: 120, 4: 240, 5: 180, 6: 150, 7: 150, 8: 120, 9: 150, 10: 90}
WEIGHTS = {1: 8, 2: 9, 3: 7, 4: 16, 5: 11, 6: 9, 7: 10, 8: 10, 9: 13, 10: 7}
# The perspective, difficulty, tool and the outcome of the case note's list of six each requirement practices.
PERSPECTIVE = {1: "auditing", 2: "auditing", 3: "auditing", 4: "auditing", 5: "auditing", 6: "auditing", 7: "auditing",
               8: "auditing", 9: "general", 10: "auditing"}
DIFFICULTY = {1: "intermediate", 2: "intermediate", 3: "intermediate", 4: "advanced", 5: "intermediate", 6: "advanced",
              7: "advanced", 8: "advanced", 9: "advanced", 10: "intermediate"}
TOOL = {1: "sql", 2: "sql", 3: "excel", 4: "sql", 5: "powerbi", 6: "excel", 7: "mixed", 8: "excel", 9: "excel", 10: "powerbi"}
OUTCOME = {1: "1", 2: "1", 3: "1, 2", 4: "2, 3", 5: "5", 6: "3, 4", 7: "3", 8: "6", 9: "6", 10: "6"}


def need(d, claim, *tables: str) -> None:
    schema(d, claim, {t: SCHEMA[t] for t in (tables or SCHEMA)})


# --- the SQL model script ---------------------------------------------------------------------------------------------

def sub(sql: str, old: str, new: str) -> str:
    """The statement with `old` replaced by `new`, ignoring differences in whitespace; the text must be there."""
    pattern = r"\s+".join(re.escape(w) for w in old.split())
    if not re.search(pattern, sql):
        raise ValueError(f"the text to change is not in the statement: {old[:70]}")
    return re.sub(pattern, lambda m: new, sql, count=1)


@lru_cache(maxsize=None)
def script() -> tuple[tuple[str, str, str], ...]:
    """(first comment line, second comment line, statement) of every query of Credits.sql, in order; the trailing result
    comment is dropped and the closing notes block is left out."""
    body = SQL.read_text(encoding="utf-8").split("\n/* Model answer", 1)[0]
    out = []
    for chunk in re.split(r"(?m)^(?=-- Requirement \d+, test [A-Z]\d+:)", body):
        if not re.match(r"-- Requirement \d+, test [A-Z]\d+:", chunk):
            continue
        lines = chunk.rstrip().split("\n")
        stmt = re.sub(r";[ \t]*-- Result:[^\n]*$", ";", "\n".join(lines[2:]), flags=re.M)
        out.append((lines[0], lines[1], stmt.strip()))
    return tuple(out)


def quote(key: str) -> str:
    """The statement of the one query whose first comment line starts with `key` ('Requirement 1, test M2')."""
    hits = [s for first, _, s in script() if first.removeprefix("-- ").startswith(key + ":")]
    if len(hits) != 1:
        raise KeyError(f"{len(hits)} queries of Credits.sql start with {key!r}")
    return hits[0]


def run(d, sql: str) -> list:
    return d.q(sql.strip().rstrip(";"))


def test(d, key: str) -> list:
    """The rows of one test of the model script ('M4', 'T6', ...), run read-only."""
    hits = [s for first, _, s in script() if re.match(rf"-- Requirement \d+, test {key}:", first)]
    if len(hits) != 1:
        raise KeyError(f"{len(hits)} queries of Credits.sql are test {key}")
    return run(d, hits[0])


# --- the Excel model workbook -----------------------------------------------------------------------------------------------

@lru_cache(maxsize=None)
def workbook():
    import openpyxl
    return openpyxl.load_workbook(XLSX, read_only=False, data_only=False)


@lru_cache(maxsize=None)
def values():
    import openpyxl
    return openpyxl.load_workbook(XLSX, read_only=False, data_only=True)


def sheets() -> list[str]:
    return list(workbook().sheetnames)


def clean(text) -> str:
    """A formula as the formula bar reads it: the file's prefixes and the stored spill references written the way a person
    types them."""
    t = str(getattr(text, "text", text))
    t = re.sub(r"_xl(?:fn|pm|ws)\.", "", t)
    t = re.sub(r"ANCHORARRAY\(((?:'[^']+'!)?\$?[A-Z]+\$?\d+)\)", r"\1#", t)
    t = re.sub(r"(\w+)\[\[#This Row\],\[([^\]]+)\]\]", r"[@\2]", t)
    return t


def formula(claim, sheet: str, cell: str) -> str:
    value = workbook()[sheet][cell].value
    ok = isinstance(value, str) and value.startswith("=") or hasattr(value, "text")
    claim(ok, f"the cell {sheet}!{cell} of the model workbook holds a formula")
    return clean(value) if ok else ""


def cell(sheet: str, address: str):
    """The cached value of one cell of the model workbook."""
    return values()[sheet][address].value


def maybe_formula(sheet: str, address: str) -> str:
    """The cell's formula as the formula bar reads it, or '' when the cell holds a constant (a model workbook built before the
    case's last revision)."""
    value = workbook()[sheet][address].value
    return clean(value) if (isinstance(value, str) and value.startswith("=")) or hasattr(value, "text") else ""


def has_name(name: str) -> bool:
    return name in workbook().defined_names


def defined(claim, name: str) -> str:
    wb = workbook()
    ok = name in wb.defined_names
    claim(ok, f"the model workbook has the defined name {name}")
    return wb.defined_names[name].attr_text if ok else ""


# --- the Power BI model answer -------------------------------------------------------------------------------------------

def _model_text(relative: str) -> str:
    return (MODEL / "definition" / relative).read_text(encoding="utf-8")


def table_m(name: str) -> list[str]:
    return _block(_model_text(f"tables/{name}.tmdl"), "source =")


def step(claim, lines: list[str], *names: str) -> str:
    """The named steps of an M query as the formula bar reads them."""
    out = []
    for n in names:
        found = [l for l in lines if l.startswith(f'#"{n}" =') or l.startswith(f"{n} =")]
        claim(len(found) == 1, f"the solution project's query has the step {n}")
        out += ["    " + l.rstrip(",") for l in found]
    return "\n".join(out)


@lru_cache(maxsize=None)
def _measures_in(model: Path) -> dict[str, str]:
    """Every measure of every table of a project as the formula bar reads it: `Name = expression`, the body dedented."""
    out: dict[str, str] = {}
    for path in sorted((model / "definition" / "tables").glob("*.tmdl")):
        lines = path.read_text(encoding="utf-8").split("\n")
        for i, line in enumerate(lines):
            m = re.match(r"^\tmeasure (?:'([^']+)'|(\S+)) =(.*)$", line)
            if not m:
                continue
            name, first = m.group(1) or m.group(2), m.group(3).strip()
            body, j = [], i + 1
            while j < len(lines) and lines[j].startswith("\t\t\t"):
                body.append(lines[j][3:].rstrip())
                j += 1
            out[name] = f"{name} = {first}" if first and not body else f"{name} =\n" + "\n".join(body)
    return out


def measure_table() -> dict[str, str]:
    """Every measure of the Credits Analytics project."""
    return _measures_in(MODEL)


def dax(claim, *names: str) -> str:
    table = measure_table()
    for n in names:
        claim(n in table, f"the solution project has the measure {n}")
    return "\n\n".join(table.get(n, "") for n in names)


def column_dax(claim, table: str, column: str) -> str:
    """The expression of a calculated column of the project, as the formula bar reads it."""
    text = _model_text(f"tables/{table}.tmdl")
    m = re.search(rf"(?m)^\tcolumn {re.escape(column)} = (.+)$", text)
    claim(m is not None, f"the solution project's table {table} has the calculated column {column}")
    return f"{column} = {m.group(1).strip()}" if m else ""


def check_expr(claim, label: str) -> str:
    """The DAX expression of one row of the project's Checks tab (DAX query view), found by its label."""
    text = (MODEL / "DAXQueries" / "Checks.dax").read_text(encoding="utf-8")
    m = re.search(r'"Check", "' + re.escape(label) + r'", "Expected", .*?, "Actual", (.*?), "Agrees"', text)
    claim(m is not None, f"the Checks tab has the row {label!r}")
    return m.group(1) if m else ""


def check_sections() -> dict[str, int]:
    text = (MODEL / "DAXQueries" / "Checks.dax").read_text(encoding="utf-8")
    counts: dict[str, int] = {}
    for sec in re.findall(r'ROW \( "Section", "([^"]*)"', text):
        counts[sec] = counts.get(sec, 0) + 1
    return counts


def title(claim, page: str, name: str, report: Path = REPORT) -> str:
    """The title of a visual of the solution project's report page, as Desktop shows it."""
    try:
        v = json.loads((report / "definition" / "pages" / page / "visuals" / name / "visual.json").read_text(encoding="utf-8"))
        text = v["visual"]["visualContainerObjects"]["title"][0]["properties"]["text"]["expr"]["Literal"]["Value"].strip("'")
        text = text.replace("''", "'")
    except (OSError, KeyError, IndexError, ValueError):
        text = ""
    claim(text != "", f"the visual {name} of page {page} in the solution project has a title")
    return text


def has_visual(claim, page: str, *names: str, report: Path = REPORT) -> None:
    for n in names:
        claim((report / "definition" / "pages" / page / "visuals" / n / "visual.json").is_file(),
              f"the page {page} of the solution project has the visual {n}")


def test_block(claim, n: int) -> str:
    """Test n of the Tests tab, comment line first."""
    text = (MODEL / "DAXQueries" / "Tests.dax").read_text(encoding="utf-8").strip()
    found = [p.strip() for p in re.split(r"(?m)^(?=// Test \d)", text) if p.startswith(f"// Test {n}:")]
    claim(len(found) == 1, f"the Tests tab holds Test {n}")
    return found[0] if found else ""


# --- the deck ------------------------------------------------------------------------------------------------------------------

def deck_titles() -> list[str]:
    titles = ["(title slide)"]
    for line in DECK.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^#{1,2} (.+?)\s*(\{[^}]*\})?\s*$", line)
        if m:
            titles.append(m.group(1).strip())
    return titles


def at(titles: list[str], first: str, nth: int = 1) -> int:
    hits = [n for n, t in enumerate(titles, 1) if t.startswith(first)]
    return hits[nth - 1]


def span(titles: list[str], first: str, after: str, nth: int = 1, after_nth: int = 1) -> str:
    start, end = at(titles, first, nth), at(titles, after, after_nth) - 1
    return str(start) if end <= start else f"{start} to {end}"


def dates_of(d) -> dict:
    """The last date of each cycle table (the first test of Requirement 1), keyed by table."""
    return {r[0]: r[4] for r in test(d, "M1") if r[4]}


# --- the case note -----------------------------------------------------------------------------------------------------------

@context("ch19.g0")
def g0(d, claim):
    need(d, claim)
    t = deck_titles()
    claim(len(t) == 69, "the case's deck has 69 slides")
    text = CHAPTER.read_text(encoding="utf-8")
    n_req = len(re.findall(r"^\*\*Requirement \d+", text, re.M))
    claim(n_req == 10, "the case has ten requirements")
    n_ms = len(re.findall(r"^\d+\. \*\*[^*]+\*\* \(Requirements? ", text, re.M))
    claim(n_ms == 3, "the case has three milestones, one for each phase")
    claim("Guided Tutorial" not in text and "## Tutorial" not in text, "the case has no tutorials")
    claim("Multiple-Choice" not in text and "Multiple Choice" not in text, "the case has no multiple-choice questions")
    downloads = DOWNLOADS.read_text(encoding="utf-8")
    zips = ("Chapter19-capstone-excel.zip", "Chapter19-capstone-pbi.zip", "Chapter19-capstone-pbi-monitoring.zip",
            "Chapter19-capstone-sql.zip")
    claim(all(z in downloads for z in zips), "the Downloads page lists the four solution files of the case")
    claim(SQL.is_file() and XLSX.is_file() and MODEL.is_dir() and REPORT.is_dir() and FLAGS_CSV.is_file()
          and MONITOR.is_dir(), "the model files of the case are in outputs/companion/solutions")
    queries = script()
    claim(len(queries) == 40, "the model script holds 40 queries")
    per: dict[int, int] = {}
    for first, _, _ in queries:
        k = int(re.match(r"-- Requirement (\d+),", first).group(1))
        per[k] = per.get(k, 0) + 1
    claim(per == {1: 8, 2: 9, 3: 2, 4: 12, 6: 2, 7: 6, 10: 1}, "the model script's queries are spread over the requirements as the guides say")
    wb = sheets()
    claim(all(s in wb for s in ("Documentation", "Cycle Map", "Reconciliation", "Materiality", "Risks", "Tests", "Sample",
                                "Timeline", "Deficiencies", "Report", "Ledger", "Credits", "Refunds", "ReturnLines",
                                "Clawbacks", "RefundFlags", "Solution Notes")),
          "the model workbook has the worksheets the guides name")
    pages = json.loads((REPORT / "definition" / "pages" / "pages.json").read_text(encoding="utf-8"))["pageOrder"]
    claim(pages == ["about", "returns", "timing", "concentration", "ranking", "refundDetail", "validation", "notes"],
          "the solution project has the eight pages the guides name, in order")
    sections = check_sections()
    claim(set(sections) == {"Requirement 5", "Requirement 5 (the flags of Requirement 4)", "Milestone 2"},
          "the Checks tab holds the rows of Requirement 5, of the flags of Requirement 4, and of Milestone 2")
    claim(sum(WEIGHTS.values()) == 100, "the rubric weights of the ten requirements add up to 100")
    claim(set(PLAN) == set(WEIGHTS) == set(range(1, 11)), "the plan and the weights cover the ten requirements")
    for path in sorted(GUIDES.glob("req-*.md")):
        k = int(re.match(r"req-(\d+)-", path.name).group(1))
        body = path.read_text(encoding="utf-8")
        mm = re.search(r"(?m)^minutes: (\d+)", body)
        claim(not mm or int(mm.group(1)) in (0, PLAN[k]), f"requirement {k}'s guide gives its student time as {PLAN[k]} minutes")
        pm = re.search(r"(?m)^perspective: (\w+)", body)
        claim(not pm or pm.group(1) in ("general", PERSPECTIVE[k]), f"requirement {k}'s guide gives its perspective as {PERSPECTIVE[k]}")
        tm = re.search(r"(?m)^tool: (\w+)", body)
        claim(not tm or tm.group(1) in ("conceptual", TOOL[k]), f"requirement {k}'s guide gives its tool as {TOOL[k]}")
    # the data conditions of the at-a-glance table
    last = dates_of(d)
    claim(last["SalesReturn"] < last["CreditMemo"] <= last["SalesCommissionAdjustment"], "the cycle's tables end on different dates: returns first, then credits and clawbacks")
    planted = n19.anomaly_tables(d)
    claim(planted is not None and not (planted & n19.CYCLE_TABLES), "the AnomalyLog plants nothing in the cycle's tables")
    claim(n19.no_preparer(d), "neither the credit memo nor the refund records a preparer")
    s = dict(
        kick=span(t, "Returns grow faster than sales", "The Engagement"), engage=span(t, "The Engagement", "The Credits Cycle and Its Risks"),
        cycle=span(t, "The Credits Cycle and Its Risks", "The Data You Will Use"), data=span(t, "The Data You Will Use", "Getting Started"),
        setup=span(t, "Getting Started", "Phase 1: Plan the Audit"), tools=at(t, "Choose the tool for each step"),
        check1=f"{at(t, 'Check your understanding', 1)} to {at(t, 'Check your understanding', 1) + 1}", p1=span(t, "Phase 1: Plan the Audit", "Phase 2: Test Controls and Balances"),
        r1=at(t, "Map the cycle from its records"), r2=at(t, "Reconcile each population"), r3=at(t, "Assess the risks and set materiality"),
        check2=f"{at(t, 'Check your understanding', 2)} to {at(t, 'Check your understanding', 2) + 1}", m1=at(t, "Milestone 1"),
        p2=span(t, "Phase 2: Test Controls and Balances", "Phase 3: Evaluate and Report"), r4=at(t, "Test the controls on every document"),
        r5=at(t, "Explain the returns and rank the refunds"), sampling=span(t, "Sampling still has a place", "Sample what the database cannot show"),
        r6=at(t, "Sample what the database cannot show"), check3=f"{at(t, 'Check your understanding', 3)} to {at(t, 'Check your understanding', 3) + 1}",
        r7=at(t, "Test the balances, and follow the largest refund"), m2=at(t, "Milestone 2"),
        p3=span(t, "Phase 3: Evaluate and Report", "Deliverables and Checklist"), deficient=span(t, "What makes a control deficient", "Evaluate the deficiencies"),
        r8=at(t, "Evaluate the deficiencies"), check4=f"{at(t, 'Check your understanding', 4)} to {at(t, 'Check your understanding', 4) + 1}",
        r9=at(t, "Report to the audit committee"), r10=at(t, "Keep the tests running"), m3=at(t, "Milestone 3"),
        deliver=span(t, "Deliverables and Checklist", "Looking Ahead"), ahead=span(t, "Looking Ahead", "About these slides"),
        watch1=at(t, "Watch out: the ledger needs care"), watch2=at(t, "Watch out: time runs through"),
        watch3=at(t, "Watch out: some evidence is not"), classes=at(t, "Three classifications"))
    return dict(plan=PLAN, weights=WEIGHTS, student_minutes=sum(PLAN.values()), student_hours=sum(PLAN.values()) // 60,
                n_req=n_req, n_ms=n_ms, slides=len(t), s=s, n_queries=len(queries), per=per, n_sheets=len(wb), n_pages=len(pages),
                checks=sum(sections.values()), per_section=sections, n_measures=len(measure_table()), dt=last,
                perspective=PERSPECTIVE, difficulty=DIFFICULTY, outcome=OUTCOME)


# --- Requirement 1: the map of the cycle ---------------------------------------------------------------------------------------

@context("ch19.g1")
def g1(d, claim):
    need(d, claim)
    m1, m2, m3 = test(d, "M1"), test(d, "M2"), test(d, "M3")
    m4, m5, m6, m7, m8 = test(d, "M4"), test(d, "M5"), test(d, "M6"), test(d, "M7"), test(d, "M8")
    claim(len(m1) == 12 and all(r[1] == r[2] and r[7] == 0 for r in m1),
          "the profile shows the twelve tables, each with as many distinct keys as records and no missing key")
    claim(len(m4) == 19 and all(r[1] == 0 for r in m4), "every one of the nineteen link tests finds no orphan")
    claim(len(m5) == 6 and all(r[1] == 0 and r[2] == 0 for r in m5), "every document of the cycle posts, and every posting has its document")
    claim(len(m6) == 12 and all(r[5] == 0 and r[6] == 0 for r in m6), "the four numbered series have no gap and no repeat")
    claim(all(r[7] in (None, 0) for r in m6), "each series continues across the years without a break")
    last = {r[0]: r[4] for r in m1 if r[4]}
    # wrong version: every credit without a refund (not only the refunded ones)
    no_refund = d.one("SELECT COUNT(*) FROM CreditMemo AS x LEFT JOIN CustomerRefund AS y ON y.CreditMemoID = x.CreditMemoID "
                      "WHERE y.CustomerRefundID IS NULL")
    not_refunded = d.one("SELECT COUNT(*) FROM CreditMemo WHERE Status <> 'Refunded'")
    claim(no_refund == not_refunded > 0, "the credits that no refund points at are exactly the credits not in status Refunded")
    # wrong version: an anti-join written as an inner join can never return a row
    inner = d.one("SELECT COUNT(*) FROM CreditMemoLine AS x INNER JOIN SalesReturnLine AS y "
                  "ON y.SalesReturnLineID = x.SalesReturnLineID WHERE y.SalesReturnLineID IS NULL")
    claim(inner == 0, "an anti-join written with INNER JOIN returns no rows")
    # wrong version: the year of the clawback numbers read as if the prefix had two letters
    scaj_zero = d.one("SELECT COUNT(*) FROM SalesCommissionAdjustment WHERE CAST(SUBSTR(AdjustmentNumber, 4, 4) AS INTEGER) = 0")
    claim(scaj_zero == d.one("SELECT COUNT(*) FROM SalesCommissionAdjustment"),
          "the clawback numbers have a four-letter prefix, so reading the year from the other series' position gives zero for every row")
    # wrong version: a gap test that expects every year to restart at 1
    restart = sum(1 for r in m6 if r[2] != 1)
    claim(restart == 8, "eight of the twelve series-years start above 1, because the numbers continue across years")
    # wrong version: the ledger joined on the document's ID alone
    ids = d.q("SELECT COUNT(*), COUNT(DISTINCT SourceDocumentType) FROM GLEntry WHERE SourceDocumentID IN (SELECT CreditMemoID FROM CreditMemo)")[0]
    own = d.one("SELECT COUNT(*) FROM GLEntry WHERE SourceDocumentType = 'CreditMemo'")
    claim(ids[0] > 10 * own and ids[1] > 5, "the credit IDs, matched without the document type, pick up the rows of many other document types")
    # wrong version: AccountID taken for the account number
    id_4060 = d.one("SELECT AccountID FROM Account WHERE AccountNumber = 4060")
    claim(d.one("SELECT COUNT(*) FROM GLEntry WHERE AccountID = 4060") == 0 and id_4060 != 4060,
          "GLEntry.AccountID is the account's surrogate key, so AccountID 4060 matches no row")
    # who records each document, by cost center (M8): the warehouse receives, customer service handles cash and credits,
    # administration approves refunds
    dept = {}
    for r in m8:
        dept.setdefault(r[0], set()).add(r[3])
    claim(dept == {"SalesReturn": {"Warehouse"}, "CreditMemo": {"Customer Service"}, "CustomerRefund": {"Administration"},
                   "SalesCommissionAdjustment": {"Customer Service"}, "CashReceipt": {"Customer Service"},
                   "CashReceiptApplication": {"Customer Service"}},
          "the warehouse receives the returns, customer service approves credits and clawbacks and handles the cash, and administration approves the refunds")
    index = [r[2] for r in d.q("PRAGMA index_info(ix_glentry_sourcedocument_trace)")]
    claim(index[:2] == ["SourceDocumentType", "SourceDocumentID"], "the ledger's trace index starts with the source document type, then the ID")
    post = {(r[0], r[1]): r for r in m7}
    both = d.one("SELECT COUNT(*) FROM (SELECT SourceDocumentID FROM GLEntry WHERE SourceDocumentType = 'CreditMemo' "
                 "AND AccountID IN (?, ?) AND Credit > 0 GROUP BY 1 HAVING COUNT(DISTINCT AccountID) = 2)",
                 d.account("1020"), d.account("2060"))
    rows_1020 = post[("CreditMemo", 1020)][5]
    rows_2060 = post[("CreditMemo", 2060)][5]
    claim(both > 0 and rows_1020 + rows_2060 > d.one("SELECT COUNT(*) FROM CreditMemo"),
          "some credits post to both 1020 and 2060, so the credit rows on the two accounts outnumber the credits")
    # the Cycle Map's control-point rows (the control points C1 to C9 in order, each with its exceptions in the records)
    ws = values()["Cycle Map"]
    points = []
    for r in range(82, 91):
        points.append(dict(point=ws.cell(r, 1).value, exceptions=ws.cell(r, 13).value, tested=ws.cell(r, 15).value))
    claim([p["point"] for p in points] == [f"C{i}" for i in range(1, 10)], "the Cycle Map holds the nine control points in order")
    exc = {p["point"]: p["exceptions"] for p in points}
    tt = n19.tests(d)
    claim(exc["C2"] == len(tt.cm_above) and exc["C4"] == len(tt.cash_credit) and exc["C6"] == len(tt.rf_above) and exc["C7"] == len(tt.before),
          "the Cycle Map's exceptions at C2, C4, C6 and C7 are the counts of Requirement 4")
    none_found = [p["point"] for p in points if isinstance(p["exceptions"], str) or p["exceptions"] == 0]
    cp_cols = [ws.cell(r, c).value for r in range(82, 91) for c in (3, 8)]
    claim(all(cp_cols), "every control point of the Cycle Map says what the records evidence and what they cannot show")
    return dict(
        m1=len(m1), m2=len(m2), m3=len(m3), m4=len(m4), m5=len(m5), m6=len(m6), m7=len(m7), m8=len(m8), last=last,
        no_refund=no_refund, inner=inner, scaj_zero=scaj_zero, restart=restart, id_rows=ids[0], id_types=ids[1], own_rows=own,
        id_4060=id_4060, both=both, rows_1020=rows_1020, rows_2060=rows_2060, none_found=none_found,
        f_count=formula(claim, "Cycle Map", "C6"), f_matrix=formula(claim, "Cycle Map", "B48"),
        f_flag=formula(claim, "Cycle Map", "M83"), n_tables=len(m1))


# --- Requirement 2: the reconciliations ------------------------------------------------------------------------------------------

@context("ch19.g2")
def g2(d, claim):
    need(d, claim)
    g1_, g2_, g3_, g4_, g5_ = (test(d, k) for k in ("G1", "G2", "G3", "G4", "G5"))
    g6_, g7_, g8_, g9_ = (test(d, k) for k in ("G6", "G7", "G8", "G9"))
    claim(all(r[11] == 0 and r[12] == 0 for r in g1_) and len(g1_) == 3, "the credits reconcile to the ledger in every year")
    claim(all(r[3] == 0 and r[6] == r[9] == r[12] == r[15] == 0 for r in g2_) and len(g2_) == 3,
          "the returns' standard cost reconciles to inventory and to each cost of goods sold account in every year")
    claim(all(r[5] == 0 and r[6] == 0 for r in g3_) and len(g3_) == 3, "the refunds reconcile to cash and to 2060 in every year")
    claim(all(r[4] == 0 for r in g4_) and len(g4_) == 3, "the clawbacks reconcile to 2034 in every year")
    claim(all(r[7] == 0 and r[8] == 0 for r in g5_) and len(g5_) == 3, "the receipts and applications reconcile to cash and receivables in every year")
    claim(len(g6_) == 3 and all(r[6] == 0 for r in g6_), "nothing but receipts, credits, applications and refunds posts to 2060")
    claim(len(g7_) == 8, "eight credits make up the year-end balances of 2060")
    claim(len(g9_) == 6, "six year-end closes post to 4060 and 6290, one a year on each")
    after = d.q("SELECT DISTINCT SourceDocumentType FROM GLEntry WHERE PostingDate > ?", f"{d.C}-12-31")
    claim(after == [("DisbursementPayment",)], "the ledger rows after the last year-end are supplier payments only")
    # wrong version: the credits tied to receivables alone
    parts_2060 = [r[10] for r in g1_]
    claim(all(p > 0 for p in parts_2060), "every year some credits post to 2060, so a tie to 1020 alone leaves a difference")
    # wrong version: 4060 netted over every source (the close credits it back)
    net_4060 = d.q("SELECT g.FiscalYear, ROUND(SUM(g.Debit - g.Credit), 2) FROM GLEntry g JOIN Account a ON a.AccountID = g.AccountID "
                   "WHERE a.AccountNumber = 4060 AND g.FiscalYear BETWEEN 2024 AND 2026 GROUP BY 1 ORDER BY 1")
    claim(len(net_4060) == 3 and all(abs(x[1]) < 0.005 for x in net_4060), "the net balance of 4060 over every source is zero in each year, because the close credits it back")
    # wrong version: the returns tied to 5010 alone
    cogs_total, cogs_5010 = g2_[-1][1], g2_[-1][4]       # the year's returned cost, and its part credited to 5010
    cogs_5010_gap = round(cogs_total - cogs_5010, 2)
    claim(cogs_5010_gap > 0.005 and cogs_5010 > 0, "the returns' standard cost is spread over four cost of goods sold accounts, so 5010 alone leaves a difference")
    claim(abs(cogs_5010_gap - sum(g2_[-1][k] for k in (7, 10, 13))) < 0.005, "the difference is what the other three accounts, 5020 to 5040, received")
    close_entries = len({r[1] for r in g9_})
    claim(close_entries == 3 and len(g9_) == 6, "the six account rows of the closes belong to three year-end closing journal entries")
    # wrong version: every credit to 6290, whatever its source
    all_6290 = d.q("SELECT g.FiscalYear, ROUND(SUM(g.Credit), 2) FROM GLEntry g JOIN Account a ON a.AccountID = g.AccountID "
                   "WHERE a.AccountNumber = 6290 AND g.FiscalYear BETWEEN 2024 AND 2026 GROUP BY 1 ORDER BY 1")
    claw = [r[1] for r in g4_]
    claim(all(a[1] > 50 * c for a, c in zip(all_6290, claw)), "the credits to 6290 from every source are far larger than the clawbacks, because the close is among them")
    # wrong version: the open items of 2060 by today's status
    issued = d.q("SELECT COUNT(*), ROUND(SUM(GrandTotal), 2) FROM CreditMemo WHERE Status = 'Issued'")[0]
    last_year_items = [r for r in g7_ if r[0] == d.C]
    claim(issued[0] == len(last_year_items) and abs(issued[1] - sum(r[4] for r in last_year_items)) < 0.005,
          "the credits still in status Issued are the open items of the last year-end, and only those")
    claim(len(g7_) > issued[0], "earlier year-ends hold items that were refunded afterward, which today's status no longer shows")
    # wrong version: the commission payable without the clawbacks
    claw_total = d.one("SELECT ROUND(SUM(CommissionAdjustmentAmount), 2) FROM SalesCommissionAdjustment WHERE AdjustmentDate <= ?", f"{d.C}-12-31")
    claim(claw_total > 0, "the clawbacks reduce commission payable, so leaving them out leaves a difference")
    # the Excel appendix: a difference that should be zero shows a floating-point residue
    resid = [abs(cell("Reconciliation", f"{c}49")) for c in "CDE"] + [abs(cell("Reconciliation", f"{c}19")) for c in "CDE"]
    claim(0 < max(resid) < 0.005, "a difference cell of the Reconciliation sheet shows a floating-point residue far below a cent")
    opening = [cell("Reconciliation", f"{c}56") for c in "CDE"]
    return dict(
        g1=len(g1_), g2=len(g2_), g3=len(g3_), g4=len(g4_), g5=len(g5_), g6=len(g6_), g7=len(g7_), g8=len(g8_), g9=len(g9_),
        parts_2060=parts_2060, parts_total=round(sum(parts_2060), 2), net_4060=[x[1] for x in net_4060],
        cogs_gap=cogs_5010_gap, cogs_total=cogs_total, cogs_5010=cogs_5010, close_entries=close_entries, all_6290=[x[1] for x in all_6290], issued=issued, last_items=len(last_year_items),
        claw_total=claw_total, resid=f"{max(resid):.1e}", not_netted=g8_[0][4], opening=opening,
        f_sub=formula(claim, "Reconciliation", "C7"), f_4060=formula(claim, "Reconciliation", "C8"),
        f_diff=formula(claim, "Reconciliation", "C9"), f_credit=formula(claim, "Reconciliation", "C19"),
        f_close=formula(claim, "Reconciliation", "C61"), f_ledger=formula(claim, "Reconciliation", "C62"),
        f_items=formula(claim, "Reconciliation", "C75"))


# --- Requirement 3: materiality and the expectation ---------------------------------------------------------------------------

# The formulas and texts the Excel builder writes since the revision of 2026-10-17. A model workbook built before it holds the older
# ones, so the guides quote these until the workbook is rebuilt; a claim checks them against the workbook once it holds the change.
NEW_CONCLUSION = ('=IF(ABS(B21)>AnalyticalFollowUp,"Investigate (Requirement 5): an unexplained difference above the follow-up trigger"'
                  '&IF(ABS(B21)>PerformanceMateriality,", above performance materiality",", below performance materiality"),'
                  '"Below the quantitative trigger; consider qualitative risk and expectation precision")')
NEW_RECON = ('=LET(mask,LEFT(Reconciliation!A5:A{last},10)="Difference",d,FILTER(Reconciliation!C5:E{last},mask),'
             'ok,FILTER(ISNUMBER(Reconciliation!C5:E{last}),mask),IF(AND(ok),SUM(ABS(d)),'
             '"Not evaluable: missing or erroneous difference"))')
NEW_RECON_CONCLUSION = '=IF(ISNUMBER(C32),IF(C32<0.005,"Reconciles to the cent","Investigate differences"),C32)'
# the builder's new wording of the report's findings and of the owners' answers (the earlier workbook holds the older)
FRESH = {
    "Refunds approved without recorded authority and dated before": "Refunds approved above the recorded general limits (effective delegation and enforcement unverified) and dated before",
    "Credits approved above recorded authority, in the department": "Credits approved above the recorded general limits, in the department",
    "The receipt shows that goods came back, not that the credit's amount or the customer's cash was right; show the credits whose approver applied the invoice's cash.":
        "The receiving record documents a return; it needs authenticating, with inspection evidence, before it shows that the goods arrived, in the stated condition, or justified the credit, and it says nothing of who handled the customer's cash; show the credits whose approver applied the invoice's cash.",
}


def fresh(text):
    """A workbook text with the builder's newer wording where the workbook is older."""
    out = str(text or "")
    for old, new in FRESH.items():
        out = out.replace(old, new)
    return out


def builder_text(claim, needle: str) -> str:
    """A line of the Excel builder that the guide quotes (the Power Query step the workbook's Ledger Table uses)."""
    text = XL_BUILDER.read_text(encoding="utf-8")
    claim(needle in text, "the Excel builder holds the Power Query step the guide quotes")
    return needle


@context("ch19.g3")
def g3(d, claim):
    need(d, claim)
    m = n19.materiality(d)
    ex = n19.expectation(d)
    claim(all(abs(cell("Materiality", a) - v) < 0.005 for a, v in (("B4", m["income"]), ("B6", m["mat"]), ("B8", m["perf"]), ("B10", m["trivial"]))),
          "the Materiality sheet holds the note's income and thresholds")
    claim(abs(cell("Materiality", "B19") - ex["exp"]) < 0.005 and abs(cell("Materiality", "B21") - ex["diff"]) < 0.005,
          "the Materiality sheet holds the note's expectation and difference")
    for name in ("Materiality", "PerformanceMateriality", "ClearlyTrivial"):
        defined(claim, name)
    # wrong version: the closes left in
    with_closes = d.q("SELECT g.FiscalYear, ROUND(SUM(g.Credit) - SUM(g.Debit), 2) FROM GLEntry g JOIN Account a ON a.AccountID = g.AccountID "
                      "WHERE a.AccountType IN ('Revenue', 'Expense') AND g.FiscalYear BETWEEN ? AND ? GROUP BY 1 ORDER BY 1", d.F, d.C)
    claim(all(abs(x[1]) < 0.005 for x in with_closes), "with the year-end closes left in, the revenue and expense accounts net to zero in every year")
    # wrong version: the difference measured against materiality
    claim(ex["diff"] < m["mat"], "the difference between the recorded and the expected returns is below overall materiality")
    claim(ex["follow_up"] < ex["diff"] < m["perf"], "the difference is above the follow-up trigger the case assumes and below performance materiality")
    claim(ex["follow_up"] != m["trivial"], "the follow-up trigger is its own assumption, not the clearly trivial threshold")
    # the model workbook's trigger cell, once the workbook holds it (the builder's rows 11 and 12 of the Materiality sheet)
    if has_name("AnalyticalFollowUp"):
        claim(abs(cell("Materiality", "B12") - ex["follow_up"]) < 0.005 and abs(cell("Materiality", "B11") - n19.FOLLOW_UP) < 1e-9,
              "the Materiality sheet holds the follow-up trigger at its rate of the expectation")
    claim(ex["exp_f"] < ex["actual"] and ex["f_pct"] > ex["diff_pct"] * 2, "the first year's rate understates the expectation by far more than the prior year's")
    # the SQL cross-check
    p1, p2 = test(d, "P1"), test(d, "P2")
    claim(len(p1) == 3 and len(p2) == 3 and abs(p1[-1][4] - ex["exp"]) < 0.005 and abs(p2[-1][1] - m["income"]) < 0.005,
          "the SQL tests P1 and P2 give the workbook's expectation and income")
    facts = [cell("Risks", f"B{r}") for r in range(6, 12)]
    ratings = [(cell("Risks", f"A{r}"), cell("Risks", f"B{r}")) for r in range(15, 19)]
    claim([x[1] for x in ratings] == ["High", "High", "Moderate", "Low"], "the Risks sheet rates refunds and credits High, cash application Moderate and clawbacks Low")
    grid = [[cell("Risks", f"{c}{r}") for c in "ABCDEFGHI"] for r in range(30, 40)]
    covered = cell("Risks", "B41")
    claim(covered == 9, "the risk-and-control matrix covers all nine control points")
    brainstorm = [cell("Risks", f"A{r}") for r in range(22, 27)]
    claim(len(brainstorm) == 5 and all(brainstorm), "the fraud brainstorm names the five schemes of the chapter's table")
    # the conclusion the sheet's formula gives (the workbook built since the revision of 2026-10-17); an earlier workbook's text is the
    # clearly-trivial wording the revision replaced, so the expected text is computed from the rule instead
    f_concl = formula(claim, "Materiality", "B25")
    if has_name("AnalyticalFollowUp"):
        claim(f_concl == NEW_CONCLUSION, "the Materiality sheet's conclusion is the builder's formula against the follow-up trigger")
    concl = (cell("Materiality", "B25") if has_name("AnalyticalFollowUp") else
             "Investigate (Requirement 5): an unexplained difference above the follow-up trigger, "
             + ("above" if ex["diff"] > m["perf"] else "below") + " performance materiality")
    ledger_rows = workbook()["Ledger"].max_row - 1
    iyec = builder_text(claim, '\'[EntryType] <> null and Text.StartsWith([EntryType], "Year-End Close")\'').strip("'")
    claw_share = facts[1]
    return dict(
        income_closes=[x[1] for x in with_closes], ratings=ratings, matrix_rows=len(grid), covered=covered, brainstorm=brainstorm,
        facts=facts, ledger_rows=ledger_rows, iyec=iyec, p1=len(p1), p2=len(p2), claw_share=claw_share,
        mat_ratio=ex["diff"] / m["mat"], follow_up=ex["follow_up"], follow_rate=n19.FOLLOW_UP, concl=concl,
        f_income=formula(claim, "Materiality", "B4"), f_mat=formula(claim, "Materiality", "B6"), f_rate=formula(claim, "Materiality", "B17"),
        f_ret=formula(claim, "Materiality", "B16"), f_exp=formula(claim, "Materiality", "B19"), f_concl=f_concl if has_name("AnalyticalFollowUp") else NEW_CONCLUSION,
        f_cover=formula(claim, "Risks", "B41"), f_rev=formula(claim, "Materiality", "B15"))


# --- Requirement 4: the control tests --------------------------------------------------------------------------------------

@context("ch19.g4")
def g4(d, claim):
    need(d, claim)
    t = n19.tests(d)
    t1, t2, t3, t4, t5 = (test(d, k) for k in ("T1", "T2", "T3", "T4", "T5"))
    t6, t7, t8, t9, t10, t11, t12 = (test(d, k) for k in ("T6", "T7", "T8", "T9", "T10", "T11", "T12"))
    claim(len(t1) == 5 and len(t2) == 0 and len(t3) == len(t.cc_rec) and len(t6) == len(t.before), "the control tests return the rows the note counts")
    claim(t4[0] == (0, 0, 0), "no refund was approved by the credit's approver, no return received by the credit's approver, none outside the warehouse")
    claim(len(t10) == len(n19.cycle(d).rf), "the refund flags hold one row for each refund")
    # wrong version: the limit read without COALESCE (a blank limit)
    null_limit = d.one("SELECT COUNT(*) FROM Employee WHERE EmployeeID IN (SELECT ApprovedByEmployeeID FROM CreditMemo UNION "
                       "SELECT ApprovedByEmployeeID FROM CustomerRefund) AND MaxApprovalAmount IS NULL")
    claim(null_limit == 0, "every approver of a credit or a refund has a recorded limit, so LimitMissing is zero today")
    # wrong version: the credit and the applications of its invoice joined, and rows counted
    rows_applied = d.one("SELECT COUNT(*) FROM CreditMemo AS cm INNER JOIN CashReceiptApplication AS cra ON cra.SalesInvoiceID = cm.OriginalSalesInvoiceID "
                         "WHERE cra.AppliedByEmployeeID = cm.ApprovedByEmployeeID")
    applied = len(t.cash_credit)
    claim(rows_applied > applied, "counting the joined rows instead of the credits counts some credits more than once")
    recorded_only = d.one("SELECT COUNT(DISTINCT cm.CreditMemoID) FROM CreditMemo AS cm INNER JOIN CashReceiptApplication AS cra "
                          "ON cra.SalesInvoiceID = cm.OriginalSalesInvoiceID INNER JOIN CashReceipt AS cr ON cr.CashReceiptID = cra.CashReceiptID "
                          "WHERE cr.RecordedByEmployeeID = cm.ApprovedByEmployeeID")
    claim(applied < recorded_only < len(t.cc_rec), "recorded-only credits fall between applied-only and applied-or-recorded")
    # wrong version: a refund before the last application tested with >= (same-day payments)
    same_day = d.one("SELECT COUNT(*) FROM CustomerRefund AS rf INNER JOIN CreditMemo AS cm ON cm.CreditMemoID = rf.CreditMemoID "
                     "INNER JOIN (SELECT SalesInvoiceID, MAX(ApplicationDate) AS lp FROM CashReceiptApplication GROUP BY 1) AS l "
                     "ON l.SalesInvoiceID = cm.OriginalSalesInvoiceID WHERE l.lp >= rf.RefundDate")
    claim(same_day > len(t.before), "some refunds are dated the day of the last application, so >= flags more refunds than >")
    # wrong version: the customer's past-due balance read on today's balances instead of the refund date's
    lifetime = d.one("""
        WITH Bal AS (
            SELECT si.SalesInvoiceID, si.CustomerID, si.DueDate, si.InvoiceDate,
                si.GrandTotal - COALESCE((SELECT SUM(AppliedAmount) FROM CashReceiptApplication c WHERE c.SalesInvoiceID = si.SalesInvoiceID), 0)
                    - COALESCE((SELECT SUM(GrandTotal) FROM CreditMemo m WHERE m.OriginalSalesInvoiceID = si.SalesInvoiceID), 0) AS Open
            FROM SalesInvoice AS si)
        SELECT COUNT(DISTINCT rf.CustomerRefundID) FROM CustomerRefund AS rf INNER JOIN Bal AS b ON b.CustomerID = rf.CustomerID
            AND b.InvoiceDate <= rf.RefundDate AND b.DueDate < rf.RefundDate WHERE b.Open > 0.005""")
    claim(lifetime == 0, "no refund's customer has an invoice that was past due on the refund date and is still open today")
    pd_any = sum(1 for r in t10 if r[11] == 1)
    pd_atleast = sum(1 for r in t10 if r[10] == 1)
    claim(pd_any > 100 and pd_atleast > 100, "past-due balances on the refund date are found for most refunds, on the balances as of that date")
    # wrong version: a refund compared with the credit's total instead of the credit's part posted to 2060
    vs_total = d.one("SELECT COUNT(*) FROM CustomerRefund AS rf INNER JOIN CreditMemo AS cm ON cm.CreditMemoID = rf.CreditMemoID "
                     "WHERE ABS(rf.Amount - cm.GrandTotal) > 0.005")
    parts = n19.cycle(d).parts
    both_credits = sum(1 for p in parts.values() if p[0] > 0 and p[1] > 0)
    claim(vs_total == both_credits, "the refunds that differ from their credit's total are the credits that posted a part to both 1020 and 2060")
    # wrong version: GLEntry.CreatedByEmployeeID read as the preparer
    prep = {r[0]: (r[2], r[3]) for r in t5}
    claim(all(prep[k][0] == prep[k][1] for k in ("CreditMemo", "CustomerRefund")), "GLEntry.CreatedByEmployeeID is the approver on every credit and refund row")
    ceo_rows = prep["SalesInvoice"]
    # wrong version: the clawback amount tested against SQLite's ROUND
    claim(len(t8) == 3 and all(r[0] == "Amount is not ROUND(base x rate, 2)" for r in t8), "the clawback test finds three amounts that differ from SQLite's ROUND by a cent")
    claim(len(t8) == len(n19.clawbacks(d)["off"]), "the three clawbacks are the half-cent products of the note")
    by_year = Counter(r[0] for r in t11)
    claim(len(t11) == len(t.cm_above) + len(t.cash_credit) + len(t.rf_above) + len(t.before), "the combined list holds the four tests' exceptions")
    # the refund flags: RefundFlags.csv is the export of this result
    import csv
    csv_rows = list(csv.DictReader(FLAGS_CSV.open(encoding="utf-8-sig")))
    claim(len(csv_rows) == len(t10) and all(a["RefundNumber"] == b[0] and int(a["Score"]) == b[12] for a, b in zip(csv_rows, t10)),
          "RefundFlags.csv of the Power BI solution is the export of test T10, row for row")
    ac = {r[0]: (r[1], r[2]) for r in t12}
    claim(all(a == b for a, b in (ac["Commission rates approved"], ac["Commission accruals created"], ac["Commission payments approved"], ac["Payroll registers approved"])),
          "the Accounting Manager made every commission rate, accrual and payment approval, and approved every payroll register")
    return dict(
        null_limit=null_limit, t1=len(t1), t3=len(t3), t5=len(t5), t6=len(t6), t7=t7[0], t10=len(t10), t11=len(t11), t12=len(t12),
        rows_applied=rows_applied, recorded_only=recorded_only, same_day=same_day, lifetime=lifetime, vs_total=vs_total,
        both_credits=both_credits, prep_rows=[prep["CreditMemo"][0], prep["CustomerRefund"][0]], ceo_rows=ceo_rows[0],
        by_year=[by_year[y] for y in d.years], pd_any=pd_any, pd_atleast=pd_atleast, am=ac, t8=len(t8),
        high=sum(1 for r in t10 if r[12] >= 4))


# --- Requirement 5: the Power BI report ---------------------------------------------------------------------------------

@context("ch19.g5")
def g5(d, claim):
    need(d, claim)
    r5 = n19.r5(d, claim)
    t = n19.tests(d)
    cm_n = d.one("SELECT COUNT(*) FROM CreditMemo")
    lines_n = d.one("SELECT COUNT(*) FROM CreditMemoLine")
    claim(lines_n > cm_n, "some credits have more than one line, so counting lines overstates the credits")
    # wrong version: the credit memo total (with tax and freight) tied to 4060
    tot = d.q("SELECT ROUND(SUM(GrandTotal) - SUM(SubTotal), 2), ROUND(SUM(FreightCreditAmount), 2), ROUND(SUM(TaxAmount), 2) FROM CreditMemo")[0]
    claim(abs(tot[0] - tot[1] - tot[2]) < 0.005 and tot[0] > 0, "the credit memos' grand total exceeds their subtotal by the freight and tax credited")
    # wrong version: the rate of the first month of the first year, by the month of the credit
    cred = dict(d.q("SELECT substr(cm.CreditMemoDate, 1, 7), SUM(l.LineTotal) FROM CreditMemoLine l JOIN CreditMemo cm ON cm.CreditMemoID = l.CreditMemoID GROUP BY 1"))
    rev = dict(d.q("SELECT substr(si.InvoiceDate, 1, 7), SUM(l.LineTotal) FROM SalesInvoiceLine l JOIN SalesInvoice si ON si.SalesInvoiceID = l.SalesInvoiceID GROUP BY 1"))
    first_month = f"{d.F}-01"
    first_rate = cred.get(first_month, 0) / rev[first_month]
    other = [cred.get(k, 0) / v for k, v in rev.items() if k != first_month]
    claim(first_rate < 0.1 * max(other) and first_rate < 0.001, "the first month of the first year has a far lower return rate than later months, by the month of the credit")
    # wrong version: the rate by the month of the shipment, with and without the lines that never ship
    c_all = d.one("SELECT SUM(LineTotal) FROM CreditMemoLine")
    r_ship = d.one("SELECT SUM(LineTotal) FROM SalesInvoiceLine WHERE ShipmentLineID IS NOT NULL")
    r_all = d.one("SELECT SUM(LineTotal) FROM SalesInvoiceLine")
    claim(r_ship < r_all and c_all / r_ship > c_all / r_all, "the lines with no shipment (design services) are left out of revenue by ship date, so the rate by ship date reads higher overall")
    # wrong version: customers ranked by credit dollars
    cc = dict(d.q("SELECT cm.CustomerID, SUM(l.LineTotal) FROM CreditMemoLine l JOIN CreditMemo cm ON cm.CreditMemoID = l.CreditMemoID GROUP BY 1"))
    cr = dict(d.q("SELECT si.CustomerID, SUM(l.LineTotal) FROM SalesInvoiceLine l JOIN SalesInvoice si ON si.SalesInvoiceID = l.SalesInvoiceID GROUP BY 1"))
    by_dollars = sorted(cc, key=lambda c: -cc[c])
    by_rate = sorted(cc, key=lambda c: -cc[c] / cr[c])
    focus = r5["focus"]
    claim(by_dollars[0] == focus and by_rate.index(focus) == len(r5["higher"]), "the customer with the largest credits is first by dollars and fourth by rate")
    # wrong version: the score with any past-due balance instead of one at least the refund
    import csv
    rows = list(csv.DictReader(FLAGS_CSV.open(encoding="utf-8-sig")))
    any_hi = sum(1 for r in rows if int(r["AboveLimit"]) + int(r["CreditAboveLimit"]) + int(r["MethodMismatch"]) + int(r["BeforePayment"]) + int(r["PastDueAny"]) >= 4)
    claim(any_hi > r5["hi"], "counting any past-due balance instead of one at least the refund puts more refunds in the top scores")
    n_rf = d.one("SELECT COUNT(*) FROM CustomerRefund")
    n_il = d.one("SELECT COUNT(*) FROM SalesInvoiceLine")
    n_gl = d.one("SELECT COUNT(*) FROM GLEntry WHERE SourceDocumentType IN ('CreditMemo', 'CustomerRefund', 'SalesInvoice')")
    n_noship = d.one("SELECT COUNT(*) FROM SalesInvoiceLine WHERE ShipmentLineID IS NULL")
    rate_year = [sum(v for k, v in cred.items() if k.startswith(str(y))) / sum(v for k, v in rev.items() if k.startswith(str(y))) for y in d.years]
    rel = re.findall(r"fromColumn: (\w+)\.(\w+)\n\s+toColumn: Date\.Date(\n\s+isActive: false)?", _model_text("relationships.tmdl"))
    rel_active = sum(1 for r in rel if not r[2])
    rel_inactive = sum(1 for r in rel if r[2])
    claim(rel_active == 5 and rel_inactive == 2, "the model relates five facts to Date by their own date and two by the shipment date, inactive")
    rel_active_tables = [r[0] for r in rel if not r[2]]
    rel_inactive_tables = [r[0] for r in rel if r[2]]
    claim(set(rel_inactive_tables) == {"CreditMemoLine", "SalesInvoiceLine"} and set(rel_inactive_tables) < set(rel_active_tables),
          "only the credit lines and the invoice lines are also related to Date by the shipment date")
    fact_titles = {p: title(claim, p, v) for p, v in (("returns", "byYear"), ("timing", "rateByMonth"), ("concentration", "customers"),
                                                       ("ranking", "refunds"), ("refundDetail", "invoice"), ("validation", "tieOut"))}
    has_visual(claim, "returns", "byYear", "byGroup", "byReason", "families", "damaged", "bySegment")
    has_visual(claim, "timing", "rateByMonth", "rateYears", "shipMonths", "lagCards")
    has_visual(claim, "concentration", "carriers", "warehouses", "overall", "customers")
    has_visual(claim, "ranking", "scores", "flagYears", "approvers", "highCards", "refunds")
    has_visual(claim, "refundDetail", "backButton", "refund", "flags", "credit", "invoice", "receipts")
    has_visual(claim, "validation", "flagCards", "tieOut")
    sections = check_sections()
    tests = [(n, test_block(claim, n)) for n in (1, 2, 3)]
    ids = dax(claim, "Return Rate", "Credits by Ship Date", "Revenue by Ship Date", "Return Rate by Ship Date", "Flagged Refunds", "High-Score Refunds",
              "Applied to the Invoice")
    return dict(
        t=fact_titles, rel_active_tables=rel_active_tables, rel_inactive_tables=rel_inactive_tables, n_rf=n_rf, n_il=n_il, n_gl=n_gl, n_noship=n_noship, rate_year=rate_year, rel_active=rel_active,
        rel_inactive=rel_inactive, c_scores=check_expr(claim, "scores different from the export (Test 3)"),
        c_exports=check_expr(claim, "refunds whose date or amount differs from the export (Test 3)"), n_lines=lines_n, n_cm=cm_n, tot_diff=tot[0], first_month=first_month, first_rate=first_rate, rate_ship=c_all / r_ship,
        rate_all=c_all / r_all, any_hi=any_hi, n_checks=sections["Requirement 5"], n_flags=sections["Requirement 5 (the flags of Requirement 4)"],
        n_ms=sections["Milestone 2"], n_tests=len(tests), dax=ids, rank_rate=by_rate.index(focus) + 1,
        d_rate=dax(claim, "Credits", "Revenue", "Return Rate"), d_ship=dax(claim, "Credits by Ship Date", "Revenue by Ship Date", "Return Rate by Ship Date"),
        d_flags=dax(claim, "Flagged Refunds"), d_high=dax(claim, "High-Score Refunds"), d_paid=dax(claim, "Applied to the Invoice"),
        d_score=column_dax(claim, "CustomerRefund", "Score"),
        m_lines=step(claim, table_m("CreditMemoLine"), "Merged Queries", "Added Custom"),
        m_ship=step(claim, table_m("CreditMemoLine"), "Merged Queries2", "Expanded ShipmentLine"),
        m_gl=step(claim, table_m("GLEntry"), "Filtered Rows"),
        c_flags=check_expr(claim, "refunds with their flags (Test 3)"),
        test1=tests[0][1], test3=tests[2][1])


# --- helpers of Requirements 6 to 10 ------------------------------------------------------------------------------------------

def row_of(sheet: str, text: str, exact: bool = True, col: int = 1, after: int = 0) -> int:
    """The row of the first cell of a column of the model workbook that is (or starts with) `text`."""
    ws = values()[sheet]
    for r in range(after + 1, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and (v == text if exact else v.startswith(text)):
            return r
    raise KeyError(f"no cell of column {col} on {sheet} {'is' if exact else 'starts with'} {text!r}")


def value_at(sheet: str, label: str, exact: bool = True, col: str = "B"):
    """The cached value next to a label of column A."""
    return cell(sheet, f"{col}{row_of(sheet, label, exact)}")


def formula_at(claim, sheet: str, label: str, exact: bool = True, col: str = "B") -> str:
    return formula(claim, sheet, f"{col}{row_of(sheet, label, exact)}")


def words_of(*texts) -> int:
    return sum(len(re.findall(r"\w+", str(t))) for t in texts if t)


def _mon_text(relative: str) -> str:
    return (MONITOR / "definition" / relative).read_text(encoding="utf-8")


def mon_table_m(name: str) -> list[str]:
    """The M of a loaded table of the Audit Monitoring project, one stripped line per step."""
    return _block(_mon_text(f"tables/{name}.tmdl"), "source =")


def mon_expression_m(name: str) -> list[str]:
    """The M of a query of Audit Monitoring that does not load, from its expressions.tmdl."""
    text = _mon_text("expressions.tmdl")
    for head in (f"expression {name} =", f"expression '{name}' ="):
        if head in text:
            return _block(text, head)
    raise KeyError(f"Audit Monitoring has no query {name}")


def mon_dax(claim, *names: str) -> str:
    table = _measures_in(MONITOR)
    for n in names:
        claim(n in table, f"the Audit Monitoring project has the measure {n}")
    return "\n\n".join(table.get(n, "") for n in names)


def mon_check_expr(claim, label: str) -> str:
    text = (MONITOR / "DAXQueries" / "Checks.dax").read_text(encoding="utf-8")
    m = re.search(r'"Check", "' + re.escape(label) + r'", "Expected", .*?, "Actual", (.*?), "Agrees"', text)
    claim(m is not None, f"the Audit Monitoring Checks tab has the row {label!r}")
    return m.group(1) if m else ""


def mon_check_sections() -> dict[str, int]:
    text = (MONITOR / "DAXQueries" / "Checks.dax").read_text(encoding="utf-8")
    counts: dict[str, int] = {}
    for sec in re.findall(r'ROW \( "Section", "([^"]*)"', text):
        counts[sec] = counts.get(sec, 0) + 1
    return counts


def mon_test_block(claim, n: int) -> str:
    text = (MONITOR / "DAXQueries" / "Tests.dax").read_text(encoding="utf-8").strip()
    found = [p.strip() for p in re.split(r"(?m)^(?=// Test \d)", text) if p.startswith(f"// Test {n}:")]
    claim(len(found) == 1, f"the Audit Monitoring Tests tab holds Test {n}")
    return found[0] if found else ""


def mon_entered(name: str) -> list[list]:
    """The rows of an Enter data table of the Audit Monitoring project (the project stores them compressed)."""
    m = re.search(r'Binary\.FromText\("([^"]+)"', _mon_text(f"tables/{name}.tmdl"))
    return json.loads(zlib.decompress(base64.b64decode(m.group(1)), -15))


# --- Requirement 6: the sample --------------------------------------------------------------------------------------------

@context("ch19.g6")
def g6(d, claim):
    need(d, claim, "CustomerRefund", "CreditMemo", "SalesReturn", "CashReceiptApplication")
    s = n19.sampling(d)
    c = n19.cycle(d)
    size, conf, tol = n19.SAMPLE, n19.CONFIDENCE, n19.TOLERABLE
    nk = len(s["key"])
    claim(cell("Sample", "B15") == size and abs(cell("Sample", "B12") - conf) < 1e-9 and abs(cell("Sample", "B13") - tol) < 1e-9
          and cell("Sample", "B14") == 0, "the Sample sheet holds the plan's parameters: 95% confidence, 5% tolerable, 0% expected, a size of 59")
    claim(cell("Sample", "B16") == 0, "the plan allows no deviation")
    # the key items, read from the sheet (the spill of the LET formula)
    r0 = row_of("Sample", "RefundNumber")
    key_rows = [[cell("Sample", f"{col}{r}") for col in "ABCDEFG"] for r in range(r0 + 1, r0 + 1 + nk)]
    claim([r[0] for r in key_rows] == [x["num"] for x in s["key"]], "the key items of the Sample sheet are the note's, in order")
    claim(all(r[3] >= n19.KEY_AMOUNT or r[4] > n19.KEY_CUSTOMER for r in key_rows), "every key item is a refund of 5,000 or more or a refund to a customer over 10,000")
    claim(cell("Sample", f"A{r0 + 1 + nk}") is None, "the key items end where the note says")
    totals: dict = {}
    for r in c.rf.values():
        totals[r["cust"]] = totals.get(r["cust"], 0.0) + r["amt"]
    focus, focus_n = Counter(r[2] for r in key_rows).most_common(1)[0]
    focus_total = max(r[4] for r in key_rows if r[2] == focus)
    other = [r for r in key_rows if r[2] != focus]
    claim(focus_total > n19.KEY_CUSTOMER and all(r[3] >= n19.KEY_AMOUNT for r in other),
          "one customer's refunds exceed 10,000 and every other key item is a refund of 5,000 or more")
    by_amount = sum(1 for r in c.rf.values() if r["amt"] >= n19.KEY_AMOUNT)
    claim(by_amount < nk, "a test of the refund's own amount alone misses the key items that come from a customer's total")
    key_if_5000 = sum(1 for r in c.rf.values() if r["amt"] >= n19.KEY_AMOUNT or totals[r["cust"]] > n19.KEY_AMOUNT)
    claim(key_if_5000 > 3 * nk, "reading 5,000 as the customer threshold more than triples the key items")
    exp_key = size * nk / len(c.rf)
    p_no_key = float(hypergeom(len(c.rf), nk, size).pmf(0))
    # the frozen draw
    first = row_of("Sample", "#") + 1
    drawn = [cell("Sample", f"B{r}") for r in range(first, first + size)]
    claim(len(set(drawn)) == size and cell("Sample", f"B{first + size}") is None, "the frozen draw holds 59 distinct refunds")
    claim(set(drawn) <= {x["num"] for x in s["rest"]}, "every refund of the frozen draw is in the remainder, so none is a key item")
    comment = workbook()["Sample"][f"B{first - 1}"].comment
    claim(comment is not None, "the header of the frozen draw carries a note with the formula that drew it")
    m = re.search(r"with\n(=.*?)\nand frozen", comment.text if comment else "", re.S)
    draw = clean(m.group(1)) if m else ""
    claim(draw.startswith("=LET(") and "RANDARRAY" in draw, "the note holds the LET formula with RANDARRAY")
    rule_row = row_of("Sample", "Deviations found in the sample")
    rule = [(cell("Sample", f"A{r}"), cell("Sample", f"B{r}"), cell("Sample", f"C{r}")) for r in range(rule_row + 1, rule_row + 5)]
    claim(all(abs(x[1] - u) < 1e-9 for x, u in zip(rule, s["udl"])) and [x[2] for x in rule] == ["Yes", "No", "No", "No"],
          "the rule of the Sample sheet is the note's: only zero deviations supports reliance")
    k_drawn = value_at("Sample", "Deviations in the sample")
    udl_drawn = value_at("Sample", "Upper deviation limit")
    dmg_drawn = value_at("Sample", "Of them, this draw holds")
    udl_dmg = value_at("Sample", "Upper deviation limit if every one of them passes")
    claim(s["lo"] <= k_drawn <= s["hi"], "the frozen draw holds a number of deviations inside the range that 93% of draws give")
    claim(abs(udl_drawn - float(beta.ppf(conf, k_drawn + 1, size - k_drawn))) < 1e-9, "the upper limit of the frozen draw is BETA.INV of its deviations")
    claim(abs(udl_dmg - float(beta.ppf(conf, 1, dmg_drawn))) < 1e-9 and udl_dmg > tol, "even if every Damaged refund of the frozen draw passed, its upper limit is above 5%")
    # the table's other columns, for the wrong size
    claim(cell("Sample", "F13") == conf and abs(cell("Sample", "F14") - 0.10) < 1e-9 and abs(cell("Sample", "G13") - 0.90) < 1e-9,
          "the Sample sheet's table columns are 95% and 5%, 95% and 10%, 90% and 5%, 90% and 10%")
    size_tol10, size_conf90 = int(cell("Sample", "F15")), int(cell("Sample", "G15"))
    udl_tol10 = float(beta.ppf(conf, 1, size_tol10))
    udl_conf90 = float(beta.ppf(conf, 1, size_conf90))
    claim(udl_tol10 > tol and udl_conf90 > tol, "a sample from the 10% column or the 90% column cannot support reliance at 5% even with no deviation")
    size_exp1 = int(cell("Sample", "E16"))                  # the table's size at a 1% expected rate, the plan a changed input would give
    udl1_exp1 = float(beta.ppf(conf, 2, size_exp1 - 1))     # one deviation, with the planned size where the tested count belongs
    claim(float(beta.ppf(conf, 2, size - 1)) > tol > udl1_exp1 or udl1_exp1 < float(beta.ppf(conf, 2, size - 1)),
          "one deviation gives a different bound at the planned size of another plan than at the tested count, a change that can cross the tolerable rate")
    size_2pct = int(cell("Sample", "E17"))
    claim(size_2pct / len(s["rest"]) > 0.8, "at 2% expected deviations the table's size is most of the remainder")
    dmg = s["dmg"]
    claim(dmg["mean"] > 10 and dmg["rest"] < size, "a draw of 59 holds about 13 refunds of Damaged returns, and the remainder holds only 48")
    claim(cell("Sample", f"H{first}") == cell("Sample", f"I{first}") == "Not tested: no document in the database",
          "the attribute columns of the frozen draw say Not tested: no document in the database")
    # the frozen plan and the guard (columns M and N of the Sample sheet, in a workbook built since the revision of 2026-10-17)
    ws_s = values()["Sample"]
    guard_row = {ws_s.cell(r, 13).value: r for r in range(1, ws_s.max_row + 1) if isinstance(ws_s.cell(r, 13).value, str)}
    guard = {label: ws_s.cell(r, 14).value for label, r in guard_row.items()}
    if guard:
        claim(guard.get("Plan changed since the freeze?") == "No" and guard.get("All intended items accounted for?") == "Yes",
              "the frozen plan agrees with the live plan, and every intended item is accounted for")
        claim(guard.get("Tested (valid, distinct, eligible, with an outcome)") == size,
              "the tested count of the frozen selection is the planned size")
    tested_row, changed_row = guard_row.get("Tested (valid, distinct, eligible, with an outcome)"), guard_row.get("Plan changed since the freeze?")
    return dict(
        guard=guard, f_tested=maybe_formula("Sample", f"N{tested_row}") if tested_row else "",
        f_changed=maybe_formula("Sample", f"N{changed_row}") if changed_row else "",
        size=size, conf=conf, tol=tol, nk=nk, key_rows=key_rows, focus=focus, focus_n=focus_n, focus_total=focus_total, by_amount=by_amount,
        key_if_5000=key_if_5000, exp_key=exp_key, p_no_key=p_no_key, n_rest=len(s["rest"]), n_all=len(c.rf), rule=rule,
        k_drawn=k_drawn, udl_drawn=udl_drawn, dmg_drawn=dmg_drawn, udl_dmg=udl_dmg, size_tol10=size_tol10, udl_tol10=udl_tol10,
        size_conf90=size_conf90, udl_conf90=udl_conf90, size_exp1=size_exp1, udl1_exp1=udl1_exp1, size_2pct=size_2pct, share_2pct=size_2pct / len(s["rest"]),
        pop_rate=len(s["dev"]) / len(s["rest"]), p_zero=f"{s['zero_draw']:.1e}", draw_rate=k_drawn / size,
        dmg_all_pass=float(beta.ppf(conf, 1, dmg["rest"])), dmg_k=dmg["k"],
        f_size=formula(claim, "Sample", "B15"), f_key=formula(claim, "Sample", f"A{r0 + 1}"), f_draw=draw,
        f_udl=formula(claim, "Sample", f"B{rule_row + 1}"), f_count=formula_at(claim, "Sample", "Deviations in the sample"),
        f_d1udl=formula_at(claim, "Sample", "Upper deviation limit"),
        f_a3=formula_at(claim, "Sample", "Refunds of Damaged returns"), f_a3udl=formula_at(claim, "Sample", "Upper deviation limit if every one of them passes"))


# --- Requirement 7: the balances and the timeline ------------------------------------------------------------------------

@context("ch19.g7")
def g7(d, claim):
    need(d, claim)
    r7 = n19.r7(d, claim)
    s1, s2, s3, s4, s5, s6 = (test(d, k) for k in ("S1", "S2", "S3", "S4", "S5", "S6"))
    lines, traced, above, qty_off, twice, price_off, total_off, lines_off = s1[0]
    claim(lines == traced and above == qty_off == twice == price_off == 0 and total_off == 1,
          "S1 finds every credit line traced and one half-cent line")
    # wrong version: the stored total against the unrounded product
    unrounded = d.one("SELECT COUNT(*) FROM CreditMemoLine WHERE ABS(LineTotal - Quantity * UnitPrice * (1 - Discount)) > 0.004")
    half_cent = d.one("SELECT COUNT(*) FROM CreditMemoLine WHERE ABS(LineTotal - Quantity * UnitPrice * (1 - Discount)) > 0.0051")
    claim(unrounded > 50 and half_cent == 0, "against the unrounded product with a tolerance of less than half a cent, many lines differ; with half a cent, none does")
    credits, invoices, rate, credit_rate, inv_tax_off, above_inv, sub_off, tax_off, tax_alone, with_freight, freight_over, parts_off = s2[0]
    claim(credits == invoices and above_inv == sub_off == tax_off == freight_over == parts_off == 0 and tax_alone == with_freight > 0,
          "S2 finds every credit recomputed, and a tax test on SubTotal alone flags the credits with freight")
    claim(inv_tax_off == 0 and abs(rate - credit_rate) < 1e-9,
          "S2 corroborates the tax of every original invoice at the rate on SubTotal plus freight, the rate the credits' own totals give")
    open_items = [(r[0], r[5]) for r in s3]
    claim(abs(sum(x[1] for x in open_items) - n19.balance_2060(d, d.C)) < 0.005, "the open items of 2060 add up to its balance at the last year-end")
    claim([x[0] for x in open_items] == [x["number"] for x in n19.open_2060(d, d.C)], "S3 lists the note's open items")
    by_reason = {r[0]: r for r in s4}
    n_off = sum(r[7] for r in s4)
    gap = round(sum(r[8] for r in s4), 2)
    claim(n_off == r7["n_frac"] and abs(gap - r7["frac_gap"]) < 0.005, "S4 finds the note's lines restocked below standard")
    dmg = by_reason["Damaged"]
    claim(abs(dmg[5] - r7["dmg_total"]) < 0.005, "S4's Damaged total is the note's")
    other_c, dep_lo, dep_hi, raw_lo, raw_hi, raw_over, raw_early, lag_lo, lag_hi, not_full, early_r, early_a, early_amt = s5[0]
    lap = n19.lapping(d)
    claim((raw_lo, raw_hi, raw_over, raw_early) == (lap["raw_lo"], lap["raw_hi"], lap["raw_over"], lap["raw_over_early"])
          and raw_over == raw_early, "S5 shows the raw delay of the note, every long one on a receipt dated before its invoice")
    claim(other_c == 0 and not_full == 0 and lag_lo >= 0 and early_r > 0, "S5 finds no application to another customer's invoice and the deposits held in 2060")
    # the raw delay, from the receipt date alone: the observation that the receipts dated before their invoice explain
    lag_receipt = d.q("SELECT MAX(julianday(cra.ApplicationDate) - julianday(cr.ReceiptDate)) FROM CashReceiptApplication cra "
                      "JOIN CashReceipt cr ON cr.CashReceiptID = cra.CashReceiptID")[0][0]
    lag_receipt_n = d.one("SELECT COUNT(*) FROM CashReceiptApplication cra JOIN CashReceipt cr ON cr.CashReceiptID = cra.CashReceiptID "
                          "WHERE julianday(cra.ApplicationDate) - julianday(cr.ReceiptDate) > ?", lag_hi)
    claim(lag_receipt > 5 * lag_hi and lag_receipt_n > 100, "measured from the receipt date alone, some applications come weeks later, which the advances explain")
    # the largest refund and its timeline
    big = n19.largest_refund(d)
    c = n19.cycle(d)
    cm = c.cm[big["cm"]]
    inv_id = cm["inv"]
    inv_lines = d.one("SELECT COUNT(*) FROM SalesInvoiceLine WHERE SalesInvoiceID = ?", inv_id)
    ship_rows = d.one("SELECT COUNT(*) FROM SalesInvoiceLine sil JOIN ShipmentLine sl ON sl.ShipmentLineID = sil.ShipmentLineID "
                      "JOIN Shipment s ON s.ShipmentID = sl.ShipmentID WHERE sil.SalesInvoiceID = ?", inv_id)
    ships = d.one("SELECT COUNT(DISTINCT sl.ShipmentID) FROM SalesInvoiceLine sil JOIN ShipmentLine sl ON sl.ShipmentLineID = sil.ShipmentLineID "
                  "WHERE sil.SalesInvoiceID = ?", inv_id)
    claim(ship_rows > ships, "the largest refund's invoice has more lines than shipments, so a join without DISTINCT repeats the shipment")
    claim(len(s6) == ships + 1 + len(c.apps[inv_id]) + 1 + 1 + sum(1 for a in c.adj if a["cm"] == cm["id"]) + 1,
          "S6 lists the shipment, invoice, receipts, return, credit, clawbacks and refund")
    weekend = d.one("SELECT COUNT(*) FROM CustomerRefund WHERE strftime('%w', RefundDate) IN ('0', '6')")
    refunds = d.one("SELECT COUNT(*) FROM CustomerRefund")
    claim(abs(weekend / refunds - 2 / 7) < 0.05, "weekend dates hold about two sevenths of the refunds, so a Sunday adds nothing")
    unpaid = value_at("Timeline", "Unpaid on the invoice on the refund date")
    claim(abs(unpaid - n19.m2(d, claim)["unpaid"]) < 0.005, "the Timeline sheet's unpaid balance is the note's")
    claim(cell("Timeline", "B7") == "Sunday" and cell("Timeline", "B15") is True and cell("Timeline", "B19") is True,
          "the live block of the Timeline sheet reads a Sunday refund of the largest credit, dated before the last application")
    t0 = row_of("Timeline", "Date")
    steps = [[cell("Timeline", f"{col}{r}") for col in "ABCDE"] for r in range(t0 + 1, t0 + 12)]
    claim(steps[-1][1] == "Invoice settled" and len(steps) == len(s6) + 1, "the Timeline sheet holds S6's steps and the settled invoice")
    rr = row_of("Timeline", "Damaged")
    return dict(
        lines=lines, traced=traced, unrounded=unrounded, tax_alone=tax_alone, with_freight=with_freight, n_open=len(open_items),
        open_items=open_items, open_total=round(sum(x[1] for x in open_items), 2), n_off=n_off, gap=gap, lag_receipt=lag_receipt,
        lag_receipt_n=lag_receipt_n, early_r=early_r, early_a=early_a, early_amt=early_amt, inv_lines=inv_lines, ship_rows=ship_rows,
        ships=ships, n_steps=len(s6), weekend=weekend, refunds=refunds, unpaid=unpaid, big_inv=c.inv[inv_id]["num"],
        n_apps=len(c.apps[inv_id]), reasons=[r[0] for r in s4], n_reasons=len(s4), big=big["num"], big_amount=big["amt"],
        open_text="; ".join(f"{x[0]} (${x[1]:,.2f})" for x in open_items),
        f_max=formula(claim, "Timeline", "B4"), f_who=formula(claim, "Timeline", "B5"), f_dow=formula(claim, "Timeline", "B7"),
        f_sum=formula(claim, "Timeline", f"B{rr}"), f_score=formula_at(claim, "Timeline", "Score on the five refund flags (RefundFlags)"))


# --- Requirement 8: the evaluation of the deficiencies -----------------------------------------------------------------------

@context("ch19.g8")
def g8(d, claim):
    need(d, claim)
    m = n19.materiality(d)
    t = n19.tests(d)
    c = n19.cycle(d)
    r8 = n19.r8(d, claim)
    ws = values()["Deficiencies"]
    head = row_of("Deficiencies", "ID")
    rows = []
    r = head + 1
    while str(ws.cell(r, 1).value or "").startswith("D"):
        rows.append(dict(id=ws.cell(r, 1).value, pattern=ws.cell(r, 2).value, point=ws.cell(r, 3).value, kind=ws.cell(r, 4).value, n3=ws.cell(r, 5).value,
                         n26=ws.cell(r, 6).value, e3=ws.cell(r, 7).value, e26=ws.cell(r, 8).value, band=ws.cell(r, 9).value,
                         cls=ws.cell(r, 13).value, disp=ws.cell(r, 14).value))
        r += 1
    claim([x["id"] for x in rows] == ["D1", "D2", "D3", "D4", "D5", "D5", "D6", "D7", "D8"], "the matrix holds D1 to D8, with D5 on two rows")
    short = {"At or above materiality": "at or above materiality",
             "Above performance materiality, below materiality": "between performance materiality and materiality",
             "Above clearly trivial, below performance materiality": "between clearly trivial and performance materiality",
             "Clearly trivial": "clearly trivial"}
    for x in rows:
        x["short"] = short[x["band"]]

    def band(x):
        return ("At or above materiality" if x >= m["mat"] else "Above performance materiality, below materiality" if x >= m["perf"]
                else "Above clearly trivial, below performance materiality" if x >= m["trivial"] else "Clearly trivial")
    claim(all(band(x["e26"]) == x["band"] for x in rows), "every band of the matrix follows from its 2026 exposure and the three thresholds")
    g0_ = row_of("Deficiencies", "Group")
    groups = []
    for r in range(g0_ + 1, g0_ + 6):
        groups.append(dict(id=ws.cell(r, 1).value, accounts=ws.cell(r, 2).value, defs=ws.cell(r, 3).value, e26=ws.cell(r, 4).value,
                           potential=ws.cell(r, 5).value, band=ws.cell(r, 7).value, cls=ws.cell(r, 10).value))
    claim([g["cls"] for g in groups] == ["Significant deficiency", "Significant deficiency", "Deficiency",
                                         "Follow up (no policy in the records)", "Valuation follow-up"],
          "the combined evaluation rates credits and refunds significant deficiencies, the concentration a deficiency, and two follow-ups")
    claim(abs(groups[0]["e26"] - r8["u12"][1]) < 0.005 and abs(groups[1]["e26"] - r8["u34"][1]) < 0.005,
          "the groups' exposures are the note's, each document counted once")
    sum12 = rows[0]["e26"] + rows[1]["e26"]
    claim(sum12 >= m["mat"] > r8["u12"][1] > m["perf"], "adding D1 and D2 reaches materiality; counting each credit once stays below it, above performance materiality")
    claim(rows[0]["e3"] >= m["mat"], "D1's three-year exposure is above overall materiality")
    count_mism = next(x for x in rows if x["id"] == "D5" and "method" in x["pattern"])["n3"]
    claim(count_mism > next(x for x in rows if x["id"] == "D4")["n3"], "the method mismatch has more exceptions than the refunds before payment, but is rated lower")
    ca = r0 = row_of("Deficiencies", "Test", col=2)
    clean = [[ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 5).value] for r in range(ca + 1, ca + 9)]
    claim(all(x[0] for x in clean) and ws.cell(ca + 9, 2).value is None, "the clean results of the matrix are eight rows")
    tot = n19.amount(n19.in_year(c.cm.values(), d.C))
    old_recon = formula(claim, "Deficiencies", "C32")
    last_row = re.search(r"A5:A(\d+)", old_recon).group(1)
    recon_formula = old_recon if maybe_formula("Deficiencies", "E32") else NEW_RECON.format(last=last_row)
    if maybe_formula("Deficiencies", "E32"):
        claim(old_recon == NEW_RECON.format(last=last_row) and maybe_formula("Deficiencies", "E32") == NEW_RECON_CONCLUSION,
              "the Deficiencies sheet's reconciliation summary is the builder's formula that blocks an invalid difference")
    resid = f"{clean[0][1]:.1e}"
    claim(float(clean[0][1]) < 0.005, "the reconciliation differences add up to far less than a cent")
    return dict(resid=resid,
        mat=m["mat"], perf=m["perf"], trivial=m["trivial"], rows=rows, groups=groups, sum12=round(sum12, 2), union12=r8["u12"][1],
        union12_n=r8["u12"][0], d1_3y=rows[0]["e3"], mism=count_mism, d4_n=next(x for x in rows if x["id"] == "D4")["n3"],
        self_all=len(c.cm) + len(c.rf), n_clean=len(clean), clean=clean, potential_credits=tot, n_sig=cell("Deficiencies", "D26"),
        n_mw=cell("Deficiencies", "D27"), overall=cell("Deficiencies", "D28"), d7_registers=d.one("SELECT COUNT(*) FROM PayrollRegister"),
        f_band=formula(claim, "Deficiencies", "I8"), f_union=formula(claim, "Deficiencies", "D20"),
        f_potential=formula(claim, "Deficiencies", "E20"), f_clean=recon_formula,
        f_overall=formula(claim, "Deficiencies", "D28"), f_exposure=formula(claim, "Deficiencies", "H8"),
        f_clean_concl=maybe_formula("Deficiencies", "E32") or NEW_RECON_CONCLUSION,
        names_ok=all(defined(claim, n) for n in ("Materiality", "PerformanceMateriality", "ClearlyTrivial")))


# --- Requirement 9: the report ------------------------------------------------------------------------------------------------

# the three action plans whose owners are expected to push back (the refund plan's text changed in the revision of 2026-10-17)
PUSHBACK = ("Have credits approved by someone", "Pay refunds to the original", "Verify the customer terms", "Review the concentration")


@context("ch19.g9")
def g9(d, claim):
    need(d, claim)
    m = n19.materiality(d)
    t = n19.tests(d)
    r8 = n19.r8(d, claim)
    ws = values()["Report"]
    claim(all(cell("Report", f"A{r}") for r in range(5, 11)), "the Report sheet has six key figures")
    keys = [(cell("Report", f"A{r}"), cell("Report", f"D{r}")) for r in range(5, 11)]
    claim(keys[0][1] == len(t.cm_above) and keys[1][1] == len(t.cash_credit) and keys[2][1] == len(t.rf_above)
          and keys[3][1] == len(t.before), "the Report sheet's key figures are the counts of Requirement 4")
    h = row_of("Report", "Finding")
    findings = [dict(id=ws.cell(r, 1).value.split(".")[0], title=fresh(ws.cell(r, 1).value), rating=ws.cell(r, 6).value)
                for r in range(h + 1, h + 7)]
    claim([f["id"] for f in findings] == [f"F{i}" for i in range(1, 7)], "the Report sheet lists six findings, F1 to F6")
    claim([f["rating"] for f in findings] == ["Significant deficiency (provisional)", "Significant deficiency (provisional)",
                                              "Deficiency (may combine entity-wide)", "Follow up", "Valuation follow-up", "Follow up"],
          "the report's ratings are the provisional ones of Requirement 8")
    a0 = row_of("Report", "Action")
    plans = []
    r = a0 + 1
    while ws.cell(r, 1).value and not str(ws.cell(r, 1).value).startswith("Action plans"):
        plans.append(dict(action=ws.cell(r, 1).value, owner=ws.cell(r, 3).value, date=ws.cell(r, 4).value,
                          response=ws.cell(r, 5).value, pushback=ws.cell(r, 6).value))
        r += 1
    claim(len(plans) == value_at("Report", "Action plans with an owner and a date", col="D"), "the Report sheet counts its action plans")
    dates = sorted(p["date"] for p in plans)
    owners = sorted({p["owner"] for p in plans})
    conclusion = ws.cell(row_of("Report", "Overall conclusion", exact=False), 1).value
    claim("no assurance" in conclusion and "lie outside the database" in conclusion and "no material weakness" in conclusion,
          "the report's conclusion says what the tests could not reach and gives no assurance on it")
    claim(sum(1 for p in plans if p["pushback"] and "None expected" not in p["pushback"]) >= 5,
          "most action plans name where the owner is expected to disagree")
    narrative = words_of(*(ws.cell(rr, 1).value for rr in (13, 14, 15, 27)),
                         *(ws.cell(rr, cc).value for rr in range(h + 1, h + 7) for cc in range(1, 8)),
                         *(p[k] for p in plans for k in ("action", "response", "pushback")))
    claim(all(any(k in p["action"].lower() for p in plans) for k in ("delegation of authority", "unpaid invoices", "concentration of duties",
                                                                     "original instrument", "damaged items", "largest refund")),
          "every finding of the report has an action plan")
    claim(sum(1 for p in plans if p["action"].startswith(PUSHBACK)) == 3,
          "the three plans whose owners are expected to push back are in the model's table")
    claim(r8["u12"][1] < m["mat"] < n19.amount(n19.in_year(n19.cycle(d).cm.values(), d.C)),
          "a report that sets the year's credits against materiality, not the credits that failed, overstates the case")
    return dict(
        figs=keys, popn=n19.m1(d, claim), findings=findings, plans=plans, n_plans=len(plans), first_date=dates[0].date().isoformat(),
        last_date=dates[-1].date().isoformat(), owners=owners, n_owners=len(owners), narrative=narrative,
        mat=m["mat"], d1_3y=n19.amount(t.cm_above), d1_26=r8["d1"], before_def=(len(t.before), len(t.before_rcpt)),
        f_key=formula(claim, "Report", "D5"), f_sig=formula(claim, "Report", "D9"), f_overall=formula(claim, "Report", "D10"),
        f_count=formula_at(claim, "Report", "Action plans with an owner and a date", col="D"),
        pushback=[(p["owner"], p["response"], fresh(p["pushback"])) for p in plans if p["action"].startswith(PUSHBACK)],
        am_registers=r8["am_largest"])


# --- Requirement 10: the continuing tests ----------------------------------------------------------------------------------

@context("ch19.g10")
def g10(d, claim):
    need(d, claim)
    r10 = n19.r10(d, claim)
    r4 = n19.r4(d, claim)
    c, t = n19.cycle(d), n19.tests(d)
    old = n19.register(d)
    reviewed = sum(1 for e in old if e["disposition"])
    new_rows = r10["new"] - r10["old"]
    tests_rows = mon_entered("Tests")
    claim(len(tests_rows) == 15 and [r[0] for r in tests_rows[-4:]] == ["CM AboveLimit", "CM CashAndCredit", "RF AboveLimit", "RF BeforePayment"]
          and {r[2] for r in tests_rows[-4:]} == {"Credit memos", "Refunds"}, "the Tests table holds Chapter 16's eleven tests and the four new ones")
    # one row for each supported cause, with its owner, dates and closure fields, and the document-level mapping (a project built
    # since the revision of 2026-10-17); the earlier project held one Design Findings row for each test
    new_findings = (MONITOR / "definition" / "tables" / "Findings.tmdl").is_file()
    if new_findings:
        findings = mon_entered("Findings")
        mapping = mon_entered("Finding Exceptions")
        claim(len(findings) == 2 and all(f[4] == n19.year_end(d.C) for f in findings),
              "the Findings table holds the two supported causes, each covering the exceptions through the last year-end")
        claim(all(f[6] is None and f[7] is None and f[8] is None and f[9] is None for f in findings),
              "no finding has an implementation date, a retest or a closure, so every finding is open")
        claim(len(mapping) == r10["new"] - r10["old"] and len({(m_[0], m_[1]) for m_ in mapping}) == len(mapping),
              "the mapping holds one row for each historical exception of the new tests")
        fix_by = findings[0][5]
    else:
        findings = mon_entered("Design Findings")
        mapping = []
        claim(len(findings) == 4 and all(f[2] == n19.year_end(d.C) for f in findings), "the Design Findings table covers each new test through the last year-end")
        fix_by = findings[0][3]
    log = mon_entered("Change Log")
    claim(len(log) == 5, "the Change Log holds five rows")
    pop = mon_dax(claim, "Population")
    claim('"Credit memos"' in pop and '"Refunds"' in pop, "the Population measure counts the credits and the refunds")
    through_dax = mon_dax(claim, "Data Through")
    claim(through_dax.count("MAX (") == 5 and "MINX" in through_dax, "Data Through is the earliest of five last dates")
    has_visual(claim, "creditsCycle", "newTests", "registerCards", "findings", "cycleAnswer", report=MON_REPORT)
    has_visual(claim, "anomalies", "anomalyYears", "anomalyAnswer", report=MON_REPORT)
    has_visual(claim, "changeLog", "log", "logAnswer", report=MON_REPORT)
    titles = {v: title(claim, p, v, MON_REPORT) for p, v in (("creditsCycle", "newTests"), ("creditsCycle", "findings"),
                                                            ("anomalies", "anomalyYears"), ("changeLog", "log"))}
    # wrong versions
    any_app = d.one("SELECT COUNT(*) FROM CreditMemo cm WHERE EXISTS (SELECT 1 FROM CashReceiptApplication a "
                    "WHERE a.SalesInvoiceID = cm.OriginalSalesInvoiceID)")
    claim(any_app > 2 * len(t.cash_credit), "merging the applications on the invoice alone flags most credits, not those whose approver applied the cash")
    same_day = d.one("SELECT COUNT(*) FROM CustomerRefund AS rf INNER JOIN CreditMemo AS cm ON cm.CreditMemoID = rf.CreditMemoID "
                     "INNER JOIN (SELECT SalesInvoiceID, MAX(ApplicationDate) AS lp FROM CashReceiptApplication GROUP BY 1) AS l "
                     "ON l.SalesInvoiceID = cm.OriginalSalesInvoiceID WHERE l.lp >= rf.RefundDate")
    claim(same_day > len(t.before), "a refund dated on or before the last application flags more refunds than one dated before it")
    open_flood = (len(old) - reviewed) + new_rows
    anomaly_flood = r10["new"] + r4["mism"]["n"] + r4["pd2"]["n"]
    last = dict(JE=max(e["date"] for e in n19.entries(d)), PO=max(o["date"] for o in n19.orders(d)), PR=r10["through"],
                CM=r10["cm_last"], RF=r10["rf_last"])
    max_through = max(last.values())
    claim(max_through > r10["through"] and min(last.values()) == r10["through"], "the latest of the five last dates is later than payroll's, which is the earliest")
    after = sum(1 for rows in (t.cm_above, t.cash_credit, t.rf_above, t.before) for r in rows if r["date"] > n19.year_end(d.C))
    claim(after == 0, "no exception of the new tests is dated after the last year-end")
    sections = mon_check_sections()
    m_cm = mon_table_m("CreditMemo")
    m_rf = mon_table_m("CustomerRefund")
    m_ex = mon_table_m("Exceptions")
    claim(sections.get("Requirement 10", 0) > 0 and sections.get("Milestone 2") == 7 and sections.get("Milestone 3") == 6,
          "the Checks tab holds the rows of Requirement 10 and of Milestones 2 and 3")
    return dict(
        reviewed=reviewed, new_rows=new_rows, open_old=len(old) - reviewed, open_flood=open_flood, anomaly_flood=anomaly_flood,
        any_app=any_app, n_credits=len(c.cm), same_day=same_day, max_through=max_through, last=last, titles=titles,
        n_tests=len(tests_rows), n_findings=len(findings), n_log=len(log), checks=sum(sections.values()), per_section=sections,
        n_r10_checks=sections["Requirement 10"] + sections["Requirement 10 (Chapter 16's checks, on the larger register)"],
        fix_by=fix_by, new_findings=new_findings, n_mapped=len(mapping), mism=r4["mism"]["n"], pd2=r4["pd2"]["n"],
        m_last="\n".join("    " + l.rstrip(",") for l in mon_expression_m("Last Payments")[2:3]),
        m_appliers="\n".join("    " + l.rstrip(",") for l in mon_expression_m("Appliers")[2:3]),
        m_cm=step(claim, m_cm, "Merged Queries1", "Added Custom1"),
        m_rf=step(claim, m_rf, "Merged Queries2", "Added Custom1"),
        m_cmx=step(claim, mon_expression_m("CM Exceptions"), "Unpivoted Only Selected Columns", "Added Custom"),
        m_cov=step(claim, m_ex, "Added Custom"),
        m_anom=step(claim, m_rf, "Added Custom2", "Added Custom3"),
        d_pop=pop, d_open=mon_dax(claim, "Open Exceptions"), d_cov=mon_dax(claim, "Design Finding Exceptions"), d_through=through_dax,
        d_status=mon_dax(claim, "Finding Status") if new_findings else "",
        d_open_findings=mon_dax(claim, "Open Findings") if new_findings else "",
        d_rate=mon_dax(claim, "Rate per 1,000"), test6=mon_test_block(claim, 6),
        test_rows=[r for r in tests_rows[-4:]], finding_rows=findings)

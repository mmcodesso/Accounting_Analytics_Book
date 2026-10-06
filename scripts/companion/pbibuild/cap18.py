"""Product Costs.pbip: the Power BI part of the solution to the Chapter 18 capstone case ("Make, Buy, or Reprice"),
the explorer for the meeting that Requirement 5 asks for.

The case starts a new Power BI file (Getting Started). As Requirement 5 says, it loads the item cost Table of the case
workbook (Make Buy Reprice.xlsx, the Excel part of this solution, built by xlbuild/cap18.py) and the invoice lines and
invoices of CharlesRiver.xlsx, adds a year column, and relates the lines to the item costs and the invoices. Design
services have no row in the item cost Table, so their lines show as a (Blank) family: a meaningful blank (Chapter 13),
left out of the Families and Concentration pages by a page filter and shown as its own row on the Validation page, so
the totals still reconcile. A Costing View table (Enter data) switches the Product Cost measure as the Steps table of
Tutorial 15.1 switched Gross to Net, with Margin and Margin % beside it.

The workbook's capacity table (Requirement 3, on its Activities sheet) is a range of formulas, not an Excel Table, so
the Capacity page rebuilds it from the two Tables it is made from: Activity (the routing hours of the year's
completions by family and work center, unpivoted) and Calendar (the hours available), at the practical share of
tbl-18-02. Four more small Tables of the workbook, each pasted there from MakeBuy.sql, give the committed cost of the
capacity (PayClass and Debits1090) and the ledger tie-outs of the Validation page (FamilyResults and Ledger5080).

The workbook is a second source. Its queries read C:\\CharlesRiver\\Make Buy Reprice.xlsx, and build.py points the
verification copy at the instructor workbook (the manifest's also_sources). Power Query types the workbook's columns
from the cached values of each Table's first 200 rows; a formula that returns "" reads as text.

Checks: every value of the instructor notes of Requirement 5 and Milestone 2 that the explorer shows, and the
Requirement 3 capacity values of its Capacity page, computed independently from CharlesRiver.sqlite by
facts/notes/ch18.py (the SQL twin of the case workbook) or by SQL here. The item costs are read from the workbook
unrounded, so the absorption margin is the note's unrounded value.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from functools import cached_property, lru_cache
from pathlib import Path

from paths import REPO
from pbibuild.ex13 import expand_all
from pbibuild.model import MEASURES_TABLE, Column, Measure, Model, Query, Table, query_table
from pbibuild.pbir import Page, Report, Visual, agg, col, lit, meas, textbox
from pbibuild.project import Project
from pbibuild.reports import COUNT, MONEY, card, claim, entered, hide, keep, money, slicer
from xlbuild import pq
from xlbuild.expected import Expected

from db import Data  # noqa: E402  (facts/ is on sys.path through pbibuild.reports)
from notes import ch18  # noqa: E402

FILE = "Product Costs"
WORKBOOK = "Make Buy Reprice.xlsx"
NEUTRAL_WB = "C:\\CharlesRiver\\" + WORKBOOK          # where the shipped project looks for the case workbook
DEFAULT_WB = REPO / "outputs" / "companion" / "solutions" / "Chapter18-capstone-excel" / WORKBOOK
PCT = "0.0%"
HOURS = "#,0.0"
KM = MEASURES_TABLE
M = lambda n: meas(KM, n)                 # noqa: E731
VIEWS = [("std", "Standard"), ("absorp", "Absorption"), ("td", "Time-driven")]
FG = ("Manufactured", "Purchased")        # the finished goods' supply modes: the item cost Table's rows


@dataclass
class Build:
    xlsx: Path
    exp: Expected
    year: int
    sources: dict = field(default_factory=dict)        # neutral path -> real path of the other sources (also_sources)
    project: Project = None

    def __post_init__(self):
        self.project = Project(FILE, Model(FILE), Report())

    @property
    def model(self) -> Model:
        return self.project.model

    @property
    def report(self) -> Report:
        return self.project.report

    @property
    def workbook(self) -> Path:
        p = Path(self.sources.get(NEUTRAL_WB, DEFAULT_WB))
        if not p.exists():
            raise SystemExit(f"build the Excel part of the case first (python scripts/companion/build.py --solutions "
                             f"--id ch18-capstone): no {p}")
        return p

    def check(self, section: str, label: str, expected, dax: str, tolerance: float = 0.005) -> None:
        self.project.check(section, label, expected, dax, tolerance)

    @cached_property
    def data(self) -> Data:
        return Data()


# --- the case workbook as a source --------------------------------------------------------------------------------

@lru_cache(maxsize=None)
def workbook_types(workbook: Path, table: str) -> tuple[tuple[str, str], ...]:
    """The types Power Query detects for each column of a Table of the case workbook from its first 200 rows, read from
    the cached values Excel saved (as pq.detected_types does for CharlesRiver.xlsx). A formula that returns "" is saved
    as an empty text, which Power Query reads as text."""
    import openpyxl
    from datetime import datetime
    values, formulas = openpyxl.load_workbook(workbook, data_only=True), openpyxl.load_workbook(workbook)
    ws = next(w for w in values.worksheets if table in dict(w.tables.items()))
    rows = ws[dict(ws.tables.items())[table]]
    header = [c.value for c in rows[0]]
    out = []
    for i, name in enumerate(header):
        vals = []
        for r in rows[1:201]:
            v = r[i].value
            if v is None and formulas[ws.title][r[i].coordinate].value is not None:
                v = ""
            if v is not None:
                vals.append(v)
        if not vals:
            kind = "type any"
        elif all(isinstance(v, datetime) for v in vals):
            kind = "type datetime"
        elif all(isinstance(v, bool) for v in vals):
            kind = "type logical"
        elif all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals):
            kind = "Int64.Type" if all(float(v).is_integer() for v in vals) else "type number"
        elif all(isinstance(v, str) for v in vals):
            kind = "type text"
        else:
            kind = "type any"
        out.append((name, kind))
    values.close()
    formulas.close()
    return tuple(out)


class WorkbookQuery(Query):
    """A query on a Table of the case workbook: Source, Navigation (<Table>_Table), and Changed Type, as the Navigator's
    Transform Data creates them. It always reads the neutral path; build.py repoints the verification copy."""

    @classmethod
    def table(cls, name: str, workbook: Path, table: str) -> "WorkbookQuery":
        q = cls(name)
        types = workbook_types(workbook, table)
        q.columns = dict(types)
        q.source = ("workbook", table, types)
        q._counts["Changed Type"] = 1
        return q

    def m(self, path: str) -> str:
        _, table, types = self.source
        nav = f"{table}_Table"
        return pq.steps_query([
            ("Source", f"Excel.Workbook(File.Contents({pq.m_string(NEUTRAL_WB)}), null, true)"),
            (nav, f"Source{{[Item={pq.m_string(table)},Kind=\"Table\"]}}[Data]"),
            ("Changed Type", f"Table.TransformColumnTypes({nav},{pq.type_list(list(types))})")] + self.steps)


def short_name(name: str) -> str:
    """The work center's name as the workbook's Activity columns spell it (the M rule of the Work Centers query):
    'Cutting Work Center' -> Cutting, 'Finishing and Test Work Center' -> Finishing, 'Quality Assurance ...' -> QA."""
    w = name.split(" ")
    return w[0] if w[1] in ("Work", "and") else w[0][0] + w[1][0]


SHORT_M = ('let words = Text.Split([WorkCenterName], " ") in if words{1} = "Work" or words{1} = "and" then words{0} '
           'else Text.Start(words{0}, 1) & Text.Start(words{1}, 1)')


# --- DAX ---------------------------------------------------------------------------------------------------------

def year_fg(y: int, extra: str = "") -> str:
    """The filters of the Families page as CALCULATE arguments: the base year, and the finished goods (no blank)."""
    return f"SalesInvoice[Year] = {y}, NOT ISBLANK ( ItemCosts[Family] )" + (f", {extra}" if extra else "")


def view(v: str) -> str:
    return f"'Costing View'[View] = \"{v}\""


def summarized(columns: list[str], filters: list[str], measures: list[tuple[str, str]]) -> str:
    """A SUMMARIZECOLUMNS table expression, laid out as the Tests tab prints it."""
    parts = columns + filters + [f'"{n}", {e}' for n, e in measures]
    return "SUMMARIZECOLUMNS (\n" + ",\n".join("    " + p for p in parts) + "\n)"


def cumulative(rank: int, v: str, y: int) -> str:
    """The cumulative margin share at a rank under a view, as a visual evaluates it (SUMMARIZECOLUMNS gives
    ALLSELECTED the filters to keep)."""
    return (f"MAXX ( SUMMARIZECOLUMNS ( 'Item Rank'[Rank], TREATAS ( {{ {rank} }}, 'Item Rank'[Rank] ), "
            f"TREATAS ( {{ \"{v}\" }}, 'Costing View'[View] ), TREATAS ( {{ {y} }}, SalesInvoice[Year] ), "
            f"TREATAS ( {{ \"{FG[0]}\", \"{FG[1]}\" }}, ItemCosts[SupplyMode] ), \"@Share\", [Cumulative Margin Share] ), "
            "[@Share] )")


CUMULATIVE = """VAR TopRank = MAX ( 'Item Rank'[Rank] )
VAR Items =
    FILTER (
        ADDCOLUMNS ( ALLSELECTED ( ItemCosts[ItemCode] ), "@Margin", [Margin] ),
        NOT ISBLANK ( [@Margin] )
    )
VAR Shown =
    IF ( TopRank <= COUNTROWS ( Items ), SUMX ( TOPN ( TopRank, Items, [@Margin], DESC ), [@Margin] ) )
RETURN
    DIVIDE ( Shown, SUMX ( Items, [@Margin] ) )"""

PLANT_PRACTICAL = "CALCULATE ( [Practical Capacity], REMOVEFILTERS ( 'Work Centers' ) )"

MEASURES = [
    # (folder, name, DAX, format, description)
    ("Views", "Revenue", "SUM ( SalesInvoiceLine[LineTotal] )", MONEY, "Invoiced revenue: the sum of LineTotal."),
    ("Views", "Units", "SUM ( SalesInvoiceLine[Quantity] )", "#,0.00", "Units invoiced (quantities are fractional)."),
    ("Views", "Standard Cost", "SUMX ( SalesInvoiceLine, SalesInvoiceLine[Quantity] * RELATED ( ItemCosts[StandardCost] ) )",
     MONEY, "Units sold at each item's standard cost."),
    ("Views", "Absorption Cost",
     "SUMX ( SalesInvoiceLine, SalesInvoiceLine[Quantity] * RELATED ( ItemCosts[AbsorptionCost] ) )", MONEY,
     "Units sold at standard cost plus the family's variance per unit completed (actual absorption)."),
    ("Views", "Time-Driven Cost",
     "SUMX ( SalesInvoiceLine, SalesInvoiceLine[Quantity] * RELATED ( ItemCosts[TimeDrivenCost] ) )", MONEY,
     "Units sold at the time-driven cost: material at the actual rate, Factory Overhead per standard hour, and the "
     "committed cost per practical hour times the routing time."),
    ("Views", "Product Cost", 'SWITCH (\n    SELECTEDVALUE ( \'Costing View\'[View], "Standard" ),\n'
                              '    "Standard", [Standard Cost],\n    "Absorption", [Absorption Cost],\n'
                              '    "Time-driven", [Time-Driven Cost]\n)', MONEY,
     "The cost of the units sold under the costing view selected; Standard when no single view is selected."),
    ("Views", "Margin", "[Revenue] - [Product Cost]", MONEY, "Revenue less the product cost of the view selected."),
    ("Views", "Margin %", "DIVIDE ( [Margin], [Revenue] )", PCT, "Margin as a share of revenue."),
    ("Views", "Cost per Unit", "DIVIDE ( [Product Cost], [Units] )", MONEY,
     "The product cost of the view selected per unit sold, at the year's mix of items."),
    ("Views", "Cumulative Margin Share", CUMULATIVE, PCT,
     "The share of the margin of the items selected that the top N items earn, N being the rank on the axis."),
    ("Capacity", "Available Hours", "SUM ( 'Work Centers'[AvailableHours] )", "#,0.00",
     "Hours available in WorkCenterCalendar in the base year."),
    ("Capacity", "Practical Share", "0.8", "0%", "Practical capacity as a share of the available hours (tbl-18-02)."),
    ("Capacity", "Practical Capacity", "[Available Hours] * [Practical Share]", HOURS, "Available hours at the practical share."),
    ("Capacity", "Routing Time Used", "SUM ( 'Routing Time'[Hours] )", HOURS,
     "Routing time of the base year's completions: run hours per unit times the units, plus setup per work order."),
    ("Capacity", "Unused Capacity",
     "[Practical Capacity] - CALCULATE ( [Routing Time Used], REMOVEFILTERS ( 'Routing Time'[Family], "
     "'Routing Time'[ItemGroup] ) )", HOURS, "Practical capacity that the year's routing time left unused."),
    ("Capacity", "Share of Available", "DIVIDE ( [Routing Time Used], [Available Hours] )", PCT,
     "Routing time used as a share of the available hours."),
    ("Capacity", "Share of Practical", "DIVIDE ( [Routing Time Used], [Practical Capacity] )", PCT,
     "Routing time used as a share of practical capacity."),
    ("Capacity", "Committed Cost",
     'SUM ( PayClass[Total] ) + CALCULATE ( SUM ( Debits1090[Debits] ), Debits1090[Source] = "Depreciation" )', MONEY,
     "The committed cost of capacity: the manufacturing crew and salaried staff with their employer costs, and the "
     "depreciation charged to 1090."),
    ("Capacity", "Capacity Cost Rate", f"DIVIDE ( [Committed Cost], {PLANT_PRACTICAL} )", "#,0.00",
     "The committed cost per hour of the plant's practical capacity."),
    ("Capacity", "Crew Rate",
     f'DIVIDE ( CALCULATE ( SUM ( PayClass[Total] ), PayClass[PayClass] = "Hourly" ), {PLANT_PRACTICAL} )', "#,0.00",
     "The hourly crew's cost per hour of the plant's practical capacity."),
    ("Capacity", "Unused Capacity Cost", "[Unused Capacity] * [Capacity Cost Rate]", MONEY,
     "Unused capacity at the capacity cost rate: a cost of the plant, not of the products."),
    ("Validation", "Invoice Lines", "COUNTROWS ( SalesInvoiceLine )", COUNT, "Invoice lines."),
    ("Validation", "Last Invoice", "MAX ( SalesInvoice[InvoiceDate] )", "yyyy-mm-dd", "The last invoice date."),
    ("Validation", "Invoices", "DISTINCTCOUNT ( SalesInvoiceLine[SalesInvoiceID] )", COUNT, "Invoices with lines."),
    ("Validation", "Workbook Revenue", "SUM ( ItemCosts[Revenue] )", MONEY,
     "The base year's revenue by item as the workbook's ItemCosts Table holds it."),
    ("Validation", "Workbook Margin", 'SWITCH (\n    SELECTEDVALUE ( \'Costing View\'[View], "Standard" ),\n'
                                     '    "Standard", SUM ( ItemCosts[MarginStandard] ),\n'
                                     '    "Absorption", SUM ( ItemCosts[MarginAbsorption] ),\n'
                                     '    "Time-driven", SUM ( ItemCosts[MarginTimeDriven] )\n)', MONEY,
     "The base year's margin of the view selected as the workbook's ItemCosts Table computes it."),
    ("Validation", "Variance Charged", "SUM ( FamilyResults[TotalVariance] )", MONEY,
     "The variance of the base year's work order closes, by family (the workbook's FamilyResults)."),
    ("Validation", "Ledger 5080", "SUM ( Ledger5080[Variance5080] )", MONEY,
     "Account 5080 Manufacturing Variance in the base year, closes left out (the workbook's Ledger5080)."),
    ("Validation", "Inputs to 1090", "SUM ( Debits1090[Debits] )", MONEY,
     "The base year's debits to 1090 other than the closes': payroll, Factory Overhead, and depreciation."),
    ("Validation", "Margin with the Year's Variance", f"CALCULATE ( [Margin], {view('Standard')} ) - [Variance Charged]",
     MONEY, "The ledger's view: margin at standard less the variance closed in the year."),
    ("Validation", "Time-Driven Margin after Unused Capacity",
     f"CALCULATE ( [Margin], {view('Time-driven')} ) - CALCULATE ( [Unused Capacity Cost], REMOVEFILTERS ( 'Work Centers' ) )",
     MONEY, "Time-driven margin less the cost of the capacity left unused."),
]


def answer(name: str, x, y, w, h, title: str, lines: list, size: int = 10) -> Visual:
    return textbox(name, x, y, w, h, [(title, True)] + lines, size=size)


def no_totals(kind: str) -> dict:
    """Totals off: a table's total row, or a matrix's column subtotals (the views are alternatives, never added)."""
    if kind == "tableEx":
        return {"total": [{"properties": {"totals": lit(False)}}]}
    return {"subTotals": [{"properties": {"columnSubtotals": lit(False)}}]}


# --- Requirement 5: the explorer ---------------------------------------------------------------------------------

def r5(b: Build) -> None:
    d, y, m, wb = b.data, b.year, b.model, b.workbook
    claim(y == d.C, "the base year is the last fiscal year of sales")
    ctx = ch18.r5(d, claim)                 # the note's claims: the blank lines are Services, the 5080 tie-out, ...
    cm = ch18.model(d)
    e = ctx["e"]

    # Queries: the item cost Table of the workbook, the invoice lines and invoices of CharlesRiver.xlsx (a year column)
    keep_ic = ["ItemID", "ItemCode", "ItemName", "Family", "ItemGroup", "SupplyMode", "StandardCost", "Units", "Revenue",
               "AbsorptionCost", "TimeDrivenCost", "MarginStandard", "MarginAbsorption", "MarginTimeDriven"]
    ic = WorkbookQuery.table("ItemCosts", wb, "ItemCosts")
    ic.select([c for c in ic.columns if c in keep_ic])
    claim(all(ic.columns[c] == "type number" for c in keep_ic[6:]), "the workbook's cost and margin columns are numbers")
    lines = Query.navigator("SalesInvoiceLine", b.xlsx, 17, "SalesInvoiceLine").select(
        ["SalesInvoiceLineID", "SalesInvoiceID", "ItemID", "Quantity", "LineTotal"])
    invoices = (Query.navigator("SalesInvoice", b.xlsx, 16, "SalesInvoice")
                .select(["SalesInvoiceID", "InvoiceNumber", "InvoiceDate"])
                .types({"InvoiceDate": "type date"})
                .custom("Year", "Date.Year([InvoiceDate])", "Int64.Type"))
    # The capacity table, rebuilt from the workbook's Activity and Calendar Tables
    act = WorkbookQuery.table("Routing Time", wb, "Activity")
    hour_cols = [c for c in act.columns if c.endswith(("Run", "Setup"))]
    centers = cm["centers"]
    claim(len(hour_cols) == 2 * len(centers), "the Activity Table has a run and a setup column for each work center")
    act.select(["Family", "ItemGroup"] + hour_cols)
    act.raw("Unpivoted Other Columns", 'Table.UnpivotOtherColumns({prev}, {"Family", "ItemGroup"}, "Attribute", "Value")',
            {"Family": "type text", "ItemGroup": "type text", "Attribute": "type text", "Value": "type number"})
    act.custom("Kind", 'if Text.EndsWith([Attribute], "Setup") then "Setup" else "Run"')
    act.custom("WorkCenter", "Text.Start([Attribute], Text.Length([Attribute]) - Text.Length([Kind]))")
    act.raw("Removed Columns", 'Table.RemoveColumns({prev},{"Attribute"})',
            {c: t for c, t in act.columns.items() if c != "Attribute"})
    act.rename({"Value": "Hours"}).types({"Kind": "type text", "WorkCenter": "type text", "Hours": "type number"})
    wcq = WorkbookQuery.table("Work Centers", wb, "Calendar").custom("WorkCenter", SHORT_M, "type text")
    names = dict(d.q("SELECT WorkCenterID, WorkCenterName FROM WorkCenter"))
    claim(sorted(short_name(n) for n in names.values()) == sorted({c[:-3] if c.endswith("Run") else c[:-5]
                                                                    for c in hour_cols}),
          "the Work Centers query's short names are the Activity Table's column prefixes")
    pay = WorkbookQuery.table("PayClass", wb, "PayClass").select(
        ["PayClass", "Employees", "GrossPay", "EmployerTax", "Benefits", "Total"])
    debits = WorkbookQuery.table("Debits1090", wb, "Debits1090")
    fam = WorkbookQuery.table("FamilyResults", wb, "FamilyResults").select(
        ["Family", "ItemGroup", "Closes", "TotalVariance", "UnitsCompleted", "VariancePerUnit"])
    ledger = WorkbookQuery.table("Ledger5080", wb, "Ledger5080")

    unit_cost = {"StandardCost": "none", "AbsorptionCost": "none", "TimeDrivenCost": "none", "ItemID": "none"}
    m.add(query_table(ic, formats={c: MONEY for c in keep_ic[6:] if c != "Units"} | {"Units": "#,0.00"},
                      hidden={"ItemID", "Units", "Revenue", "MarginStandard", "MarginAbsorption", "MarginTimeDriven"},
                      summarize=unit_cost))
    li = m.add(query_table(lines, formats={"LineTotal": MONEY, "Quantity": "#,0.00"}))
    hide(li, "SalesInvoiceLineID", "SalesInvoiceID", "ItemID")
    inv = m.add(query_table(invoices, summarize={"Year": "none"}))
    hide(inv, "SalesInvoiceID")
    m.add(query_table(act, formats={"Hours": HOURS}))
    m.add(query_table(wcq, formats={"AvailableHours": "#,0.00"}, summarize={"WorkCenterID": "none", "WorkingDays": "none"},
                      sort_by={"WorkCenterName": "WorkCenterID", "WorkCenter": "WorkCenterID"}))   # in routing order
    m.add(query_table(pay, formats={c: MONEY for c in ("GrossPay", "EmployerTax", "Benefits", "Total")}))
    m.add(query_table(debits, formats={"Debits": MONEY}))
    m.add(query_table(fam, formats={"TotalVariance": MONEY, "VariancePerUnit": MONEY, "UnitsCompleted": "#,0.00"},
                      summarize={"VariancePerUnit": "none"}))
    m.add(query_table(ledger, formats={"Variance5080": MONEY}))
    m.relate("SalesInvoiceLine.SalesInvoiceID", "SalesInvoice.SalesInvoiceID")
    m.relate("SalesInvoiceLine.ItemID", "ItemCosts.ItemID")
    m.relate("Routing Time.WorkCenter", "Work Centers.WorkCenter")
    # Enter data: the Costing View table (View sorted by Order), and the measures table
    cv = m.add(entered("Costing View", [("View", "string"), ("Order", "int64")],
                       [[name, str(i + 1)] for i, (_, name) in enumerate(VIEWS)]))
    cv.column("View").sort_by = "Order"
    hide(cv, "Order")
    m.query_order.append("Costing View")
    m.add(Table("Item Rank", [Column("Rank", "int64", "[Rank]", fmt="0", summarize="none")],
                dax='SELECTCOLUMNS ( GENERATESERIES ( 1, COUNTROWS ( ItemCosts ), 1 ), "Rank", [Value] )'))
    mt = m.add(entered(KM, [("Column1", "string")], []))
    m.query_order.append(KM)
    hide(mt, "Column1")
    mt.measures += [Measure(n, dax, fmt, desc, display_folder=folder) for folder, n, dax, fmt, desc in MEASURES]

    fams = sorted(cm["fam"].values(), key=lambda f: f["name"])
    groups = cm["groups"]
    example = cm["fam"][ctx["example"]]
    exact = e["exact"]
    unused_cost = cm["unused_h"] * cm["td_rate"]
    td_unused = exact["td"] - unused_cost
    last_invoice = d.one("SELECT MAX(InvoiceDate) FROM SalesInvoice")
    last_close = d.one("SELECT MAX(CloseDate) FROM WorkOrderClose")
    years = d.q("SELECT CAST(substr(si.InvoiceDate, 1, 4) AS INTEGER), COUNT(*), COUNT(DISTINCT l.SalesInvoiceID), "
                "SUM(l.LineTotal), SUM(l.Quantity * i.StandardCost) FROM SalesInvoiceLine l JOIN SalesInvoice si ON "
                "si.SalesInvoiceID = l.SalesInvoiceID JOIN Item i ON i.ItemID = l.ItemID GROUP BY 1 ORDER BY 1")
    cur = next(r for r in years if r[0] == y)
    groups_y = d.q("SELECT CASE WHEN i.ItemType = 'Finished Good' THEN i.ItemGroup END, COUNT(*), SUM(l.LineTotal), "
                   "SUM(l.Quantity * i.StandardCost) FROM SalesInvoiceLine l JOIN SalesInvoice si ON si.SalesInvoiceID = "
                   "l.SalesInvoiceID JOIN Item i ON i.ItemID = l.ItemID WHERE si.InvoiceDate BETWEEN ? AND ? GROUP BY 1",
                   f"{y}-01-01", f"{y}-12-31")
    year_filter = lambda name: keep(name, col("SalesInvoice", "Year"), [y])          # noqa: E731
    not_blank = lambda name: keep(name, col("ItemCosts", "Family"), [None], exclude=True)  # noqa: E731

    # About
    about = b.report.add(Page("about", "About"))
    about.add(textbox("aboutText", 20, 20, 1240, 470, [
        (FILE, True),
        "Purpose: the explorer for the CFO and the production manager. It shows what each family and item of finished "
        "goods earns under three costing views, and what the plant had and used of its capacity, in the base year, "
        f"fiscal {y}.",
        "Views: Standard is each item's standard cost. Absorption adds its family's variance per unit completed in the "
        "year, so the year's whole cost is spread over the units made. Time-driven charges material at the year's "
        "actual rate, the Factory Overhead entries per standard hour, and the committed cost of the crew, supervision, "
        "and equipment per hour of practical capacity times the item's routing time; the capacity left unused is a "
        "cost of the plant, shown on the Capacity page. Purchased finished goods carry their standard cost in every "
        "view. Choose a view in the slicer of the Families page; with none selected, the measures use Standard.",
        f"Sources: the ItemCosts Table of {WORKBOOK} (one row per finished good, from Requirement 4), and its Activity, "
        "Calendar, PayClass, Debits1090, FamilyResults, and Ledger5080 Tables (pasted there from MakeBuy.sql); the "
        f"invoice lines and invoices of CharlesRiver.xlsx. The invoices run through {last_invoice}, and the work order "
        f"closes through {last_close}.",
        "The blank: design services have no row in the item cost Table, so their lines show as a (Blank) family. They "
        "carry no standard cost, so they are left out of the Families and Concentration pages (a page filter keeps "
        "families that are not blank), and the Validation page shows them as the (Blank) row, so the totals still "
        "reconcile.",
        "Checks: the Validation page reproduces control totals of Chapters 6, 10, 13, and 14 and compares each view "
        "with the workbook; the Tests tab in DAX query view tests the same totals."], size=12))
    about.add(card("lastInvoice", 20, 510, 300, 90, [(M("Last Invoice"), "Last invoice")]))
    about.add(Visual("pageNavigator", "pageNavigator", 340, 520, 920, 60, objects={"pages": [{"properties": {
        "showHiddenPages": lit(False), "showTooltipPages": lit(False)}}]}))
    about.expect = [FILE, "Families", "Capacity", "Validation", last_invoice]

    # Families
    page = b.report.add(Page("families", "Families"))
    page.extra = {"filterConfig": {"filters": [year_filter("famYear"), not_blank("famNotBlank")]},
                  "visualInteractions": [{"source": "viewSlicer", "target": t, "type": "NoFilter"}
                                         for t in ("familyMatrix", "viewTotals")]}
    v_col, fam_col, grp_col = col("Costing View", "View"), col("ItemCosts", "Family"), col("ItemCosts", "ItemGroup")
    page.add(slicer("viewSlicer", 20, 20, 360, 80, v_col, ["Time-driven"], style="tile", single=True, title="Costing view"))
    page.add(card("viewCards", 400, 20, 860, 80, [M("Revenue"), M("Product Cost"), M("Margin"), M("Margin %")],
                  units_none=True))
    levels = [grp_col, fam_col]
    page.add(Visual("familyMatrix", "pivotTable", 20, 115, 700, 590,
                    {"Rows": levels, "Columns": [v_col], "Values": [M("Margin %"), M("Cost per Unit")]},
                    title=f"Margin % and cost per unit by family and costing view, manufactured items, fiscal {y}",
                    objects=no_totals("pivotTable"), filters=[keep("matrixMade", col("ItemCosts", "SupplyMode"),
                                                                   ["Manufactured"])],
                    extra=expand_all(levels), active=("Rows",)))
    page.add(Visual("viewTotals", "tableEx", 740, 115, 520, 150,
                    {"Values": [v_col, M("Revenue"), M("Product Cost"), M("Margin"), M("Margin %")]},
                    title=f"Every finished good sold in fiscal {y}, by costing view", objects=no_totals("tableEx")))
    page.add(Visual("itemsChart", "clusteredBarChart", 740, 280, 520, 425,
                    {"Category": [col("ItemCosts", "ItemCode")], "Y": [M("Margin")]}, sort=[(M("Margin"), "Descending")],
                    title="Items ranked by margin under the costing view selected",
                    alt_text=f"Finished goods sold in fiscal {y}, ranked by their margin under the costing view chosen "
                             "in the slicer, from the largest."))
    page.expect = ["Margin %", "Cost per Unit", money(exact["std"]), money(exact["absorp"]), money(exact["td"]),
                   money(e["rev"]), ctx["example"], f"{example['unit']['m_td']:.1%}", f"{example['unit']['td']:,.2f}"]

    # Concentration (the optional curve of the note)
    conc = b.report.add(Page("concentration", "Concentration"))
    conc.extra = {"filterConfig": {"filters": [year_filter("concYear"), not_blank("concNotBlank")]}}
    conc.add(Visual("curve", "lineChart", 20, 20, 1240, 520,
                    {"Category": [col("Item Rank", "Rank")], "Y": [M("Cumulative Margin Share")], "Series": [v_col]},
                    title=f"Cumulative share of margin by item rank, each costing view, finished goods sold in fiscal {y}",
                    alt_text="Three nearly identical curves rise steeply over the first items and flatten: about a "
                             "fifth of the items earn half of the margin under every view."))
    tops = e["tops"]
    conc.add(answer("concAnswer", 20, 560, 1240, 150, "Model answer: how concentrated the margin is", [
        f"The top fifth of the {e['n_items']} finished goods sold ({e['top_n']} items) earns "
        + ", ".join(f"{tops[k]['share']:.1%} of the margin {lab}" for k, lab in
                    (("std", "at standard"), ("absorp", "under absorption"), ("td", "time-driven"))) +
        f"; every curve reaches 100% at rank {e['n_items']} and stops there. The ranking barely depends on the view, "
        "because the views move each family's costs, not the order of its items. No item loses money under any view "
        f"(the lowest margin under absorption is {tops['absorp']['low']}, {tops['absorp']['low_margin']:.1%})."]))
    conc.expect = ["Cumulative share of margin by item rank", "Cumulative Margin Share"]

    # Capacity
    cap = b.report.add(Page("capacity", "Capacity"))
    wn, wshort = col("Work Centers", "WorkCenterName"), col("Work Centers", "WorkCenter")
    cap.add(card("capacityCards", 20, 20, 1240, 90, [M("Practical Capacity"), M("Routing Time Used"),
                                                     M("Unused Capacity"), M("Share of Practical"), M("Committed Cost"),
                                                     M("Capacity Cost Rate"), M("Unused Capacity Cost")],
                 units_none=True))
    cap.add(Visual("capacityTable", "tableEx", 20, 125, 800, 230,
                   {"Values": [wn, M("Available Hours"), M("Practical Capacity"), M("Routing Time Used"),
                               M("Share of Available"), M("Share of Practical"), M("Unused Capacity")]},
                   title=f"Capacity and its use by work center, fiscal {y} (the total is the plant)"))
    cap.add(Visual("capacityChart", "clusteredBarChart", 840, 125, 420, 230,
                   {"Category": [wn], "Y": [M("Practical Capacity"), M("Routing Time Used")]},
                   title="Practical capacity and routing time used (hours)",
                   alt_text="For each work center, practical capacity beside the routing time used: Assembly and "
                            "Finishing use the most of theirs, Quality Assurance very little."))
    fam_levels = [col("Routing Time", "ItemGroup"), col("Routing Time", "Family")]
    cap.add(Visual("familyHours", "pivotTable", 20, 370, 1240, 335,
                   {"Rows": fam_levels, "Columns": [wshort], "Values": [M("Routing Time Used")]},
                   title=f"Routing time used by family and work center, fiscal {y} (hours)",
                   extra=expand_all(fam_levels), active=("Rows",)))
    plant = dict(avail=cm["avail"], pract=cm["pract"], used=cm["used"])
    cap.expect = ["Routing Time Used", "Share of Practical", f"{cm['pract']:,.1f}", f"{cm['used']:,.1f}",
                  f"{cm['unused_h']:,.1f}", f"{cm['used'] / cm['pract']:.1%}", money(cm["committed"]),
                  money(unused_cost)]

    # Validation
    val = b.report.add(Page("validation", "Validation"))
    val.add(Visual("byYear", "tableEx", 20, 20, 760, 170, {"Values": [
        col("SalesInvoice", "Year"), M("Invoice Lines"), M("Invoices"), M("Revenue"), M("Standard Cost")]},
        title="Every invoice line by year (Chapters 13 and 14)"))
    val.add(Visual("byGroup", "tableEx", 800, 20, 460, 170, {"Values": [
        grp_col, M("Invoice Lines"), M("Revenue"), M("Standard Cost")]},
        title=f"Fiscal {y} by item group: (Blank) is the design services (Chapter 6)", filters=[year_filter("bgYear")]))
    val.add(Visual("viewsWorkbook", "tableEx", 20, 210, 760, 150, {"Values": [
        v_col, M("Revenue"), M("Workbook Revenue"), M("Margin"), M("Workbook Margin")]},
        title=f"Each view against the workbook's ItemCosts Table, finished goods, fiscal {y}",
        filters=[year_filter("vwYear"), not_blank("vwNotBlank")], objects=no_totals("tableEx")))
    val.add(Visual("inputs1090", "tableEx", 800, 210, 460, 150, {"Values": [
        col("Debits1090", "Source"), (agg("Debits1090", "Debits"), "Debits")]},
        title=f"Inputs to 1090, fiscal {y} (Chapter 10)"))
    val.add(card("ledgerCards", 20, 380, 1240, 110, [
        M("Variance Charged"), M("Ledger 5080"), M("Margin with the Year's Variance"),
        M("Time-Driven Margin after Unused Capacity")], units_none=True,
        filters=[year_filter("lcYear"), not_blank("lcNotBlank")]))
    val.add(answer("valAnswer", 20, 510, 1240, 200, "What the totals mean", [
        f"Revenue of the finished goods sold in fiscal {y} is {money(e['rev'])}: the year's {money(e['all_rev'])} less "
        f"the design services, {money(e['services'])}, the (Blank) row. The standard cost of the year's lines is "
        f"Chapter 14's {money(e['all_std'])}, because the services carry none.",
        f"Margin at standard is {money(exact['std'])}; charging the year's variance instead ({money(cm['totals']['total'])}"
        f", equal to account 5080) gives {money(e['year'])}, the ledger's view. The totals differ from absorption "
        f"({money(exact['absorp'])}) through the inventory build and timing: the variance per unit completed is charged "
        "only on the units sold.",
        f"Time-driven margin is {money(exact['td'])}, and {money(td_unused)} after the unused capacity "
        f"({money(unused_cost)}), which is a cost of the plant, not of the products."]))
    val.expect = [f"{cur[1]:,}", money(cur[3]), money(e["services"]), money(cm["totals"]["total"]), money(e["year"]),
                  money(td_unused), money(cm["inputs"])]
    b.report.active = "about"

    # The Tests tab
    fg_filter = f'TREATAS ( {{ "{FG[0]}", "{FG[1]}" }}, ItemCosts[SupplyMode] )'
    yr = f"TREATAS ( {{ {y} }}, SalesInvoice[Year] )"
    tests = [
        (f"Test 1: each view's fiscal {y} margin must equal the workbook's (finished goods; design services left out)",
         summarized(["'Costing View'[Order]", "'Costing View'[View]"], [yr, fg_filter], [
             ("Revenue", "ROUND ( [Revenue], 2 )"), ("Workbook Revenue", "ROUND ( [Workbook Revenue], 2 )"),
             ("Margin", "ROUND ( [Margin], 2 )"), ("Workbook Margin", "ROUND ( [Workbook Margin], 2 )"),
             ("Difference", "ROUND ( [Margin] - [Workbook Margin], 2 )")]), "'Costing View'[Order]", len(VIEWS)),
        ("Test 2: lines, invoices, revenue, and standard cost by year must equal Chapters 13 and 14",
         summarized(["SalesInvoice[Year]"], [], [
             ("Lines", "[Invoice Lines]"), ("Invoices", "[Invoices]"), ("Revenue", "ROUND ( [Revenue], 2 )"),
             ("Standard Cost", "ROUND ( [Standard Cost], 2 )")]), "SalesInvoice[Year]", len(years)),
        ("Test 3: the variance charged to the families must equal account 5080, and the inputs to 1090 Chapter 10's",
         "ROW (\n    \"Variance Charged\", ROUND ( [Variance Charged], 2 ),\n    \"Ledger 5080\", ROUND ( [Ledger 5080], 2 ),"
         "\n    \"Difference\", ROUND ( [Variance Charged] - [Ledger 5080], 2 ),\n"
         "    \"Inputs to 1090\", ROUND ( [Inputs to 1090], 2 )\n)", None, 1),
        ("Test 4: the capacity by work center must equal the workbook's capacity table (Requirement 3)",
         summarized(["'Work Centers'[WorkCenterID]", "'Work Centers'[WorkCenterName]"], [], [
             ("Available", "ROUND ( [Available Hours], 2 )"), ("Practical", "ROUND ( [Practical Capacity], 2 )"),
             ("Used", "ROUND ( [Routing Time Used], 2 )"), ("Share of Practical", "ROUND ( [Share of Practical], 4 )")]),
         "'Work Centers'[WorkCenterID]", len(centers)),
        (f"Test 5: the design services are the only lines with no item cost: the (Blank) row of fiscal {y}",
         summarized(["ItemCosts[ItemGroup]"], [yr], [
             ("Lines", "[Invoice Lines]"), ("Revenue", "ROUND ( [Revenue], 2 )"),
             ("Standard Cost", "ROUND ( [Standard Cost], 2 )")]), "ItemCosts[ItemGroup]", len(groups_y)),
        (f"Test 6: the top fifth of the items sold in fiscal {y} earns about half of the margin under every view",
         summarized(["'Costing View'[Order]", "'Costing View'[View]", "'Item Rank'[Rank]"],
                    [f"TREATAS ( {{ {e['top_n']}, {e['n_items']} }}, 'Item Rank'[Rank] )", yr, fg_filter],
                    [("Share", "ROUND ( [Cumulative Margin Share], 4 )")]),
         "'Costing View'[Order], 'Item Rank'[Rank]", 2 * len(VIEWS)),
    ]
    b.project.queries["Tests"] = "\n\n".join(
        f"// {c}\nEVALUATE\n{expr}" + (f"\nORDER BY {order}" if order else "") for c, expr, order, _ in tests) + "\n"

    t = "Requirement 5"
    n_fg = d.one("SELECT COUNT(*) FROM Item WHERE ItemType = 'Finished Good'")
    n_lines, n_inv = d.one("SELECT COUNT(*) FROM SalesInvoiceLine"), d.one("SELECT COUNT(*) FROM SalesInvoice")
    blank_lines = d.one("SELECT COUNT(*) FROM SalesInvoiceLine l JOIN Item i ON i.ItemID = l.ItemID "
                        "WHERE i.ItemType <> 'Finished Good'")
    b.check(t, "ItemCosts rows (every finished good)", n_fg, "COUNTROWS ( ItemCosts )", 0)
    b.check(t, "SalesInvoiceLine rows", n_lines, "COUNTROWS ( SalesInvoiceLine )", 0)
    b.check(t, "SalesInvoice rows", n_inv, "COUNTROWS ( SalesInvoice )", 0)
    b.check(t, "invoice lines with no invoice", 0,
            "COUNTROWS ( FILTER ( SalesInvoiceLine, ISBLANK ( RELATED ( SalesInvoice[Year] ) ) ) )", 0)
    b.check(t, "lines with no item cost row (the design services, every year)", blank_lines,
            "COUNTROWS ( FILTER ( SalesInvoiceLine, ISBLANK ( RELATED ( ItemCosts[ItemID] ) ) ) )", 0)
    b.check(t, f"{y} revenue of the (Blank) row (design services)", round(e["services"], 2),
            f"SUMX ( FILTER ( SalesInvoiceLine, ISBLANK ( RELATED ( ItemCosts[ItemID] ) ) && "
            f"RELATED ( SalesInvoice[Year] ) = {y} ), SalesInvoiceLine[LineTotal] )")
    b.check(t, f"{y} revenue, every line", round(e["all_rev"], 2), f"CALCULATE ( [Revenue], SalesInvoice[Year] = {y} )")
    b.check(t, f"{y} standard cost, every line (Chapter 14's)", round(e["all_std"], 2),
            f"CALCULATE ( [Standard Cost], SalesInvoice[Year] = {y} )", 0.01)
    b.check(t, f"{y} revenue, finished goods", round(e["rev"], 2), f"CALCULATE ( [Revenue], {year_fg(y)} )")
    b.check(t, f"{y} finished goods sold", e["n_items"],
            f"CALCULATE ( COUNTROWS ( FILTER ( VALUES ( ItemCosts[ItemCode] ), NOT ISBLANK ( [Revenue] ) ) ), {year_fg(y)} )", 0)
    b.check(t, f"{y} finished goods sold, by the workbook (Units > 0)", e["n_items"],
            "COUNTROWS ( FILTER ( ItemCosts, ItemCosts[Units] > 0 ) )", 0)
    for k, name in VIEWS:
        b.check(t, f"{y} product cost, {name}", round(e["rev"] - exact[k], 2),
                f"CALCULATE ( [Product Cost], {year_fg(y, view(name))} )", 0.01)
        b.check(t, f"{y} margin, {name} (item costs unrounded)", round(exact[k], 2),
                f"CALCULATE ( [Margin], {year_fg(y, view(name))} )", 0.01)
        b.check(t, f"{y} Margin %, {name}", round(exact[k] / e["rev"], 9),
                f"CALCULATE ( [Margin %], {year_fg(y, view(name))} )", 1e-6)
        b.check(t, f"{y} margin, {name}, in the workbook's ItemCosts Table", round(exact[k], 2),
                f"CALCULATE ( [Workbook Margin], {view(name)} )", 0.01)
        b.check(t, f"{y} margin less the workbook's, {name}", 0,
                f"CALCULATE ( [Margin] - [Workbook Margin], {year_fg(y, view(name))} )", 0.01)
    b.check(t, f"{y} margin with no view selected (falls back to Standard)", round(exact["std"], 2),
            f"CALCULATE ( [Margin], {year_fg(y)} )", 0.01)
    b.check(t, f"{y} revenue in the workbook's ItemCosts Table", round(e["rev"], 2), "[Workbook Revenue]", 0.01)
    b.check(t, f"{y} time-driven margin after unused capacity", round(td_unused, 2),
            f"CALCULATE ( [Time-Driven Margin after Unused Capacity], {year_fg(y)} )", 0.01)
    b.check(t, f"{y} margin with the year's variance charged (the ledger's view)", round(e["year"], 2),
            f"CALCULATE ( [Margin with the Year's Variance], {year_fg(y)} )", 0.01)
    b.check(t, "variance charged to the families", round(cm["totals"]["total"], 2), "[Variance Charged]")
    b.check(t, "account 5080 (closes left out)", round(ch18.ledger_5080(d), 2), "[Ledger 5080]")
    b.check(t, "variance charged less 5080", 0, "[Variance Charged] - [Ledger 5080]")
    b.check(t, "inputs to 1090", round(cm["inputs"], 2), "[Inputs to 1090]")
    for src, amount in sorted(cm["src"].items()):
        b.check(t, f"inputs to 1090: {src}", round(amount, 2),
                f"CALCULATE ( [Inputs to 1090], Debits1090[Source] = \"{src}\" )")
    for f in fams:
        fy = year_fg(y, f"ItemCosts[Family] = \"{f['name']}\", ItemCosts[SupplyMode] = \"Manufactured\"")
        for k, name in VIEWS:
            pct = {"std": f["m_std"], "absorp": f["m_unit"], "td": f["unit"]["m_td"]}[k]
            b.check(t, f"{f['name']} (manufactured) Margin %, {name}", round(pct, 9),
                    f"CALCULATE ( [Margin %], {fy}, {view(name)} )", 1e-6)
            b.check(t, f"{f['name']} (manufactured) cost per unit, {name}", round(f["unit"][k], 6),
                    f"CALCULATE ( [Cost per Unit], {fy}, {view(name)} )", 0.0001)
    for g in groups:
        gy = year_fg(y, f"ItemCosts[ItemGroup] = \"{g['name']}\", ItemCosts[SupplyMode] = \"Manufactured\"")
        for k, name in VIEWS[:2]:
            b.check(t, f"{g['name']} (manufactured) Margin %, {name}", round(g["m_std" if k == "std" else "m_unit"], 9),
                    f"CALCULATE ( [Margin %], {gy}, {view(name)} )", 1e-6)
    for k, name in VIEWS:
        b.check(t, f"top fifth ({e['top_n']} items) share of margin, {name}", round(tops[k]["share"], 9),
                cumulative(e["top_n"], name, y), 1e-6)
        b.check(t, f"share of margin at rank {e['n_items']}, {name}", 1, cumulative(e["n_items"], name, y), 1e-9)
    b.check(t, f"share of margin beyond rank {e['n_items']} is blank", 1,
            f"IF ( ISBLANK ( {cumulative(e['n_items'] + 1, 'Standard', y)} ), 1, 0 )", 0)
    b.check(t, "Item Rank rows (one per finished good)", n_fg, "COUNTROWS ( 'Item Rank' )", 0)
    b.check(t, "Costing View rows", len(VIEWS), "COUNTROWS ( 'Costing View' )", 0)
    # the Capacity page (the values of Requirement 3's capacity table)
    b.check(t, "routing time rows with no work center", 0,
            "COUNTROWS ( FILTER ( 'Routing Time', ISBLANK ( RELATED ( 'Work Centers'[WorkCenterID] ) ) ) )", 0)
    for c in centers:
        w = f"'Work Centers'[WorkCenterID] = {c['id']}"
        b.check(t, f"{c['name']}: available hours", round(c["avail"], 2), f"CALCULATE ( [Available Hours], {w} )")
        b.check(t, f"{c['name']}: practical capacity", round(c["pract"], 4), f"CALCULATE ( [Practical Capacity], {w} )", 0.0005)
        b.check(t, f"{c['name']}: routing time used", round(c["used"], 4), f"CALCULATE ( [Routing Time Used], {w} )", 0.0005)
        b.check(t, f"{c['name']}: share of available", round(c["of_avail"], 9), f"CALCULATE ( [Share of Available], {w} )", 1e-6)
        b.check(t, f"{c['name']}: share of practical", round(c["of_pract"], 9), f"CALCULATE ( [Share of Practical], {w} )", 1e-6)
    b.check(t, "plant: available hours", round(plant["avail"], 2), "[Available Hours]")
    b.check(t, "plant: practical capacity", round(plant["pract"], 4), "[Practical Capacity]", 0.0005)
    b.check(t, "plant: routing time used", round(plant["used"], 4), "[Routing Time Used]", 0.0005)
    b.check(t, "plant: share of available", round(plant["used"] / plant["avail"], 9), "[Share of Available]", 1e-6)
    b.check(t, "plant: share of practical", round(plant["used"] / plant["pract"], 9), "[Share of Practical]", 1e-6)
    b.check(t, "plant: run hours equal the standard labor hours completed", round(cm["run_total"], 4),
            "CALCULATE ( [Routing Time Used], 'Routing Time'[Kind] = \"Run\" )", 0.0005)
    b.check(t, "plant: setup hours", round(cm["setup_total"], 4),
            "CALCULATE ( [Routing Time Used], 'Routing Time'[Kind] = \"Setup\" )", 0.0005)
    # the Validation page's control totals by year
    for yy, n, ninv, rev, std in years:
        f_ = f"SalesInvoice[Year] = {yy}"
        b.check(t, f"{yy} invoice lines", n, f"CALCULATE ( [Invoice Lines], {f_} )", 0)
        b.check(t, f"{yy} invoices", ninv, f"CALCULATE ( [Invoices], {f_} )", 0)
        b.check(t, f"{yy} revenue", round(rev, 2), f"CALCULATE ( [Revenue], {f_} )")
        b.check(t, f"{yy} standard cost", round(std, 2), f"CALCULATE ( [Standard Cost], {f_} )", 0.01)
    for g, n, rev, std in groups_y:
        gf = f"ItemCosts[ItemGroup] = \"{g}\"" if g else "ISBLANK ( ItemCosts[ItemGroup] )"
        label = g or "(Blank)"
        b.check(t, f"{y} {label}: lines", n, f"CALCULATE ( [Invoice Lines], SalesInvoice[Year] = {y}, {gf} )"
                if g else f"COUNTROWS ( FILTER ( SalesInvoiceLine, ISBLANK ( RELATED ( ItemCosts[ItemGroup] ) ) && "
                          f"RELATED ( SalesInvoice[Year] ) = {y} ) )", 0)
        if g:
            b.check(t, f"{y} {label}: revenue", round(rev, 2), f"CALCULATE ( [Revenue], SalesInvoice[Year] = {y}, {gf} )")
    b.check(t, "last invoice date", last_invoice, "FORMAT ( [Last Invoice], \"yyyy-mm-dd\" )")
    # every query of the Tests tab runs and returns its rows
    for c, expr, _, n in tests:
        b.check(t, f"Tests tab: {c.split(':')[0]} returns {n} rows", n, f"COUNTROWS ( {' '.join(expr.split())} )", 0)


# --- Milestone 2: what a product costs -------------------------------------------------------------------------

def m2(b: Build) -> None:
    d, y = b.data, b.year
    ctx = ch18.m2(d, claim)
    t = "Milestone 2"
    b.check(t, "committed cost of capacity", round(ctx["committed"], 2), "[Committed Cost]")
    b.check(t, "capacity cost rate (per practical hour)", round(ctx["td_rate"], 9), "[Capacity Cost Rate]", 1e-6)
    b.check(t, "crew rate (per practical hour)", round(ctx["crew_rate"], 9), "[Crew Rate]", 1e-6)
    b.check(t, "unused capacity (hours)", round(ctx["unused_h"], 4), "[Unused Capacity]", 0.0005)
    b.check(t, "unused capacity cost", round(ctx["unused"], 2), "[Unused Capacity Cost]", 0.01)
    for g in ctx["groups"]:
        gy = year_fg(y, f"ItemCosts[ItemGroup] = \"{g['name']}\", ItemCosts[SupplyMode] = \"Manufactured\"")
        b.check(t, f"{g['name']} (manufactured) time-driven margin", round(g["m_td"], 9),
                f"CALCULATE ( [Margin %], {gy}, {view('Time-driven')} )", 1e-6)
    e = ctx["e"]
    for k, name in VIEWS:
        b.check(t, f"explorer margin, {name}", round(e["exact"][k], 2),
                f"CALCULATE ( [Margin], {year_fg(y, view(name))} )", 0.01)
    b.check(t, "explorer revenue", round(e["rev"], 2), f"CALCULATE ( [Revenue], {year_fg(y)} )")


EXERCISES = [("Requirement 5", r5), ("Milestone 2", m2)]

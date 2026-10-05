"""Charles River Lender Review.pbip: the Power BI part of the solution to the Chapter 17 capstone case ("Financial
Statements for a Lender"), the Review file of Requirements 2 and 6.

The case starts a new Power BI file (Getting Started). Requirement 2 builds it from CharlesRiver.xlsx as Tutorial 14.3
built its ledger star (GLEntry with the EntryType of its journal entries and the closes flagged by IsYearEndClose,
Account, CostCenter, JournalEntry with Enable load cleared, and the Date table of Tutorial 14.1), with Exercise 14.1's
Balance measure, which leaves out only the closing entries dated on the as-of date. The statement lines come from the
package workbook (Charles River Lender Package.xlsx, the Excel part of this solution, built by xlbuild/cap17.py): its
WorkingTB Table maps every account to a line by AccountSubType and account number (Requirement 1), so the Review file
reads that mapping instead of typing it again, and WorkingTB filters Account as an outrigger of the star. The review
pages are the matrix of statement lines and accounts with the year-end balances and the change from the prior year to
the current one, flagged above a what-if threshold that defaults to the controller's materiality (as in Exercise 15.4);
a drill-through page to the postings; balances by SourceDocumentType with the opening entry's lines; the Questions
table (Enter data, as Chapter 16 kept its dispositions); and a Validation page with Requirement 1's control totals.

Requirement 6 loads the AdjustmentLines Table of the workbook beside the ledger (related to Account and to Date by its
year-end) and tests the package on the Package Tests page and in the Tests tab: each entry balances, no entry touches
cash, adjusted net income (the ledger's net income less the adjustment lines on revenue and expense accounts) agrees
with the workbook's adjusted trial balance, and recorded total assets agree with the workbook's trial balances of
Requirement 1, account by account.

The workbook is a second source. Its queries read C:\\CharlesRiver\\Charles River Lender Package.xlsx, and build.py points
the verification copy at the instructor workbook (the manifest's also_sources). No query combines the two sources, so
Power Query's privacy firewall never applies; the model relates them. Power Query types the workbook's columns from the
cached values of each Table's first 200 rows. Power Query reads a Table's totals row as one more row (Desktop rejected
the blank AccountNumber of WorkingTB's totals row on the one side of its relationship), so each query filters it out.

Checks: every value of the instructor notes of Requirement 2, of the Review tests of Requirement 6, and of Milestones 1
and 3 that the Review file can compute, from facts/notes/ch17.py (computed read-only from CharlesRiver.sqlite) or SQL
here; the workbook's own values are compared with the ledger by the tests, never typed.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from functools import cached_property, lru_cache
from pathlib import Path

from paths import REPO
from pbibuild.audit_monitoring import drillthrough
from pbibuild.case4 import rows_expanded, text_measure
from pbibuild.model import DATE_FORMAT, MEASURES_TABLE, Column, Measure, Model, Query, Table, query_table
from pbibuild.pbir import Page, Report, Visual, col, lit, meas, textbox
from pbibuild.project import Project
from pbibuild.reports import CH14, COUNT, MONEY, card, claim, code_after, dax_blocks, definitions, entered, hide, keep, \
    money, one, rows, slicer
from xlbuild import pq
from xlbuild.expected import Expected

from db import Data  # noqa: E402  (facts/ is on sys.path through pbibuild.reports)
from notes import ch17  # noqa: E402

FILE = "Charles River Lender Review"
WORKBOOK = "Charles River Lender Package.xlsx"
NEUTRAL_WB = "C:\\CharlesRiver\\" + WORKBOOK          # where the shipped project looks for the package workbook
DEFAULT_WB = REPO / "outputs" / "companion" / "instructor-2026" / "Chapter17-capstone-excel" / WORKBOOK
KM = MEASURES_TABLE
M = lambda n: meas(KM, n)                 # noqa: E731
CH14_EX = CH14 / "_exercises.qmd"
PL = '{ "Revenue", "Expense" }'
ABOVE = "Above threshold"
# The statement lines of the workbook's mapping (Requirement 1), in the order the statements show them.
LINES = ["Cash", "Receivables", "Inventories", "Prepaid expenses and other current assets", "Property and equipment, net",
         "Accounts payable", "Goods received not invoiced", "Accrued liabilities", "Sales tax payable", "Customer credits",
         "Notes payable, less current portion", "Common stock", "Retained earnings", "Operating revenue",
         "Sales returns and allowances", "Cost of goods sold", "Operating expenses", "Interest expense",
         "Loss on disposal of equipment", "Other income"]
PL_LINES = LINES[LINES.index("Operating revenue"):]


@dataclass
class Build:
    xlsx: Path
    exp: Expected
    year: int
    sources: dict = field(default_factory=dict)        # neutral path -> real path of the other sources (also_sources)
    project: Project = None
    tests: list = field(default_factory=list)

    def __post_init__(self):
        self.project = Project(FILE, Model(FILE), Report())

    @property
    def model(self) -> Model:
        return self.project.model

    @property
    def report(self) -> Report:
        return self.project.report

    @property
    def workbook(self) -> Path:
        p = Path(self.sources.get(NEUTRAL_WB, DEFAULT_WB))
        if not p.exists():
            raise SystemExit(f"build the Excel part of the case first (python scripts/companion/build.py --solutions "
                             f"--id ch17-capstone): no {p}")
        return p

    @cached_property
    def d(self) -> Data:
        return Data()

    @cached_property
    def ctx(self) -> dict:
        return {k: getattr(ch17, k)(self.d, claim) for k in ("r1", "r2", "r6", "m1", "m3")}

    def check(self, section: str, label: str, expected, dax: str, tolerance: float = 0.005) -> None:
        self.project.check(section, label, expected, dax, tolerance)

    def measure(self, folder: str, name: str, dax: str, fmt: str | None, description: str) -> None:
        mt = self.model.tables[KM]
        assert all(m.name != name for m in mt.measures), f"measure {name} exists already"
        mt.measures.append(Measure(name, dax.strip("\n"), fmt, description, display_folder=folder))

    def add_test(self, comment: str, query: str) -> None:
        """A query typed at the end of the Tests tab, under its comment line (the tab kept before Checks)."""
        self.tests.append(f"// {comment}\n{query.strip()}")
        self.project.queries["Tests"] = "\n\n".join(self.tests) + "\n"
        if "Checks" in self.project.queries:
            self.project.queries["Checks"] = self.project.queries.pop("Checks")


# --- the package workbook as a source ----------------------------------------------------------------------------

@lru_cache(maxsize=None)
def workbook_table(workbook: Path, table: str) -> tuple[tuple[tuple[str, str], ...], tuple[tuple, ...]]:
    """The types Power Query detects for each column of a Table of the workbook from its first 200 rows, read from the
    cached values Excel saved (as pq.detected_types does for CharlesRiver.xlsx), and the Table's data rows. Power
    Query reads the totals row as a row too, so it takes part in the detection, but not in the rows returned. A formula
    that returns "" is saved as an empty text."""
    import openpyxl
    from datetime import datetime
    values, formulas = openpyxl.load_workbook(workbook, data_only=True), openpyxl.load_workbook(workbook)
    ws = next(w for w in values.worksheets if table in dict(w.tables.items()))
    t = ws.tables[table]
    cells = ws[t.ref]
    header = [c.value for c in cells[0]]
    body = cells[1:len(cells) - (t.totalsRowCount or 0)]
    out = []
    for i, name in enumerate(header):
        vals = []
        for r in cells[1:201]:
            v = r[i].value
            if v is None and formulas[ws.title][r[i].coordinate].value is not None:
                v = ""
            if v is not None:
                vals.append(v)
        if not vals:
            kind = "type any"
        elif all(isinstance(v, datetime) for v in vals):
            kind = "type datetime"
        elif all(isinstance(v, bool) for v in vals):
            kind = "type logical"
        elif all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals):
            kind = "Int64.Type" if all(float(v).is_integer() for v in vals) else "type number"
        elif all(isinstance(v, str) for v in vals):
            kind = "type text"
        else:
            kind = "type any"
        out.append((name, kind))
    data = tuple(tuple(c.value for c in r) for r in body)
    values.close()
    formulas.close()
    return tuple(out), tuple(dict(zip(header, r)) for r in data)


class PackageQuery(Query):
    """A query on a Table of the package workbook: Source, Navigation (<Table>_Table), and Changed Type, as the
    Navigator's Transform Data creates them. It always reads the neutral path; build.py repoints the verification copy."""

    @classmethod
    def table(cls, name: str, workbook: Path, table: str) -> "PackageQuery":
        q = cls(name)
        types = workbook_table(workbook, table)[0]
        q.columns = dict(types)
        q.source = ("workbook", table, types)
        q._counts["Changed Type"] = 1
        return q

    def m(self, path: str) -> str:
        _, table, types = self.source
        nav = f"{table}_Table"
        return pq.steps_query([
            ("Source", f"Excel.Workbook(File.Contents({pq.m_string(NEUTRAL_WB)}), null, true)"),
            (nav, f"Source{{[Item={pq.m_string(table)},Kind=\"Table\"]}}[Data]"),
            ("Changed Type", f"Table.TransformColumnTypes({nav},{pq.type_list(list(types))})")] + self.steps)


# --- small helpers ---------------------------------------------------------------------------------------------------

def signed(x: float) -> str:
    return ("+" if x > 0 else "") + money(x)


def answer(name: str, x, y, w, h, lines: list, size: int = 10) -> Visual:
    return textbox(name, x, y, w, h, [("Model answer", True)] + lines, size=size)


def above_zero(name: str, f: dict) -> dict:
    """Filters on this visual: show items when the measure is greater than 0 (a 1/0 flag)."""
    kind = next(iter(f))
    ref = {kind: {"Expression": {"SourceRef": {"Source": "k"}}, "Property": f[kind]["Property"]}}
    return {"name": name, "field": f, "type": "Advanced", "howCreated": "User",
            "filter": {"Version": 2, "From": [{"Name": "k", "Entity": f[kind]["Expression"]["SourceRef"]["Entity"],
                                               "Type": 0}],
                       "Where": [{"Condition": {"Comparison": {"ComparisonKind": 1, "Left": ref,
                                                                 "Right": {"Literal": {"Value": "0D"}}}}}]}}


def by_year(column: str, years: list[int]) -> str:
    """SWITCH on the year in the filter context: the workbook's column of that year-end (Recorded2026, ...)."""
    cases = ",\n".join(f"    {y}, SUM ( WorkingTB[{column}{y}] )" for y in years)
    return f"SWITCH (\n    SELECTEDVALUE ( 'Date'[Year] ),\n{cases}\n)"


def normal_sign(d: Data, number: str) -> int:
    """+1 for a debit-balance account, -1 for a credit-balance one (the notes state amounts on the normal side)."""
    return -1 if ch17.chart(d)[number]["normal"] == "Credit" else 1


def label_of(d: Data, number: str) -> str:
    return f"{number} {ch17.chart(d)[number]['name']}"


# --- Requirement 2: the Review file --------------------------------------------------------------------------------

def r2(b: Build) -> None:
    d, m, x, wb = b.d, b.model, b.xlsx, b.workbook
    F, P, C = d.F, d.P, d.C
    years = [F, P, C]
    c2, c1 = b.ctx["r2"], b.ctx["r1"]
    claim(b.year == C, "the last fiscal year of sales is the current year of the case")
    t14 = CH14 / "_tutorial-03.qmd"

    # The ledger star of Tutorial 14.3 (Steps 1-4): GLEntry (Transform Data) with the EntryType of its journal entries
    # and the closes flagged, Account, CostCenter, JournalEntry with Enable load cleared; the Date table of Tutorial 14.1
    gl = (Query.navigator("GLEntry", x, 3, "GLEntry")
          .select(["GLEntryID", "PostingDate", "AccountID", "Debit", "Credit", "VoucherNumber", "SourceDocumentType",
                   "CostCenterID"]).types({"PostingDate": "type date"}))
    journal = Query.navigator("JournalEntry", x, 2, "JournalEntry").select(["EntryNumber", "EntryType"])
    gl.merge(journal, "VoucherNumber", "EntryNumber", ["EntryType"])
    gl.custom("IsYearEndClose", code_after(t14, "name the column `IsYearEndClose`"), "type logical")
    account = Query.navigator("Account", x, 1, "Account").select(
        ["AccountID", "AccountNumber", "AccountName", "AccountType", "AccountSubType", "NormalBalance"])
    account.custom("AccountLabel", 'Text.From([AccountNumber]) & " " & [AccountName]', "type text")
    center = Query.navigator("CostCenter", x, 75, "CostCenter").select(["CostCenterID", "CostCenterName"])
    # the FixedAsset register, to see which postings it supports (1110's single posting)
    assets = (Query.navigator("FixedAsset", x, 40, "FixedAsset")
              .select(["FixedAssetID", "AssetCode", "AssetDescription", "AssetCategory", "AssetAccountID", "InServiceDate",
                       "OriginalCost", "Status", "DisposalDate"]).types({"InServiceDate": "type date", "DisposalDate": "type date"}))
    # The workbook's mapping and trial balances (WorkingTB of Requirement 1, with the adjustments of Requirement 5)
    wt = PackageQuery.table("WorkingTB", wb, "WorkingTB")
    amounts = [f"{k}{y}" for k in ("Recorded", "Adj", "Adjusted") for y in years]
    missing = [c for c in ["AccountNumber", "AccountType", "StatementLine", "Statement"] + amounts if c not in wt.columns]
    claim(not missing, f"the workbook's WorkingTB Table has the columns the Review file reads (missing {missing})")
    wt.select(["AccountNumber", "AccountName", "AccountType", "StatementLine", "Statement"] + amounts)
    wt.filter("[AccountNumber] <> null")                  # the Table's totals row
    found = sorted({r["StatementLine"] for r in workbook_table(wb, "WorkingTB")[1]})
    claim(set(found) <= set(LINES), f"every statement line of the workbook is a line of the statements ({found})")
    order = "{" + ", ".join(pq.m_string(s) for s in LINES) + "}"
    wt.custom("LineOrder", f"List.PositionOf({order}, [StatementLine]) + 1", "Int64.Type")
    m.add(query_table(account, summarize={"AccountNumber": "none"}))
    m.stage(journal)
    ledger = m.add(query_table(gl, formats={"Debit": MONEY, "Credit": MONEY}))
    m.add(query_table(center))
    hide(ledger, "GLEntryID", "AccountID", "CostCenterID")
    hide(m.tables["Account"], "AccountID")
    ((name, expr),) = definitions(dax_blocks(CH14 / "_tutorial-01.qmd", "table")[0])
    assert name == "Date"
    m.add(Table("Date", [
        Column("Date", "dateTime", "[Date]", DATE_FORMAT, extra=("isKey",)),
        Column("Year", "int64", "[Year]"), Column("Quarter", "string", "[Quarter]"),
        Column("YearQuarter", "string", "[YearQuarter]"), Column("YearMonth", "string", "[YearMonth]"),
        Column("MonthNumber", "int64", "[MonthNumber]"),
        Column("MonthName", "string", "[MonthName]", sort_by="MonthNumber")], dax=expr, props=("dataCategory: Time",)))
    m.tables["Date"].column("Year").summarize = "none"
    fa = m.add(query_table(assets, formats={"OriginalCost": MONEY},
                           summarize={"FixedAssetID": "none", "AssetAccountID": "none"}))
    hide(fa, "FixedAssetID", "AssetAccountID")
    w = m.add(query_table(wt, formats={c: MONEY for c in amounts}, summarize={"AccountNumber": "none", "LineOrder": "none"},
                          sort_by={"StatementLine": "LineOrder"}))
    hide(w, "LineOrder")
    for many, one_ in [("GLEntry.AccountID", "Account.AccountID"), ("GLEntry.CostCenterID", "CostCenter.CostCenterID"),
                       ("GLEntry.PostingDate", "Date.Date"), ("Account.AccountNumber", "WorkingTB.AccountNumber"),
                       ("FixedAsset.AssetAccountID", "Account.AccountID")]:
        m.relate(many, one_)
    mt = m.add(entered(KM, [("Column1", "string")], []))
    m.query_order.append(KM)
    hide(mt, "Column1")

    # The measures
    L, R, S, V = "Ledger", "Review", "Sources", "Validation"
    b.measure(L, "GL Amount", text_measure(t14, "measure", "GL Amount"), MONEY,
              "Debits less credits, without the year-end closing entries (Tutorial 14.3).")
    b.measure(L, "P&L Amount", text_measure(t14, "measure", "P&L Amount"), MONEY,
              "GL Amount with its sign reversed: revenue positive, expenses negative (Tutorial 14.3).")
    b.measure(L, "Net Income", text_measure(t14, "measure", "Net Income"), MONEY,
              "P&L Amount over the revenue and expense accounts, without the closes: no income tax is recorded, so this "
              "is also income before income taxes (Tutorial 14.3).")
    b.measure(L, "Balance", text_measure(CH14_EX, "measure", "Balance"), MONEY,
              "Exercise 14.1's balance at the last date in the filter context: every posting up to that date, debits less "
              "credits, with a closing entry left out only when it is dated on that day. At a year-end a revenue or "
              "expense account shows that year's activity, because the prior year's close zeroed it.")
    b.measure(L, "Total Assets", 'CALCULATE ( [Balance], Account[AccountType] = "Asset" )', MONEY,
              "Balance of the Asset accounts at the last date in the filter context (contra assets reduce it).")
    for y in years:
        b.measure(R, f"Balance {y}", f"CALCULATE ( [Balance], REMOVEFILTERS ( 'Date' ), 'Date'[Date] = DATE ( {y}, 12, 31 ) )",
                  MONEY, f"Balance at the end of fiscal {y}, before that day's closing entries.")
    chg = f"Change {P} to {C}"
    b.measure(R, chg, f"[Balance {C}] - [Balance {P}]", MONEY,
              f"The change from the end of fiscal {P} to the end of {C}, debits less credits. For a revenue or expense "
              f"account it is {C}'s activity less {P}'s.")
    b.measure(R, "Absolute Change", f"ABS ( [{chg}] )", MONEY, "The size of the change, whatever its direction.")
    b.measure(R, "Materiality", "ROUND ( 0.05 * [Net Income], 2 )", MONEY,
              "The controller's materiality: 5% of income before income taxes for the year, as the ledger records it.")
    for y in (P, C):
        b.measure(R, f"Materiality {y}", f"CALCULATE ( [Materiality], REMOVEFILTERS (), 'Date'[Year] = {y} )",
                  MONEY, f"The controller's materiality for fiscal {y}, whatever the filters (an account or a statement "
                         "line in a row of a visual would otherwise change it).")
    # Modeling > New parameter > Numeric range (Exercise 15.4): the GENERATESERIES table and its Value measure; the
    # default is the controller's materiality, read from its measure rather than typed, so it follows the ledger
    m.add(Table("Threshold", [Column("Threshold", "int64", "[Value]", fmt=COUNT, summarize="none",
                                     extra=('extendedProperty ParameterMetadata = {"version":0}',))],
                measures=[Measure("Threshold Value", f"SELECTEDVALUE ( 'Threshold'[Threshold], [Materiality {C}] )", MONEY,
                                  "The threshold the slider sets; until it is moved, the controller's materiality for "
                                  f"fiscal {C}.")],
                dax="GENERATESERIES ( 0, 1000000, 25000 )"))
    b.measure(R, "Is Above Threshold", "IF ( [Absolute Change] > [Threshold Value], 1, 0 )", COUNT,
              "1 when the change is larger than the threshold, else 0 (the filter of the flagged list).")
    b.measure(R, "Flag", f'IF ( [Absolute Change] > [Threshold Value], "{ABOVE}" )', None,
              "Above threshold when the change is larger than the threshold the slider sets.")
    b.measure(S, "Opening Entry Amount", 'CALCULATE ( SUM ( GLEntry[Debit] ) - SUM ( GLEntry[Credit] ), '
                                         'GLEntry[EntryType] = "Opening" )', MONEY,
              "The opening entry's line to the account (JE-2024-000001, EntryType Opening), debits less credits.")
    b.measure(S, "In Opening Entry", "IF ( NOT ISBLANK ( [Opening Entry Amount] ), 1, 0 )", COUNT,
              "1 when the opening entry has a line to the account.")
    b.measure(S, "Postings", "CALCULATE ( COUNTROWS ( GLEntry ), REMOVEFILTERS ( 'Date' ) ) + 0", COUNT,
              "Every posting to the account, in any year.")
    b.measure(S, "Postings after Opening", 'CALCULATE ( COUNTROWS ( GLEntry ), REMOVEFILTERS ( \'Date\' ), '
                                           'GLEntry[EntryType] <> "Opening" || ISBLANK ( GLEntry[EntryType] ) ) + 0', COUNT,
              "Postings other than the opening entry's line.")
    b.measure(S, "Moved since Opening", f"[Balance {C}] - CALCULATE ( [Opening Entry Amount], REMOVEFILTERS ( 'Date' ) )",
              MONEY, f"What the postings after the opening entry added to the balance by the end of fiscal {C}.")
    b.measure(S, "Document Postings", 'CALCULATE ( COUNTROWS ( GLEntry ), REMOVEFILTERS ( \'Date\' ), '
                                      'GLEntry[SourceDocumentType] <> "JournalEntry" ) + 0', COUNT,
              "Postings from a source document (anything but a journal entry).")
    b.measure(S, "Journal Entries Only", "IF ( [Postings] > 0 && [Document Postings] = 0, 1, 0 )", COUNT,
              "1 when every posting to the account is a journal entry: a balance no source document supports.")
    b.measure(V, "Trial Balance Accounts", "COUNTROWS ( FILTER ( VALUES ( Account[AccountNumber] ), NOT ISBLANK ( [Balance] ) ) )",
              COUNT, "Accounts with postings on or before the as-of date (the rows of Requirement 1's trial balance).")
    b.measure(V, "Accounts at Zero", "COUNTROWS ( FILTER ( VALUES ( Account[AccountNumber] ), VAR B = [Balance] RETURN "
                                     "NOT ISBLANK ( B ) && ABS ( B ) < 0.005 ) )", COUNT,
              "Trial balance accounts whose balance is zero.")
    b.measure(V, "Debit Balances", "SUMX ( VALUES ( Account[AccountNumber] ), VAR B = [Balance] RETURN IF ( B > 0, B ) )",
              MONEY, "The sum of the debit balances of the trial balance.")
    b.measure(V, "Credit Balances", "SUMX ( VALUES ( Account[AccountNumber] ), VAR B = [Balance] RETURN IF ( B < 0, -B ) )",
              MONEY, "The sum of the credit balances of the trial balance.")
    b.measure(V, "Accounts Never Posted", 'CALCULATE ( COUNTROWS ( FILTER ( Account, Account[AccountSubType] <> "Header" '
                                          '&& ISEMPTY ( RELATEDTABLE ( GLEntry ) ) ) ), REMOVEFILTERS ( \'Date\' ) )', COUNT,
              "Accounts of the chart, headers aside, with no posting in any year (the anti-join of Requirement 1).")
    b.measure(V, "Income Tax Accounts", 'COUNTROWS ( FILTER ( ALL ( Account ), SEARCH ( "Income Tax", '
                                        'Account[AccountName], 1, 0 ) > 0 ) ) + 0', COUNT,
              "Accounts whose name mentions income tax: the chart has neither an expense nor a payable.")

    # Review page: the threshold, the matrix of statement lines and accounts, and the flagged changes
    yr = col("Date", "Year")
    line, lab = col("WorkingTB", "StatementLine"), col("Account", "AccountLabel")
    page = b.report.add(Page("review", "Review", height=1000))
    page.add(slicer("threshold", 20, 20, 300, 110, col("Threshold", "Threshold"), title="Threshold (what-if)"))
    page.visuals[-1].objects = {"data": [{"properties": {"mode": lit("Single")}}]}
    page.add(card("materiality", 340, 20, 920, 110, [meas("Threshold", "Threshold Value"), M(f"Materiality {C}"),
                                                     M(f"Materiality {P}")], units_none=True, title="Threshold and materiality"))
    levels = [line, lab]
    page.add(Visual("reviewMatrix", "pivotTable", 20, 150, 820, 560, {
        "Rows": levels, "Values": [M(f"Balance {y}") for y in years] + [M(chg), M("Flag")]},
        title=f"Statement lines and accounts: year-end balances (debits less credits) and the change from {P} to {C}",
        extra=rows_expanded(levels), active=("Rows",)))
    page.add(Visual("aboveTable", "tableEx", 860, 150, 400, 560, {
        "Values": [(lab, "Account"), M(chg), M("Absolute Change")]},
        title="Changes above the threshold (right-click > Drill through > Postings)",
        filters=[above_zero("aboveFlag", M("Is Above Threshold"))], sort=[(M("Absolute Change"), "Descending")]))
    a = c2["above"]
    tb_p, tb_c = ch17.tb(d, ch17.ye(P)), ch17.tb(d, ch17.ye(C))
    named = lambda n: f"{n} {ch17.chart(d)[n]['name']}"                  # noqa: E731
    re_change = ch17.xr(-(tb_c["3030"] - tb_p["3030"]))
    page.add(answer("reviewAnswer", 20, 730, 1240, 250, [
        f"Materiality is 5% of income before income taxes as the ledger records it: {money(c2['mat_c'])} for {C} and "
        f"{money(c2['mat_p'])} for {P} (no income tax is recorded, so it is 5% of net income). With the slider at its "
        f"default, the changes from the end of {P} to the end of {C} above {money(c2['mat_c'])} are: " +
        "; ".join(f"{named(x['n'])} {signed(x['v'])}" + (" (credit)" if x["credit"] else "") for x in a) +
        f"; and 3030 Retained Earnings, {signed(re_change)} (credit), the {P} net income closed into it. The next largest "
        "are " + " and ".join(f"{named(x['n'])} {money(x['v'])}" for x in c2["following"]) + ".",
        "A revenue or expense account's balance at a year-end is that year's activity, because the prior close zeroed "
        f"it, so its change is {C}'s activity less {P}'s. Drill through from any flagged row to its postings; the "
        "Sources page shows which balances no document supports, and the Questions page records what the review "
        "asks and which requirement answers it."], size=10))
    top = a[0]
    page.expect = ["Threshold Value", money(c2["mat_c"]), money(c2["mat_p"]), "Absolute Change",
                   money(normal_sign(d, top["n"]) * top["v"]), ABOVE, "Model answer"]

    # Postings page: the drill-through from a flagged account, left as after a drill-through on the largest change
    detail = b.report.add(Page("postings", "Postings", hidden=True))
    drillthrough(detail, lab, label_of(d, top["n"]))
    detail.add(card("changeCards", 80, 20, 1180, 100, [M(f"Balance {P}"), M(f"Balance {C}"), M(chg)], units_none=True,
                    title="The account drilled through"))
    src = col("GLEntry", "SourceDocumentType")
    detail.add(Visual("bySource", "pivotTable", 20, 140, 560, 560, {"Rows": [src], "Columns": [yr], "Values": [M("GL Amount")]},
                      title=f"Activity by source document, {P} and {C} (closes left out)",
                      filters=[keep("bySourceYears", yr, [P, C])]))
    detail.add(Visual("postingList", "tableEx", 600, 140, 660, 560, {"Values": [
        col("GLEntry", "PostingDate"), col("GLEntry", "VoucherNumber"), src, col("GLEntry", "EntryType"),
        col("GLEntry", "Debit"), col("GLEntry", "Credit")]}, title=f"Postings in fiscal {C}",
        filters=[keep("postingYear", yr, [C])], sort=[(col("GLEntry", "PostingDate"), "Ascending")]))
    top_src = rows("SELECT g.SourceDocumentType, ROUND(SUM(g.Debit - g.Credit), 2) FROM GLEntry g JOIN Account a ON "
                   "a.AccountID = g.AccountID WHERE a.AccountNumber = ? AND g.FiscalYear = ? GROUP BY 1 ORDER BY "
                   "ABS(SUM(g.Debit - g.Credit)) DESC LIMIT 1", int(top["n"]), C)[0]
    detail.expect = ["VoucherNumber", label_of(d, top["n"]), money(normal_sign(d, top["n"]) * top["v"]), top_src[0],
                     money(top_src[1])]

    # Sources page: balances by SourceDocumentType, the opening entry's lines, and the balances of journal entries alone
    page = b.report.add(Page("sources", "Sources", height=1220))
    bs_types = ["Asset", "Liability", "Equity"]
    page.add(Visual("sourceMatrix", "pivotTable", 20, 20, 1240, 420, {
        "Rows": [lab], "Columns": [src], "Values": [M("Balance")]},
        title=f"Balance-sheet accounts at the end of fiscal {C}, by source document",
        filters=[keep("smYear", yr, [C]), keep("smTypes", col("Account", "AccountType"), bs_types)]))
    page.add(Visual("openingTable", "tableEx", 20, 460, 760, 520, {"Values": [
        (lab, "Account"), M("Opening Entry Amount"), M(f"Balance {C}"), M("Moved since Opening"), M("Postings after Opening")]},
        title="The opening entry's lines (debits less credits)", filters=[above_zero("openingFlag", M("In Opening Entry"))]))
    page.add(Visual("journalOnly", "tableEx", 800, 460, 460, 520, {"Values": [
        (lab, "Account"), M(f"Balance {C}"), M("Postings")]}, title="Balance-sheet accounts posted by journal entries alone",
        filters=[above_zero("journalOnlyFlag", M("Journal Entries Only")),
                 keep("journalOnlyTypes", col("Account", "AccountType"), bs_types)]))
    opening = {n: ch17.opening_line(d, n) * normal_sign(d, n) for n in
               [str(n) for (n,) in rows("SELECT DISTINCT a.AccountNumber FROM GLEntry g JOIN Account a ON a.AccountID = "
                                        "g.AccountID WHERE g.VoucherNumber = ?", c2["entry"])]}
    after = dict((str(n), k) for n, k in rows(
        "SELECT a.AccountNumber, COUNT(g.GLEntryID) FROM Account a LEFT JOIN GLEntry g ON g.AccountID = a.AccountID AND "
        "g.VoucherNumber <> ? GROUP BY a.AccountNumber", c2["entry"]))
    only_je = [str(n) for (n,) in rows(
        "SELECT a.AccountNumber FROM Account a JOIN GLEntry g ON g.AccountID = a.AccountID WHERE a.AccountType IN "
        "('Asset', 'Liability', 'Equity') GROUP BY a.AccountID HAVING SUM(g.SourceDocumentType <> 'JournalEntry') = 0 "
        "ORDER BY 1")]
    unmoved = [n for n in opening if abs(tb_c.get(n, 0.0) - opening[n]) < 0.005]
    fa_1110 = rows("SELECT f.AssetCode, f.OriginalCost FROM FixedAsset f JOIN Account a ON a.AccountID = f.AssetAccountID "
                   "WHERE a.AccountNumber = 1110")
    claim(fa_1110 == [(c2["asset"], c2["asset_cost"])], "the FixedAsset register's asset on 1110 is the note's")
    claim(set(unmoved) == {"1050", "1110", "2030", "3010"}, f"the balances unchanged since the opening entry ({unmoved})")
    page.add(answer("sourcesAnswer", 20, 1000, 1240, 200, [
        f"The opening entry {c2['entry']} has {c2['lines']} lines. Balances that have not moved since it: "
        f"{named('1050')} {money(c2['prepaid'])}, its only posting; {named('2030')} {money(c2['payroll'])}, because "
        f"every payroll posting to it nets to zero ({after['2030']:,} postings); {named('1110')} "
        f"{money(opening['1110'])}, a single posting that the FixedAsset register supports ({c2['asset']}, "
        f"{money(c2['asset_cost'])}); and {named('3010')}. Parts of three more rest on it with no document: "
        f"1020 {money(c2['receivable'])}, 2010 {money(c2['payable'])} (equal to the opening cash line to the cent), and "
        f"2040 {money(c2['accrued'])}.",
        "Balance-sheet accounts posted by journal entries alone: " + ", ".join(named(n) for n in only_je) + ". No "
        "source document supports them: 1050 rests on the opening entry alone; the equipment, its depreciation, and the "
        "notes rest on entries that the FixedAsset and DebtAgreement records support; equity moves only by the opening "
        "entry and the closes."], size=10))
    page.expect = ["JournalEntry", "Opening Entry Amount", money(c2["prepaid"]), money(-c2["payroll"]),
                   money(-c2["payable"]), "Model answer"]

    # Questions page: the Questions table (Enter data), each with the requirement that answers it
    cash_c = tb_c["1010"]
    notes_c = -tb_c["2110"]
    qrows = [
        [1, f"Why did 2040 Accrued Expenses grow by {money(c2['grew'])} in {C}, and which of its accruals are still owed?",
         "2040", "Requirement 3"],
        [2, f"Was revenue recorded in the year of delivery at each year-end, and were any December {C} shipments never "
            "invoiced?", "4010-4040, 1020", "Requirement 4"],
        [3, f"2030 has held {money(c2['payroll'])} since the opening entry while every payroll posting to it nets to zero: "
            "is the payroll of each year's last days, paid in January, owed at the year-end?", "2030-2033, 1090",
         "Requirement 4"],
        [4, "2080 Interest Payable has never been used: is the interest on the notes accrued at the year-ends?",
         "2080, 7030", "Requirement 4"],
        [5, f"Why does 1090 carry a credit balance of {money(-tb_c['1090'])} inside current assets?", "1090, 1046",
         "Requirement 5"],
        [6, "Is an allowance for credit losses needed? 1030 has never been used.", "1030, 1020", "Requirement 5"],
        [7, f"Which part of the notes payable ({money(notes_c)} at the end of {C}) is due within a year? All of 2110 is "
            "classified long-term.", "2110", "Requirement 5"],
        [8, "2050 has no debits except credit memos: has sales tax ever been remitted?", "2050", "Requirement 9"],
        [9, "Rent is paid every month from cash while 2120 and 1140 have never been used: are there leases to record?",
         "6070, 6080, 2120, 1140", "Requirement 9"],
        [10, "The chart has no income tax account: what is the company's tax status?", "none", "Requirement 9"],
        [11, f"Why does 7010 show no interest income on more than {c2['millions']} million of cash?", "7010, 1010",
         "Requirement 9"],
        [12, "Do the opening balances no document supports (1020, 1050, 2010, 2030, 2040) exist?",
         "1020, 1050, 2010, 2030, 2040", "Requirement 9"]]
    m.add(Table("Questions", [Column("No", "int64"), Column("Question", "string"), Column("Accounts", "string"),
                              Column("Requirement", "string")], entered=[[str(v) for v in r] for r in qrows]))
    m.query_order.append("Questions")
    page = b.report.add(Page("questions", "Questions"))
    page.add(Visual("questionTable", "tableEx", 20, 20, 1240, 680, {"Values": [
        col("Questions", "No"), col("Questions", "Question"), col("Questions", "Accounts"), col("Questions", "Requirement")]},
        title="The questions the review raises, and the requirement that answers each",
        sort=[(col("Questions", "No"), "Ascending")]))
    page.expect = ["Question", "Requirement 3", "Requirement 9", qrows[0][2]]

    # Validation page (Requirement 1's control totals, Milestone 1)
    page = b.report.add(Page("validation", "Validation"))
    page.add(Visual("tbTotals", "tableEx", 20, 20, 1240, 170, {"Values": [
        yr, M("Trial Balance Accounts"), M("Accounts at Zero"), M("Debit Balances"), M("Credit Balances"), M("Net Income"),
        M("Total Assets")]}, title="Trial balances at each fiscal year-end, before that day's closing entries",
        filters=[keep("tbYears", yr, years)]))
    page.add(card("chartCards", 20, 210, 620, 110, [M("Accounts Never Posted"), M("Income Tax Accounts")],
                  title="The chart of accounts"))
    rec = c1["rec"]
    page.add(answer("validationAnswer", 20, 340, 1240, 300, [
        "The Balance measure reproduces Requirement 1's trial balances: " + "; ".join(
            f"{y} {r['n']} accounts ({r['zero']} at zero), debit and credit balances {money(r['debits'])}"
            for y, r in zip(years, rec)) + f". The last is Chapter 6's pre-closing trial balance.",
        "Net income " + ", ".join(f"{y} {money(n)}" for y, n in zip(years, c1["ni"])) + "; total assets " +
        ", ".join(f"{y} {money(r['ta'])}" for y, r in zip(years, rec)) + f". {c1['n_never']} accounts of the chart, "
        "headers aside, have never been posted, and no account is named for income tax."], size=10))
    page.expect = ["Trial Balance Accounts", money(rec[-1]["debits"]), money(c1["ni"][-1]), money(rec[-1]["ta"])]
    b.report.active = "review"

    # Test 1: the trial balances of Requirement 1
    yl = ", ".join(map(str, years))
    b.add_test("Test 1: the trial balance at each fiscal year-end (Requirement 1): accounts, debit and credit balances, "
               "net income, and total assets",
               "EVALUATE\nSUMMARIZECOLUMNS (\n    'Date'[Year],\n    TREATAS ( { " + yl + " }, 'Date'[Year] ),\n"
               "    \"Accounts\", [Trial Balance Accounts],\n    \"Debit Balances\", ROUND ( [Debit Balances], 2 ),\n"
               "    \"Credit Balances\", ROUND ( [Credit Balances], 2 ),\n    \"Net Income\", ROUND ( [Net Income], 2 ),\n"
               "    \"Total Assets\", ROUND ( [Total Assets], 2 )\n)\nORDER BY 'Date'[Year]")
    b.add_test(f"Test 2: the changes from {P} to {C} above the threshold (the controller's materiality by default)",
               "EVALUATE\nFILTER (\n    SUMMARIZECOLUMNS (\n        Account[AccountLabel],\n"
               f"        \"Change\", ROUND ( [{chg}], 2 ),\n        \"Threshold\", [Threshold Value]\n    ),\n"
               "    ABS ( [Change] ) > [Threshold]\n)\nORDER BY ABS ( [Change] ) DESC")

    # checks
    s = "Requirement 2"
    b.check(s, "GLEntry rows (all of them)", one("SELECT COUNT(*) FROM GLEntry"), "COUNTROWS ( GLEntry )", 0)
    b.check(s, "postings flagged IsYearEndClose", one(
        "SELECT COUNT(*) FROM GLEntry g JOIN JournalEntry j ON j.EntryNumber = g.VoucherNumber WHERE "
        "g.SourceDocumentType = 'JournalEntry' AND j.EntryType LIKE 'Year-End Close%'"),
        "COUNTROWS ( FILTER ( GLEntry, GLEntry[IsYearEndClose] ) )", 0)
    b.check(s, "the closing entries", ", ".join(sorted(d.closes)),
            "CONCATENATEX ( CALCULATETABLE ( VALUES ( GLEntry[VoucherNumber] ), GLEntry[IsYearEndClose] ), "
            "GLEntry[VoucherNumber], \", \", GLEntry[VoucherNumber], ASC )")
    b.check(s, "accounts of the workbook's mapping (WorkingTB rows)", one("SELECT COUNT(*) FROM Account WHERE "
                                                                        "AccountSubType <> 'Header'"),
            "COUNTROWS ( WorkingTB )", 0)
    b.check(s, "accounts with postings and no statement line", 0,
            "COUNTROWS ( FILTER ( Account, [Postings] > 0 && ISBLANK ( RELATED ( WorkingTB[StatementLine] ) ) ) ) + 0", 0)
    b.check(s, f"materiality {C}", c2["mat_c"], f"[Materiality {C}]")
    b.check(s, f"materiality {P}", c2["mat_p"], f"[Materiality {P}]")
    b.check(s, "the threshold until the slider is moved (the controller's materiality)", c2["mat_c"],
            "[Threshold Value]")
    flagged = [x["n"] for x in a] + ["3030"]
    for x in a:
        n = x["n"]
        b.check(s, f"{named(n)}: change {P} to {C} (debits less credits)", normal_sign(d, n) * x["v"],
                f"CALCULATE ( [{chg}], Account[AccountNumber] = {n} )", 0.01)
        b.check(s, f"{n} flagged", ABOVE, f"CALCULATE ( [Flag], Account[AccountNumber] = {n} )")
    b.check(s, f"3030 Retained Earnings: change {P} to {C}, the {P} net income closed into it",
            -ch17.net_income(d, P), f"CALCULATE ( [{chg}], Account[AccountNumber] = 3030 )", 0.01)
    b.check(s, "accounts above the threshold (with 3030)", len(flagged),
            "COUNTROWS ( FILTER ( VALUES ( Account[AccountNumber] ), [Is Above Threshold] = 1 ) )", 0)
    for x in c2["following"]:
        b.check(s, f"next largest: {named(x['n'])}, absolute change", x["v"],
                f"CALCULATE ( [Absolute Change], Account[AccountNumber] = {x['n']} )", 0.01)
        b.check(s, f"{x['n']} not flagged", 0, f"CALCULATE ( [Is Above Threshold], Account[AccountNumber] = {x['n']} )", 0)
    for n in ("1010", "1040", "1045", "2010", "2050"):
        b.check(s, f"{named(n)}: balance at {ch17.ye(C)}", ch17.xr(tb_c[n]), f"CALCULATE ( [Balance {C}], "
                f"Account[AccountNumber] = {n} )", 0.01)
    for n in ("4020", "5020", "5080"):
        b.check(s, f"{named(n)}: balance at {ch17.ye(C)} = the year's activity", ch17.xr(ch17.activity(d, [n], C)),
                f"CALCULATE ( [Balance {C}], Account[AccountNumber] = {n} )", 0.01)
    b.check(s, "opening entry: lines", c2["lines"], 'CALCULATE ( COUNTROWS ( GLEntry ), GLEntry[EntryType] = "Opening" )', 0)
    b.check(s, "opening entry: its number", c2["entry"], 'CALCULATE ( MAX ( GLEntry[VoucherNumber] ), '
                                                       'GLEntry[EntryType] = "Opening" )')
    for n, key in (("1050", "prepaid"), ("2030", "payroll"), ("1020", "receivable"), ("2010", "payable"),
                   ("2040", "accrued")):
        b.check(s, f"opening entry line to {named(n)}", normal_sign(d, n) * c2[key],
                f"CALCULATE ( [Opening Entry Amount], Account[AccountNumber] = {n} )", 0.005)
    b.check(s, "opening line to 2010 equals the opening cash line", 0,
            "ROUND ( CALCULATE ( [Opening Entry Amount], Account[AccountNumber] = 1010 ) + CALCULATE ( [Opening Entry Amount], "
            "Account[AccountNumber] = 2010 ), 2 )")
    b.check(s, "postings to 1050 after the opening entry", 0,
            "CALCULATE ( [Postings after Opening], Account[AccountNumber] = 1050 )", 0)
    b.check(s, "2030: what the postings after the opening entry added", 0,
            "CALCULATE ( [Moved since Opening], Account[AccountNumber] = 2030 )")
    b.check(s, "2030: postings after the opening entry", after["2030"],
            "CALCULATE ( [Postings after Opening], Account[AccountNumber] = 2030 )", 0)
    b.check(s, "accounts unchanged since the opening entry", ", ".join(sorted(unmoved)),
            "CONCATENATEX ( FILTER ( VALUES ( Account[AccountNumber] ), [In Opening Entry] = 1 && ABS ( "
            "[Moved since Opening] ) < 0.005 ), Account[AccountNumber], \", \", Account[AccountNumber], ASC )")
    b.check(s, "balance-sheet accounts posted by journal entries alone", ", ".join(only_je),
            "CONCATENATEX ( FILTER ( VALUES ( Account[AccountNumber] ), [Journal Entries Only] = 1 && CALCULATE ( "
            "SELECTEDVALUE ( Account[AccountType] ) ) IN { \"Asset\", \"Liability\", \"Equity\" } ), "
            "Account[AccountNumber], \", \", Account[AccountNumber], ASC )")
    b.check(s, "1110: postings", 1, "CALCULATE ( [Postings], Account[AccountNumber] = 1110 )", 0)
    b.check(s, "1110: the FixedAsset register's asset", c2["asset"],
            "CALCULATE ( CONCATENATEX ( FixedAsset, FixedAsset[AssetCode], \", \" ), Account[AccountNumber] = 1110 )")
    b.check(s, "1110: its original cost", c2["asset_cost"],
            "CALCULATE ( SUM ( FixedAsset[OriginalCost] ), Account[AccountNumber] = 1110 )")
    b.check(s, f"2040: change {P} to {C} (it grew)", -c2["grew"], f"CALCULATE ( [{chg}], Account[AccountNumber] = 2040 )",
            0.01)
    b.check(s, f"1090: balance at {ch17.ye(C)} (a credit inside current assets)", ch17.xr(tb_c["1090"]),
            f"CALCULATE ( [Balance {C}], Account[AccountNumber] = 1090 )", 0.01)
    for n in ("1030", "1140", "2080", "2120", "7010"):
        b.check(s, f"postings to {named(n)} (unused)", 0, f"CALCULATE ( [Postings], Account[AccountNumber] = {n} )", 0)
    b.check(s, "debits to 2050 from anything but credit memos", 0,
            'CALCULATE ( SUM ( GLEntry[Debit] ), Account[AccountNumber] = 2050, GLEntry[SourceDocumentType] <> "CreditMemo" ) + 0')
    b.check(s, "debits to 2050 from credit memos", ch17.xr(one(
        "SELECT SUM(g.Debit) FROM GLEntry g WHERE g.AccountID = ? AND g.SourceDocumentType = 'CreditMemo'", d.account("2050"))),
        'CALCULATE ( SUM ( GLEntry[Debit] ), Account[AccountNumber] = 2050, GLEntry[SourceDocumentType] = "CreditMemo" )', 0.01)
    months = one("SELECT COUNT(DISTINCT substr(PostingDate, 1, 7)) FROM JournalEntry WHERE EntryType = 'Rent'")
    claim(months == 12 * len(years), "rent is posted every month of the window")
    b.check(s, "months with rent entries", months,
            'COUNTROWS ( SUMMARIZE ( FILTER ( GLEntry, GLEntry[EntryType] = "Rent" ), \'Date\'[YearMonth] ) )', 0)
    b.check(s, "income tax accounts", 0, "[Income Tax Accounts]", 0)
    b.check(s, f"cash at {ch17.ye(C)} (more than {c2['millions']} million)", ch17.xr(cash_c),
            f"CALCULATE ( [Balance {C}], Account[AccountNumber] = 1010 )", 0.01)
    b.check(s, "Questions rows", len(qrows), "COUNTROWS ( Questions )", 0)
    for req in sorted({r[3] for r in qrows}):
        b.check(s, f"questions answered by {req}", sum(r[3] == req for r in qrows),
                f"CALCULATE ( COUNTROWS ( Questions ), Questions[Requirement] = \"{req}\" )", 0)
    b.check(s, "Threshold rows (GENERATESERIES from 0 to 1,000,000 by 25,000)", 1_000_000 // 25_000 + 1,
            "COUNTROWS ( 'Threshold' )", 0)


# --- Milestone 1: the control totals of the Validation page ----------------------------------------------------------

def m1(b: Build) -> None:
    d, f = b.d, b.ctx["m1"]
    s = "Milestone 1"
    for y, r, ni in zip(d.years, f["rec"], f["ni"]):
        cy = f"'Date'[Year] = {y}"
        b.check(s, f"trial balance accounts at {ch17.ye(y)}", r["n"], f"CALCULATE ( [Trial Balance Accounts], {cy} )", 0)
        b.check(s, f"accounts at zero at {ch17.ye(y)}", r["zero"], f"CALCULATE ( [Accounts at Zero], {cy} )", 0)
        b.check(s, f"debit balances at {ch17.ye(y)}", r["debits"], f"CALCULATE ( [Debit Balances], {cy} )", 0.01)
        b.check(s, f"credit balances at {ch17.ye(y)}", r["credits"], f"CALCULATE ( [Credit Balances], {cy} )", 0.01)
        b.check(s, f"recorded net income {y}", ni, f"CALCULATE ( [Net Income], {cy} )", 0.01)
        b.check(s, f"recorded total assets at {ch17.ye(y)}", r["ta"], f"CALCULATE ( [Total Assets], {cy} )", 0.01)
    b.check(s, "accounts never posted", f["never"], "[Accounts Never Posted]", 0)
    b.check(s, "income tax accounts", 0, "[Income Tax Accounts]", 0)


# --- Requirement 6: the tests of the package -------------------------------------------------------------------------

def r6(b: Build) -> None:
    d, m, wb = b.d, b.model, b.workbook
    F, P, C = d.F, d.P, d.C
    years = [F, P, C]
    c6, c1 = b.ctx["r6"], b.ctx["r1"]
    # The AdjustmentLines Table of the workbook (Requirement 5), beside the ledger
    adj = PackageQuery.table("AdjustmentLines", wb, "AdjustmentLines")
    keep_cols = ["No", "Adjustment", "YearEnd", "Account", "AccountName", "StatementLine", "Amount", "Debit", "Credit",
                 "Type", "Evidence"]
    missing = [c for c in keep_cols if c not in adj.columns]
    claim(not missing, f"the workbook's AdjustmentLines Table has the columns the Review file reads (missing {missing})")
    adj.select(keep_cols).filter("[Account] <> null").types({"YearEnd": "type date"})      # the totals row out
    claim(adj.columns["Account"] == "Int64.Type" and adj.columns["Amount"] == "type number",
          "the workbook's account numbers are whole numbers and its amounts decimals")
    lines = workbook_table(wb, "AdjustmentLines")[1]
    t = m.add(query_table(adj, formats={"Amount": MONEY, "Debit": MONEY, "Credit": MONEY},
                          summarize={"Account": "none"}))
    t.column("YearEnd").fmt = DATE_FORMAT
    m.relate("AdjustmentLines.Account", "Account.AccountNumber")
    m.relate("AdjustmentLines.YearEnd", "Date.Date")

    P6 = "Package"
    b.measure(P6, "Adjustment Amount", "SUM ( AdjustmentLines[Amount] )", MONEY,
              "The package's adjustment lines, debits less credits, at the year-ends in the filter context.")
    b.measure(P6, "Adjustment Debits", "SUM ( AdjustmentLines[Debit] )", MONEY, "The debits of the adjustment lines.")
    b.measure(P6, "Adjustment Credits", "SUM ( AdjustmentLines[Credit] )", MONEY, "The credits of the adjustment lines.")
    b.measure(P6, "Adjustment Lines", "COUNTROWS ( AdjustmentLines ) + 0", COUNT, "The number of adjustment lines.")
    b.measure(P6, "Entries", "COUNTROWS ( SUMMARIZE ( AdjustmentLines, AdjustmentLines[No], AdjustmentLines[YearEnd] ) )",
              COUNT, "Entries: each numbered adjustment at each year-end.")
    b.measure(P6, "Largest Entry Imbalance", "MAXX ( SUMMARIZE ( AdjustmentLines, AdjustmentLines[No], "
                                             "AdjustmentLines[YearEnd] ), ABS ( ROUND ( CALCULATE ( SUM ( "
                                             "AdjustmentLines[Amount] ) ), 2 ) ) ) + 0", MONEY,
              "The largest debits less credits of any entry: zero when every entry balances.")
    b.measure(P6, "Entries out of Balance", "COUNTROWS ( FILTER ( SUMMARIZE ( AdjustmentLines, AdjustmentLines[No], "
                                            "AdjustmentLines[YearEnd] ), ABS ( CALCULATE ( SUM ( AdjustmentLines[Amount] ) ) ) "
                                            "> 0.005 ) ) + 0", COUNT, "Entries whose debits and credits differ.")
    b.measure(P6, "Lines on Cash", "CALCULATE ( COUNTROWS ( AdjustmentLines ), AdjustmentLines[Account] = 1010 ) + 0",
              COUNT, "Adjustment lines to 1010 Cash and Cash Equivalents: an adjustment never moves cash.")
    b.measure(P6, "Adjusted P&L Amount", "[P&L Amount] - [Adjustment Amount]", MONEY,
              "P&L Amount of the year plus the adjustment lines at its year-end (revenue positive, expenses negative).")
    b.measure(P6, "Adjusted Net Income", f"CALCULATE ( [Adjusted P&L Amount], Account[AccountType] IN {PL} )", MONEY,
              "The ledger's net income of the year plus the adjustment lines on revenue and expense accounts at its "
              "year-end.")
    b.measure(P6, "Workbook Recorded Balance", by_year("Recorded", years), MONEY,
              "The workbook's recorded balance (WorkingTB, from the trial balances of Requirement 1) at the year-end in "
              "the filter context.")
    b.measure(P6, "Workbook Adjusted Balance", by_year("Adjusted", years), MONEY,
              "The workbook's adjusted balance (WorkingTB) at the year-end in the filter context.")
    b.measure(P6, "Workbook Adjusted Net Income", f"- CALCULATE ( [Workbook Adjusted Balance], WorkingTB[AccountType] IN {PL} )",
              MONEY, "Net income from the workbook's adjusted trial balance: its revenue and expense accounts, sign reversed.")
    b.measure(P6, "Net Income Difference", "ROUND ( [Adjusted Net Income] - [Workbook Adjusted Net Income], 2 )", MONEY,
              "Adjusted net income in this file less the workbook's: zero when they agree.")
    b.measure(P6, "Workbook Recorded Total Assets", 'CALCULATE ( [Workbook Recorded Balance], WorkingTB[AccountType] = "Asset" )',
              MONEY, "Total assets of the workbook's trial balance at the year-end in the filter context.")
    b.measure(P6, "Total Assets Difference", "ROUND ( [Total Assets] - [Workbook Recorded Total Assets], 2 )", MONEY,
              "Total assets from the ledger less the workbook's trial balance: zero when they agree.")
    b.measure(P6, "Accounts Disagreeing", "COUNTROWS ( FILTER ( VALUES ( WorkingTB[AccountNumber] ), "
                                          "ABS ( [Balance] - [Workbook Recorded Balance] ) > 0.005 ) ) + 0", COUNT,
              "Accounts whose ledger balance at the year-end differs from the workbook's trial balance.")
    b.measure(P6, "Adjusted Total Assets", 'CALCULATE ( [Balance] + [Adjustment Amount], Account[AccountType] = "Asset" )',
              MONEY, "Total assets after the adjustment lines at the year-end in the filter context.")

    # Package Tests page
    yr, no = col("Date", "Year"), col("AdjustmentLines", "No")
    page = b.report.add(Page("packageTests", "Package Tests", height=1100))
    page.add(card("testCards", 20, 20, 1240, 110, [M("Adjustment Lines"), M("Entries"), M("Entries out of Balance"),
                                                   M("Largest Entry Imbalance"), M("Lines on Cash")],
                  title="The adjustment lines of the package workbook"))
    page.add(Visual("entryMatrix", "pivotTable", 20, 150, 560, 400, {
        "Rows": [no], "Columns": [yr], "Values": [M("Adjustment Debits"), M("Adjustment Credits"), M("Adjustment Amount")]},
        title="Each entry balances: debits, credits, and debits less credits by adjustment and year-end"))
    page.add(Visual("agreeTable", "tableEx", 600, 150, 660, 190, {"Values": [
        yr, M("Adjusted Net Income"), M("Workbook Adjusted Net Income"), M("Net Income Difference")]},
        title="Adjusted net income: the ledger plus the adjustment lines, against the workbook",
        filters=[keep("agreeYears", yr, [P, C])]))
    page.add(Visual("assetTable", "tableEx", 600, 360, 660, 190, {"Values": [
        yr, M("Total Assets"), M("Workbook Recorded Total Assets"), M("Total Assets Difference"), M("Accounts Disagreeing"),
        M("Adjusted Total Assets")]}, title="Recorded total assets: the ledger against the workbook's trial balances",
        filters=[keep("assetYears", yr, years)]))
    pl_types = ["Revenue", "Expense"]
    page.add(Visual("incomeMatrix", "pivotTable", 20, 570, 760, 330, {
        "Rows": [col("WorkingTB", "StatementLine")], "Columns": [yr], "Values": [M("Adjusted P&L Amount")]},
        title="The adjusted income statement by line (revenue positive, expenses negative)",
        filters=[keep("imYears", yr, [P, C]), keep("imTypes", col("Account", "AccountType"), pl_types)]))
    i_p, i_c = c6["i_p"], c6["i_c"]
    rec = c1["rec"]
    page.add(answer("packageAnswer", 800, 570, 460, 510, [
        f"The workbook's {len(lines)} adjustment lines form {len({(r['No'], r['YearEnd']) for r in lines})} entries "
        "(each numbered adjustment at each year-end); every entry balances and none touches 1010.",
        f"Adjusted net income: {P} {money(i_p['ni'])} and {C} {money(i_c['ni'])}, the ledger's net income plus the "
        "lines on revenue and expense accounts, equal to the workbook's adjusted trial balance.",
        "Recorded total assets equal the trial balances of Requirement 1 account by account: " +
        ", ".join(f"{y} {money(r['ta'])}" for y, r in zip(years, rec)) + ". After the adjustments: " +
        f"{P} {money(c6['s_p']['ta'])}, {C} {money(c6['s_c']['ta'])}; only the revenue cutoff moves an asset, and the "
        "1090 reclassification stays within inventories.",
        "The income statement here is recomputed from the ledger and the lines; its lines are the workbook's: "
        f"operating revenue {money(i_c['opr'])}, cost of goods sold {money(i_c['cogs'])}, operating expenses "
        f"{money(i_c['opx'])} in {C}."], size=10))
    page.expect = ["Adjusted Net Income", money(i_c["ni"]), money(i_p["ni"]), money(rec[-1]["ta"]), "Model answer"]

    # Tests 3 to 5
    yl = ", ".join(map(str, years))
    b.add_test("Test 3: each entry of the package balances (debits less credits by adjustment and year-end)",
               "EVALUATE\nSUMMARIZECOLUMNS (\n    AdjustmentLines[No],\n    AdjustmentLines[YearEnd],\n"
               "    \"Debits\", ROUND ( [Adjustment Debits], 2 ),\n    \"Credits\", ROUND ( [Adjustment Credits], 2 ),\n"
               "    \"Debits less Credits\", ROUND ( [Adjustment Amount], 2 )\n)\n"
               "ORDER BY AdjustmentLines[YearEnd], AdjustmentLines[No]")
    b.add_test("Test 4: no adjustment touches cash, and no entry is out of balance",
               "EVALUATE\nROW (\n    \"Lines on Cash\", [Lines on Cash],\n    \"Entries\", [Entries],\n"
               "    \"Entries out of Balance\", [Entries out of Balance]\n)")
    b.add_test("Test 5: adjusted net income agrees with the workbook, and recorded total assets agree with the trial "
               "balances of Requirement 1",
               "EVALUATE\nSUMMARIZECOLUMNS (\n    'Date'[Year],\n    TREATAS ( { " + yl + " }, 'Date'[Year] ),\n"
               "    \"Adjusted Net Income\", ROUND ( [Adjusted Net Income], 2 ),\n"
               "    \"Workbook Adjusted Net Income\", ROUND ( [Workbook Adjusted Net Income], 2 ),\n"
               "    \"Total Assets\", ROUND ( [Total Assets], 2 ),\n"
               "    \"Workbook Recorded Total Assets\", ROUND ( [Workbook Recorded Total Assets], 2 ),\n"
               "    \"Accounts Disagreeing\", [Accounts Disagreeing]\n)\nORDER BY 'Date'[Year]")

    s = "Requirement 6"
    n_entries = len({(r["No"], r["YearEnd"]) for r in lines})
    claim(n_entries == 4 * len(years) + 2 * 2, "A1-A4 at each year-end and R1-R2 at the two year-ends presented")
    b.check(s, "adjustment lines loaded (the workbook's Table, without its totals row)", len(lines),
            "[Adjustment Lines]", 0)
    b.check(s, "entries (A1-A4 at each year-end, R1 and R2 at the two presented)", n_entries, "[Entries]", 0)
    b.check(s, "every entry nets to zero: entries out of balance", 0, "[Entries out of Balance]", 0)
    b.check(s, "every entry nets to zero: the largest imbalance", 0, "[Largest Entry Imbalance]")
    b.check(s, "lines on 1010 Cash", 0, "[Lines on Cash]", 0)
    b.check(s, "lines whose account is not in the chart", 0,
            "COUNTROWS ( FILTER ( AdjustmentLines, ISBLANK ( RELATED ( Account[AccountID] ) ) ) ) + 0", 0)
    b.check(s, "lines dated on a year-end of the window", len(lines),
            f"CALCULATE ( [Adjustment Lines], 'Date'[Date] IN {{ {', '.join(f'DATE ( {y}, 12, 31 )' for y in years)} }} )", 0)
    for y, i in ((P, i_p), (C, i_c)):
        cy = f"'Date'[Year] = {y}"
        b.check(s, f"adjusted net income {y} (ledger plus lines)", i["ni"], f"CALCULATE ( [Adjusted Net Income], {cy} )", 0.01)
        b.check(s, f"adjusted net income {y}: the workbook's", i["ni"], f"CALCULATE ( [Workbook Adjusted Net Income], {cy} )",
                0.01)
        b.check(s, f"adjusted net income {y}: difference", 0, f"CALCULATE ( [Net Income Difference], {cy} )")
        for key, ln, sign in (("opr", "Operating revenue", 1), ("ret", "Sales returns and allowances", -1),
                              ("cogs", "Cost of goods sold", -1), ("opx", "Operating expenses", -1),
                              ("intr", "Interest expense", -1), ("loss", "Loss on disposal of equipment", -1)):
            b.check(s, f"adjusted {y}: {ln.lower()}", sign * i[key],
                    f"CALCULATE ( [Adjusted P&L Amount], {cy}, WorkingTB[StatementLine] = \"{ln}\" )", 0.01)
    for y, r in zip(years, rec):
        cy = f"'Date'[Year] = {y}"
        b.check(s, f"recorded total assets at {ch17.ye(y)} (ledger)", r["ta"], f"CALCULATE ( [Total Assets], {cy} )", 0.01)
        b.check(s, f"recorded total assets at {ch17.ye(y)}: the workbook's trial balance", r["ta"],
                f"CALCULATE ( [Workbook Recorded Total Assets], {cy} )", 0.01)
        b.check(s, f"recorded total assets at {ch17.ye(y)}: difference", 0, f"CALCULATE ( [Total Assets Difference], {cy} )")
        b.check(s, f"accounts whose balance at {ch17.ye(y)} differs from the workbook's", 0,
                f"CALCULATE ( [Accounts Disagreeing], {cy} )", 0)
    b.check(s, f"the workbook's adjusted trial balance at {ch17.ye(C)} balances", 0,
            f"ROUND ( SUM ( WorkingTB[Adjusted{C}] ), 2 )")


# --- Milestone 3: the statements' totals the Review file reproduces --------------------------------------------------

def m3(b: Build) -> None:
    d, f = b.d, b.ctx["m3"]
    s = "Milestone 3"
    for y in (d.P, d.C):
        cy = f"'Date'[Year] = {y}"
        st = f["s_c"] if y == d.C else f["s_p"]
        b.check(s, f"adjusted net income {y}", f["adj_ni"][y], f"CALCULATE ( [Adjusted Net Income], {cy} )", 0.01)
        b.check(s, f"total assets at {ch17.ye(y)}, adjusted", st["ta"], f"CALCULATE ( [Adjusted Total Assets], {cy} )", 0.01)
    b.check(s, f"retained earnings at 1 January {d.P}, as previously reported (3030 at {ch17.ye(d.F)} plus {d.F} net "
               "income)", f["reported"],
            f"- CALCULATE ( [Balance {d.F}], Account[AccountNumber] = 3030 ) + CALCULATE ( [Net Income], 'Date'[Year] = {d.F} )",
            0.01)
    b.check(s, f"the correction of retained earnings at 1 January {d.P} (lines to 3030 at {ch17.ye(d.P)})",
            ch17.xr(f["restated"] - f["reported"]),
            f"- CALCULATE ( [Adjustment Amount], Account[AccountNumber] = 3030, 'Date'[Year] = {d.P} )", 0.01)
    b.check(s, f"property and equipment cost at {ch17.ye(d.C)}", f["cost"],
            f'CALCULATE ( [Balance {d.C}], Account[AccountSubType] = "Fixed Asset" )', 0.01)
    b.check(s, f"accumulated depreciation at {ch17.ye(d.C)}", -f["accumulated"],
            f'CALCULATE ( [Balance {d.C}], Account[AccountSubType] = "Contra Fixed Asset" )', 0.01)
    b.check(s, f"notes payable at {ch17.ye(d.C)}", -f["s_c"]["notes"],
            f"CALCULATE ( [Balance {d.C}], Account[AccountNumber] = 2110 )", 0.01)
    b.check(s, f"current portion at {ch17.ye(d.C)} (R1)", f["s_c"]["cur"],
            f"CALCULATE ( [Adjustment Amount], AdjustmentLines[No] = \"R1\", AdjustmentLines[Amount] > 0, "
            f"'Date'[Year] = {d.C} )", 0.01)


EXERCISES = [("Requirement 2", r2), ("Milestone 1", m1), ("Requirement 6", r6), ("Milestone 3", m3)]

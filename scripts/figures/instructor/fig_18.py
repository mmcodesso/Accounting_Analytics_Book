"""Instructor figures of Chapter 18 (guides in facts/instructor/guides/chapter-18-make-buy-or-reprice/).

Chapter 18, "Make, Buy, or Reprice", is the second capstone case.  Its model answer is three files: the script MakeBuy.sql, the
workbook Make Buy Reprice.xlsx and the Power BI project Product Costs (outputs/companion/solutions/Chapter18-capstone-*).  The SQL
mock quotes its query from the script, at the editor line where it sits, and runs it read-only, so the grid and the message pane
come from the data.  The Excel mock draws a block of the workbook: its values come from the guides' own contexts (ch18.r4) and
each is asserted against the value the workbook holds in the cell it draws, so a stale workbook fails the build.  The Power BI
mocks draw the Families page and the Tests tab from the project's definition (page, visuals, titles, query tabs, the test's
text) with values from the registered notes, asserted against the workbook's ItemCosts Table and the project's Checks tab.  The
diagrams take their structure from the case text and from the guides' tables, and the charts from the same contexts and from the
query results the guides name.

The figures of Requirements 6 to 10 (IF-18-6-A and later) are drawn by the second half of this module; the helpers above the
first figure (the Excel sheet, the SQL mock, the notes, the workbook reader) are written to be reused by it.
"""

from __future__ import annotations

import json
import re
import statistics
from functools import lru_cache

import ctx
import dbbrowser as db
import excel as xl
import kit
import powerbi as pbi
from check_figures import text_width
from data import connection
from drawio import (AMBER, AMBER_TINT, BLUE, BLUE_TINT, CORAL, CORAL_TINT, FONT, GRAY, GRAY_TINT, INK, ROW_H, RULE, SMALL, TEAL,
                    TEAL_TINT, WHITE, Diagram, esc)

REPO = ctx.REPO
SOLUTIONS = REPO / "outputs" / "companion" / "solutions"
SQL = SOLUTIONS / "Chapter18-capstone-sql" / "MakeBuy.sql"
XLSX = SOLUTIONS / "Chapter18-capstone-excel" / "Make Buy Reprice.xlsx"
PROJECT = SOLUTIONS / "Chapter18-capstone-pbi" / "Product Costs"
MODEL = PROJECT / "Product Costs.SemanticModel"
REPORT = PROJECT / "Product Costs.Report" / "definition"
CHAPTER = REPO / "chapters" / "18-make-buy-or-reprice" / "chapter.qmd"
GUIDES = REPO / "facts" / "instructor" / "guides" / "chapter-18-make-buy-or-reprice"
GRID_W = 820            # the 860 px frame less the row-number column of the result grid
CELL_PAD = 20           # a cell wraps its text within about 12 px of its width, so leave room
MINUS = "−"
STRAIGHT = "edgeStyle=none;"


# -- shared helpers -----------------------------------------------------------------------------------------------------

def _note(d: Diagram, text: str, y: float, x: float = 0, w: float = 860) -> float:
    """A short italic gray note sized to its text; returns its bottom."""
    h = kit.height_for(text, w, SMALL, pad=8)
    d.text(f"<i>{text}</i>", x, y, w, h, size=SMALL, color=GRAY)
    return y + h


def _money(x: float) -> str:
    return f"{x:,.2f}"


def _g():
    """guide_contexts.ch18: the helpers that read the model answers."""
    from guide_contexts import ch18
    return ch18


def _n():
    from notes import ch18
    return ch18


def _model() -> dict:
    """The case's twin of every value (facts/notes/ch18.py), for the values a guide context does not hold."""
    return _n().model(ctx.data())


def _clip(text: str, width: float, size: int = SMALL, margin: float = 14, bold: bool = False) -> str:
    """A value cut to its column with an ellipsis, as a narrow column shows it."""
    if not text or text_width(esc(text), size, bold) <= width - margin:
        return text
    while text and text_width(esc(text + "…"), size, bold) > width - margin:
        text = text[:-1]
    return text.rstrip() + "…"


def _nowrap(text: str) -> str:
    return f'<span style="white-space:nowrap">{esc(text)}</span>'


def _wrap(text: str, width: float, size: int = SMALL, bold: bool = False) -> list[str]:
    """A greedy word wrap measured with text_width."""
    out, cur = [], ""
    for word in text.split(" "):
        trial = (cur + " " + word).strip()
        if not cur or text_width(trial, size, bold) <= width:
            cur = trial
        else:
            out.append(cur)
            cur = word
    out.append(cur)
    return out


def _polyline(d: Diagram, pts: list[tuple[float, float]], color: str = BLUE, width: float = 3, dashed: bool = False) -> None:
    """One edge through the points, drawn straight between them (separate edges leave gaps at their joints)."""
    a, b = d.anchor(*pts[0]), d.anchor(*pts[-1])
    d.edge(a, b, color=color, width=width, dashed=dashed, points=pts[1:-1], meta="edgeStyle=none;rounded=0;")


def _label(d: Diagram, text: str, x: float, y: float, w: float, h: float = 18, color: str = INK, align: str = "left",
           bold: bool = False) -> None:
    """A white-backed text label, so a gridline does not strike through it."""
    shown = f"<b>{text}</b>" if bold else text
    d.vertex(shown, f"text;html=1;whiteSpace=wrap;align={align};verticalAlign=middle;fontSize={SMALL};fontColor={color};{FONT}"
                    f"spacing=2;fillColor={WHITE};strokeColor=none;", x, y, w, h)


# -- the case text -----------------------------------------------------------------------------------------------------

def _case_text() -> str:
    return CHAPTER.read_text(encoding="utf-8")


def _requirements() -> list[dict]:
    """Each requirement of the case text: number, title, chapters (empty for the memo) and the phase it belongs to."""
    found, phase = [], None
    for line in _case_text().split("\n"):
        m = re.match(r"^### Phase (\d): (.+)$", line)
        if m:
            phase = (int(m.group(1)), m.group(2).strip())
        m = re.match(r"^\*\*Requirement (\d+): (.+?)(?: \((Chapters? [^)]+)\))?\.\*\*", line)
        if m:
            found.append(dict(n=int(m.group(1)), title=m.group(2), chapters=m.group(3) or "", phase=phase))
    assert [r["n"] for r in found] == list(range(1, 11)), found
    assert [r["phase"][0] for r in found] == [1, 1, 1, 2, 2, 3, 3, 3, 4, 4], found
    return found


def _tools() -> dict[int, list[str]]:
    """The tools of each requirement as the case notes' table 'Using the requirements' lists them (Power Query is Excel's)."""
    text = (GUIDES / "00-case-notes.md").read_text(encoding="utf-8")
    out = {}
    for m in re.finditer(r"^\| (\d+)\. [^|]+\| (?:Managerial|Financial|Auditing|General) \| ([^|]+?) \|", text, re.M):
        tags: list[str] = []
        for part in re.split(r", | and ", m.group(2)):
            tag = {"Power Query": "Excel", "a written memo": "Memo"}.get(part.strip(), part.strip())
            if tag not in tags:
                tags.append(tag)
        out[int(m.group(1))] = tags
    assert sorted(out) == list(range(1, 11)), out
    return out


def _milestones() -> list[dict]:
    found = [dict(n=int(m.group(1)), title=m.group(2), reqs=m.group(3), body=m.group(4))
             for m in re.finditer(r"^(\d)\. \*\*([^*]+)\*\* \(Requirements? ([^)]+)\): (.+?)\.$", _case_text(), re.M)]
    assert [x["n"] for x in found] == [1, 2, 3, 4], found
    return found


def _tool_files() -> list[dict]:
    """The rows of the case's table of tools and files: tool, file, what it holds and how it hands results on."""
    rows = []
    for m in re.finditer(r"^\| (SQL|Excel|Power BI) \| ([^|]+?) \| ([^|]+?) \| ([^|]+?) \|$", _case_text(), re.M):
        rows.append(dict(tool=m.group(1), file=m.group(2), holds=m.group(3), handoff=m.group(4)))
    assert [r["tool"] for r in rows] == ["SQL", "Excel", "Power BI"], rows
    return rows


TOOL_COLORS = {"SQL": (TEAL, TEAL_TINT), "Excel": (BLUE, BLUE_TINT), "Power BI": (AMBER, AMBER_TINT), "Memo": (GRAY, GRAY_TINT)}


# -- IF-18-0-A: the requirements, the files and the milestones ---------------------------------------------------------------

def if_18_0_a() -> Diagram:
    """Ten requirement cards in four phase columns with their tool tags, a milestone under each column, the three files with the
    hand-offs between them, and the memo they all feed."""
    reqs, tools, miles, files = _requirements(), _tools(), _milestones(), _tool_files()
    assert [m["reqs"] for m in miles] == ["1 to 3", "4 and 5", "6 to 8", "9 and 10"], miles
    text = _case_text()
    assert "no more than two pages" in text and "no more than three pages" in text, "the memo's page limits are the case's"
    d = Diagram("How the Requirements, Files and Milestones Fit Together")
    W, GAP = 200.0, 20.0
    xs = [i * (W + GAP) for i in range(4)]
    bottoms = []
    for p in range(1, 5):
        phase = [r for r in reqs if r["phase"][0] == p]
        d.box(f"<b>Phase {p}</b><br><b>{esc(phase[0]['phase'][1])}</b>", xs[p - 1], 0, W, 40, fill=GRAY, stroke=GRAY, color=WHITE, size=SMALL,
              rounded=False)
        y = 48.0
        for r in phase:
            label = f"<b>Requirement {r['n']}</b><br>{esc(r['title'])}" + (f"<br><i>{esc(r['chapters'])}</i>" if r["chapters"] else "")
            h = kit.height_for(label, W - 4, SMALL, pad=10) + 4 + 30
            d.box(label, xs[p - 1], y, W, h, fill=WHITE, stroke=GRAY, stroke_width=1.5, size=SMALL, align="left", valign="top", rounded=False)
            tx = xs[p - 1] + 6
            for tag in tools[r["n"]]:
                fill, _ = TOOL_COLORS[tag]
                tw = text_width(f"<b>{tag}</b>", SMALL, True) + 16
                d.box(f"<b>{tag}</b>", tx, y + h - 27, tw, 22, fill=fill, stroke=fill, color=WHITE, size=SMALL)
                tx += tw + 6
            y += h + 8
        bottoms.append(y - 8)
    top = max(bottoms) + 14
    # the milestones, one under each column
    labels = [f"<b>Milestone {m['n']}</b>: {esc(m['title'])}<br>{esc(m['body'][0].upper() + m['body'][1:])}" for m in miles]
    mh = max(kit.height_for(t, W - 10, SMALL, pad=12) for t in labels)
    for m, label in zip(miles, labels):
        d.box(label, xs[m["n"] - 1], top, W, mh, fill=TEAL_TINT, stroke=TEAL, stroke_width=1.5, size=SMALL, align="left", valign="top", rounded=False)
        d.box("", xs[m["n"] - 1], top, 5, mh, fill=TEAL, stroke=TEAL, rounded=False)
    # the files and their hand-offs
    fy = top + mh + 40
    FW, FG = 220.0, 100.0
    fxs = [i * (FW + FG) for i in range(3)]
    fh = max(kit.height_for(f"<b>{esc(f['file'])}</b><br>{esc(f['holds'])}", FW, SMALL, pad=12) for f in files) + 8
    fids = []
    for f, fx in zip(files, fxs):
        fill, tint_ = TOOL_COLORS[f["tool"]]
        fids.append(d.box(f"<br><b>{esc(f['file'])}</b><br>{esc(f['holds'])}", fx, fy, FW, fh, fill=tint_, stroke=fill, stroke_width=1.5, size=SMALL,
                          align="left", valign="top", rounded=False))
        tw = text_width(f"<b>{f['tool']}</b>", SMALL, True) + 16
        d.box(f"<b>{f['tool']}</b>", fx + 6, fy - 11, tw, 22, fill=fill, stroke=fill, color=WHITE, size=SMALL)
    for i in (0, 1):
        d.arrow(fids[i], fids[i + 1], exit=(1, 0.5), entry=(0, 0.5), meta=STRAIGHT)
        d.text(f"<i>{esc(files[i]['handoff'])}</i>", fxs[i] + FW + 2, fy + fh / 2 - 62, FG - 4, 58, size=SMALL, color=GRAY, align="center",
               valign="bottom")
    # the memo
    my = fy + fh + 50
    memo = d.box("<b>The memo to the CFO and the production manager</b> (Requirement 10): at most two pages, with the decision table, and at most three pages of exhibits.<br>"
                 "Every number in it traces to one of the three files.", 0, my, 860, 52, fill=GRAY_TINT, stroke=GRAY, stroke_width=1.5, size=SMALL,
                 rounded=False)
    for fid, fx in zip(fids, fxs):
        d.arrow(fid, memo, exit=(0.5, 1), entry=((fx + FW / 2) / 860, 0), meta=STRAIGHT)
    _note(d, "Each card is a requirement of the case, in the order of its four phases; the tags name the tools it uses (Requirement 4 builds its item costs "
             "with Power Query inside Excel). The teal boxes are the four milestones an instructor may collect and check. Arrows show how results move: "
             "Copy with Headers or Export to CSV from the script into the workbook, the workbook's Table of item costs into the explorer, and all three "
             "files into the memo.", my + 62)
    return d


# -- IF-18-0-B: the same inputs to the clearing account cut three ways ---------------------------------------------------------

def if_18_0_b() -> Diagram:
    """Three bars of equal length, each the year's inputs to account 1090: the standard view (released at standard, and the rest, the
    variance), actual absorption (every dollar spread over the units made) and time-driven (the products' conversion and the unused
    capacity, with the crew's part of it marked)."""
    data = ctx.data()
    r2, r4, m1 = ctx.values("ch18.r2"), ctx.values("ch18.r4"), ctx.values("ch18.m1")
    C = data.C
    inputs, released = r4["total"], m1["released"]
    conv, unused, crew, other = r4["conv"], r4["unused"], r4["unused_crew"], r4["unused"] - r4["unused_crew"]
    rest = inputs - released
    assert abs(m1["inputs"] - inputs) < 0.005 and abs(r2["inputs"] - inputs) < 0.005, "the inputs are the guides' own"
    assert abs(conv + unused - inputs) < 0.01, "the products' conversion and the unused capacity are the inputs"
    assert 0 < rest < inputs and 0 < crew < unused, (rest, crew, unused)
    wb = _cached()["Costs"]
    for ref, v in (("B24", conv), ("B26", unused), ("B27", crew), ("B28", inputs), ("B29", inputs), ("B30", 0)):
        _agree(wb, ref, v)
    # the part of the inputs not released is what the closes carry to 5080 as the conversion variance, apart from the year's net change in 1090
    first, last = f"{C}-01-01", f"{C}-12-31"
    conv_var = data.one("SELECT SUM(DirectLaborVarianceAmount + OverheadVarianceAmount) FROM WorkOrderClose WHERE CloseDate BETWEEN ? AND ?", first, last)
    net_1090 = data.one("SELECT SUM(Debit) - SUM(Credit) FROM GLEntry WHERE AccountID = ? AND FiscalYear = ?", data.account("1090"), C)
    assert abs(rest - conv_var - net_1090) < 0.01, (rest, conv_var, net_1090)
    td_rate, unused_h = r4["td_rate"], r4["unused_h"]
    d = Diagram("The Same Inputs, Three Costing Views")
    X0, BAR_W, BAR_H = 138.0, 720.0, 54.0
    scale = BAR_W / inputs
    wd = lambda v: v * scale                                                                   # noqa: E731
    # the total, as a bracket over the bars
    d.text(f"<b>The year's inputs to account 1090: ${_money(inputs)}</b>", X0, 0, BAR_W, 20, size=SMALL, align="center")
    for x, w in ((X0, 1.5), (X0 + BAR_W - 1.5, 1.5)):
        d.box("", x, 24, w, 12, fill=INK, stroke=INK, rounded=False)
    d.box("", X0, 24, BAR_W, 1.5, fill=INK, stroke=INK, rounded=False)

    def seg(x: float, y: float, value: float, label: str, fill: str, color: str = WHITE, stroke: str | None = None) -> None:
        d.box(label, x, y, wd(value), BAR_H, fill=fill, stroke=stroke or fill, color=color, size=SMALL, rounded=False, stroke_width=1.5)

    rows = [("Standard", "What the standards release to the products"),
            ("Actual absorption", "In concept, every dollar of the inputs on the units made"),
            ("Time-driven", "What the products' routing time used, and the capacity left over")]
    pitch = 120.0
    tops = [50.0 + i * pitch for i in range(3)]
    for (name, sub), y in zip(rows, tops):
        d.text(f"<b>{name}</b>", 0, y, X0 - 10, 22, size=SMALL)
        d.text(esc(sub), 0, y + 22, X0 - 10, 52, size=SMALL, color=GRAY)
    # standard
    y = tops[0]
    seg(X0, y, released, f"<b>Released at standard</b><br>${_money(released)}", BLUE)
    seg(X0 + wd(released), y, rest, f"<b>The rest: the variance</b><br>${_money(rest)}", CORAL)
    d.text("<i>The standards charge the products only what they release; the rest waits in 1090 until each work order's close sends it to account 5080.</i>",
           X0, y + BAR_H + 4, BAR_W, 34, size=SMALL, color=GRAY)
    # actual absorption
    y = tops[1]
    seg(X0, y, inputs, f"<b>Every dollar spread over the units made</b><br>${_money(inputs)}, or ${r2['actual_rate']:.2f} a standard hour over {r2['stdh']:,.0f} standard hours completed", BLUE)
    d.text("<i>In concept nothing is left over, so the idle capacity and the supervision sit in the products' cost. The case approximates it with each family's closed variance per unit completed.</i>",
           X0, y + BAR_H + 4, BAR_W, 34, size=SMALL, color=GRAY)
    # time-driven
    y = tops[2]
    seg(X0, y, conv, f"<b>The products' conversion</b><br>${_money(conv)}", BLUE)
    ux = X0 + wd(conv)
    seg(ux, y, crew, f"<b>The crew</b><br>${_money(crew)}", AMBER)
    seg(ux + wd(crew), y, other, "", AMBER_TINT, stroke=AMBER)
    bx = ux
    d.box("", bx, y - 12, wd(unused), 1.5, fill=INK, stroke=INK, rounded=False)
    for x in (bx, bx + wd(unused) - 1.5):
        d.box("", x, y - 12, 1.5, 9, fill=INK, stroke=INK, rounded=False)
    tw = 470
    d.text(f"<b>Unused capacity, the plant's cost: ${_money(unused)}</b>", X0 + BAR_W - tw, y - 34, tw, 20, size=SMALL, align="right")
    d.text(f"<i>{unused_h:,.1f} practical hours the routings did not use, at ${td_rate:.2f} an hour. The crew's part is the darker segment; the light "
           f"segment, ${_money(other)}, is supervision and equipment.</i>", X0, y + BAR_H + 4, BAR_W, 34, size=SMALL, color=GRAY)
    bottom = tops[2] + BAR_H + 44
    kit.legend(d, 0, bottom + 6, [("Charged to the products", BLUE, "box"), ("The variance, outside the product costs", CORAL, "box"),
                                  ("Unused capacity (the crew's part darker)", AMBER, "box")])
    _note(d, f"The three bars are the same ${_money(inputs)} of fiscal {C} inputs to 1090 (payroll, the Factory Overhead entries and depreciation), "
             f"drawn to one scale. The standard view's rest is the conversion variance the closes carry to 5080, ${_money(conv_var)}, plus the year's net change "
             f"in 1090, ${_money(net_1090)}.", bottom + 34)
    return d


# -- IF-18-2-B: the two scatter plots of Requirement 2 --------------------------------------------------------------------------

def _hollow(d: Diagram, cx: float, cy: float, color: str = GRAY, size: float = 11) -> None:
    d.vertex("", f"ellipse;html=1;fillColor={WHITE};strokeColor={color};strokeWidth=2;", cx - size / 2, cy - size / 2, size, size)


def _diamond(d: Diagram, cx: float, cy: float, color: str = CORAL, size: float = 12) -> None:
    d.vertex("", f"rhombus;html=1;fillColor={color};strokeColor={WHITE};strokeWidth=1;", cx - size / 2, cy - size / 2, size, size)


def _dot(d: Diagram, cx: float, cy: float, color: str = BLUE, size: float = 8) -> None:
    d.vertex("", f"ellipse;html=1;fillColor={color};strokeColor={WHITE};strokeWidth=1;", cx - size / 2, cy - size / 2, size, size)


def if_18_2_b() -> Diagram:
    """Two panels with the standard hours completed each month on the horizontal axis: the Factory Overhead entries with the
    fitted line rising, and the payroll to 1090 with a flat fitted line and the months with three pay dates marked; the start-up
    month sits apart from the line in both."""
    data = ctx.data()
    g = _g()
    m = _model()
    F, C = data.F, data.C
    rows = g.run(data, g.quote("Requirement 2, Q14"))
    assert len(rows) == 36 and rows[0][0] == f"{F}-01" and rows[-1][0] == f"{C}-12", (len(rows), rows[0][0], rows[-1][0])
    start, test = rows[0], rows[1:]
    xs = [r[1] for r in test]
    fit = {}
    for key, col in (("fo", 3), ("pay", 4)):
        ys = [r[col] for r in test]
        lr = statistics.linear_regression(xs, ys)
        r = statistics.correlation(xs, ys)
        f = m["fits"][key]
        assert abs(lr.slope - f["slope"]) < 1e-3 and abs(lr.intercept - f["intercept"]) < 10 and abs(r - f["r"]) < 1e-4, (key, lr, r, f)
        fit[key] = dict(slope=lr.slope, intercept=lr.intercept, r=r, r2=r * r)
    three = [r[0] for r in test if r[6] == 3]
    assert three == m["three"], three
    assert start[3] == 0 and start[4] == 0 and abs(start[1] - m["startup"]["stdh"]) < 0.01, "the start-up month posted no payroll and no overhead entry to 1090"
    assert m["startup"]["to_5080"] > 0, "its one payroll went to 5080"
    assert fit["fo"]["r"] > 0.9 and abs(fit["pay"]["r"]) < 0.2, "the overhead entries follow output and the payroll does not"
    names = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
    pretty = lambda ym: f"{names[int(ym[5:7]) - 1]} {ym[:4]}"                                    # noqa: E731
    d = Diagram("Overhead Entries and Payroll Against Output")
    PW, GAP = 410.0, 20.0
    top = 40.0
    PH = 338.0
    x_ticks = [0, 2000, 4000, 6000, 8000, 10000]
    comma = lambda v: f"{v:,.0f}"                                                                # noqa: E731
    panels = [("fo", 3, "Factory Overhead entries to 1090", [0, 50000, 100000, 150000, 200000], 0),
              ("pay", 4, "Payroll to 1090", [0, 100000, 200000, 300000, 400000, 500000], PW + GAP)]
    plots = {}
    for key, col, title, y_ticks, px in panels:
        f = fit[key]
        d.text(f"<b>{esc(title)}</b>", px, 0, PW, 20, size=SMALL)
        d.text(f"r = {f['r']:.3f}, R² = {f['r2']:.3f}, slope ${f['slope']:.2f} a standard hour", px, 18, PW, 20, size=SMALL, color=GRAY)
        plot = kit.Plot(d, px, top, PW, PH, y_ticks=y_ticks, x_ticks=x_ticks, x_fmt=comma, x_title="Standard hours completed in the month",
                        y_title="$ a month", left=60)
        plots[key] = plot
        lo, hi = min(xs), max(xs)
        plot.line([(lo, f["intercept"] + f["slope"] * lo), (hi, f["intercept"] + f["slope"] * hi)], color=AMBER, width=3)
        for r in test:
            if key == "pay" and r[6] == 3:
                continue
            _dot(d, plot.x_of(r[1]), plot.y_of(r[col]))
        if key == "pay":
            for r in test:
                if r[6] == 3:
                    _diamond(d, plot.x_of(r[1]), plot.y_of(r[col]))
        _hollow(d, plot.x_of(start[1]), plot.y_of(start[col]))
    # the start-up month's label, in the clear space above the axis
    text = "January 2024, the start-up month, left out of the fit"
    for key in ("fo", "pay"):
        p = plots[key]
        _label(d, text, p.x_of(start[1]) + 12, p.y_of(0) - 24, text_width(text, SMALL, False) + 12)
    bottom = top + PH + 10
    # the key
    cx = 0.0
    for shape, label in (("dot", "A month"), ("diamond", f"A month with three pay dates ({len(three)})"),
                         ("ring", "The start-up month"), ("line", f"The fitted line, {pretty(test[0][0])} to {pretty(test[-1][0])}")):
        if shape == "dot":
            _dot(d, cx + 6, bottom + 10)
            cx += 18
        elif shape == "diamond":
            _diamond(d, cx + 7, bottom + 10)
            cx += 20
        elif shape == "ring":
            _hollow(d, cx + 7, bottom + 10)
            cx += 20
        else:
            d.line_sample(cx, bottom + 10, 28, color=AMBER, width=3)
            cx += 34
        width = text_width(label, SMALL, False) + 8
        d.text(label, cx, bottom, width, 20, size=SMALL, valign="middle")
        cx += width + 20
    _note(d, f"Each point is a month, {pretty(test[0][0])} to {pretty(test[-1][0])}; the fit leaves out {pretty(start[0])}, which has {start[1]:,.0f} standard hours and no payroll "
             f"or overhead entry to 1090 (its one payroll, of salaries, went to 5080). The overhead entries move with output at about the applied rate; the payroll "
             f"does not, and the months with three pay dates ({', '.join(pretty(x) for x in three)}) carry a third payroll with no more output.", bottom + 30)
    return d


# -- SQL mocks: the model script, a read-only run and the Execute SQL tab -----------------------------------------------

def _script_lines() -> list[str]:
    if not SQL.is_file():
        raise SystemExit(f"Model answer not found: {SQL} (build the solutions first)")
    return SQL.read_text(encoding="utf-8").split("\n")


def _solution(key: str) -> tuple[list[str], str, int]:
    """The two comment lines, the SQL and the editor line of the one query of MakeBuy.sql whose first comment line starts with
    `-- key` (for example "Requirement 3, Q20").  The `-- Result: ...` remark that the model answer appends to a query's last line
    is left out: a student types the query without it."""
    lines = _script_lines()
    hits = [i for i, line in enumerate(lines) if line.startswith(f"-- {key}:")]
    if len(hits) != 1:
        raise SystemExit(f"MakeBuy.sql: {len(hits)} queries start with {key!r} (expected one)")
    i = hits[0]
    assert lines[i + 1].startswith("-- Population"), lines[i + 1]
    j = i + 2
    while not re.search(r";\s*(--.*)?$", lines[j]):
        j += 1
    sql = lines[i + 2:j + 1]
    sql[-1] = re.sub(r"\s*--\s*Result:.*$", "", sql[-1])
    return lines[i:i + 2], "\n".join(sql), i + 1


def _temp_view(name: str) -> str:
    """The CREATE VIEW statement of MakeBuy.sql for `name`, as a TEMP view: the model answer writes the view into the working
    copy, the figure's read-only run keeps it in the connection's temporary database."""
    m = re.search(rf"CREATE VIEW {name} AS\n(.*?);\s*--\s*Result", "\n".join(_script_lines()), re.S)
    if not m:
        raise SystemExit(f"MakeBuy.sql: no CREATE VIEW {name}")
    return f"CREATE TEMP VIEW {name} AS\n{m.group(1)}"


class _Views:
    """The saved views of MakeBuy.sql created as TEMP views for the length of a block, and dropped again (the connection is
    shared with other units' figures, which create views of the same names)."""

    def __init__(self, *names: str):
        self.names = names

    def __enter__(self):
        con = connection()
        for name in self.names:
            con.execute(f"DROP VIEW IF EXISTS temp.{name}")
            con.execute(_temp_view(name))
        return self

    def __exit__(self, *exc):
        con = connection()
        for name in reversed(self.names):
            con.execute(f"DROP VIEW IF EXISTS temp.{name}")


def _run(sql: str) -> tuple[list[str], list[tuple]]:
    cursor = connection().execute(sql)
    return [c[0] for c in cursor.description], cursor.fetchall()


GUTTER = 42             # the gutter of the book's editor is 34 px, too narrow for a four-digit line number


def _editor(d: Diagram, x: float, y: float, w: float, sql: str, tab: str, first_line: int, selection: slice | None) -> float:
    """db.editor with a gutter wide enough for four-digit line numbers and a tab as wide as the file's name.  Returns the bottom y."""
    bx = x
    for label, width, primary in [("Open SQL file(s)", 124, False), ("Save SQL file ▾", 124, False), ("▶ Execute all/selected SQL", 196, True)]:
        xl.button(d, bx, y, width, label, primary=primary)
        bx += width + 8
    y += 32
    d.box(f"<b>{esc(tab)}</b>", x, y, max(110, text_width(f"<b>{esc(tab)}</b>", SMALL, True) + 30), 24, fill=WHITE, stroke=RULE, size=SMALL,
          align="left", rounded=False)
    y += 24
    lines = sql.strip("\n").split("\n")
    height = len(lines) * db.LINE_H + 8
    d.box("", x, y, GUTTER, height, fill=GRAY_TINT, stroke=RULE, rounded=False)
    d.box("", x + GUTTER, y, w - GUTTER, height, fill=WHITE, stroke=RULE, rounded=False)
    if selection:
        first, last = selection.start or 0, (selection.stop or len(lines)) - 1
        d.box("", x + GUTTER + 1, y + 4 + first * db.LINE_H, w - GUTTER - 2, (last - first + 1) * db.LINE_H, fill=BLUE_TINT, stroke=BLUE_TINT,
              rounded=False)
    in_comment = False
    for i, line in enumerate(lines):
        ly = y + 4 + i * db.LINE_H
        d.text(str(first_line + i), x, ly, GUTTER - 4, db.LINE_H, size=SMALL, color=GRAY, align="right", valign="middle")
        label, in_comment = db.highlight(line, in_comment)
        assert text_width(label, SMALL, False) <= w - GUTTER - 12, f"SQL line too wide: {line}"
        d.text(label or " ", x + GUTTER + 4, ly, w - GUTTER - 8, db.LINE_H, size=SMALL, valign="middle")
    return y + height


def _fit_comment(line: str) -> str:
    """A long comment line cut at the editor's edge with an ellipsis (the real editor scrolls sideways; the mock shows its first
    part)."""
    limit = 860 - GUTTER - 12 - 8
    if text_width(db.highlight(line)[0], SMALL, False) <= limit:
        return line
    while line and text_width(db.highlight(line.rstrip() + "…")[0], SMALL, False) > limit:
        line = line[:-1]
    return line.rstrip() + "…"


def _widths(headers: list[str], rows: list[tuple], total: float = GRID_W, stretch: float = 1.5, pad: float = CELL_PAD + 2) -> list[float]:
    """Column widths that fit each heading and value, shrunk (the widest column first, its text then clipped) to fit the frame,
    or stretched a little when the result is narrow."""
    natural = []
    for c, head in enumerate(headers):
        w = text_width(f"<b>{esc(head)}</b>", SMALL, True)
        for row in rows:
            w = max(w, text_width(db.cell(row[c])[0], SMALL, False))
        natural.append(w + pad + (6 if any(isinstance(row[c], str) for row in rows) else 0))      # a text column keeps room to avoid a wrap
    if sum(natural) <= total:
        scale = min(total / sum(natural), stretch)
        return _fit([round(w * scale) for w in natural], total)
    widths = list(natural)
    floor = [text_width(f"<b>{esc(h)}</b>", SMALL, True) + pad for h in headers]
    while sum(widths) > total:
        i = max(range(len(widths)), key=lambda c: widths[c] - floor[c])
        assert widths[i] > floor[i], "the result does not fit the frame"
        widths[i] = max(floor[i], widths[i] - min(8, sum(widths) - total))
    return _fit([round(w) for w in widths], total)


def _fit(widths: list[int], total: float) -> list[int]:
    """Take any excess left by rounding off the widest column, so the grid never passes the frame."""
    excess = sum(widths) - total
    if excess > 0:
        widths[max(range(len(widths)), key=lambda c: widths[c])] -= int(excess + 0.999)
    return widths


def _cell(d: Diagram, label: str, x: float, y: float, w: float, h: float, fill: str = WHITE, color: str = INK, align: str = "left",
          spacing: float = 6) -> None:
    """A grid cell with a chosen margin beside its text (Diagram.box keeps 6 px, so a cell wraps within about 12 px of its width; a
    dense grid passes a smaller spacing and keeps its columns narrow)."""
    style = (f"rounded=0;whiteSpace=wrap;html=1;fillColor={fill};strokeColor={RULE};strokeWidth=1;fontColor={color};"
             f"fontSize={SMALL};fontStyle=0;align={align};verticalAlign=middle;{FONT}spacingLeft={spacing:g};"
             f"spacingRight={spacing:g};")
    d.vertex(label, style, x, y, w, h)


def _results(d: Diagram, y: float, headers: list[str], widths: list[float], rows: list[tuple], first: int = 1, spacing: float = 6,
             number_w: float = 40) -> dict:
    """db.results with a smaller margin in each cell and a narrower row-number column.  Returns the geometry (x, y, w, h) of every
    cell, keyed (row, column) from 0, with row -1 for the headings."""
    geometry: dict[tuple[int, int], tuple] = {}
    xs = [number_w]
    for width in widths[:-1]:
        xs.append(xs[-1] + width)
    _cell(d, "", 0, y, number_w, ROW_H, fill=GRAY_TINT, spacing=spacing)
    for c, (head, cx, width) in enumerate(zip(headers, xs, widths)):
        _cell(d, f"<b>{esc(head)}</b>", cx, y, width, ROW_H, fill=GRAY_TINT, spacing=spacing)
        geometry[(-1, c)] = (cx, y, width, ROW_H)
    for r, row in enumerate(rows):
        ry = y + (r + 1) * ROW_H
        _cell(d, str(first + r), 0, ry, number_w, ROW_H, fill=GRAY_TINT, color=GRAY, align="center", spacing=spacing)
        for c, (value, cx, width) in enumerate(zip(row, xs, widths)):
            label, align, color = db.cell(value)
            _cell(d, label, cx, ry, width, ROW_H, color=color, align=align, spacing=spacing)
            geometry[(r, c)] = (cx, ry, width, ROW_H)
    return geometry


def _mock(d: Diagram, key: str, lines: slice | None = None, shown: slice | None = None, tight: bool = False, pad: float | None = None,
          cols: slice | None = None) -> dict:
    """The Execute SQL tab with the query of MakeBuy.sql whose first comment line starts with `-- key`, as a reader sees it after
    selecting the query and clicking Execute all/selected SQL.  The editor shows the comment lines and the query at their line
    numbers in the script, or only `lines` of them (an editor scrolled to part of a long query, the line numbers kept); `shown`
    draws only some rows of the result, as a grid scrolled to them (the rows keep their numbers); `tight` keeps a smaller margin in
    each cell, for a result of many columns (`pad` narrows the columns further: a text value is then cut with an ellipsis); `cols` draws a
    grid scrolled to the right, only those columns (the row numbers stay).  Returns the grid and editor geometry (`headers` are the drawn
    columns' names, `all_headers` every column's, `rows` every row with every column)."""
    comment, sql, line = _solution(key)
    all_headers, rows = _run(sql)
    headers = all_headers[cols] if cols else all_headers
    grid_rows = [tuple(r[cols]) for r in rows] if cols else rows
    spacing, pad0, number_w = (3, 10, 32) if tight else (6, CELL_PAD + 2, 40)
    pad = pad0 if pad is None else pad
    saved, db.PATH = db.PATH, "C:\\Users\\you\\Documents\\CharlesRiver_Capstone.sqlite"    # the case works in a copy of that name
    try:
        y = db.window(d)
    finally:
        db.PATH = saved
    editor_top = y
    full = [_fit_comment(c) for c in comment] + sql.split("\n")
    n_comment = len(comment)
    start = (lines.start or 0) if lines else 0
    drawn = full[lines] if lines else full
    y = _editor(d, 0, y, 860, "\n".join(drawn), "MakeBuy.sql", line + start, slice(max(n_comment - start, 0), None)) + 8
    visible = grid_rows[shown] if shown else grid_rows
    first = (shown.start or 0) + 1 if shown else 1
    widths = _widths(headers, visible, total=860 - number_w, pad=pad)
    clipped = [tuple(_clip(v, w, margin=max(pad, 14)) if isinstance(v, str) else v for v, w in zip(row, widths)) for row in visible]
    if tight:
        geometry = _results(d, y, headers, widths, clipped, first=first, spacing=spacing, number_w=number_w)
    else:
        geometry = db.results(d, 0, y, headers, widths, clipped, first=first)
    y += (len(visible) + 1) * ROW_H + 8
    bottom = db.message(d, 0, y, 860, len(rows), sql.split("\n")[0], line=line + n_comment)
    return dict(headers=headers, all_headers=all_headers, rows=rows, geometry=geometry, bottom=bottom, sql_lines=drawn, sql=sql, n_comment=n_comment,
                editor_top=editor_top, first_line=line + start, widths=widths)


def _line_top(mock: dict, i: int) -> float:
    """The y of line i (from 0) of the text drawn in the editor, below the editor's toolbar and tab."""
    return mock["editor_top"] + 32 + 24 + 4 + i * db.LINE_H


def _find(mock: dict, text: str) -> int:
    hits = [i for i, line in enumerate(mock["sql_lines"]) if text in line]
    assert len(hits) == 1, (text, hits)
    return hits[0]


def _outline_lines(d: Diagram, mock: dict, first: int, last: int, color: str = TEAL) -> None:
    """An outline around lines first..last of the drawn editor text, inset so that neighbors do not touch."""
    top = _line_top(mock, first)
    d.outline(GUTTER + 3, top + 1.5, 860 - GUTTER - 6, (last - first + 1) * db.LINE_H - 3, color, width=2.5, front=True)


def _outline_row(d: Diagram, mock: dict, first: int, last: int | None = None, color: str = CORAL, inset: float = 0) -> None:
    """Outline rows first..last of the grid (indexes into the drawn rows), every column."""
    last = first if last is None else last
    geo, c_last = mock["geometry"], len(mock["headers"]) - 1
    x0, y0, _, _ = geo[(first, 0)]
    x1, y1, w1, h1 = geo[(last, c_last)]
    left, right = x0 - 1 + inset, min(860.0, x1 + w1 + 1 - inset)
    d.outline(left, y0 - 1 + inset, right - left, y1 + h1 - y0 + 2 - 2 * inset, color, width=2.5, front=True)


def _outline_column(d: Diagram, mock: dict, column: int, rows: int, color: str = CORAL) -> None:
    """Outline one column of the grid, from its heading down through `rows` drawn rows."""
    geo = mock["geometry"]
    x, y, w, _ = geo[(-1, column)]
    _, y_last, _, h_last = geo[(rows - 1, column)]
    right = min(860.0, x + w + 1)
    d.outline(x - 1, y - 1, right - x + 1, y_last + h_last - y + 2, color, width=2.5, front=True)


# -- IF-18-3-A: the capacity query and its six rows -------------------------------------------------------------------------

def if_18_3_a() -> Diagram:
    """Requirement 3's capacity query (Q20), scrolled to its last CTE and the final SELECT, with its six rows: one for each work
    center and a Plant row, which is outlined."""
    r3 = ctx.values("ch18.r3")
    d = Diagram("The Capacity Query and Its Result")
    key = "Requirement 3, Q20"
    comment, sql, line = _solution(key)
    n = len(comment) + len(sql.split("\n"))
    start = next(i for i, text in enumerate(comment + sql.split("\n")) if text.startswith("Rows AS ("))
    mock = _mock(d, key, lines=slice(start, n), tight=True, cols=slice(1, None))
    rows = mock["rows"]
    assert mock["all_headers"] == ["WorkCenterID", "WorkCenterName", "AvailableHours", "PracticalHours", "RunHours", "SetupHours", "RoutingHours",
                               "ShareOfAvailable", "ShareOfPractical"], mock["headers"]
    assert len(rows) == 6 and rows[-1][0] is None and rows[-1][1] == "Plant" and all(r[0] is not None for r in rows[:5]), rows
    plant = rows[-1]
    assert abs(plant[2] - r3["avail"]) < 0.005 and abs(plant[3] - r3["pract"]) < 0.005 and abs(plant[6] - r3["used"]) < 0.05, "the plant row is the guide's"
    assert abs(plant[7] - r3["of_avail"]) < 0.0005 and abs(plant[8] - r3["of_pract"]) < 0.0005, "the plant's shares are the guide's"
    assert abs(sum(r[2] for r in rows[:5]) - plant[2]) < 0.01 and abs(sum(r[6] for r in rows[:5]) - plant[6]) < 0.3, "the plant row adds the centers"
    assert abs(plant[3] - 0.80 * plant[2]) < 0.01, "practical capacity is 80% of the available hours"
    union = _find(mock, "UNION ALL")
    assert "'Plant'" in mock["sql_lines"][union + 1]
    _outline_row(d, mock, 5)
    _outline_lines(d, mock, union, union + 3, TEAL)
    _note(d, "Outlined in coral: the Plant row, the available, practical and routing hours of the five centers added, and the two shares of the plant. Outlined in "
             "teal: the lines of the last CTE that build it, a UNION ALL of one row that sums the centers (its WorkCenterID is NULL). The editor is scrolled to the "
             "last CTE and the final SELECT, and the grid one column to the right: WorkCenterID, the first column, is out of view. The CTEs above sum the "
             "completions, the routings and the calendar.", mock["bottom"] + 6)
    return d


# -- IF-18-3-B: where the plant's hours go -------------------------------------------------------------------------------------

def if_18_3_b() -> Diagram:
    """A clustered bar chart: for each of the five work centers its available hours, its practical capacity and the routing time
    used, with the Plant total in a label and a note on the Quality Assurance bar."""
    data = ctx.data()
    r3 = ctx.values("ch18.r3")
    g = _g()
    C = data.C
    rows = g.run(data, g.quote("Requirement 3, Q20"))
    centers, plant = rows[:5], rows[5]
    assert plant[1] == "Plant" and len(centers) == len(r3["centers"]) == 5
    names = [re.sub(r" Work Center$", "", r[1]) for r in centers]
    avail, pract = [r[2] for r in centers], [r[3] for r in centers]
    used = [r[4] + r[5] for r in centers]
    for c, a, p, u in zip(r3["centers"], avail, pract, used):
        assert abs(c["avail"] - a) < 0.01 and abs(c["pract"] - p) < 0.01 and abs(c["used"] - u) < 0.05, c
    qa = len(centers) - 1
    assert names[qa] == "Quality Assurance" and used[qa] / pract[qa] < 0.15 and all(used[i] / pract[i] > 0.5 for i in range(qa)), "only QA uses a small fraction"
    total_used = sum(used)
    assert abs(total_used - r3["used"]) < 0.05 and abs(plant[2] - r3["avail"]) < 0.005
    d = Diagram("Where the Plant's Hours Go")
    d.text(f"<b>Hours by work center, fiscal {C}</b>", 0, 0, 500, 20, size=SMALL)
    ticks = list(range(0, 60001, 10000))
    plot = kit.Plot(d, 0, 24, 860, 340, y_ticks=ticks, categories=names, y_fmt=lambda v: f"{v:,.0f}", y_title="Hours", left=66)
    fills = [("Available hours", RULE, GRAY, avail), ("Practical capacity (80% of available)", GRAY, GRAY, pract), ("Routing time used", BLUE, BLUE, used)]
    slot = plot.pw / len(names)
    bw, gap = slot * 0.22, 5
    group = 3 * bw + 2 * gap
    tops: dict[tuple[int, int], tuple[float, float]] = {}
    for i in range(len(names)):
        x0 = plot.px + slot * i + (slot - group) / 2
        for k, (_, fill, stroke, vals) in enumerate(fills):
            top, base = plot.y_of(vals[i]), plot.y_of(0)
            d.box("", x0 + k * (bw + gap), top, bw, base - top, fill=fill, stroke=stroke, stroke_width=1.5, rounded=False)
            tops[(i, k)] = (x0 + k * (bw + gap) + bw / 2, top)
    # the plant total, over the first group
    total = (f"<b>Plant total</b><br>Available {_money(plant[2])}<br>Practical {_money(plant[3])}<br>Routing time {plant[6]:,.1f}")
    _label(d, total, plot.px + 8, plot.y_of(58000), 150, 66, align="left")
    # the note on the Quality Assurance bar
    share = used[qa] / pract[qa]
    nx, ny = plot.px + plot.pw - 296, plot.y_of(36000)
    note = f"Quality Assurance used {share * 100:.1f}% of its practical hours: {used[qa]:,.0f} of {pract[qa]:,.0f}."
    nh = kit.height_for(note, 280, SMALL, pad=12)
    d.box(note, nx, ny, 280, nh, fill=AMBER_TINT, stroke=AMBER, size=SMALL, rounded=False, align="left", stroke_width=1.5)
    bx, by = tops[(qa, 2)]
    a1, a2 = d.anchor(bx, ny + nh), d.anchor(bx, by - 3)
    d.edge(a1, a2, end="blockThin", color=AMBER, width=1.5, meta=STRAIGHT)
    bottom = 24 + 340 + 8
    cx = 0.0
    for label, fill, stroke, _ in fills:
        d.box("", cx, bottom + 3, 14, 14, fill=fill, stroke=stroke, stroke_width=1.5, rounded=False)
        cx += 20
        width = text_width(label, SMALL, False) + 8
        d.text(label, cx, bottom, width, 20, size=SMALL, valign="middle")
        cx += width + 22
    _note(d, "Each group is one work center of the calendar; the three bars are its hours available in the year, its practical capacity at the share of the "
             "available hours the case assumes, and the routing time the year's completions needed (run hours per unit times the units, plus the setup of each "
             f"work order). The plant used {plant[8] * 100:.1f}% of its practical capacity; Quality Assurance's bar for routing time is the smallest by far.", bottom + 28)
    return d


# -- the model workbook ---------------------------------------------------------------------------------------------------------

@lru_cache(maxsize=None)
def _cached():
    """The model workbook with its saved values (the guides' own reader reads formulas; the figures draw what Excel shows)."""
    import openpyxl
    return openpyxl.load_workbook(XLSX, read_only=False, data_only=True)


def _agree(ws, ref: str, expected, tol: float = 0.006) -> None:
    """The model workbook holds in this cell the value the figure draws."""
    value = ws[ref].value
    if hasattr(value, "date"):
        value = value.date().isoformat()
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        assert isinstance(value, (int, float)) and abs(value - expected) < tol, f"{ws.title}!{ref} holds {value!r}, the figure draws {expected!r}"
    else:
        assert value == expected, f"{ws.title}!{ref} holds {value!r}, the figure draws {expected!r}"


def _acct(x: float | None, places: int = 2) -> str:
    """The workbook's number format #,##0.00: a blank cell empty."""
    return "" if x is None else f"{x:,.{places}f}"


def _book(sheet: str):
    """One worksheet of the model workbook with its saved values."""
    return _cached()[sheet]


def _gap(d: Diagram, x: float, y: float, w: float, text: str, h: float = 22) -> float:
    """A band that stands for the rows of a worksheet that the figure leaves out."""
    d.box(f"<i>{esc(text)}</i>", x, y, w, h, fill=WHITE, stroke=RULE, color=GRAY, size=SMALL, rounded=False, dashed=True)
    return y + h


def _wrapped(text: str, width: float, bold: bool = False, **kw) -> dict:
    """A cell that wraps its text over as many lines as the column needs (the cell's lines are kept in "lines")."""
    lines = _wrap(text, width - 16, bold=bold)
    html = "<br>".join(esc(line) for line in lines)
    return dict(html=f"<b>{html}</b>" if bold else html, lines=len(lines), **kw)


def _ws(d: Diagram, x: float, y: float, letters: list[str], widths: list[float], rows: list[dict], number_color: str = GRAY,
        show_letters: bool = True, row_h: float = 22) -> dict:
    """A worksheet with a letter row and rows of any height.  widths[0] is the row-number column; a row is {"n": row number,
    "cells": [cell, ...], "h": height}, its cells taken in order from the first column (a cell with span uses that many columns); a
    cell is a string or a dict: text, bold, italic, align, color, fill, stroke, html (a ready label), span, clip (False keeps a long
    text whole), valign.  A cell made by _wrapped() sets the row's height when the row has no "h".  Returns {(row index, column):
    (x, y, w, h)} and "bottom"."""
    xs = [x]
    for w in widths[:-1]:
        xs.append(xs[-1] + w)
    top = 22 if show_letters else 0
    if show_letters:
        for cx, w, letter in zip(xs, widths, ["", *letters]):
            d.box(esc(letter), cx, y, w, 22, fill=GRAY_TINT, stroke=RULE, color=GRAY, size=SMALL, rounded=False)
    geo: dict = {}
    ry = y + top
    for i, row in enumerate(rows):
        lines = max([c.get("lines", 1) for c in row["cells"] if isinstance(c, dict)] or [1])
        h = row.get("h", max(row_h, lines * 15 + 8))
        d.box(esc(row["n"]), xs[0], ry, widths[0], h, fill=GRAY_TINT, stroke=RULE, color=number_color, size=SMALL, rounded=False)
        cells = list(row["cells"])
        c = idx = 0
        while c < len(widths) - 1:
            cell = cells[idx] if idx < len(cells) else ""
            idx += 1
            if isinstance(cell, str):
                cell = {"text": cell}
            span = cell.get("span", 1)
            cx = xs[c + 1]
            cw = sum(widths[c + 1:c + 1 + span])
            text = cell.get("text", "")
            numeric = text[:1].isdigit() or text[:1] in "-$(" and (text[1:2].isdigit() or text == "-")
            align = cell.get("align", "right" if numeric else "left")
            bold = cell.get("bold", False)
            if "html" in cell:
                label = cell["html"]
            else:
                lead = len(text) - len(text.lstrip(" "))                      # leading spaces are the workbook's indent
                body = text.lstrip(" ")
                shown = _clip(body, cw - 3.5 * lead, bold=bold, margin=cell.get("margin", 26)) if cell.get("clip", not numeric) else body
                label = "&nbsp;" * lead + esc(shown)
                if bold:
                    label = f"<b>{label}</b>"
                if cell.get("italic"):
                    label = f"<i>{label}</i>"
            d.box(label, cx, ry, cw, h, fill=cell.get("fill", WHITE), stroke=cell.get("stroke", RULE), color=cell.get("color", INK), size=SMALL,
                  rounded=False, align=align, valign=cell.get("valign", "middle"))
            geo[(i, c)] = (cx, ry, cw, h)
            c += span
        ry += h
    geo["bottom"] = ry
    return geo


def _outline(d: Diagram, geo: dict, first: tuple[int, int], last: tuple[int, int], color: str = CORAL, inset: float = 0) -> None:
    """An outline around the cells from first to last (row index, column) of a worksheet drawn by _ws, inset so that outlines on
    neighboring cells do not touch."""
    x0, y0, _, _ = geo[first]
    x1, y1, w1, h1 = geo[last]
    left, top = max(0, x0 - 2 + inset), max(0, y0 - 2 + inset)
    right = min(860, x1 + w1 + 2 - inset)
    d.outline(left, top, right - left, y1 + h1 + 2 - inset - top, color, width=2.5, front=True)


def _formula_bar(d: Diagram, x: float, y: float, w: float, name: str, formula: str) -> None:
    """The Name Box, the fx label and the formula bar."""
    d.box(esc(name), x, y, 90, 28, fill=WHITE, stroke=RULE, size=SMALL, align="left", rounded=False)
    d.box("<i>fx</i>", x + 94, y, 30, 28, fill=WHITE, stroke=WHITE, color=GRAY, size=SMALL, rounded=False)
    d.box(esc(formula), x + 128, y, w - 128, 28, fill=WHITE, stroke=RULE, size=SMALL, align="left", rounded=False)


def _tabs(d: Diagram, y: float, active: str, names: list[str]) -> float:
    """The sheet tabs of the model workbook around the active one."""
    wb = _cached().sheetnames
    assert all(n in wb for n in names) and active in names, names
    xl.sheet_tabs(d, 0, y, names, active)
    return y + 26


# -- IF-18-4-A: the reconciliation block of the Costs sheet ---------------------------------------------------------------------

def if_18_4_a() -> Diagram:
    """Rows 22 to 36 of Costs: the routing time used, the products' conversion, the unused capacity and its crew part, the products
    plus the unused capacity beside the inputs to 1090 with a difference of zero, and the three lines of the setup-allocation bridge
    with its own zero difference."""
    r4 = ctx.values("ch18.r4")
    ws = _book("Costs")
    g = _g()
    cells = [("B23", ctx.values("ch18.r3")["total"]), ("B24", r4["conv"]), ("B25", r4["unused_h"]), ("B26", r4["unused"]), ("B27", r4["unused_crew"]),
             ("B28", r4["total"]), ("B29", r4["total"]), ("B30", 0.0), ("B31", r4["setup"]), ("B32", r4["alloc_setup"]), ("B33", r4["item_conv"]),
             ("B34", r4["bridge"]), ("B35", r4["total"]), ("B36", 0.0)]
    for ref, v in cells:
        _agree(ws, ref, v, tol=0.01)
    assert ws["A22"].value.startswith("Products and unused capacity") and ws["A30"].value == "Difference" and ws["A36"].value.strip().startswith("difference from the inputs")
    widths = [28, 648, 132]                                           # row numbers, then A and B (column A widened to show its labels)
    d = Diagram("The Reconciliation Block of the Costs Sheet")
    formula = g.clean(g.workbook()["Costs"]["B30"].value)
    assert formula == "=ProductsConversion+UnusedCost-Inputs1090", formula
    _formula_bar(d, 0, 0, 808, "B30", formula)
    sy = 40.0
    _ws(d, 0, sy, ["A", "B"], widths, [], show_letters=True)
    y = _gap(d, 0, sy + 22, sum(widths), "Rows 1 to 21: the committed cost, the time-driven rate and the Factory Overhead entries per standard hour")
    rows = []
    for r in range(22, 37):
        label = ws.cell(r, 1).value
        value = ws.cell(r, 2).value
        bold = r in (22, 28, 30, 36)
        cell_b = dict(text=_acct(value) if isinstance(value, (int, float)) else "", bold=bold, align="right", clip=False)
        rows.append(dict(n=str(r), cells=[dict(text=label, bold=r == 22, clip=True), cell_b]))
    geo = _ws(d, 0, y, [], widths, rows, show_letters=False)
    idx = {r: r - 22 for r in range(22, 37)}
    for r in (24, 26):
        _outline(d, geo, (idx[r], 1), (idx[r], 1), TEAL, inset=1.5)
    for r in (30, 36):
        _outline(d, geo, (idx[r], 1), (idx[r], 1), CORAL, inset=1.5)
    bottom = _gap(d, 0, geo["bottom"], sum(widths), "Rows 37 to 101: the family summary, the top fifth of the items and the Data Table of practical capacity")
    bottom = _tabs(d, bottom + 8, "Costs", ["Activities", "Costs", "Make or Buy"])
    _note(d, "Outlined in teal: the products' conversion and the unused capacity cost, the two parts of the inputs to 1090. Outlined in coral: the two differences, "
             "both zero: products plus unused capacity against the inputs (B30, whose formula is in the bar) and, after the setup-allocation bridge in rows "
             "31 to 34, the item costs' conversion plus the setup difference plus the unused capacity against the same inputs. Column A is widened here to show each "
             "label on one line; in the workbook a long label wraps within the column.", bottom + 4)
    return d


# -- Power BI pieces: the Product Costs project --------------------------------------------------------------------------------

def _page_ids() -> list[str]:
    return json.loads((REPORT / "pages" / "pages.json").read_text(encoding="utf-8"))["pageOrder"]


def _page(page: str) -> dict:
    return json.loads((REPORT / "pages" / page / "page.json").read_text(encoding="utf-8"))


def _page_name(page: str) -> str:
    return _page(page)["displayName"]


def _visual(page: str, name: str) -> dict:
    return json.loads((REPORT / "pages" / page / "visuals" / name / "visual.json").read_text(encoding="utf-8"))


def _title(page: str, name: str) -> str:
    """The title the model answer gives a visual (a doubled apostrophe in the file is one in the title)."""
    value = _visual(page, name)["visual"]["visualContainerObjects"]["title"][0]["properties"]["text"]["expr"]["Literal"]["Value"]
    return value.strip("'").replace("''", "'")


def _projections(page: str, name: str, role: str) -> list[str]:
    return [p["nativeQueryRef"] for p in _visual(page, name)["visual"]["query"]["queryState"][role]["projections"]]


def _query_tabs() -> list[str]:
    """The DAX query tabs of the project in their order (Tests, Checks)."""
    return json.loads((MODEL / ".pbi" / "daxQueries.json").read_text(encoding="utf-8"))["tabOrder"]


def _expected(label: str) -> float:
    """The Expected value of one row of the project's Checks tab (DAX query view), found by its label."""
    text = (MODEL / "DAXQueries" / "Checks.dax").read_text(encoding="utf-8")
    m = re.search(r'"Check", "' + re.escape(label) + r'", "Expected", ([-0-9.eE]+),', text)
    assert m, f"the Checks tab has no row {label!r}"
    return float(m.group(1))


VIEWS = [("Standard", "MarginStandard"), ("Absorption", "MarginAbsorption"), ("Time-driven", "MarginTimeDriven")]


@lru_cache(maxsize=None)
def _item_costs() -> tuple[dict, ...]:
    """The rows of the workbook's ItemCosts Table with their saved values: what the Product Costs model loads."""
    ws = _book("ItemCosts")
    head = [c.value for c in ws[1]]
    return tuple(dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None)


def _frame(d: Diagram, x: float, y: float, w: float, h: float, title: str | None, selected: bool = False,
           lines: int | None = None) -> tuple[float, float, float, float]:
    """A visual's container.  The title may take two lines, as a long title does in Power BI.  Returns the inner area."""
    d.box("", x, y, w, h, fill=WHITE, stroke=BLUE if selected else RULE, stroke_width=2 if selected else 1, rounded=False)
    if not title:
        return x + 6, y + 6, w - 12, h - 12
    label = f"<b>{esc(title)}</b>"
    lines = lines or kit.lines_for(label, w - 12, SMALL, bold=True)
    th = 16 * lines + 4
    d.text(label, x + 6, y + 4, w - 12, th, size=SMALL)
    return x + 6, y + 4 + th + 2, w - 12, h - th - 12


def _is_num(value: str) -> bool:
    return value[:1].isdigit() or value[:1] in "-−" and value[1:2].isdigit()


def _heads(head: str, width: float, bold: bool = True) -> str:
    """A column heading as HTML: an explicit <br> breaks it, and a heading too wide for its column is cut with an ellipsis."""
    if "<br>" in head:
        return "<br>".join(esc(p) for p in head.split("<br>"))
    return esc(_clip(head, width, margin=8, bold=bold))


def _grid_line(d: Diagram, x: float, y: float, w: float, h: float) -> None:
    """A one-pixel gridline (a box with a stroke would draw two)."""
    d.vertex("", f"rounded=0;html=1;fillColor={RULE};strokeColor=none;", x, y, w, h)


def _scrollbar(d: Diagram, right: float, top: float, height: float, shown: int, total: int) -> None:
    """A vertical scroll bar at a visual's right edge: the thumb is the share of the rows in view."""
    d.box("", right - 10, top, 6, height, fill=GRAY_TINT, stroke=GRAY_TINT, rounded=False)
    d.box("", right - 10, top, 6, max(height * shown / total, 12), fill=GRAY, stroke=GRAY, rounded=False)


def _table_visual(d: Diagram, x: float, y: float, title: str, headers: list[str], widths: list[float], rows: list[list[str]],
                  aligns: list[str] | None = None, title_lines: int | None = None, head_lines: int = 1, row_h: float = 22,
                  total: list[str] | None = None) -> dict:
    """A table visual: a title, a header row, the rows (banded) and an optional bold total row.  aligns names each column's side.
    Returns the cell geometry keyed (row, column) with row -1 for the headers, and the visual's bottom and right edges."""
    w = sum(widths) + 12
    label = f"<b>{esc(title)}</b>"
    th = 16 * (title_lines or kit.lines_for(label, w - 12, SMALL, bold=True)) + 4
    hh = row_h if head_lines == 1 else 15 * head_lines + 8
    n = len(rows) + (1 if total else 0)
    h = th + 12 + hh + n * row_h
    ix, iy, _, _ = _frame(d, x, y, w, h, title, lines=title_lines)
    xs = [ix]
    for wd in widths[:-1]:
        xs.append(xs[-1] + wd)
    aligns = aligns or ["right" if any(_is_num(r[c]) for r in rows) else "left" for c in range(len(headers))]
    geometry = {}
    for c, (head, cx, wd) in enumerate(zip(headers, xs, widths)):
        d.box(f"<b>{_heads(head, wd)}</b>", cx, iy, wd, hh, fill=WHITE, stroke=WHITE, size=SMALL, align=aligns[c], rounded=False)
        geometry[(-1, c)] = (cx, iy, wd, hh)
    d.box("", ix, iy + hh - 1, sum(widths), 1, fill=INK, stroke=INK, rounded=False)
    for r, row in enumerate(rows + ([total] if total else [])):
        ry = iy + hh + row_h * r
        is_total = total is not None and r == len(rows)
        if is_total:
            d.box("", ix, ry, sum(widths), 1, fill=INK, stroke=INK, rounded=False)
        for c, (value, cx, wd) in enumerate(zip(row, xs, widths)):
            plain = r % 2 == 0 or is_total
            shown = esc(_clip(value, wd, margin=8 if aligns[c] == "right" else 16, bold=is_total))
            d.box(f"<b>{shown}</b>" if is_total else shown, cx, ry, wd, row_h,
                  fill=WHITE if plain else GRAY_TINT, stroke=WHITE if plain else GRAY_TINT, size=SMALL, align=aligns[c], rounded=False)
            geometry[(r, c)] = (cx, ry, wd, row_h)
    return dict(geometry=geometry, bottom=y + h, right=x + w, left=ix, width=sum(widths))


def _matrix(d: Diagram, x: float, y: float, title: str | None, row_w: float, groups: list[tuple[str | None, list[tuple[str, float]]]],
            row_heads: list[str], cells: list[list[str]], row_h: float = 22, head_lines: int = 1, bold_rows: set[int] | None = None,
            levels: list[int] | None = None, title_lines: int | None = None, expanded: set[int] | None = None) -> dict:
    """A matrix visual: a title, group headings (the view over its measures), the column headings (<br> where one breaks), the row
    headings (a level indents a row under its parent, an expanded parent starts with a minus sign) and the values.  groups are
    (heading, [(column heading, width), ...]); bold_rows are drawn bold.  Returns the cell geometry keyed (row, column) over the
    flattened value columns."""
    cols = [c for _, hs in groups for c in hs]
    inner = row_w + sum(wd for _, wd in cols)
    w = inner + 12
    gh = 22
    hh = row_h if head_lines == 1 else 15 * head_lines + 8
    label = f"<b>{esc(title)}</b>"
    th = 16 * (title_lines or kit.lines_for(label, w - 12, SMALL, bold=True)) + 4
    h = th + 12 + gh + hh + len(row_heads) * row_h
    ix, iy, _, _ = _frame(d, x, y, w, h, title, lines=title_lines)
    bold_rows, levels = bold_rows or set(), levels or [0] * len(row_heads)
    cy = iy
    gx = ix + row_w
    for g, hs in groups:
        gw = sum(wd for _, wd in hs)
        d.text(f"<b>{esc(g)}</b>", gx, cy, gw, gh, size=SMALL, align="center", valign="middle")
        gx += gw
    cy += gh
    xs, cx = [], ix + row_w
    for head, wd in cols:
        xs.append(cx)
        d.text(f"<b>{_heads(head, wd)}</b>", cx, cy, wd - 4, hh, size=SMALL, align="right", valign="bottom")
        cx += wd
    rule_y = cy + hh - 1
    cy += hh
    geometry = {}
    for r, head in enumerate(row_heads):
        ry = cy + r * row_h
        bold = r in bold_rows
        text = "&nbsp;" * (5 * levels[r]) + esc(("– " if r in (expanded or set()) else "") + head)
        d.text(f"<b>{text}</b>" if bold else text, ix + 4, ry, row_w - 4, row_h, size=SMALL, align="left", valign="middle")
        for c, ((_, wd), vx) in enumerate(zip(cols, xs)):
            value = cells[r][c]
            assert not value or text_width(value, SMALL, bold) + 6 <= wd, f"{title}: {value!r} does not fit its {wd:.0f} px column"
            d.box(f"<b>{esc(value)}</b>" if bold else esc(value), vx, ry, wd, row_h, fill=WHITE, stroke=WHITE, size=SMALL, align="right", rounded=False)
            geometry[(r, c)] = (vx, ry, wd, row_h)
    d.box("", ix, rule_y, inner, 1, fill=INK, stroke=INK, rounded=False)
    return dict(geometry=geometry, bottom=y + h, right=x + w, left=ix, width=inner, top_rows=cy, row_h=row_h, head_top=iy)


def _cards_visual(d: Diagram, x: float, y: float, h: float, items: list[tuple[str, str]], widths: list[float]) -> dict:
    """A card visual, one card for each field: the value above the field's name.  Returns the cards' geometry."""
    w = sum(widths) + 6 * (len(items) - 1) + 12
    d.box("", x, y, w, h, fill=WHITE, stroke=RULE, rounded=False)
    geo, cx = [], x + 6
    for (label, value), cw in zip(items, widths):
        d.box("", cx, y + 6, cw, h - 12, fill=WHITE, stroke=RULE, rounded=True)
        d.text(f"<b>{esc(value)}</b>", cx + 2, y + 9, cw - 4, 22, size=14, align="center")
        d.text(esc(label), cx + 2, y + 31, cw - 4, 18, size=SMALL, color=GRAY, align="center")
        geo.append((cx, y + 6, cw, h - 12))
        cx += cw + 6
    return dict(bottom=y + h, right=x + w, cards=geo)


def _items_chart(d: Diagram, x: float, y: float, w: float, h: float, title: str, cats: list[str], values: list[float], total: int,
                 title_lines: int = 2) -> dict:
    """A clustered bar chart sorted as given (largest first), as many bars as fit, with a scroll bar for the rest."""
    ix, iy, iw, ih = _frame(d, x, y, w, h, title, lines=title_lines)
    label_w, axis_h, slot = 106, 22, 20
    n = int((ih - axis_h - 4) // slot)
    assert n <= len(cats), (n, len(cats))
    left, top = ix + label_w, iy + 2
    plot_w = iw - label_w - 22
    ticks = pbi.nice_ticks(0, max(values[:n]), 4)
    x_of = lambda v: left + v * plot_w / ticks[-1]                                                        # noqa: E731
    for t in ticks:
        _grid_line(d, x_of(t), top, 1, n * slot)
        d.text(kit.compact(t), x_of(t) - 26, top + n * slot + 2, 52, 18, size=SMALL, color=GRAY, align="center")
    for i in range(n):
        by = top + i * slot
        d.text(esc(cats[i]), ix, by, label_w - 4, slot, size=SMALL, align="right", valign="middle")
        d.box("", left, by + 3, max(x_of(values[i]) - left, 1), slot - 6, fill=BLUE, stroke=BLUE, rounded=False)
    _scrollbar(d, x + w, top, n * slot, n, total)
    return dict(n=n, left=left, top=top, slot=slot)


def _page_tabs(active: str, d: Diagram, y: float) -> None:
    """The page tabs under the canvas: the pages around the active one, as the report orders them."""
    order = _page_ids()
    i = order.index(active)
    shown = order[max(i - 1, 0):i + 2]
    pbi.page_tabs(d, 4, y, [_page_name(p) for p in shown], _page_name(active))


# -- IF-18-5-A: the Families page ------------------------------------------------------------------------------------------------

def if_18_5_a() -> Diagram:
    """The Families page with the Time-driven view selected: the slicer of the three views, the cards (revenue, product cost, margin and
    margin percent of that view), the matrix of item groups and families by view, the bar chart of items ranked by margin and the table of
    totals by view."""
    g5, r5 = ctx.values("ch18.g5"), ctx.values("ch18.r5")
    m = _model()
    e = r5["e"]
    page = "families"
    titles = dict(slicer=_title(page, "viewSlicer"), matrix=_title(page, "familyMatrix"), chart=_title(page, "itemsChart"), totals=_title(page, "viewTotals"))
    assert titles == dict(slicer=g5["t_slicer"], matrix=g5["t_matrix"], chart=g5["t_chart"], totals=g5["t_totals"]), titles
    cards = _projections(page, "viewCards", "Data")
    assert cards == ["Revenue", "Product Cost", "Margin", "Margin %"], cards
    mt = _g().measure_table()
    assert all(c in mt for c in cards) and all(f'"{v}"' in mt["Product Cost"] for v, _ in VIEWS), "the measures the page uses are the model's"
    assert [p["target"] for p in _page(page)["visualInteractions"]] == ["familyMatrix", "viewTotals"] and \
        all(p["source"] == "viewSlicer" and p["type"] == "NoFilter" for p in _page(page)["visualInteractions"]), "the slicer leaves the matrix and the totals table alone"
    assert _projections(page, "familyMatrix", "Rows") == ["ItemGroup", "Family"] and _projections(page, "familyMatrix", "Columns") == ["View"]
    assert _projections(page, "familyMatrix", "Values") == ["Margin %", "Cost per Unit"], "the matrix shows Margin % and Cost per Unit"
    assert _projections(page, "viewTotals", "Values") == ["View", "Revenue", "Product Cost", "Margin", "Margin %"]
    # the items sold in the year (the page is filtered to the year and to a family that is not blank), from the workbook's Table
    sold = [r for r in _item_costs() if (r["Revenue"] or 0) > 0]
    assert len(sold) == e["n_items"], (len(sold), e["n_items"])
    rev = sum(r["Revenue"] for r in sold)
    margin = {v: sum(r[col] for r in sold) for v, col in VIEWS}
    exact = dict(zip((v for v, _ in VIEWS), (e["exact"]["std"], e["exact"]["absorp"], e["exact"]["td"])))
    assert abs(rev - e["rev"]) < 0.01 and all(abs(margin[v] - exact[v]) < 0.01 for v, _ in VIEWS), "the workbook's Table is the note's"
    td_cost = rev - margin["Time-driven"]
    assert abs(td_cost - _expected("2026 product cost, Time-driven")) < 0.01 and abs(margin["Time-driven"] / rev - _expected("2026 Margin %, Time-driven")) < 1e-6
    # the matrix: manufactured items by group and family
    made = [r for r in sold if r["SupplyMode"] == "Manufactured"]

    def agg(items: list[dict]) -> list[str]:
        units, rv = sum(r["Units"] for r in items), sum(r["Revenue"] for r in items)
        out = []
        for _, col in VIEWS:
            mg = sum(r[col] for r in items)
            out += [f"{mg / rv * 100:.1f}%", f"{(rv - mg) / units:,.2f}"]
        return out
    heads, cells, levels, bold, expanded = [], [], [], set(), set()
    for grp in sorted({r["ItemGroup"] for r in made}):
        items = [r for r in made if r["ItemGroup"] == grp]
        bold.add(len(heads))
        expanded.add(len(heads))
        heads.append(grp)
        cells.append(agg(items))
        levels.append(0)
        for fam in sorted({r["Family"] for r in items}):
            heads.append(fam)
            cells.append(agg([r for r in items if r["Family"] == fam]))
            levels.append(1)
    assert sum(1 for lv in levels if lv == 1) == len(m["fam"]) == 16 and len(bold) == len(m["groups"]) == 3
    for fam, v in m["fam"].items():                                                 # each family against the note's own per-unit costs
        row = cells[heads.index(fam)]
        for k, key in enumerate(("std", "absorp", "td")):
            assert abs(float(row[2 * k + 1].replace(",", "")) - v["unit"][key]) < 0.01, (fam, key)
            assert abs(float(row[2 * k].rstrip("%")) / 100 - (1 - v["unit"][key] / v["unit"]["price"])) < 0.0006, (fam, key)
    assert abs(float(cells[heads.index("FUR-BKC")][4].rstrip("%")) / 100 - _expected("FUR-BKC (manufactured) Margin %, Time-driven")) < 0.0006
    assert abs(float(cells[heads.index("Furniture")][0].rstrip("%")) / 100 - _expected("Furniture (manufactured) Margin %, Standard")) < 0.0006
    # the items ranked by margin under the view selected
    ranked = sorted(sold, key=lambda r: r["MarginTimeDriven"], reverse=True)
    cat, val = [r["ItemCode"] for r in ranked], [r["MarginTimeDriven"] for r in ranked]
    totals = [[v, f"{rev:,.2f}", f"{rev - margin[v]:,.2f}", f"{margin[v]:,.2f}", f"{margin[v] / rev * 100:.1f}%"] for v, _ in VIEWS]
    card_items = [("Revenue", f"{rev:,.2f}"), ("Product Cost", f"{td_cost:,.2f}"), ("Margin", f"{margin['Time-driven']:,.2f}"),
                  ("Margin %", f"{margin['Time-driven'] / rev * 100:.1f}%")]
    pad = 16
    groups = [(v, [("Margin %", 64), ("Cost per<br>Unit", 60)]) for v, _ in VIEWS]

    def draw(d: Diagram) -> tuple[dict, dict, dict, dict, float]:
        right = pbi.slicer_tiles(d, pad, pad, titles["slicer"], [v for v, _ in VIEWS], "Time-driven", tile_w=92)
        cv = _cards_visual(d, right + 12, pad, 56, card_items, [132, 132, 132, 84])
        y2 = pad + 56 + 12
        mx = _matrix(d, pad, y2, titles["matrix"], 90, groups, heads, cells, row_h=20, head_lines=2, bold_rows=bold, levels=levels, title_lines=2,
                     expanded=expanded)
        cx = mx["right"] + 12
        ch = _items_chart(d, cx, y2, 860 - pad - cx, mx["bottom"] - y2, titles["chart"], cat, val, len(cat))
        tb = _table_visual(d, pad, mx["bottom"] + 12, titles["totals"], ["View", "Revenue", "Product Cost", "Margin", "Margin %"], [156, 172, 172, 172, 144], totals,
                           aligns=["left", "right", "right", "right", "right"], title_lines=1)
        return cv, mx, ch, tb, tb["bottom"] + pad

    canvas_h = draw(Diagram("probe"))[4]
    d = Diagram("The Families Page")
    pbi.canvas(d, 0, 0, 860, canvas_h)
    cv, mx, ch, tb, _ = draw(d)
    geo = mx["geometry"]
    x0, y0, _, _ = geo[(0, 4)]
    x1, y1, w1, h1 = geo[(len(heads) - 1, 5)]
    d.outline(x0 - 1, mx["head_top"] - 1, x1 + w1 - x0 + 2, y1 + h1 - mx["head_top"] + 2, TEAL, width=2.5, front=True)
    tx, ty, tw, th_ = tb["geometry"][(2, 0)]
    right = tb["geometry"][(2, 4)]
    d.outline(tx - 1, ty - 1, right[0] + right[2] - tx + 2, th_ + 2, TEAL, width=2.5, front=True)
    _page_tabs(page, d, canvas_h + 4)
    _note(d, f"Report view, page {_page_name(page)}, with the slicer saved on {VIEWS[2][0]} and the page filtered to fiscal {ctx.data().C} and to items with a family. Drawn here in "
             f"a different arrangement: the totals table is below the matrix in place of beside it. The slicer filters the cards and the bar chart only (Edit interactions "
             f"set to None on the matrix and the totals table), so those two show every view at once; outlined in teal: the {VIEWS[2][0]} columns and row, which the cards and "
             f"the chart follow. The matrix has the manufactured items only; the other items are in the totals and the chart. The chart shows the {ch['n']} items of the "
             f"{len(cat)} with the largest margin, and scrolls.", canvas_h + 34)
    return d


# -- IF-18-5-B: the Tests tab in DAX query view ---------------------------------------------------------------------------------

def _dax_editor(d: Diagram, x: float, y: float, w: float, rows: list[tuple[int | None, str]], tab: str, updates: int,
                selected: range | None = None) -> dict:
    """DAX query view's Run button, query tabs and editor.  rows are (line number, text): the editor shows long lines cut at its edge,
    and a row with no number is a gap whose text says which lines the picture leaves out.  selected holds the line numbers drawn
    selected.  Returns the editor's top and bottom."""
    xl.button(d, x, y, 70, "▶ Run", primary=True)
    xl.button(d, x + 78, y, 214, f"Update model with changes ({updates})")
    y += 32
    cx = x
    for name in _query_tabs():
        on = name == tab
        width = text_width(esc(name), SMALL, on) + 24
        d.box(f"<b>{esc(name)}</b>" if on else esc(name), cx, y, width, 24, fill=WHITE, stroke=BLUE if on else RULE,
              stroke_width=2 if on else 1, color=BLUE if on else INK, size=SMALL, rounded=False)
        cx += width + 2
    d.box("<b>+</b>", cx, y, 26, 24, fill=WHITE, stroke=RULE, size=SMALL, rounded=False)
    y += 24
    gutter, line_h = 34, 20
    height = len(rows) * line_h + 8
    d.box("", x, y, gutter, height, fill=GRAY_TINT, stroke=RULE, rounded=False)
    d.box("", x + gutter, y, w - gutter, height, fill=WHITE, stroke=RULE, rounded=False)
    ys = {}
    for i, (number, _) in enumerate(rows):
        if number is not None:
            ys[number] = y + 4 + i * line_h
    if selected:
        first, last = min(selected), max(selected)
        d.box("", x + gutter + 1, ys[first], w - gutter - 2, ys[last] + line_h - ys[first], fill=BLUE_TINT, stroke=BLUE_TINT, rounded=False)
    for i, (number, text) in enumerate(rows):
        ly = y + 4 + i * line_h
        if number is None:
            d.box(f"<i>{esc(text)}</i>", x + gutter + 4, ly + 1, w - gutter - 8, line_h - 2, fill=GRAY_TINT, stroke=GRAY_TINT, color=GRAY,
                  size=SMALL, align="left", rounded=False)
            continue
        d.text(str(number), x, ly, gutter - 4, line_h, size=SMALL, color=GRAY, align="right", valign="middle")
        label = pbi.highlight_dax(text)
        assert text_width(label, SMALL, False) <= w - gutter - 12, f"DAX line too wide: {text}"
        d.text(label or " ", x + gutter + 4, ly, w - gutter - 8, line_h, size=SMALL, valign="middle")
    return dict(top=y, bottom=y + height, ys=ys, left=x + gutter, width=w - gutter, line_h=line_h)


def _dax_num(x: float) -> str:
    """A rounded value as DAX query view's grid writes it: no thousands separator, no trailing zero."""
    text = f"{x:.2f}".rstrip("0").rstrip(".")
    return "0" if text in ("", "-0") else text


def if_18_5_b() -> Diagram:
    """DAX query view with the tab Tests: Test 1 selected in the editor and its three-row result beneath, one row for each view with the
    model's revenue and margin beside the workbook's and a difference of zero."""
    g5, r5 = ctx.values("ch18.g5"), ctx.values("ch18.r5")
    e = r5["e"]
    lines = (MODEL / "DAXQueries" / "Tests.dax").read_text(encoding="utf-8").rstrip("\n").split("\n")
    end = next(i for i, text in enumerate(lines) if text.startswith("// Test 2:"))
    assert lines[0].startswith("// Test 1:") and lines[end - 1] == "" and lines[end - 2] == "ORDER BY 'Costing View'[Order]", lines[end - 2]
    test1 = lines[:end - 1]
    assert "\n".join(test1) == _g().test_block(lambda ok, msg: None, 1), "Test 1 is the project's"
    n_tests = g5["n_tests"]
    assert n_tests == sum(1 for text in lines if text.startswith("// Test ")) == 6
    sold = [r for r in _item_costs() if (r["Revenue"] or 0) > 0]
    wb_rev = round(sum(r["Revenue"] for r in sold), 2)
    rows = []
    for (view, col), key in zip(VIEWS, ("std", "absorp", "td")):
        wb_margin = round(sum(r[col] for r in sold), 2)
        margin = round(e["exact"][key], 2)
        rev = round(e["rev"], 2)
        assert rev == wb_rev and margin == wb_margin, (view, rev, wb_rev, margin, wb_margin)
        rows.append([view, _dax_num(rev), _dax_num(wb_rev), _dax_num(margin), _dax_num(wb_margin), _dax_num(margin - wb_margin)])
    headers = ["Costing View[View]", "[Revenue]", "[Workbook Revenue]", "[Margin]", "[Workbook Margin]", "[Difference]"]
    # the grid is scrolled one column to the right: Costing View[Order] (1, 2, 3) is out of view, the status still counts all seven
    widths = [max(text_width(f"<b>{esc(h)}</b>", SMALL, True), max(text_width(esc(r[c]), SMALL, False) for r in rows)) + 24 for c, h in enumerate(headers)]
    win_h = 96 + 10 + 32 + 24 + (len(test1) + 3) * 20 + 8 + 14 + 24 + pbi.ROW + 2 + len(rows) * pbi.ROW + 44 + 12
    d = Diagram("The Tests Tab and Test 1")
    win = pbi.window(d, win_h, view="DAX query", tab="Home", tabs=pbi.QUERY_TABS, panes=None,
                     buttons=["Format", "Comment", "Uncomment", "Find", "Replace", "Command palette"], file="Product Costs")
    x, w = win.canvas_x + 8, win.canvas_w - 16
    assert sum(widths) <= w, (sum(widths), w)
    editor_rows: list[tuple[int | None, str]] = [(i + 1, text) for i, text in enumerate(test1)]
    editor_rows.append((end, ""))
    editor_rows.append((end + 1, lines[end]))
    editor_rows.append((None, f"Lines {end + 2} to {len(lines)}: the rest of Test 2, and Tests 3 to {n_tests}"))
    ed = _dax_editor(d, x, win.top + 10, w, editor_rows, "Tests", 0, selected=range(1, len(test1) + 1))
    gy = ed["bottom"] + 14
    d.text("<b>Results</b>", x, gy, 120, 22, size=SMALL)
    gy += 24
    geo = {}
    cx = x
    for c, (head, wd) in enumerate(zip(headers, widths)):
        d.box(f"<b>{esc(head)}</b>", cx, gy, wd, pbi.ROW + 2, fill=GRAY_TINT, stroke=RULE, size=SMALL, align="left", rounded=False)
        geo[(-1, c)] = (cx, gy, wd, pbi.ROW + 2)
        for r, row in enumerate(rows):
            d.box(esc(row[c]), cx, gy + pbi.ROW + 2 + r * pbi.ROW, wd, pbi.ROW, fill=WHITE, stroke=RULE, size=SMALL, align="left", rounded=False)
            geo[(r, c)] = (cx, gy + pbi.ROW + 2 + r * pbi.ROW, wd, pbi.ROW)
        cx += wd
    grid_bottom = gy + pbi.ROW + 2 + len(rows) * pbi.ROW
    d.box("", x, grid_bottom + 4, w, 6, fill=GRAY_TINT, stroke=GRAY_TINT, rounded=False)                  # the horizontal scroll bar
    d.box("", x + w * 0.12, grid_bottom + 4, w * 0.88, 6, fill=GRAY, stroke=GRAY, rounded=False)
    d.text(esc(f"{len(headers) + 1} columns, {len(rows)} rows"), x, grid_bottom + 14, 220, 20, size=SMALL, color=GRAY)
    for c0, c1, color in ((3, 4, TEAL), (5, 5, CORAL)):
        x0, y0, _, _ = geo[(-1, c0)]
        xa, ya, wa, ha = geo[(len(rows) - 1, c1)]
        d.outline(x0 - 1, y0 - 1, xa + wa - x0 + 2, ya + ha - y0 + 2, color, width=2.5, front=True)
    _note(d, f"The tab Tests of Product Costs with Test 1 selected and run: the three views of fiscal 2026 for the finished goods, the model's revenue and margin beside the "
             f"workbook's. Outlined in teal: the margin and the workbook's margin; in coral: the difference, zero for each view. The grid is scrolled one column to the right "
             f"(Costing View[Order] is out of view). The tab holds {n_tests} tests; the tab Checks holds a row for each value of Requirement 5 ({g5['n_checks']}) and of Milestone 2 ({g5['n_milestone']}).", win.bottom + 10)
    return d


# -- the second half: Requirements 6 to 10 ---------------------------------------------------------------------------------------
# These figures draw the sheets Make or Buy, Capacity, Growth 40%, Prices, Standards and Memo of the model workbook.  Each value
# is read from the workbook as Excel shows it (its own number format) after the notes' model has been asserted equal to the cell.

def _excel_format(v: float, fmt: str) -> str:
    """A number as Excel shows it in the number formats the model workbook uses (#,##0.00; 0.0%; \\+#,##0;\\-#,##0; ...)."""
    sections = fmt.split(";")
    negative = v < 0
    if negative:
        sec, v = sections[1 if len(sections) > 1 else 0], -v
    elif v == 0 and len(sections) > 2:
        sec = sections[2]
    else:
        sec = sections[0]
    sec = sec.replace("\\", "")
    m = re.search(r"([#0,]+(?:\.0+)?)(%?)", sec)
    assert m, f"number format {fmt!r} is not one the figure knows"
    body, pct = m.groups()
    places = len(body.split(".")[1]) if "." in body else 0
    x = xl.xround(v * 100 if pct else v, places)
    text = f"{x:,.{places}f}" if "," in body else f"{x:.{places}f}"
    out = sec[:m.start()] + text + pct + sec[m.end():]
    return "-" + out if negative and len(sections) == 1 else out


def _show(ws, ref: str) -> str:
    """What the cell shows: text as it is, a number in its number format, TRUE and FALSE, an empty cell empty."""
    cell = ws[ref]
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return str(int(v)) if cell.number_format == "General" and float(v).is_integer() else _excel_format(float(v), cell.number_format)
    return str(v)


def _num(text: str, **kw) -> dict:
    """A numeric cell: right-aligned and never clipped (a sign in front of the digits would make _ws read it as text)."""
    return dict(text=text, align="right", clip=False, **kw)


def _head(text: str, width: float, **kw) -> dict:
    """A column heading: bold, wrapped over as many lines as the column needs, sitting at the foot of its row."""
    return _wrapped(text, width, bold=True, valign="bottom", **kw)


def _row_of(ws, text: str, col: int = 1) -> int:
    """The row of a worksheet whose cell in `col` starts with `text` (exactly one)."""
    hits = [r for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, col).value, str) and ws.cell(r, col).value.startswith(text)]
    assert len(hits) == 1, (ws.title, text, hits)
    return hits[0]


def _formula_of(sheet: str, ref: str, width: float = 732) -> str:
    """A cell's formula as the formula bar reads it, cut at the bar's edge as a long formula is."""
    return _clip(_g().clean(_g().workbook()[sheet][ref].value), width, size=SMALL, margin=36).replace(" ", " ")   # no break: one line


def _key(d: Diagram, x: float, y: float, items: list[tuple[str, str, str]], w: float = 860) -> float:
    """A key of swatches: (label, fill, stroke), the swatch a 22 x 14 box with a 2 px stroke.  Wraps to a new row at the frame's edge.
    Returns the bottom y."""
    cx, cy = x, y
    for label, fill, stroke in items:
        width = text_width(esc(label), SMALL, False) + 8
        if cx > x and cx + 28 + width > x + w:
            cx, cy = x, cy + 22
        d.box("", cx, cy + 3, 22, 14, fill=fill, stroke=stroke, stroke_width=2, rounded=False)
        d.text(esc(label), cx + 28, cy, width, 20, size=SMALL, valign="middle")
        cx += 28 + width + 18
    return cy + 22


# -- IF-18-6-A: the Make or Buy table ------------------------------------------------------------------------------------------

def if_18_6_a() -> Diagram:
    """The six quoted families of the Make or Buy sheet (rows 10 to 15): the delivered quote, the relevant cost with the crew kept and
    reduced, the annual effect of buying under the absorption view, the crew kept and the crew reduced, and the decision; the families
    to buy are outlined."""
    m, n = _model(), _n()
    ws = _book("Make or Buy")
    quoted, buy = m["quoted"], m["buy"]
    first = 10
    assert ws["A9"].value == "Family" and [ws.cell(first + i, 1).value for i in range(len(quoted))] == [q["name"] for q in quoted]
    decisions = []
    for i, q in enumerate(quoted):
        r = first + i
        for col, key in (("G", "delivered"), ("I", "kept"), ("J", "reduced"), ("L", "absorp"), ("M", "a_kept"), ("N", "a_reduced")):
            _agree(ws, f"{col}{r}", q[key], tol=0.05)
        # the decision cell also asks whether the two years the quote is fixed cover the family's one-time costs (MOD18-04)
        decision = ("Buy" if q["a_kept"] > 0 else
                    ("Buy only with the crew reduced" if q["two_year"] > 0
                     else "Make: the two years do not cover the one-time costs") if q["a_reduced"] > 0 else "Make")
        assert ws[f"O{r}"].value == decision, (r, ws[f"O{r}"].value, decision)
        decisions.append(decision)
    to_buy = [i for i, dec in enumerate(decisions) if dec != "Make"]
    assert [quoted[i]["name"] for i in to_buy] == [b["name"] for b in buy] and 0 < len(to_buy) < len(quoted)
    assert all(q["a_kept"] <= 0 for q in quoted), "with the crew kept, no family is worth buying"
    absorp_buy = sum(1 for q in quoted if q["absorp"] > 0)
    assert 0 < absorp_buy and absorp_buy > len(buy), "the absorption view would buy more families than the relevant costs do"
    letters = ["A", "G", "I", "J", "L", "M", "N", "O"]
    widths = [28, 80, 88, 88, 88, 96, 96, 96, 200]
    d = Diagram("The Make or Buy Table")
    formula = _g().clean(_g().workbook()["Make or Buy"]["O10"].value)
    assert formula == ('=IF(M10>0,"Buy",IF(N10>0,IF(Y10>0,"Buy only with the crew reduced",'
                       '"Make: the two years do not cover the one-time costs"),"Make"))'), formula
    _formula_bar(d, 0, 0, 860, "O10", formula)
    sy = 40.0
    _ws(d, 0, sy, letters, widths, [], show_letters=True)
    y = _gap(d, 0, sy + 22, sum(widths), "Rows 1 to 8: the title, the method, and the two sensitivity inputs (B5 and B6)")
    head = dict(n="9", cells=[_head(ws[f"{c}9"].value, widths[k + 1]) for k, c in enumerate(letters)])
    rows = [head]
    for i in range(len(quoted)):
        r = first + i
        rows.append(dict(n=str(r), cells=[dict(text=ws[f"A{r}"].value, clip=False)] + [_num(_show(ws, f"{c}{r}")) for c in "GIJLMN"]
                         + [dict(text=decisions[i])]))
    geo = _ws(d, 0, y, [], widths, rows, show_letters=False)
    for i in to_buy:
        _outline(d, geo, (1 + i, 0), (1 + i, 7), TEAL, inset=1.5)
    bottom = _gap(d, 0, geo["bottom"], sum(widths), "Rows 16 to 56: the Goal Seek block, the families to buy, the two-way Data Table and the materials the families draw")
    bottom = _tabs(d, bottom + 8, "Make or Buy", ["Costs", "Make or Buy", "Capacity"])
    _note(d, f"Outlined in teal: the {n.word(len(buy))} families whose decision is to buy only with the crew reduced. A positive effect means buying saves money; "
             f"with the crew kept (column M) every quoted family costs money to buy, and the CFO's absorption view (column L) is positive for {n.word(absorp_buy)} of "
             f"the {n.word(len(quoted))}. Columns B to F, H, K and P to Z are hidden here, and the columns are narrower than the workbook's; the formula bar "
             f"shows the rule of the decision cells, which also asks (column Y) whether the two years the quote is fixed cover the family's one-time costs.", bottom + 4)
    return d


# -- IF-18-6-B: the two-way Data Table ----------------------------------------------------------------------------------------

def if_18_6_b() -> Diagram:
    """The sensitivity inputs (B5 and B6), the corner cell, the quote adjustments across, the crew shares down and the nine results of the
    two-way Data Table (A41 to D44), with the cell for all the crew time removed at the quote as it is outlined."""
    m = _model()
    ws = _book("Make or Buy")
    buy, two_way = m["buy"], m["two_way"]
    adjustments, shares = (-0.10, 0.0, 0.10), (0.0, 0.5, 1.0)
    assert ws["A4"].value.startswith("Sensitivity inputs") and ws["A5"].value == "Quote adjustment" and ws["A6"].value.startswith("Share of the freed crew time")
    _agree(ws, "B5", 0.0)
    _agree(ws, "B6", 1.0)
    table = _g().workbook()["Make or Buy"]["B42"].value
    assert getattr(table, "ref", "") == "B42:D44" and table.r1 == "B5" and table.r2 == "B6", "the Data Table's inputs are B5 (row input) and B6 (column input)"
    _agree(ws, "A41", m["gain"], tol=0.05)
    for j, adj in enumerate(adjustments):
        _agree(ws, f"{'BCD'[j]}41", adj, tol=1e-9)
    for i, share in enumerate(shares):
        _agree(ws, f"A{42 + i}", share, tol=1e-9)
        for j in range(3):
            _agree(ws, f"{'BCD'[j]}{42 + i}", two_way[i][j], tol=0.05)
    kept_cost = -sum(x["a_kept"] for x in buy)
    assert abs(two_way[2][1] - m["gain"]) < 0.01 and abs(two_way[0][1] + kept_cost) < 0.01, "the centre of the bottom row is the annual effect, the top row's the cost with the crew kept"
    assert ws["A40"].value.startswith("Two-way Data Table") and "(across)" in ws["A40"].value and "(down)" in ws["A40"].value
    widths = [28, 308, 174, 174, 174]
    letters = ["A", "B", "C", "D"]
    d = Diagram("The Two-Way Data Table")
    _formula_bar(d, 0, 0, 860, "C44", "{=TABLE(B5,B6)}")
    sy = 40.0
    _ws(d, 0, sy, letters, widths, [], show_letters=True)
    y = _gap(d, 0, sy + 22, sum(widths), "Rows 1 to 3: the title and the method of the sheet")
    rows = [dict(n="4", cells=[dict(text=ws["A4"].value, bold=True, span=4, clip=False)]),
            dict(n="5", cells=[dict(text=ws["A5"].value, clip=False), _num(_show(ws, "B5"))]),
            dict(n="6", cells=[dict(text=ws["A6"].value, clip=False), _num(_show(ws, "B6"))])]
    geo1 = _ws(d, 0, y, [], widths, rows, show_letters=False)
    y = _gap(d, 0, geo1["bottom"], sum(widths), "Rows 7 to 39: the table of the six families, the Goal Seek block and the summary of the families to buy")
    rows = [dict(n="40", cells=[dict(**_wrapped(ws["A40"].value, sum(widths[1:]), bold=True), span=4)]),
            dict(n="41", cells=[_num(_show(ws, "A41"))] + [_num(_show(ws, f"{c}41")) for c in "BCD"])]
    for i in range(3):
        r = 42 + i
        rows.append(dict(n=str(r), cells=[_num(_show(ws, f"A{r}"))] + [_num(_show(ws, f"{c}{r}")) for c in "BCD"]))
    geo2 = _ws(d, 0, y, [], widths, rows, show_letters=False)
    _outline(d, geo1, (1, 1), (2, 1), TEAL, inset=1.5)
    _outline(d, geo2, (1, 0), (1, 0), BLUE, inset=1.5)
    _outline(d, geo2, (4, 2), (4, 2), CORAL, inset=1.5)
    bottom = _gap(d, 0, geo2["bottom"], sum(widths), "Rows 45 to 56: the materials the families draw and the model answer")
    bottom = _tabs(d, bottom + 8, "Make or Buy", ["Costs", "Make or Buy", "Capacity"])
    names = ", ".join(x["name"] for x in buy)
    _note(d, f"Outlined in teal: the two input cells, B5 (the row input, across) and B6 (the column input, down), which sit on this sheet because a Data Table's input "
             f"cells must. In blue: the corner cell A41, the sum of column V over the families to buy ({names}). In coral: C44, all the crew time removed at the quote as "
             f"it is, which equals the annual effect of buying them with the crew reduced, {m['gain']:,.0f}; the row of none removed, C42, is the {kept_cost:,.0f} they "
             f"cost with the crew kept, shown with a minus sign. Every grid cell is the one array formula {{=TABLE(B5,B6)}}, shown in the bar.", bottom + 4)
    return d


# -- IF-18-7-A: the Capacity sheet's table of hours freed ---------------------------------------------------------------------

def if_18_7_a() -> Diagram:
    """Rows 4 to 11 of Capacity: for each work center and the plant the routing time used, the hours freed by buying, the hours used
    after buying, the practical capacity and the use of it before and after; the Plant row is outlined, and the Quality Assurance row
    shows nothing freed."""
    m = _model()
    ws = _book("Capacity")
    centers, freed, buy = m["centers"], m["freed"], m["buy"]
    n = len(centers)
    assert ws["A5"].value == "Work center" and [ws.cell(6 + i, 1).value for i in range(n)] == [c["name"] for c in centers] and ws.cell(6 + n, 1).value == "Plant"
    assert all(b["name"] in ws["A4"].value for b in buy), "the heading names the families bought"
    tot_freed, used, pract = sum(freed.values()), m["used"], m["pract"]
    for i, c in enumerate(centers):
        r = 6 + i
        f = freed.get(c["id"], 0.0)
        _agree(ws, f"B{r}", c["used"], 0.05)
        _agree(ws, f"C{r}", f, 0.05)
        _agree(ws, f"D{r}", c["used"] - f, 0.05)
        _agree(ws, f"E{r}", c["pract"], 0.05)
        _agree(ws, f"F{r}", c["used"] / c["pract"], 1e-6)
        _agree(ws, f"G{r}", (c["used"] - f) / c["pract"], 1e-6)
    p = 6 + n
    for col, v in (("B", used), ("C", tot_freed), ("D", used - tot_freed), ("E", pract)):
        _agree(ws, f"{col}{p}", v, 0.05)
    _agree(ws, f"F{p}", used / pract, 1e-6)
    _agree(ws, f"G{p}", (used - tot_freed) / pract, 1e-6)
    r7 = ctx.values("ch18.r7")
    assert abs(tot_freed - r7["total"]) < 0.01 and abs(used / pract - r7["use_before"]) < 1e-6 and abs((used - tot_freed) / pract - r7["use_after"]) < 1e-6
    none = [i for i, c in enumerate(centers) if freed.get(c["id"], 0.0) == 0.0]
    assert len(none) == 1 and centers[none[0]]["name"] == "Quality Assurance", "only Quality Assurance frees nothing"
    g40 = _book("Growth 40%")
    assert all(g40.cell(14 + i, 11).value == 0 for i in range(6) if g40.cell(14 + i, 1).value in {b["name"] for b in buy}), "no family bought is routed through Quality Assurance"
    widths = [28, 182, 108, 108, 108, 108, 108, 108]
    letters = list("ABCDEFG")
    d = Diagram("The Capacity Sheet's Table of Hours Freed")
    _formula_bar(d, 0, 0, 860, "C6", _formula_of("Capacity", "C6"))
    sy = 40.0
    _ws(d, 0, sy, letters, widths, [], show_letters=True)
    y = _gap(d, 0, sy + 22, sum(widths), "Rows 1 to 3: the title and the note on the freed hours")
    rows = [dict(n="4", cells=[dict(text=ws["A4"].value, bold=True, span=7, clip=False)]),
            dict(n="5", cells=[_head(ws[f"{c}5"].value, widths[k + 1]) for k, c in enumerate(letters)])]
    for r in range(6, p + 1):
        rows.append(dict(n=str(r), cells=[dict(text=ws[f"A{r}"].value, bold=r == p, clip=False)]
                         + [_num(_show(ws, f"{c}{r}"), bold=r == p) for c in "BCDEFG"]))
    geo = _ws(d, 0, y, [], widths, rows, show_letters=False)
    _outline(d, geo, (2 + n, 0), (2 + n, 6), CORAL, inset=1.5)
    _outline(d, geo, (2 + none[0], 2), (2 + none[0], 2), TEAL, inset=1.5)
    bottom = _gap(d, 0, geo["bottom"], sum(widths), "Rows 13 to 42: the comparison with the unused capacity, the crew-kept case, the headroom table and the model answer")
    bottom = _tabs(d, bottom + 8, "Capacity", ["Make or Buy", "Capacity", "Growth 40%"])
    names = ", ".join(b["name"] for b in buy)
    _note(d, f"Outlined in coral: the Plant row, whose use of practical capacity falls from {_show(ws, f'F{p}')} to {_show(ws, f'G{p}')} when {names} are bought. "
             f"Outlined in teal: Quality Assurance, which frees nothing, because none of the three families is routed through it. The formula bar shows the first "
             f"freed cell, a SUMPRODUCT over the Activity Table weighted by whether the family is bought; the use columns divide by the practical capacity of column E.", bottom + 4)
    return d


# -- IF-18-7-B: the growth sheet that holds Solver's model ---------------------------------------------------------------------

def if_18_7_b() -> Diagram:
    """The sheet Growth 40%: the work-center block (practical capacity, load at growth, hours used after buying, slack, status) with the
    binding row outlined, the family block with the units bought as the shaded variable cells and their added cost, and the total
    added cost, the objective."""
    m, n = _model(), _n()
    ws = _book("Growth 40%")
    lp, centers, quoted = m["lp"][0], m["centers"], m["quoted"]
    assert abs(ws["B3"].value - n.GROWTH[0]) < 1e-9 and ws["A3"].value.startswith("Growth in volume")
    nc, nq = len(centers), len(quoted)
    assert [ws.cell(6 + i, 1).value for i in range(nc)] == [c["name"] for c in centers] and [ws.cell(14 + i, 1).value for i in range(nq)] == [q["name"] for q in quoted]
    binding = [i for i in range(nc) if ws[f"G{6 + i}"].value == "Binding"]
    assert len(binding) == 1 and centers[binding[0]]["short"] == lp["binding"][0]["short"], "one work center binds, the one the note names"
    _agree(ws, f"I{6 + binding[0]}", lp["binding"][0]["price"], 1e-4)
    assert abs(ws[f"F{6 + binding[0]}"].value) < 0.01
    for i, c in enumerate(centers):
        _agree(ws, f"B{6 + i}", c["pract"], 0.05)
    bought = {b["short"]: b["units"] for b in lp["bought"]}
    for i, q in enumerate(quoted):
        r = 14 + i
        _agree(ws, f"B{r}", q["units"], 0.05)
        _agree(ws, f"C{r}", q["units"] * (1 + n.GROWTH[0]), 0.05)
        _agree(ws, f"D{r}", q["delivered"], 0.005)
        _agree(ws, f"L{r}", bought.get(q["short"], 0.0), 1e-3)
        assert ws[f"L{r}"].fill.fill_type == "solid", "the variable cells are shaded in the workbook"
    _agree(ws, "B21", lp["cost"], 0.01)
    assert ws["A21"].value.startswith("Total added cost") and "minimize" in ws["A21"].value
    assert abs(sum(ws[f"M{14 + i}"].value for i in range(nq)) - lp["cost"]) < 0.01, "the objective adds the added cost column"
    formula = _g().clean(_g().workbook()["Growth 40%"]["E6"].value)
    assert formula == "=D6-SUMPRODUCT(G$14:G$19,$L$14:$L$19)", formula
    letters = ["A", "B", "C", "D", "E", "F", "G", "L", "M"]
    widths = [28, 196, 78, 86, 80, 80, 66, 80, 90, 76]
    d = Diagram("The Growth Sheet That Holds Solver's Model")
    _formula_bar(d, 0, 0, 860, "E6", formula)
    sy = 40.0
    _ws(d, 0, sy, letters, widths, [], show_letters=True)
    y = _gap(d, 0, sy + 22, sum(widths), "Rows 1 and 2: the title and the note on the model")
    wc_fill = dict(fill=TEAL_TINT, stroke=TEAL)

    rows = [dict(n="3", cells=[dict(text=ws["A3"].value, clip=False), _num(_show(ws, "B3"))]), dict(n="4", cells=[])]
    rows.append(dict(n="5", cells=[_head(ws[f"{c}5"].value, widths[k + 1]) for k, c in enumerate(letters[:7])]))
    for i in range(nc):
        r = 6 + i
        rows.append(dict(n=str(r), cells=[dict(text=ws[f"A{r}"].value, clip=False), _num(_show(ws, f"B{r}"), **wc_fill), _num(_show(ws, f"C{r}")),
                                           _num(_show(ws, f"D{r}")), _num(_show(ws, f"E{r}"), **wc_fill), _num(_show(ws, f"F{r}")),
                                           dict(text=ws[f"G{r}"].value, bold=i == binding[0], clip=False)]))
    rows += [dict(n="11", cells=[]), dict(n="12", cells=[])]
    rows.append(dict(n="13", cells=[_head(ws[f"{c}13"].value, widths[k + 1]) for k, c in enumerate(letters)]))
    for i in range(nq):
        r = 14 + i
        rows.append(dict(n=str(r), cells=[dict(text=ws[f"A{r}"].value, clip=False), _num(_show(ws, f"B{r}")), _num(_show(ws, f"C{r}"), **wc_fill),
                                           _num(_show(ws, f"D{r}")), _num(_show(ws, f"E{r}")), _num(_show(ws, f"F{r}")), _num(_show(ws, f"G{r}")),
                                           _num(_show(ws, f"L{r}"), fill=AMBER_TINT), _num(_show(ws, f"M{r}"))]))
    rows.append(dict(n="20", cells=[]))
    rows.append(dict(n="21", cells=[_wrapped(ws["A21"].value, widths[1]), _num(_show(ws, "B21"))]))
    geo = _ws(d, 0, y, [], widths, rows, show_letters=False)
    wc0, f0, last = 3, 3 + nc + 2 + 1, 3 + nc + 2 + 1 + nq - 1          # row indexes in `rows`: row 6, row 14, row 19
    _outline(d, geo, (wc0 + binding[0], 0), (wc0 + binding[0], 6), CORAL, inset=1.5)
    _outline(d, geo, (f0, 7), (last, 7), AMBER, inset=1.5)
    _outline(d, geo, (last + 2, 1), (last + 2, 1), BLUE, inset=1.5)
    bottom = _gap(d, 0, geo["bottom"], sum(widths), "Rows 22 to 32: the binding work centers, the families bought, the checks, and the plan the supplier would accept")
    y = _key(d, 0, bottom + 8, [("Variable cells, L14:L19", AMBER_TINT, AMBER), ("Constraints: E6:E10 not above B6:B10, L14:L19 not above C14:C19", TEAL_TINT, TEAL),
                                ("Objective, B21", WHITE, BLUE), ("The binding work center", WHITE, CORAL)])
    y = _tabs(d, y + 6, "Growth 40%", ["Capacity", "Growth 40%", "Growth 40% Sensitivity"])
    bw = centers[binding[0]]["name"]
    _note(d, f"The model of Requirement 7's growth test at {n.GROWTH[0]:.0%}: minimize the total added cost, B21, by changing the units bought, L14:L19, with the hours used after "
             f"buying (E6:E10) not above the practical capacity (B6:B10) and the units bought not above the units at growth (C14:C19). {bw} binds: its slack is 0 and an "
             f"hour of its capacity is worth {lp['binding'][0]['price']:.2f} (column I, hidden). The values in L14:L19 are the optimum computed by the simplex method and "
             f"written when the workbook was built; Solver itself was not run. The formula bar shows the hours used after buying at the first work center. Columns H to K and "
             f"N to R are hidden here: the shadow price, the hours per unit at each work center and the plan the supplier would accept.", y + 4)
    return d


# -- IF-18-8-A: the Prices sheet's table of floor prices ------------------------------------------------------------------------

def _runs(flags: list[bool]) -> list[tuple[int, int]]:
    """The first and last index of each run of consecutive True flags."""
    out, start = [], None
    for i, f in enumerate(flags + [False]):
        if f and start is None:
            start = i
        elif not f and start is not None:
            out.append((start, i - 1))
            start = None
    return out


def if_18_8_a() -> Diagram:
    """The first table of the Prices sheet (rows 4 to 20): each manufactured family's units, price, time-driven cost and margin, the
    Below the floor flag, the target price, the increase, the contribution per unit, the volume the increase can lose and the revenue
    a year; the families below the floor are outlined, and the totals block follows."""
    m, n = _model(), _n()
    ws = _book("Prices")
    fam, reprice, buy = m["fam"], m["reprice"], m["buy"]
    names = sorted(fam)
    assert ws["A4"].value == "Family" and [ws.cell(5 + i, 1).value for i in range(len(names))] == names
    inc = {r["name"]: r for r in reprice}
    below = [fam[f]["unit"]["m_td"] < n.FLOOR for f in names]
    assert sorted(f for f, b in zip(names, below) if b) == sorted(inc) and 0 < sum(below) < len(names)
    for i, f in enumerate(names):
        r, w = 5 + i, fam[f]["unit"]
        _agree(ws, f"B{r}", fam[f]["units"], 0.05)
        _agree(ws, f"C{r}", w["price"], 0.01)
        _agree(ws, f"D{r}", w["td"], 0.01)
        _agree(ws, f"E{r}", w["m_td"], 1e-5)
        _agree(ws, f"F{r}", below[i])
        _agree(ws, f"L{r}", w["price"] - w["kept"], 0.01)
        if below[i]:
            x = inc[f]
            _agree(ws, f"G{r}", w["td"] / (1 - n.FLOOR), 0.01)
            _agree(ws, f"J{r}", x["increase"], 1e-6)
            _agree(ws, f"M{r}", x["loss"], 1e-6)
            _agree(ws, f"N{r}", x["revenue"], 0.05)
            assert ws[f"G{r}"].fill.fill_type == "solid", "Goal Seek's changing cells are shaded in the workbook"
        else:
            assert all(ws[f"{c}{r}"].value is None for c in "GJMN"), "a family at or above the floor has no target price"
    buy_names = {b["name"] for b in buy}
    need = [f for f, b in zip(names, below) if b and f not in buy_names]
    t_rows = [_row_of(ws, "Families below the floor"), _row_of(ws, "  which"), _row_of(ws, "Revenue a year if volume holds"),
              _row_of(ws, "If the families of Requirement 6 are bought"), _row_of(ws, "  their revenue a year")]
    assert t_rows == [22, 23, 24, 25, 26], t_rows
    _agree(ws, "B22", sum(below))
    assert ws["B23"].value == ", ".join(f for f, b in zip(names, below) if b) and ws["B25"].value == ", ".join(need)
    _agree(ws, "B24", sum(r["revenue"] for r in reprice), 0.05)
    _agree(ws, "B26", sum(inc[f]["revenue"] for f in need), 0.05)
    letters = ["A", "B", "C", "D", "E", "F", "G", "J", "L", "M", "N"]
    widths = [24, 72, 56, 72, 84, 84, 58, 84, 66, 90, 90, 80]
    assert abs(sum(widths) - 860) < 1
    d = Diagram("The Prices Sheet's Table of Floor Prices")
    formula = _g().clean(_g().workbook()["Prices"]["M6"].value)
    assert formula == '=IF(F6,1-L6/(L6+G6-C6),"")', formula
    _formula_bar(d, 0, 0, 860, "M6", formula)
    sy = 40.0
    _ws(d, 0, sy, letters, widths, [], show_letters=True)
    y = _gap(d, 0, sy + 22, sum(widths), "Rows 1 to 3: the title and the note on the floor")
    rows = [dict(n="4", cells=[_head(ws[f"{c}4"].value, widths[k + 1]) for k, c in enumerate(letters)])]
    for i, f in enumerate(names):
        r = 5 + i
        g_fill = dict(fill=AMBER_TINT) if ws[f"G{r}"].fill.fill_type == "solid" else {}
        rows.append(dict(n=str(r), cells=[dict(text=f, clip=False), _num(_show(ws, f"B{r}")), _num(_show(ws, f"C{r}")), _num(_show(ws, f"D{r}")),
                                           _num(_show(ws, f"E{r}")), dict(text=_show(ws, f"F{r}"), align="center", clip=False),
                                           _num(_show(ws, f"G{r}"), **g_fill), _num(_show(ws, f"J{r}")), _num(_show(ws, f"L{r}")),
                                           _num(_show(ws, f"M{r}")), _num(_show(ws, f"N{r}"))]))
    geo = _ws(d, 0, y, [], widths, rows, show_letters=False)
    for a, b in _runs(below):
        _outline(d, geo, (1 + a, 0), (1 + b, 10), CORAL, inset=1.5)
    y = _gap(d, 0, geo["bottom"], sum(widths), "Row 21 is blank. The totals block, rows 22 to 26, is drawn with column A widened to show its labels", h=24)
    w2 = [24, 400, 100, 110, 110, 116]
    rows2 = [dict(n="22", cells=[dict(text=ws["A22"].value, clip=False), _num(_show(ws, "B22"))]),
             dict(n="23", cells=[dict(text=ws["A23"].value, clip=False), dict(text=_show(ws, "B23"), span=4, clip=False)]),
             dict(n="24", cells=[dict(text=ws["A24"].value, clip=False), _num(_show(ws, "B24"))]),
             dict(n="25", cells=[dict(text=ws["A25"].value, clip=False), dict(text=_show(ws, "B25"), span=4, clip=False)]),
             dict(n="26", cells=[dict(text=ws["A26"].value, clip=False), _num(_show(ws, "B26"))])]
    geo2 = _ws(d, 0, y, [], w2, rows2, show_letters=False)
    bottom = _tabs(d, geo2["bottom"] + 10, "Prices", ["Growth 50% Sensitivity", "Prices", "Standards"])
    _note(d, f"Outlined in coral: the {n.word(sum(below))} families whose margin on time-driven cost is below the {n.FLOOR:.0%} floor; only they have a target price, an increase, "
             f"a volume they can lose and a revenue. The shaded cells of column G are Goal Seek's changing cells. The volume the increase can lose is the formula in the bar: one "
             f"minus the contribution per unit now, over that contribution plus the increase per unit. Columns H, I and K are hidden here, and the columns are narrower than the "
             f"workbook's. The totals read {_show(ws, 'B24')} a year if volume holds. Of the families bought in Requirement 6, {n.word(len(buy_names & {f for f, b in zip(names, below) if b}))} "
             f"are among those below the floor, so {_show(ws, 'B26')} a year is what the other {n.word(len(need))} bring.", bottom + 4)
    return d


# -- IF-18-9-A: the top of the Standards sheet --------------------------------------------------------------------------------------

def if_18_9_a() -> Diagram:
    """Rows 4 to 16 of Standards: the five components of the proposed rate with their basis, the proposed rate, the labor-content case, the
    current standard, the time-driven rate, the base year's actual cost per standard hour and the two rates at the ends of the
    normal-capacity range; the proposed rate is outlined."""
    m = _model()
    ws = _book("Standards")
    p, r9 = m["parts"], ctx.values("ch18.r9")
    expected = [("B5", p["direct"]), ("B6", p["variable"]), ("B7", p["indirect"]), ("B8", p["supervision"]), ("B9", p["depreciation"]),
                ("B10", m["new_rate"]), ("B11", m["one_rate"]), ("B13", m["td_std"]), ("B14", m["actual_rate"]), ("B15", m["normal"][0]["rate"]),
                ("B16", m["normal"][1]["rate"])]
    for ref, v in expected:
        _agree(ws, ref, v, 1e-5)
    _agree(ws, "B12", r9["current"], 1e-5)
    assert abs(sum(p.values()) - m["new_rate"]) < 1e-9, "the five components add to the proposed rate"
    assert m["normal"][0]["share"] < 1 < m["normal"][1]["share"] and m["normal"][0]["rate"] > m["new_rate"] > m["normal"][1]["rate"], "the range brackets the proposed rate"
    assert ws["A4"].value.startswith("Proposed standard conversion rate") and ws["A10"].value.startswith("Proposed rate") and ws["A15"].value.startswith("Normal capacity at the lower share")
    assert ws["A16"].value.startswith("Normal capacity at the higher share") and ws["A11"].value.startswith("Labor-content case") and ws["A12"].value.startswith("Current standard")
    formula = _g().clean(_g().workbook()["Standards"]["B10"].value)
    assert formula == "=SUM(B5:B9)", formula
    letters = list("ABCDEFGHI")
    widths = [28, 336, 96] + [57] * 7
    d = Diagram("The Top of the Standards Sheet")
    _formula_bar(d, 0, 0, 860, "B10", formula)
    sy = 40.0
    _ws(d, 0, sy, letters, widths, [], show_letters=True)
    y = _gap(d, 0, sy + 22, sum(widths), "Rows 1 to 3: the title and the note on normal capacity")
    rows = [dict(n="4", cells=[_wrapped(ws["A4"].value, widths[1], bold=True), dict(text=ws["B4"].value, bold=True), dict(text=ws["C4"].value, bold=True, span=7)])]
    for r in range(5, 17):
        cells = [_wrapped(ws[f"A{r}"].value, widths[1]), _num(_show(ws, f"B{r}"))]
        if r <= 9:
            cells.append(dict(text=ws[f"C{r}"].value, span=7, clip=True, align="left"))
        rows.append(dict(n=str(r), cells=cells))
    geo = _ws(d, 0, y, [], widths, rows, show_letters=False)
    _outline(d, geo, (1, 1), (5, 1), TEAL, inset=1.5)
    _outline(d, geo, (6, 0), (6, 1), CORAL, inset=1.5)
    bottom = _gap(d, 0, geo["bottom"], sum(widths), "Rows 17 and below: the plan's figures against the records, the revised standards, the stock and its revaluation, and the model answer", h=24)
    bottom = _tabs(d, bottom + 8, "Standards", ["Prices", "Standards", "Memo"])
    lo, hi = m["normal"]
    _note(d, f"Outlined in teal: the five components. Outlined in coral: the proposed rate, B10, their sum (its formula is in the bar). Below it come the labor-content case with no indirect "
             f"hours, the current standard (the conversion released per standard hour), the time-driven rate, the base year's actual cost per standard hour, and the proposed rate "
             f"with supervision and depreciation spread over {lo['share']:.0%} and {hi['share']:.0%} of the plan volume, the ends of the normal-capacity range. Labels are wrapped here, "
             f"as the workbook wraps a long label within column A. Columns J to Q hold the Table OnHand and are not shown.", bottom + 4)
    return d


# -- IF-18-9-B: the revaluation against each year's screen ---------------------------------------------------------------------

def _bar(d: Diagram, x: float, y: float, w: float, h: float, label: str, fill: str, stroke: str, color: str) -> None:
    """A bar that carries its value at the top, inside it (the label's colors are the bar's own pair, so they keep their contrast)."""
    style = (f"rounded=0;whiteSpace=wrap;html=1;fillColor={fill};strokeColor={stroke};strokeWidth=2;fontColor={color};fontSize={SMALL};"
             f"fontStyle=0;align=center;verticalAlign=top;{FONT}spacingTop=4;spacingLeft=1;spacingRight=1;")
    d.vertex(label, style, x, y, w, h)


def if_18_9_b() -> Diagram:
    """Two groups of bars, one for each year-end of the statements issued: the revaluation of the finished goods on hand at the proposed
    rate, with no indirect hours and at each end of the normal-capacity range, with that year's materiality screen drawn across the
    group; the one bar below its screen is marked."""
    data = ctx.data()
    m, n = _model(), _n()
    ws = _book("Standards")
    P, C = data.P, data.C
    years = [P, C]
    stock = {s["year"]: s for s in m["stock"]}
    lo, hi = m["normal"]
    series = [("Proposed rate", "plan"), ("No indirect hours", "one"), (f"Normal capacity at {lo['share']:.0%}", 0), (f"Normal capacity at {hi['share']:.0%}", 1)]

    def value(year: int, key) -> float:
        s = stock[year]
        return s["normal"][key] if isinstance(key, int) else s[key]
    screens = {y: n.materiality(data, y) for y in years}
    col = {y: "BCD"[list(data.years).index(y)] for y in years}                    # the sheet's columns run over the three year-ends
    rows = {"plan": _row_of(ws, "Revaluation at the proposed rate"), "one": _row_of(ws, "Revaluation with no indirect hours"),
            0: _row_of(ws, "Revaluation with normal capacity at the lower share"), 1: _row_of(ws, "Revaluation with normal capacity at the higher share")}
    for y in years:
        assert ws[f"{col[y]}34"].value == f"End of fiscal {y}", (y, ws[f"{col[y]}34"].value)
        for _, key in series:
            _agree(ws, f"{col[y]}{rows[key]}", value(y, key), 0.05)
        r = _row_of(ws, f"Materiality screen for fiscal {y}")
        _agree(ws, f"B{r}", screens[y], 0.006)
    below = [(y, label) for y in years for label, key in series if value(y, key) < screens[y]]
    assert below == [(P, "No indirect hours")], below                              # the earlier year with no indirect hours
    short = screens[P] - value(P, "one")
    assert 0 < short < 0.1 * screens[P]
    top = max(value(y, key) for y in years for _, key in series)
    assert top < 600000
    d = Diagram("The Revaluation of the Stock Against the Screens")
    d.text("<b>The revaluation of the finished goods on hand at each year-end, against that year's materiality screen</b>", 0, 0, 860, 20, size=SMALL)
    ticks = [0, 100000, 200000, 300000, 400000, 500000, 600000]
    plot = kit.Plot(d, 0, 26, 860, 482, y_ticks=ticks, categories=["", ""], y_title="Dollars", left=66, bottom=94)
    bw, gap = 64.0, 8.0
    group = 4 * bw + 3 * gap
    slot = plot.pw / 2
    base = plot.y_of(0)
    bars: dict[tuple[int, int], tuple[float, float]] = {}
    for g, y in enumerate(years):
        left = plot.px + slot * g
        x0 = left + (slot - group) / 2
        for k, (label, key) in enumerate(series):
            v = value(y, key)
            bx, by = x0 + k * (bw + gap), plot.y_of(v)
            lowbar = (y, label) in below
            _bar(d, bx, by, bw, base - by, f"{v:,.0f}", BLUE_TINT if lowbar else BLUE, BLUE, INK if lowbar else WHITE)
            bars[(g, k)] = (bx + bw / 2, by)
            d.text(esc(label), bx - 2, base + 6, bw + 4, 48, size=SMALL, align="center")
        d.text(f"<b>End of fiscal {y}</b>", left, base + 58, slot, 20, size=SMALL, align="center")
    for g, y in enumerate(years):                                                       # the screens, over the bars
        left = plot.px + slot * g
        ly = plot.y_of(screens[y])
        _polyline(d, [(left + 4, ly), (left + slot - 4, ly)], color=CORAL, width=2.5)
        _label(d, f"Screen<br>{screens[y]:,.0f}", left + 2, ly - 34, 50, 32, color=CORAL, bold=True)
    # the one bar below its screen
    cx, cy = bars[(0, 1)]
    note = f"{short:,.0f} short of the screen: the one revaluation below it"
    nw = 200
    nh = kit.height_for(note, nw, SMALL, pad=12)
    ny = plot.y_of(560000)
    d.box(note, cx - nw / 2, ny, nw, nh, fill=AMBER_TINT, stroke=AMBER, size=SMALL, rounded=False, align="left", stroke_width=1.5)
    a1, a2 = d.anchor(cx, ny + nh), d.anchor(cx, cy - 2)
    d.edge(a1, a2, end="blockThin", color=AMBER, width=1.5, meta=STRAIGHT)
    y0 = 26 + 482 + 10
    _swatch = lambda x, fill, stroke, label: (d.box("", x, y0 + 3, 22, 14, fill=fill, stroke=stroke, stroke_width=2, rounded=False),      # noqa: E731
                                                d.text(esc(label), x + 28, y0, text_width(esc(label), SMALL, False) + 8, 20, size=SMALL, valign="middle"))
    x = 0.0
    for fill, stroke, label in ((BLUE, BLUE, "A revaluation above its screen"), (BLUE_TINT, BLUE, "A revaluation below its screen")):
        _swatch(x, fill, stroke, label)
        x += 28 + text_width(esc(label), SMALL, False) + 30
    d.line_sample(0, y0 + 32, 22, color=CORAL, width=2.5)
    d.text(esc(f"The year's materiality screen, {n.MATERIALITY:.0%} of income before income taxes as recorded"), 28, y0 + 22, 520, 20, size=SMALL, valign="middle")
    _note(d, f"Each bar is the revaluation of the manufactured finished goods made since the ledger opened and on hand at the year-end, as the movements rebuild them: at the proposed "
             f"rate, with no indirect hours (the labor-content case) and with supervision and depreciation spread over {lo['share']:.0%} and {hi['share']:.0%} of the plan volume. "
             f"The screens are {screens[P]:,.0f} for fiscal {P} and {screens[C]:,.0f} for fiscal {C}; the revaluation is a sensitivity on next year's plan, not an established error.",
          y0 + 50)
    return d


# -- IF-18-10-A: the Memo sheet's decision table -------------------------------------------------------------------------------

def if_18_10_a() -> Diagram:
    """The decision table of the Memo sheet (rows 4 to 15): a row for each quoted family and each other family below the floor, with the
    action, the annual effect, the one-time cost and what the decision depends on, and the two totals; the families to buy come first,
    those to make next, the other families below the floor last."""
    m, n = _model(), _n()
    ws = _book("Memo")
    g10 = ctx.values("ch18.g10")
    paths = n.paths_compared(m)
    rows_ctx = g10["rows"]
    nb, nm, no = g10["n_buy_rows"], g10["n_make_rows"], g10["n_only"]
    total = len(rows_ctx)
    assert total == nb + nm + no == g10["n_rows"] and total == 9 and [ws.cell(4, c).value for c in range(1, 6)] == ["Family", "What to do", "Annual effect", "One-time cost", "What the decision depends on"]
    for i, r in enumerate(rows_ctx):
        row = 5 + i
        assert ws[f"A{row}"].value.endswith(r["short"]), (row, ws[f"A{row}"].value, r["short"])
        _agree(ws, f"C{row}", r["effect"], 0.05)
        _agree(ws, f"D{row}", r["once"], 0.05)
        text = ws[f"B{row}"].value
        assert text.startswith("Buy") if i < nb else text.startswith("Make") if i < nb + nm else text.startswith("Raise"), (row, text)
    assert ws["A14"].value == "Total annual effect" and ws["A15"].value == "Total one-time cost"
    _agree(ws, "B14", paths["buy"], 0.05)
    _agree(ws, "B15", paths["once"], 0.05)
    assert [r["short"] for r in rows_ctx[nb:nb + nm]] == [x["short"] for x in m["lose"]], "the families to make come after those to buy"
    formula = _g().clean(_g().workbook()["Memo"]["B14"].value)
    assert formula == "=SUM(C5:C13)", formula
    idle = [r for r in rows_ctx[nb:nb + nm] if r["effect"] == 0]
    letters = list("ABCDE")
    widths = [28, 132, 232, 80, 80, 308]
    d = Diagram("The Memo Sheet's Decision Table")
    _formula_bar(d, 0, 0, 860, "B14", formula)
    sy = 40.0
    _ws(d, 0, sy, letters, widths, [], show_letters=True)
    y = _gap(d, 0, sy + 22, sum(widths), "Rows 1 to 3: the title and the note on the memo")
    cells = [_head(ws[f"{c}4"].value, widths[k + 1]) for k, c in enumerate(letters)]
    rows = [dict(n="4", cells=cells)]
    for i in range(total):
        r = 5 + i
        rows.append(dict(n=str(r), cells=[dict(text=ws[f"A{r}"].value, clip=False, valign="top"), _wrapped(ws[f"B{r}"].value, widths[2], valign="top"),
                                           _num(_show(ws, f"C{r}"), valign="top"), _num(_show(ws, f"D{r}"), valign="top"),
                                           _wrapped(ws[f"E{r}"].value, widths[5], valign="top")]))
    for r in (14, 15):
        rows.append(dict(n=str(r), cells=[dict(text=ws[f"A{r}"].value, clip=False), _num(_show(ws, f"B{r}"))]))
    geo = _ws(d, 0, y, [], widths, rows, show_letters=False)
    for (a, b), color in (((0, nb - 1), TEAL), ((nb, nb + nm - 1), BLUE), ((nb + nm, total - 1), CORAL)):
        _outline(d, geo, (1 + a, 0), (1 + b, 4), color, inset=1.5)
    bottom = _gap(d, 0, geo["bottom"], sum(widths), "Rows 16 to 23: the model memo, in five paragraphs")
    bottom = _tabs(d, bottom + 8, "Memo", ["Standards", "Memo", "ItemCosts"])
    _note(d, f"Outlined in teal: the {n.word(nb)} families to buy, only with the crew reduced. In blue: the {n.word(nm)} to make. In coral: the other {n.word(no)} families below the floor, "
             f"which only need a price. The totals are the path that buys and reprices the rest: an annual effect of {_show(ws, 'B14')} and one-time costs of {_show(ws, 'B15')} "
             f"(severance and tooling); the formula bar adds column C. {n.word(len(idle)).capitalize()} of the {n.word(nm)} make rows carry no effect, because their families are at or above the floor.", bottom + 4)
    return d


# -- IF-18-10-B: the page plan of the memo and its appendix ------------------------------------------------------------------------

def if_18_10_b() -> Diagram:
    """Two columns: the memo of two pages with its seven short sections in order, and the appendix of three pages with three exhibits, each
    tagged with the sheet of the workbook it comes from."""
    m = _model()
    text = _case_text()
    assert "no more than two pages" in text and "no more than three pages" in text and "decision table" in text, "the page limits are the case's"
    wb = _cached().sheetnames
    memo = _book("Memo")
    finding = ("On the routings, the plant supplied", "position-equivalents of practical hours")
    for key, row_start, needle in (("finding", *finding), ("views", finding[0], "Both of you are right"),
                                   ("decisions", "Make ", "only with a committed plan to reduce the crew"), ("paths", "Make ", "Compare the whole paths"),
                                   ("prices", "Raise ", "price lists"), ("standards", "Standards:", "normal capacity"), ("manager", "The one decision only the production manager can make", "")):
        row = _row_of(memo, row_start, 1)
        assert needle in memo.cell(row, 1).value, (key, memo.cell(row, 1).value[:80])
    assert m["unused_h"] > sum(m["freed"].values()), "on the same base, the plant's unused capacity is larger than what buying would free"
    sections = [
        ("The finding that changes the question", "On the routings the plant leaves more hours unused than buying would free, so fitting the crew to the work, week by week, is the larger question.", "Requirements<br>3, 4 and 7"),
        ("The two views of the cost", "The CFO is right that the products cost more than their standards say; the production manager is right that idle capacity belongs to the plant.", "Requirement 4"),
        ("Each decision with its condition", "Make, buy only with a committed plan to reduce the crew, or raise the price, with the decision table of every quoted family and every family below the floor.", "Requirements<br>6 and 8"),
        ("The whole paths compared", "Make and reprice against buy and reprice the rest, with the one-time costs shown.", "The Memo sheet and<br>Requirement 8"),
        ("The prices", "New, approved price lists with end dates for the families below the floor.", "Requirement 8"),
        ("The standards and the statements", "A rate on normal capacity for inventory, and the evidence the base year's statements need before anything is restated.", "Requirement 9"),
        ("The production manager's decision", "Whether, and how fast, to reduce the crew to the work the plant has: left to the production manager.", ""),
    ]
    exhibits = [("Exhibit 1", "The family costs under the three views", ["Costs"]),
                ("Exhibit 2", "The make-or-buy table, with the break-even quotes and the two-way table", ["Make or Buy"]),
                ("Exhibit 3", "The capacity, price and standards results", ["Capacity", "Prices", "Standards"])]
    assert all(s in wb for _, _, tags in exhibits for s in tags) and "Memo" in wb
    d = Diagram("The Page Plan of the Memo and Its Appendix")
    LW, RX, RW = 540.0, 580.0, 280.0
    d.box("<b>The memo</b>: no more than two pages, to the CFO and the production manager", 0, 0, LW, 34, fill=GRAY, stroke=GRAY, color=WHITE, size=SMALL, rounded=False)
    d.box("<b>The appendix</b>: no more than three pages", RX, 0, RW, 34, fill=GRAY, stroke=GRAY, color=WHITE, size=SMALL, rounded=False)
    H, GAP = 64.0, 8.0
    y = 44.0
    for i, (title, body, evidence) in enumerate(sections):
        d.box("", 0, y, LW, H, fill=WHITE, stroke=GRAY, stroke_width=1.5, rounded=False)
        d.marker(str(i + 1), 8, y + (H - 24) / 2)
        label = f"<b>{esc(title)}</b><br>{esc(body)}"
        tw = LW - 44 - (170 if evidence else 12)
        assert kit.lines_for(label, tw, SMALL) <= 4, (title, kit.lines_for(label, tw, SMALL))
        d.text(label, 40, y + 2, tw, H - 4, size=SMALL, valign="middle")
        if evidence:
            d.box(evidence, LW - 150, y + 14, 140, H - 28, fill=GRAY_TINT, stroke=RULE, size=SMALL, rounded=True)
        y += H + GAP
    left_bottom = y - GAP
    PH = 118.0
    py = 44.0
    for title, body, tags in exhibits:
        d.box("", RX, py, RW, PH, fill=WHITE, stroke=GRAY, stroke_width=1.5, rounded=False)
        d.text(f"<b>{esc(title)}</b><br>{esc(body)}", RX + 6, py + 6, RW - 12, PH - 44, size=SMALL)
        tx = RX + 8
        d.text("Sheet:", tx, py + PH - 31, 44, 22, size=SMALL, color=INK, valign="middle")
        tx += 48
        for tag in tags:
            tw2 = text_width(f"<b>{esc(tag)}</b>", SMALL, True) + 16
            d.box(f"<b>{esc(tag)}</b>", tx, py + PH - 30, tw2, 22, fill=BLUE, stroke=BLUE, color=WHITE, size=SMALL)
            tx += tw2 + 6
        py += PH + 12
    kit.callout(d, "Each exhibit is cited in the memo by name, with its sheet and the date of its assumptions, and supports a number the memo states. Keep the summaries on the last page legible and the detailed sheets for traceability. The decision table of the Memo sheet goes with the memo, in its third section.",
                RX, py + 2, RW, "teal")
    _note(d, "Left: the seven sections of the model memo in their order, each with the requirements that supply its evidence. Right: the appendix, at most three pages, each exhibit tagged "
             "with the sheet of Make Buy Reprice.xlsx it comes from.", left_bottom + 10)
    return d


FIGURES = {
    "IF-18-0-A": if_18_0_a,
    "IF-18-0-B": if_18_0_b,
    "IF-18-2-B": if_18_2_b,
    "IF-18-3-A": if_18_3_a,
    "IF-18-3-B": if_18_3_b,
    "IF-18-4-A": if_18_4_a,
    "IF-18-5-A": if_18_5_a,
    "IF-18-5-B": if_18_5_b,
    "IF-18-6-A": if_18_6_a,
    "IF-18-6-B": if_18_6_b,
    "IF-18-7-A": if_18_7_a,
    "IF-18-7-B": if_18_7_b,
    "IF-18-8-A": if_18_8_a,
    "IF-18-9-A": if_18_9_a,
    "IF-18-9-B": if_18_9_b,
    "IF-18-10-A": if_18_10_a,
    "IF-18-10-B": if_18_10_b,
}

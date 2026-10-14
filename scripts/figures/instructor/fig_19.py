"""Instructor figures of Chapter 19 (guides in facts/instructor/guides/chapter-19-customer-credits-audit/).

Chapter 19, "Auditing the Customer Credits Cycle", is the third capstone case.  Its model answer is four files: the script Credits.sql,
the workbook Credits Audit.xlsx, the Power BI project Credits Analytics (with RefundFlags.csv) and Audit Monitoring updated
(outputs/companion/solutions/Chapter19-capstone-*).  The Excel mocks draw a block of the workbook as Excel shows it (each value is read
from the saved workbook and asserted against the notes' own values), the Power BI mocks draw the Timing and Refund Ranking pages from
the project's definition (visuals, titles, projections) with values from the registered notes asserted against the dataset, and the
diagrams take their structure from the case text and from the workbook's own sheets (Cycle Map, Risks, Deficiencies).

The first half of the module holds the helpers (workbook reader, Excel sheet, Power BI visuals) and the figures of the case note and of
Requirements 1 to 5; the second half (IF-19-6-A to IF-19-10-B) reuses them: the Sample sheet and the distribution behind it, the largest
refund's timeline, the Deficiencies sheet, the report's page budget, and the Credits Cycle page and the flow of the continuing tests in the
Audit Monitoring project (read from its files, with the Checks tab's expected values asserted against the figure).
"""

from __future__ import annotations

import json
import re
import shutil
import tempfile
from datetime import date
from functools import lru_cache
from pathlib import Path

import ctx
import excel as xl
import kit
import powerbi as pbi
from check_figures import text_width
from drawio import (AMBER, AMBER_TINT, BLUE, BLUE_TINT, CORAL, CORAL_TINT, FONT, GRAY, GRAY_TINT, INK, ROW_H, RULE, SMALL, TEAL,
                    TEAL_TINT, WHITE, Diagram, esc)

REPO = ctx.REPO
SOLUTIONS = REPO / "outputs" / "companion" / "solutions"
XLSX = SOLUTIONS / "Chapter19-capstone-excel" / "Credits Audit.xlsx"
PROJECT = SOLUTIONS / "Chapter19-capstone-pbi" / "Credits Analytics"
MODEL = PROJECT / "Credits Analytics.SemanticModel"
REPORT = PROJECT / "Credits Analytics.Report" / "definition"
FLAGS_CSV = SOLUTIONS / "Chapter19-capstone-pbi" / "RefundFlags.csv"
CHAPTER = REPO / "chapters" / "19-customer-credits-audit" / "chapter.qmd"
GUIDES = REPO / "facts" / "instructor" / "guides" / "chapter-19-customer-credits-audit"
MINUS = "−"
STRAIGHT = "edgeStyle=none;"


# -- shared helpers -----------------------------------------------------------------------------------------------------

def _note(d: Diagram, text: str, y: float, x: float = 0, w: float = 860) -> float:
    """A short italic gray note sized to its text; returns its bottom."""
    h = kit.height_for(text, w, SMALL, pad=8)
    d.text(f"<i>{text}</i>", x, y, w, h, size=SMALL, color=GRAY)
    return y + h


def _money(x: float) -> str:
    return f"{x:,.2f}"


def _g():
    """guide_contexts.ch19: the helpers that read the model answers."""
    from guide_contexts import ch19
    return ch19


def _n():
    from notes import ch19
    return ch19


def _clip(text: str, width: float, size: int = SMALL, margin: float = 14, bold: bool = False) -> str:
    """A value cut to its column with an ellipsis, as a narrow column shows it."""
    if not text or text_width(esc(text), size, bold) <= width - margin:
        return text
    while text and text_width(esc(text + "…"), size, bold) > width - margin:
        text = text[:-1]
    return text.rstrip() + "…"


def _nowrap(text: str) -> str:
    return f'<span style="white-space:nowrap">{esc(text)}</span>'


def _wrap(text: str, width: float, size: int = SMALL, bold: bool = False) -> list[str]:
    """A greedy word wrap measured with text_width."""
    out, cur = [], ""
    for word in text.split(" "):
        trial = (cur + " " + word).strip()
        if not cur or text_width(trial, size, bold) <= width:
            cur = trial
        else:
            out.append(cur)
            cur = word
    out.append(cur)
    return out


def _polyline(d: Diagram, pts: list[tuple[float, float]], color: str = BLUE, width: float = 3, dashed: bool = False) -> None:
    """One edge through the points, drawn straight between them (separate edges leave gaps at their joints)."""
    a, b = d.anchor(*pts[0]), d.anchor(*pts[-1])
    d.edge(a, b, color=color, width=width, dashed=dashed, points=pts[1:-1], meta="edgeStyle=none;rounded=0;")


def _label(d: Diagram, text: str, x: float, y: float, w: float, h: float = 18, color: str = INK, align: str = "left",
           bold: bool = False) -> None:
    """A white-backed text label, so a gridline does not strike through it."""
    shown = f"<b>{text}</b>" if bold else text
    d.vertex(shown, f"text;html=1;whiteSpace=wrap;align={align};verticalAlign=middle;fontSize={SMALL};fontColor={color};{FONT}"
                    f"spacing=2;fillColor={WHITE};strokeColor=none;", x, y, w, h)


def _date(iso: str) -> str:
    """An ISO date that does not break at its hyphens."""
    return f'<span style="white-space:nowrap">{iso}</span>'


# -- the case text -----------------------------------------------------------------------------------------------------

def _case_text() -> str:
    return CHAPTER.read_text(encoding="utf-8")


def _requirements() -> list[dict]:
    """Each requirement of the case text: number, title, chapters (empty for the report) and the phase it belongs to."""
    found, phase = [], None
    for line in _case_text().split("\n"):
        m = re.match(r"^### Phase (\d): (.+)$", line)
        if m:
            phase = (int(m.group(1)), m.group(2).strip())
        m = re.match(r"^\*\*Requirement (\d+): (.+?)(?: \((Chapters? [^)]+)\))?\.\*\*", line)
        if m:
            found.append(dict(n=int(m.group(1)), title=m.group(2), chapters=m.group(3) or "", phase=phase))
    assert [r["n"] for r in found] == list(range(1, 11)), found
    assert [r["phase"][0] for r in found] == [1, 1, 1, 2, 2, 2, 2, 3, 3, 3], found
    return found


def _tools() -> dict[int, list[str]]:
    """The tools of each requirement as the case notes' table 'Using the requirements' lists them (a written report is its own tag)."""
    text = (GUIDES / "00-case-notes.md").read_text(encoding="utf-8")
    out = {}
    for m in re.finditer(r"^\| (\d+)\. [^|]+\| (?:Auditing|General) \| ([^|]+?) \|", text, re.M):
        tags: list[str] = []
        for part in re.split(r", | and ", m.group(2)):
            tag = {"a written report": "Report"}.get(part.strip(), part.strip())
            if tag not in tags:
                tags.append(tag)
        out[int(m.group(1))] = tags
    assert sorted(out) == list(range(1, 11)), out
    return out


def _milestones() -> list[dict]:
    found = [dict(n=int(m.group(1)), title=m.group(2), reqs=m.group(3), body=m.group(4))
             for m in re.finditer(r"^(\d)\. \*\*([^*]+)\*\* \(Requirements? ([^)]+)\): (.+?)\.$", _case_text(), re.M)]
    assert [x["n"] for x in found] == [1, 2, 3], found
    return found


def _tool_files() -> list[dict]:
    """The rows of the case's table of tools and files: tool, file, what it holds and how it hands results on."""
    rows = []
    for m in re.finditer(r"^\| (SQL|Excel|Power BI) \| ([^|]+?) \| ([^|]+?) \| ([^|]+?) \|$", _case_text(), re.M):
        rows.append(dict(tool=m.group(1), file=m.group(2), holds=m.group(3), handoff=m.group(4)))
    assert [r["tool"] for r in rows] == ["SQL", "Excel", "Power BI"], rows
    return rows


TOOL_COLORS = {"SQL": (TEAL, TEAL_TINT), "Excel": (BLUE, BLUE_TINT), "Power BI": (AMBER, AMBER_TINT), "Report": (GRAY, GRAY_TINT)}


def _tag(d: Diagram, label: str, x: float, y: float) -> float:
    """A filled tool tag; returns its width."""
    fill, _ = TOOL_COLORS[label]
    tw = text_width(f"<b>{label}</b>", SMALL, True) + 16
    d.box(f"<b>{label}</b>", x, y, tw, 22, fill=fill, stroke=fill, color=WHITE, size=SMALL)
    return tw


# -- the model workbook -------------------------------------------------------------------------------------------------

@lru_cache(maxsize=None)
def _cached():
    """The model workbook with its saved values (openpyxl on a copy; the figures draw what Excel shows)."""
    import openpyxl
    folder = Path(tempfile.mkdtemp(prefix="fig19-"))
    copy = folder / "Credits Audit.xlsx"
    shutil.copyfile(XLSX, copy)
    return openpyxl.load_workbook(copy, read_only=False, data_only=True)


def _agree(ws, ref: str, expected, tol: float = 0.006) -> None:
    """The model workbook holds in this cell the value the figure draws."""
    value = ws[ref].value
    if hasattr(value, "date"):
        value = value.date().isoformat()
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        assert isinstance(value, (int, float)) and abs(value - expected) < tol, f"{ws.title}!{ref} holds {value!r}, the figure draws {expected!r}"
    else:
        assert value == expected, f"{ws.title}!{ref} holds {value!r}, the figure draws {expected!r}"


def _book(sheet: str):
    """One worksheet of the model workbook with its saved values."""
    return _cached()[sheet]


def _acct(x: float | None, places: int = 2) -> str:
    """The workbook's number format #,##0.00: a blank cell empty."""
    return "" if x is None else f"{x:,.{places}f}"


def _gap(d: Diagram, x: float, y: float, w: float, text: str, h: float = 22) -> float:
    """A band that stands for the rows of a worksheet that the figure leaves out."""
    d.box(f"<i>{esc(text)}</i>", x, y, w, h, fill=WHITE, stroke=RULE, color=GRAY, size=SMALL, rounded=False, dashed=True)
    return y + h


def _wrapped(text: str, width: float, bold: bool = False, **kw) -> dict:
    """A cell that wraps its text over as many lines as the column needs (the cell's lines are kept in "lines")."""
    lines = _wrap(text, width - 16, bold=bold)
    html = "<br>".join(esc(line) for line in lines)
    return dict(html=f"<b>{html}</b>" if bold else html, lines=len(lines), **kw)


def _ws(d: Diagram, x: float, y: float, letters: list[str], widths: list[float], rows: list[dict], number_color: str = GRAY,
        show_letters: bool = True, row_h: float = 22) -> dict:
    """A worksheet with a letter row and rows of any height.  widths[0] is the row-number column; a row is {"n": row number,
    "cells": [cell, ...], "h": height}, its cells taken in order from the first column (a cell with span uses that many columns); a
    cell is a string or a dict: text, bold, italic, align, color, fill, stroke, html (a ready label), span, clip (False keeps a long
    text whole), valign.  A cell made by _wrapped() sets the row's height when the row has no "h".  Returns {(row index, column):
    (x, y, w, h)} and "bottom"."""
    xs = [x]
    for w in widths[:-1]:
        xs.append(xs[-1] + w)
    top = 22 if show_letters else 0
    if show_letters:
        for cx, w, letter in zip(xs, widths, ["", *letters]):
            d.box(esc(letter), cx, y, w, 22, fill=GRAY_TINT, stroke=RULE, color=GRAY, size=SMALL, rounded=False)
    geo: dict = {}
    ry = y + top
    for i, row in enumerate(rows):
        lines = max([c.get("lines", 1) for c in row["cells"] if isinstance(c, dict)] or [1])
        h = row.get("h", max(row_h, lines * 15 + 8))
        d.box(esc(row["n"]), xs[0], ry, widths[0], h, fill=GRAY_TINT, stroke=RULE, color=number_color, size=SMALL, rounded=False)
        cells = list(row["cells"])
        c = idx = 0
        while c < len(widths) - 1:
            cell = cells[idx] if idx < len(cells) else ""
            idx += 1
            if isinstance(cell, str):
                cell = {"text": cell}
            span = cell.get("span", 1)
            cx = xs[c + 1]
            cw = sum(widths[c + 1:c + 1 + span])
            text = cell.get("text", "")
            numeric = text[:1].isdigit() or text[:1] in "-$(" and (text[1:2].isdigit() or text == "-")
            align = cell.get("align", "right" if numeric else "left")
            bold = cell.get("bold", False)
            if "html" in cell:
                label = cell["html"]
            else:
                lead = len(text) - len(text.lstrip(" "))                      # leading spaces are the workbook's indent
                body = text.lstrip(" ")
                shown = _clip(body, cw - 3.5 * lead, bold=bold, margin=cell.get("margin", 26)) if cell.get("clip", not numeric) else body
                label = "&nbsp;" * lead + esc(shown)
                if bold:
                    label = f"<b>{label}</b>"
                if cell.get("italic"):
                    label = f"<i>{label}</i>"
            d.box(label, cx, ry, cw, h, fill=cell.get("fill", WHITE), stroke=cell.get("stroke", RULE), color=cell.get("color", INK), size=SMALL,
                  rounded=False, align=align, valign=cell.get("valign", "middle"))
            geo[(i, c)] = (cx, ry, cw, h)
            c += span
        ry += h
    geo["bottom"] = ry
    return geo


def _outline(d: Diagram, geo: dict, first: tuple[int, int], last: tuple[int, int], color: str = CORAL, inset: float = 0) -> None:
    """An outline around the cells from first to last (row index, column) of a worksheet drawn by _ws, inset so that outlines on
    neighboring cells do not touch."""
    x0, y0, _, _ = geo[first]
    x1, y1, w1, h1 = geo[last]
    left, top = max(0, x0 - 2 + inset), max(0, y0 - 2 + inset)
    right = min(860, x1 + w1 + 2 - inset)
    d.outline(left, top, right - left, y1 + h1 + 2 - inset - top, color, width=2.5, front=True)


def _formula_bar(d: Diagram, x: float, y: float, w: float, name: str, formula: str) -> None:
    """The Name Box, the fx label and the formula bar."""
    d.box(esc(name), x, y, 140 if len(name) > 12 else 90, 28, fill=WHITE, stroke=RULE, size=SMALL, align="left", rounded=False)
    off = 140 if len(name) > 12 else 90
    d.box("<i>fx</i>", x + off + 4, y, 30, 28, fill=WHITE, stroke=WHITE, color=GRAY, size=SMALL, rounded=False)
    d.box(esc(formula), x + off + 38, y, w - off - 38, 28, fill=WHITE, stroke=RULE, size=SMALL, align="left", rounded=False)


def _tabs(d: Diagram, y: float, active: str, names: list[str]) -> float:
    """The sheet tabs of the model workbook around the active one."""
    wb = _cached().sheetnames
    assert all(n in wb for n in names) and active in names, names
    xl.sheet_tabs(d, 0, y, names, active)
    return y + 26


def _excel_format(v: float, fmt: str) -> str:
    """A number as Excel shows it in the number formats the model workbook uses (#,##0.00; 0.0%; 0.000%; #,##0; 0%)."""
    sections = fmt.split(";")
    negative = v < 0
    if negative:
        sec, v = sections[1 if len(sections) > 1 else 0], -v
    elif v == 0 and len(sections) > 2:
        sec = sections[2]
    else:
        sec = sections[0]
    sec = sec.replace("\\", "")
    m = re.search(r"([#0,]+(?:\.0+)?)(%?)", sec)
    assert m, f"number format {fmt!r} is not one the figure knows"
    body, pct = m.groups()
    places = len(body.split(".")[1]) if "." in body else 0
    x = xl.xround(v * 100 if pct else v, places)
    text = f"{x:,.{places}f}" if "," in body else f"{x:.{places}f}"
    out = sec[:m.start()] + text + pct + sec[m.end():]
    return "-" + out if negative and len(sections) == 1 else out


def _show(ws, ref: str) -> str:
    """What the cell shows: text as it is, a number in its number format, TRUE and FALSE, an empty cell empty."""
    cell = ws[ref]
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return str(int(v)) if cell.number_format == "General" and float(v).is_integer() else _excel_format(float(v), cell.number_format)
    return str(v)


def _num(text: str, **kw) -> dict:
    """A numeric cell: right-aligned and never clipped (a sign in front of the digits would make _ws read it as text)."""
    return dict(text=text, align="right", clip=False, **kw)


def _row_of(ws, text: str, col: int = 1) -> int:
    """The row of a worksheet whose cell in `col` starts with `text` (exactly one)."""
    hits = [r for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, col).value, str) and ws.cell(r, col).value.startswith(text)]
    assert len(hits) == 1, (ws.title, text, hits)
    return hits[0]


def _formula_of(sheet: str, ref: str, width: float = 732) -> str:
    """A cell's formula as the formula bar reads it, cut at the bar's edge as a long formula is."""
    return _clip(_g().clean(_g().workbook()[sheet][ref].value), width, size=SMALL, margin=36).replace(" ", " ")


def _key(d: Diagram, x: float, y: float, items: list[tuple[str, str, str]], w: float = 860) -> float:
    """A key of swatches: (label, fill, stroke), the swatch a 22 x 14 box with a 2 px stroke.  Wraps to a new row at the frame's edge.
    Returns the bottom y."""
    cx, cy = x, y
    for label, fill, stroke in items:
        width = text_width(esc(label), SMALL, False) + 8
        if cx > x and cx + 28 + width > x + w:
            cx, cy = x, cy + 22
        d.box("", cx, cy + 3, 22, 14, fill=fill, stroke=stroke, stroke_width=2, rounded=False)
        d.text(esc(label), cx + 28, cy, width, 20, size=SMALL, valign="middle")
        cx += 28 + width + 18
    return cy + 22


# -- IF-19-0-A: the requirements, the files and the milestones ----------------------------------------------------------------

def if_19_0_a() -> Diagram:
    """Ten requirement cards in three phase columns with their tool tags, a milestone under each column, the four model files with the
    hand-offs between them (the query results into the workbook, the refund flags into the report, Chapter 16's register into the
    updated monitoring file) and the report to the committee they all feed."""
    reqs, tools, miles, files = _requirements(), _tools(), _milestones(), _tool_files()
    assert [m["reqs"] for m in miles] == ["1 to 3", "4 to 7", "8 to 10"], miles
    text = _case_text()
    assert "no more than three pages" in text and "appendix of exhibits" in text, "the report's page limits are the case's"
    # the four model files: the Power BI row of the case's table names two
    sql_f, xl_f, pbi_f = files
    names = [sql_f["file"].split(",")[0], xl_f["file"], *[x.strip() for x in pbi_f["file"].split(", and ")]]
    holds = [sql_f["holds"], xl_f["holds"], *[x.strip() for x in pbi_f["holds"].split("; ")]]
    assert names == ["Credits.sql", "Credits Audit.xlsx", "Credits Analytics.pbix", "Audit Monitoring.pbix updated"], names
    assert len(holds) == 4 and "Export to CSV" in sql_f["handoff"] and "refund flags" in pbi_f["handoff"], (holds, files)
    assert "Audit Monitoring.pbix from Chapter 16" in text and "RefundFlags.csv" in text
    d = Diagram("How the Requirements, Files and Milestones Fit Together")
    W, GAP = 273.0, 20.5
    xs = [i * (W + GAP) for i in range(3)]
    bottoms = []
    for p in range(1, 4):
        phase = [r for r in reqs if r["phase"][0] == p]
        d.box(f"<b>Phase {p}</b><br><b>{esc(phase[0]['phase'][1])}</b>", xs[p - 1], 0, W, 40, fill=GRAY, stroke=GRAY, color=WHITE, size=SMALL,
              rounded=False)
        y = 48.0
        for r in phase:
            label = f"<b>Requirement {r['n']}</b><br>{esc(r['title'])}" + (f"<br><i>{esc(r['chapters'])}</i>" if r["chapters"] else "")
            h = kit.height_for(label, W - 4, SMALL, pad=10) + 4 + 30
            d.box(label, xs[p - 1], y, W, h, fill=WHITE, stroke=GRAY, stroke_width=1.5, size=SMALL, align="left", valign="top", rounded=False)
            tx = xs[p - 1] + 6
            for tag in tools[r["n"]]:
                tx += _tag(d, tag, tx, y + h - 27) + 6
            y += h + 8
        bottoms.append(y - 8)
    top = max(bottoms) + 14
    # the milestones, one under each column
    labels = [f"<b>Milestone {m['n']}</b>: {esc(m['title'])}<br>{esc(m['body'][0].upper() + m['body'][1:])}" for m in miles]
    mh = max(kit.height_for(t, W - 10, SMALL, pad=12) for t in labels)
    for m, label in zip(miles, labels):
        d.box(label, xs[m["n"] - 1], top, W, mh, fill=TEAL_TINT, stroke=TEAL, stroke_width=1.5, size=SMALL, align="left", valign="top", rounded=False)
        d.box("", xs[m["n"] - 1], top, 5, mh, fill=TEAL, stroke=TEAL, rounded=False)
    m_bottom = top + mh
    # the four model files, with the hand-offs
    FW, FG = 170.0, 60.0
    fxs = [i * (FW + FG) for i in range(4)]
    fy = m_bottom + 104
    fh = max(kit.height_for(f"<br><b>{esc(n)}</b><br>{esc(h_)}", FW, SMALL, pad=12) for n, h_ in zip(names, holds)) + 4
    tool_of = ["SQL", "Excel", "Power BI", "Power BI"]
    fids = []
    for name, held, fx, tool in zip(names, holds, fxs, tool_of):
        fill, tint_ = TOOL_COLORS[tool]
        fids.append(d.box(f"<br><b>{esc(name)}</b><br>{esc(held)}", fx, fy, FW, fh, fill=tint_, stroke=fill, stroke_width=1.5, size=SMALL,
                          align="left", valign="top", rounded=False))
        _tag(d, tool, fx + 6, fy - 11)
    # Chapter 16's register feeds the updated monitoring file
    cx0 = fxs[3]
    ch16 = d.box("<b>Chapter 16</b>: the register of Audit Monitoring.pbix, if built", cx0, m_bottom + 14, FW, 50, fill=WHITE,
                 stroke=GRAY, stroke_width=1.5, size=SMALL, rounded=False, dashed=True)
    d.arrow(ch16, fids[3], exit=(0.5, 1), entry=(0.5, 0), meta=STRAIGHT)
    d.text("<i>the register</i>", cx0 + FW / 2 + 6, m_bottom + 72, 860 - (cx0 + FW / 2 + 6), 20, size=SMALL, color=GRAY)
    # the query results into the workbook, the refund flags into the report (over the top of the workbook)
    d.arrow(fids[0], fids[1], exit=(1, 0.5), entry=(0, 0.5), meta=STRAIGHT)
    d.text("<i>Copy with Headers</i>", fxs[0] + FW + 2, fy + fh / 2 - 36, FG - 4, 34, size=SMALL, color=GRAY, align="center", valign="bottom")
    run_y = fy - 30
    d.edge(fids[0], fids[2], end="blockThin", color=GRAY, exit=(0.5, 0), entry=(0.5, 0), points=[(fxs[0] + FW / 2, run_y), (fxs[2] + FW / 2, run_y)],
           meta="edgeStyle=none;rounded=0;")
    mid = (fxs[0] + fxs[2] + FW) / 2
    _label(d, "<i>Export to CSV: RefundFlags.csv, the refund flags of Requirement 4</i>", mid - 190, run_y - 28, 380, 20, color=GRAY, align="center")
    # the report to the committee
    ry = fy + fh + 64
    report = d.box("<b>The report to the audit committee</b> (Requirement 9): three pages, with an appendix of three. Every number in it traces to "
                   "one of the four files.", 0, ry, 860, 52, fill=GRAY_TINT, stroke=GRAY, stroke_width=1.5, size=SMALL, rounded=False)
    for fid, fx, lab in zip(fids[1:], fxs[1:], ("evaluation and appendix", "pages and Tests tab", "the continuing tests")):
        d.arrow(fid, report, exit=(0.5, 1), entry=((fx + FW / 2) / 860, 0), meta=STRAIGHT)
        d.text(f"<i>{esc(lab)}</i>", fx + FW / 2 + 6, fy + fh + 4, min(FW / 2 + 22, 860 - (fx + FW / 2 + 6)), 44, size=SMALL, color=GRAY)
    _note(d, "Each card is a requirement of the case, in the order of its three phases; the tags name the tools it uses (Requirement 9 is a written "
             "report built on the workbook). The teal boxes are the three milestones an instructor may collect and check. Arrows show how results move: Copy with "
             "Headers from the script into the workbook, the refund flags exported to CSV from the script into the Power BI report, Chapter 16's register into the "
             "updated monitoring file, and the workbook, the report and the register into the committee's report.", ry + 62)
    return d


# -- IF-19-0-B: the exposure of each deficiency against the three thresholds ---------------------------------------------------

def if_19_0_b() -> Diagram:
    """Horizontal bars to one scale for the last fiscal year: the credits above authority, the credits that fail the authority or the
    cash test counted once, the refunds above authority, the refunds before the payment, the refunds that fail either counted once, and
    all the year's credits; three vertical lines mark the clearly trivial threshold, performance materiality and overall materiality."""
    data = ctx.data()
    n = _n()
    m1, r8 = ctx.values("ch19.m1"), ctx.values("ch19.r8")
    c, t = n.cycle(data), n.tests(data)
    C = data.C
    d1, d3, d4 = n.in_year(t.cm_above, C), n.in_year(t.rf_above, C), n.in_year(t.before, C)
    all_credits = n.in_year(c.cm.values(), C)
    u12_n, u12 = r8["u12"]
    u34_n, u34 = r8["u34"]
    mat, perf, trivial = m1["mat"], m1["perf"], m1["trivial"]
    ws = _book("Deficiencies")
    for ref, v in (("D3", mat), ("D4", perf), ("D5", trivial), ("H8", n.amount(d1)), ("F8", len(d1)), ("H10", n.amount(d3, "amt")), ("F10", len(d3)),
                   ("H11", n.amount(d4, "amt")), ("F11", len(d4)), ("D20", u12), ("E20", n.amount(all_credits)), ("D21", u34)):
        _agree(ws, ref, v, tol=0.01)
    d2 = n.in_year(t.cash_credit, C)
    assert abs(n.amount(d1) + n.amount(d2) - r8["d1"] - n.amount(d2)) < 0.005 and r8["d1"] == n.amount(d1)
    assert trivial < n.amount(d4, "amt") < n.amount(d3, "amt") < u34 < perf < n.amount(d1) < u12 < mat < n.amount(all_credits), \
        "the exposures stand between the thresholds as the guide says"
    _agree(ws, "H9", n.amount(d2), tol=0.01)
    _agree(ws, "F9", len(d2), tol=0.01)
    summed = n.amount(d1) + n.amount(d2)
    assert u12 < summed and summed >= mat > u12 > perf, "added, D1 and D2 reach materiality; counted once they do not"
    bars = [("D1 Credits above the approver's authority", n.amount(d1), f"{len(d1)} credits", "pattern"),
            ("Credits failing D1 or D2, each counted once", u12, f"{u12_n} credits", "once"),
            ("D3 Refunds above the approver's authority", n.amount(d3, "amt"), f"{len(d3)} refunds", "pattern"),
            ("D4 Refunds before the payment", n.amount(d4, "amt"), f"{len(d4)} refunds", "pattern"),
            ("Refunds failing D3 or D4, each counted once", u34, f"{u34_n} refunds", "once"),
            (f"All credits of fiscal {C}: the potential magnitude", n.amount(all_credits), f"{len(all_credits)} credits", "potential")]
    colors = {"pattern": BLUE, "once": TEAL, "potential": GRAY}
    d = Diagram("The Exposure of Each Deficiency Against the Thresholds")
    X0, X1 = 270.0, 820.0
    top_ticks = 300000.0
    x_of = lambda v: X0 + v * (X1 - X0) / top_ticks                                                       # noqa: E731
    PITCH, BH = 56.0, 26.0
    y0 = 76.0                                                                                                # the first bar's row
    plot_h = PITCH * len(bars)
    # gridlines and the axis
    for tick in range(0, 300001, 50000):
        d.box("", x_of(tick) - 0.5, y0 - 8, 1, plot_h + 8, fill=RULE, stroke=RULE, rounded=False)
        d.text(kit.compact(tick), x_of(tick) - 30, y0 + plot_h + 2, 60, 18, size=SMALL, color=GRAY, align="center")
    d.text(f"<i>Fiscal {C} exposure, $</i>", X0, y0 + plot_h + 22, X1 - X0, 18, size=SMALL, color=GRAY, align="center")
    # the bars
    for i, (name, value, count, kind) in enumerate(bars):
        by = y0 + i * PITCH
        d.text(f"<b>{esc(name)}</b><br>{esc(count)}", 0, by - 12, X0 - 14, 50, size=SMALL, valign="middle")
        d.box("", X0, by, x_of(value) - X0, BH, fill=colors[kind], stroke=colors[kind], rounded=False)
    # the thresholds
    lines = [(trivial, "Clearly trivial", GRAY, 1, "dotted", 1), (perf, "Performance materiality", AMBER, 1, "dashed", 1),
             (mat, "Overall materiality", CORAL, 0, "solid", 0)]
    for value, label, color, level, style, _ in lines:
        x = x_of(value)
        meta = {"dotted": "dashed=1;dashPattern=2 4;", "dashed": "dashed=1;dashPattern=8 4;", "solid": ""}[style]
        a, b = d.anchor(x, 20 + 22 * level + 20), d.anchor(x, y0 + plot_h)
        d.edge(a, b, color=color, width=2.5, meta="edgeStyle=none;" + meta)
    for value, label, color, level, _, _ in lines:
        x = x_of(value)
        text = f"{label} {_money(value)}"
        tw = text_width(f"<b>{esc(text)}</b>", SMALL, True) + 16
        lx = min(max(x - tw / 2, X0 - 10), 860 - tw)
        d.box(f"<b>{esc(text)}</b>", lx, 20 + 22 * level, tw, 20, fill=WHITE, stroke=color, stroke_width=1.5, size=SMALL, rounded=False, color=INK)
    # the dollar labels, white-backed so a threshold line does not strike through them
    for i, (name, value, count, kind) in enumerate(bars):
        by = y0 + i * PITCH
        lab = _money(value)
        lw = text_width(lab, SMALL, False) + 16
        _label(d, f"<b>{lab}</b>", x_of(value) + 4, by + 3, lw, 20)
    bottom = y0 + plot_h + 46
    kit.legend(d, 0, bottom, [("Exposure of one pattern", colors["pattern"], "box"), ("Counted once where patterns overlap", colors["once"], "box"),
                              ("Potential magnitude", colors["potential"], "box")])
    _note(d, f"Fiscal {C}, to one scale. An exposure is the amount of the documents that failed a control, not a misstatement. Counted together, D1 and D2 "
             f"(D2 is the {len(d2)} credits whose approver also applied the invoice's cash, {_money(n.amount(d2))}) would add to {_money(summed)}, which reaches materiality; "
             f"counted once, the {u12_n} credits that fail either test are {_money(u12)}, between performance materiality and materiality. The refund patterns stay "
             f"between the clearly trivial threshold and performance materiality, while the year's {len(all_credits)} credits, {_money(n.amount(all_credits))}, are the "
             f"potential magnitude of the credit patterns: a rating rests on likelihood, magnitude and compensating controls, not on how many exceptions a test finds.", bottom + 28)
    return d


# -- IF-19-1-A: the map of the cycle, marked up -----------------------------------------------------------------------------------

def _badge(d: Diagram, label: str, x: float, y: float, style: str, size: float = 26) -> str:
    """A control point's badge: filled teal (the records evidence it), tinted with an amber ring (they evidence part of it), white
    with a dashed coral ring (the evidence lies outside the database)."""
    fill, stroke, color, dash = {"full": (TEAL, TEAL, WHITE, ""), "partial": (AMBER_TINT, AMBER, INK, ""),
                                 "outside": (WHITE, CORAL, INK, "dashed=1;dashPattern=3 2;")}[style]
    return d.vertex(f"<b>{esc(label)}</b>", f"ellipse;whiteSpace=wrap;html=1;fillColor={fill};strokeColor={stroke};strokeWidth=2.5;{dash}fontColor={color};"
                    f"fontSize={SMALL};{FONT}align=center;verticalAlign=middle;spacing=0;", x, y, size, size)


def _control_points() -> list[dict]:
    """The nine control points of the Cycle Map sheet, each with what a well-designed cycle checks, what the records evidence, what they
    cannot show and its style: 'full' when the sheet says there is nothing material outside the records, 'outside' when every row of the
    risk-and-control matrix that cites the point says the evidence lies (mostly) outside the database, 'partial' otherwise."""
    ws, risks = _book("Cycle Map"), _book("Risks")
    first = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Control point") + 1
    matrix = _row_of(risks, "Risk-and-control matrix") + 2
    out = []
    for r in range(first, first + 9):
        point = ws.cell(r, 1).value
        assert re.fullmatch(r"C[1-9]", point), point
        where = [str(risks.cell(k, 7).value) for k in range(matrix, risks.max_row + 1)
                 if risks.cell(k, 3).value and point in re.split(r",\s*", str(risks.cell(k, 3).value))]
        assert where, f"the risk-and-control matrix cites {point}"
        cannot = str(ws.cell(r, 8).value)
        style = ("full" if cannot.startswith("Nothing material") else
                 "outside" if all(w.startswith(("Mostly outside", "Outside")) for w in where) else "partial")
        out.append(dict(point=point, checks=ws.cell(r, 2).value, evidence=ws.cell(r, 3).value, cannot=cannot, style=style, where=where))
    return out


def if_19_1_a() -> Diagram:
    """The book's map of the cycle: five documents with their departments, the postings and the hub, and the nine control points drawn in
    three styles (the records evidence it, they evidence part of it, the evidence lies outside the database), with what the records
    cannot show written beside each badge in the list below."""
    data = ctx.data()
    n = _n()
    c = n.cycle(data)
    points = _control_points()
    styles = {p["point"]: p["style"] for p in points}
    assert styles == {"C1": "partial", "C2": "partial", "C3": "full", "C4": "partial", "C5": "partial", "C6": "partial", "C7": "partial",
                      "C8": "outside", "C9": "partial"}, styles
    # the departments and the postings, asserted against the records
    cc = lambda ids: sorted({c.emp[e]["cc"] for e in ids})                                                     # noqa: E731
    dept = {"CashReceipt": cc(r["rec"] for r in c.rcpt.values()), "SalesReturn": cc(s["recv"] for s in c.sr.values()),
            "CreditMemo": cc(m["appr"] for m in c.cm.values()), "SalesCommissionAdjustment": cc(a["appr"] for a in c.adj),
            "CustomerRefund": cc(r["appr"] for r in c.rf.values())}
    assert dept == {"CashReceipt": ["Customer Service"], "SalesReturn": ["Warehouse"], "CreditMemo": ["Customer Service"],
                    "SalesCommissionAdjustment": ["Customer Service"], "CustomerRefund": ["Administration"]}, dept
    assert cc(a[2] for v in c.apps.values() for a in v) == ["Customer Service"], "customer service applies every payment"
    post = {k: n.sides(data, k) for k in ("CashReceipt", "CashReceiptApplication", "SalesReturn", "CreditMemo", "SalesCommissionAdjustment", "CustomerRefund")}
    assert post["CashReceipt"] == {("Dr", "1010"), ("Cr", "2060")} and post["CashReceiptApplication"] == {("Dr", "2060"), ("Cr", "1020")}
    assert post["SalesReturn"] == {("Dr", "1040"), ("Cr", "5010"), ("Cr", "5020"), ("Cr", "5030"), ("Cr", "5040")}
    assert post["CreditMemo"] == {("Dr", "4060"), ("Dr", "4050"), ("Dr", "2050"), ("Cr", "1020"), ("Cr", "2060")}
    assert post["SalesCommissionAdjustment"] == {("Dr", "2034"), ("Cr", "6290")} and post["CustomerRefund"] == {("Dr", "2060"), ("Cr", "1010")}
    docs = [("CashReceipt and its applications", dept["CashReceipt"][0], "cash received and applied to invoices", ["C9"],
             "Receipt: Dr 1010, Cr <b>2060</b><br>Applied: Dr <b>2060</b>, Cr 1020"),
            ("SalesReturn", dept["SalesReturn"][0], "goods received back", ["C1"], "Dr 1040 Inventory<br>Cr 5010–5040 Cost of goods sold"),
            ("CreditMemo", dept["CreditMemo"][0], "credit approved", ["C2", "C3", "C4"], "Dr 4060, 4050, 2050<br>Cr 1020 and/or <b>2060</b>"),
            ("SalesCommission<br>Adjustment", dept["SalesCommissionAdjustment"][0], "commission clawed back", ["C5"],
             "Dr 2034 Commission payable<br>Cr 6290 Commission expense"),
            ("CustomerRefund", dept["CustomerRefund"][0], "refund approved and paid", ["C6", "C7", "C8"], "Dr <b>2060</b><br>Cr 1010 Cash")]
    assert sorted(p for _, _, _, ps, _ in docs for p in ps) == [p["point"] for p in points]
    d = Diagram("The Customer Credits Cycle, Marked")
    BOTTOM, TOP, LEFT, RIGHT = (0.5, 1), (0.5, 0), (0, 0.5), (1, 0.5)
    w, gap, y, h = 156.0, 20.0, 40.0, 104.0
    boxes, posts = [], []
    for i, (name, who, what, ctrl, entry) in enumerate(docs):
        x = i * (w + gap)
        boxes.append(d.box(f"<b>{name}</b><br><i>{what}</i>", x, y, w, h, fill=BLUE_TINT, stroke=BLUE, size=SMALL, valign="top"))
        d.text(f"<b>{esc(who)}</b>", x, y - 22, w, 20, size=SMALL, color=GRAY, align="left", valign="bottom")
        for j, p in enumerate(ctrl):
            _badge(d, p, x + 8 + 32 * j, y + h - 34, styles[p])
        posts.append(d.box(entry, x, 232, w, 70, fill=AMBER_TINT, stroke=AMBER, size=SMALL))
        d.arrow(boxes[-1], posts[-1], color=AMBER, dashed=True, exit=BOTTOM, entry=TOP)
    d.arrow(boxes[1], boxes[2], exit=RIGHT, entry=LEFT)
    d.arrow(boxes[2], boxes[3], exit=RIGHT, entry=LEFT)
    d.arrow(boxes[2], boxes[4], exit=(0.8, 0), entry=(0.88, 0), points=[(2 * (w + gap) + 0.8 * w, 12), (4 * (w + gap) + 0.88 * w, 12)])
    d.arrow(boxes[0], boxes[2], exit=(0.25, 1), entry=(0.2, 1), points=[(0.25 * w, 172), (2 * (w + gap) + 0.2 * w, 172)])
    d.text("<i>has the invoice been paid?</i>", 0.5 * w + 14, 174, 190, 20, size=SMALL, color=GRAY)
    d.text("<b>Account 2060 Customer Deposits and Unapplied Cash</b> joins receipts, credits on paid invoices and refunds: every receipt passes "
           "through it, and every refund is paid from it.", 0, 312, 860, 40, align="left", size=SMALL)
    # the control points and what the records cannot show of each
    y0 = 366.0
    d.text("<b>Control points: what a well-designed cycle checks, and what the records cannot show</b>", 0, y0, 860, 22, size=SMALL)
    CW, EX = 316.0, 372.0
    ry = y0 + 28
    for p in points:
        lines = max(len(_wrap(p["checks"], CW - 12)), len(_wrap(p["cannot"], 860 - EX - 12)), 1)
        rh = lines * 15 + 10
        _badge(d, p["point"], 0, ry + (rh - 26) / 2, p["style"])
        d.text(esc(p["checks"]), 34, ry, CW, rh, size=SMALL, valign="middle")
        d.text(f"<i>{esc(p['cannot'])}</i>", EX, ry, 860 - EX, rh, size=SMALL, color=GRAY, valign="middle")
        d.box("", 34, ry + rh, 826, 1, fill=RULE, stroke=RULE, rounded=False)
        ry += rh + 2
    ly = ry + 14
    cx = 0.0
    for style, label in (("full", "The records evidence it"), ("partial", "The records evidence part of it: the rest is written beside the badge"),
                         ("outside", "The evidence lies outside the database")):
        _badge(d, "C", cx, ly, style, 22)
        width = text_width(label, SMALL, False) + 12
        d.text(esc(label), cx + 28, ly, width, 22, size=SMALL, valign="middle")
        cx += 28 + width + 14
    d.line_sample(0, ly + 44, 44, color=GRAY, end="blockThin")
    d.text("Document flow", 52, ly + 34, 120, 20, size=SMALL, valign="middle")
    d.line_sample(190, ly + 44, 44, color=AMBER, dashed=True, end="blockThin")
    d.text("Posting to the ledger", 242, ly + 34, 180, 20, size=SMALL, valign="middle")
    counts = {s: sum(1 for p in points if p["style"] == s) for s in ("full", "partial", "outside")}
    _note(d, f"The model answer's marking, from the Cycle Map and Risks sheets: {counts['full']} control point is evidenced by the records (C3, whose recomputation "
             f"is complete in them), {counts['partial']} are evidenced in part, and {counts['outside']} (C8) rests mostly on evidence in no table. The badge style "
             f"carries the reading with its fill and its outline, not its color alone. Departments are the cost centers of the people each document names.", ly + 66)
    return d


# -- IF-19-2-A: the roll-forward of account 2060 --------------------------------------------------------------------------------

def if_19_2_a() -> Diagram:
    """Rows 4 and 55 to 63 of the Reconciliation sheet: the heading row with the three years and the block for account 2060 (the opening
    balance, the receipts held, the applications, the credits on paid invoices, the refunds, the closing balance rolled forward, the
    ledger balance and the difference), with the flows, the balances and the difference outlined apart."""
    data = ctx.data()
    m1 = ctx.values("ch19.m1")
    ws = _book("Reconciliation")
    g6 = _g().test(data, "G6")
    years = list(data.years)
    assert [r[0] for r in g6] == years and [ws[f"{c}4"].value for c in "CDE"] == years, "the three fiscal years head the columns"
    first = _row_of(ws, "Account 2060 Customer Deposits and Unapplied Cash, rolled forward")
    expect = {first + 1: ("Opening balance", [r[1] for r in g6]), first + 2: ("Receipts held", [r[2] for r in g6]),
              first + 3: ("Applied to invoices", [r[4] for r in g6]), first + 4: ("Credits on paid invoices", [r[3] for r in g6]),
              first + 5: ("Refunds paid", [r[5] for r in g6]), first + 6: ("Closing balance, rolled forward", [r[7] for r in g6]),
              first + 7: ("Ledger balance of 2060", list(m1["b2060"])), first + 8: ("Difference, rolled forward less ledger", [0.0] * 3)}
    for r, (label, vals) in expect.items():
        assert str(ws[f"A{r}"].value).startswith(label), (r, ws[f"A{r}"].value)
        for col, v in zip("CDE", vals):
            _agree(ws, f"{col}{r}", v, tol=0.01)
    assert all(r[6] == 0 for r in g6), "nothing but the four flows posts to 2060"
    flows = [abs(ws[f"E{r}"].value) for r in range(first + 2, first + 6)]
    balances = [abs(ws[f"{c}{first + 6}"].value) for c in "CDE"]
    assert min(flows[:2]) > 1000 * max(balances), "the flows are thousands of times the balances"
    assert all(abs(ws[f"{c}{first + 8}"].value) < 0.005 for c in "CDE")
    formula = _g().clean(_g().workbook()["Reconciliation"][f"C{first + 8}"].value)
    assert formula == f"=C{first + 6}-C{first + 7}", formula
    widths = [28, 400, 124, 124, 124]
    letters = ["A", "C", "D", "E"]
    d = Diagram("The Roll-Forward of Account 2060")
    _formula_bar(d, 0, 0, 800, f"C{first + 8}", formula)
    sy = 40.0
    _ws(d, 0, sy, letters, widths, [], show_letters=True)
    head = dict(n="4", cells=[dict(text="Item", bold=True)] + [dict(text=str(y), bold=True, align="right", clip=False) for y in years])
    geo0 = _ws(d, 0, sy + 22, letters, widths, [head], show_letters=False)
    y = _gap(d, 0, geo0["bottom"], sum(widths), f"Rows 5 to {first - 1}: each population reconciled to the ledger by year, with the year-end closes excluded")
    rows = [dict(n=str(first), cells=[dict(text=ws[f"A{first}"].value, bold=True, span=4, clip=False)])]
    for r in range(first + 1, first + 9):
        bold = r >= first + 6
        rows.append(dict(n=str(r), cells=[dict(text=ws[f"A{r}"].value, bold=bold, clip=True)]
                         + [_num(_show(ws, f"{c}{r}"), bold=bold) for c in "CDE"]))
    geo = _ws(d, 0, y, letters, widths, rows, show_letters=False)
    idx = lambda r: r - first                                                                        # noqa: E731
    _outline(d, geo, (idx(first + 2), 0), (idx(first + 5), 3), TEAL, inset=1.5)
    _outline(d, geo, (idx(first + 6), 0), (idx(first + 7), 3), BLUE, inset=1.5)
    _outline(d, geo, (idx(first + 8), 0), (idx(first + 8), 3), CORAL, inset=1.5)
    bottom = _gap(d, 0, geo["bottom"], sum(widths), f"Rows {first + 9} to {ws.max_row}: the items in 2060 at each year-end, the reconciliation of 2034, the exclusions and the model answer")
    bottom = _tabs(d, bottom + 8, "Reconciliation", ["Documentation", "Cycle Map", "Reconciliation", "Materiality", "Risks"])
    top_flow = max(abs(r[2]) for r in g6)
    top_bal = max(abs(b) for b in m1["b2060"])
    _note(d, f"Outlined in teal: the four flows through 2060, the receipts held and the applications to invoices, each about {top_flow / 1e6:.0f} million a year, then the credits "
             f"on paid invoices and the refunds, tens of thousands. In blue: the closing balance rolled forward and the ledger's, a few thousand dollars at most ({_money(top_bal)}). In coral: "
             f"the difference, which Excel shows as -0.00 because each cell is a floating-point residue far below a cent; the model answer tests it with ABS(x)&lt;0.005. The opening "
             f"balance of each year is the closing balance above it (D{first + 1} is =C{first + 6}). Columns B (Source, empty in this block) and F (Total) are hidden, and column A is "
             f"narrower than the workbook's.", bottom + 4)
    return d


# -- IF-19-3-A: the Materiality sheet ----------------------------------------------------------------------------------------

def if_19_3_a() -> Diagram:
    """Rows 1 to 25 of the Materiality sheet: the parameter block with the income, the three rates and the three thresholds (the named
    cells), the three-year block of revenue, returns and return rate, then the expectation, the recorded returns, the difference, the
    two comparisons and the conclusion."""
    data = ctx.data()
    n = _n()
    m1 = ctx.values("ch19.m1")
    g = _g()
    ws = _book("Materiality")
    years = list(data.years)
    ex = n.expectation(data)
    for ref, v in (("B4", m1["income"]), ("B5", 0.05), ("B6", m1["mat"]), ("B7", 0.75), ("B8", m1["perf"]), ("B9", 0.05), ("B10", m1["trivial"]),
                   ("B19", ex["exp"]), ("B20", ex["actual"]), ("B21", ex["diff"]), ("B22", ex["diff_pct"])):
        _agree(ws, ref, v, tol=0.01 if abs(v) > 1 else 1e-9)
    for col, k in zip("BCD", range(3)):
        _agree(ws, f"{col}14", years[k], tol=0.5)
        _agree(ws, f"{col}15", ex["revenue"][k], tol=0.01)
        _agree(ws, f"{col}16", n.returns_4060(data)[k], tol=0.01)
        _agree(ws, f"{col}17", ex["rates"][k], tol=1e-9)
    assert [ws[f"B{r}"].value for r in (23, 24)] == ["Above clearly trivial", "Below performance materiality"]
    concl = ws["B25"].value
    assert concl.startswith("Investigate (Requirement 5)"), concl
    ok = lambda cond, msg: None                                                                      # noqa: E731
    for name, ref in (("Materiality", "$B$6"), ("PerformanceMateriality", "$B$8"), ("ClearlyTrivial", "$B$10")):
        assert g.defined(ok, name) == f"Materiality!{ref}", name
    assert g.formula(ok, "Materiality", "B6") == "=B4*B5"
    assert m1["trivial"] < ex["diff"] < m1["perf"], "the difference is above clearly trivial and below performance materiality"
    # the columns: B to E as wide as the conclusion needs (it spills over C to E), column A cut where its neighbor begins
    need = text_width(f"<b>{esc(concl)}</b>", SMALL, True) + 24
    w = max(118, int(need / 4) + 1)
    widths = [28, 860 - 28 - 4 * w, w, w, w, w]
    letters = ["A", "B", "C", "D", "E"]
    d = Diagram("The Materiality Sheet")
    _formula_bar(d, 0, 0, 860, "Materiality", "=B4*B5")
    sy = 40.0
    _ws(d, 0, sy, letters, widths, [], show_letters=True)

    def val(ref: str, **kw) -> dict:
        return _num(_show(ws, ref), **kw)

    def lab(r: int, **kw) -> dict:
        return dict(text=ws[f"A{r}"].value, clip=True, **kw)
    rows = [dict(n="1", cells=[dict(text=ws["A1"].value, bold=True, span=5, clip=False)]),
            dict(n="2", cells=[dict(text=ws["A2"].value, span=5, clip=True)]),
            dict(n="3", cells=[dict(text="Parameter", bold=True), dict(text="Value", bold=True, align="right", clip=False), dict(text="Basis", bold=True)]),
            dict(n="4", cells=[lab(4), val("B4"), dict(text=ws["C4"].value, span=3, clip=False)]),
            dict(n="5", cells=[lab(5), val("B5"), dict(text=ws["C5"].value, span=3, clip=False)]),
            dict(n="6", cells=[lab(6, bold=True), val("B6", bold=True)]),
            dict(n="7", cells=[lab(7), val("B7"), dict(text=ws["C7"].value, span=3, clip=False, align="left")]),
            dict(n="8", cells=[lab(8, bold=True), val("B8", bold=True)]),
            dict(n="9", cells=[lab(9), val("B9"), dict(text=ws["C9"].value, span=3, clip=False, align="left")]),
            dict(n="10", cells=[lab(10, bold=True), val("B10", bold=True)]),
            dict(n="11", cells=[dict(text=ws["A11"].value, span=5, clip=True)]),
            dict(n="12", cells=[]),
            dict(n="13", cells=[dict(text=ws["A13"].value, bold=True, span=5, clip=False)]),
            dict(n="14", cells=[dict(text="Item", bold=True)] + [_num(str(y), bold=True) for y in years]),
            dict(n="15", cells=[lab(15)] + [val(f"{c}15") for c in "BCD"]),
            dict(n="16", cells=[lab(16)] + [val(f"{c}16") for c in "BCD"]),
            dict(n="17", cells=[lab(17)] + [val(f"{c}17") for c in "BCD"]),
            dict(n="18", cells=[]),
            dict(n="19", cells=[lab(19), val("B19")]),
            dict(n="20", cells=[lab(20), val("B20")]),
            dict(n="21", cells=[lab(21), val("B21")]),
            dict(n="22", cells=[lab(22), val("B22")]),
            dict(n="23", cells=[lab(23), dict(text=ws["B23"].value, span=2, clip=False)]),
            dict(n="24", cells=[lab(24), dict(text=ws["B24"].value, span=2, clip=False)]),
            dict(n="25", cells=[lab(25, bold=True), dict(text=concl, bold=True, span=4, clip=False)])]
    geo = _ws(d, 0, sy + 22, letters, widths, rows, show_letters=False)
    idx = lambda r: r - 1                                                                        # noqa: E731
    # the three named cells outlined, each with its name beside it (the tags are labels, not cells)
    for r, name in ((6, "Materiality"), (8, "PerformanceMateriality"), (10, "ClearlyTrivial")):
        _outline(d, geo, (idx(r), 1), (idx(r), 1), TEAL, inset=1.5)
        x0, y0, w0, h0 = geo[(idx(r), 2)]
        tw = text_width(f"<b>Name: {name}</b>", SMALL, True) + 28
        d.box(f"<b>Name: {name}</b>", x0 + 6, y0 + 2, tw, h0 - 4, fill=TEAL, stroke=TEAL, color=WHITE, size=SMALL)
    _outline(d, geo, (idx(19), 0), (idx(22), 1), BLUE, inset=1.5)
    _outline(d, geo, (idx(23), 0), (idx(25), 1), CORAL, inset=1.5)
    bottom = _gap(d, 0, geo["bottom"], sum(widths), f"Rows 26 to {ws.max_row}: the expectation at the {years[0]} rate and the model answer")
    bottom = _tabs(d, bottom + 8, "Materiality", ["Cycle Map", "Reconciliation", "Materiality", "Risks", "Tests"])
    _note(d, f"Outlined in teal: the three threshold cells, each with its name written beside it as a label (in the workbook the name shows in the Name Box, here that of B6); "
             f"every later formula reads the name, not the cell. In blue: the expectation at the {years[1]} rate, the recorded returns, the difference "
             f"({_money(ex['diff'])}, {ex['diff_pct'] * 100:.1f}% of the expectation) and its percentage. In coral: the two comparisons and the conclusion, which says to investigate "
             f"because the difference is above the clearly trivial threshold and below performance materiality. Column A is narrower than the workbook's, so a long label is cut "
             f"where its neighbor's value begins.", bottom + 4)
    return d


# -- IF-19-4-A: the dates around one refund ----------------------------------------------------------------------------------

def _day(iso: str) -> int:
    return date.fromisoformat(iso).toordinal()


def _pick_refund(data) -> dict:
    """The first refund the two readings of 'before the payment' part on (before the last application, not before the last receipt)
    whose invoice has two receipts, each applied one or two days after its receipt date."""
    n = _n()
    c, t = n.cycle(data), n.tests(data)
    by_receipt = {r["id"] for r in t.before_rcpt}
    parted = sorted((r for r in t.before if r["id"] not in by_receipt), key=lambda r: r["id"])
    for r in parted:
        inv = c.cm[r["cm"]]["inv"]
        apps = sorted(c.apps[inv])
        if len(apps) == 2 and all(1 <= _day(a[0]) - _day(c.rcpt[a[3]]["date"]) <= 2 for a in apps):
            return dict(refund=r, parted=len(parted))
    raise AssertionError("no refund of the BeforePayment test has two receipts applied a day or two after each")


def if_19_4_a() -> Diagram:
    """One invoice's dates in three lanes (the invoice and its due date, the cash with each application, the credit and the refund) on one
    scale with the refund date marked through them, and, enlarged below, the four days around the refund with three brackets for the
    three readings of 'before the payment'."""
    data = ctx.data()
    n = _n()
    c, t = n.cycle(data), n.tests(data)
    r4 = ctx.values("ch19.r4")
    pick = _pick_refund(data)
    rf, parted = pick["refund"], pick["parted"]
    cm = c.cm[rf["cm"]]
    inv = c.inv[cm["inv"]]
    apps = sorted(c.apps[inv["id"]])
    rcpt = [c.rcpt[a[3]] for a in apps]
    row = next(x for x in _g().test(data, "T6") if x[0] == rf["num"])
    last_app, last_rcpt = apps[-1][0], max(r["date"] for r in rcpt)
    assert row[4] == inv["num"] and row[5] == inv["due"] and row[6] == last_app and row[7] == last_rcpt and row[10] == 0, row
    assert last_rcpt == rf["date"] < last_app, "the refund carries the date of the last receipt and comes before the last application"
    assert all(len(x) == 10 for x in (rf["date"], last_app, last_rcpt)), "the records hold dates, not times"
    assert len(t.before) - len(t.before_rcpt) == parted == r4["same"]["n"], "the refunds the definitions part on"
    assert rcpt[0]["date"] < inv["due"] < cm["date"] < rf["date"] and inv["date"] < rcpt[0]["date"] < apps[0][0] < inv["due"]
    gap_days = _day(last_app) - _day(rf["date"])
    d = Diagram("The Dates Around One Refund")
    x0, x1 = 128.0, 800.0
    start, end = _day(inv["date"]), _day(last_app)
    xt = lambda iso: x0 + (_day(iso) - start) * (x1 - x0) / (end - start)                                     # noqa: E731
    ref_x = xt(rf["date"])
    lanes = [("The invoice", 70.0), ("The cash", 164.0), ("Credit and refund", 252.0)]
    for name, ay in lanes:
        d.text(f"<b>{esc(name)}</b>", 0, ay - 11, 116, 22, size=SMALL, valign="middle")
        d.box("", x0 - 8, ay - 1, x1 - x0 + 16, 2, fill=RULE, stroke=RULE, rounded=False)

    def mark(x: float, ay: float, kind: str, color: str, size: float = 14) -> None:
        if kind == "dot":
            d.vertex("", f"ellipse;html=1;fillColor={color};strokeColor={WHITE};strokeWidth=1.5;", x - size / 2, ay - size / 2, size, size)
        elif kind == "diamond":
            d.vertex("", f"rhombus;html=1;fillColor={WHITE};strokeColor={color};strokeWidth=2.5;", x - size / 2 - 1, ay - size / 2 - 1, size + 2, size + 2)
        else:
            d.box("", x - size / 2, ay - size / 2, size, size, fill=color, stroke=WHITE, stroke_width=1.5, rounded=False)

    def tag(lines: list[str], x: float, ay: float, w: float, align: str = "left", below: bool = False) -> None:
        text = "<br>".join((f"<b>{esc(ln)}</b>" if i == 0 else esc(ln)) for i, ln in enumerate(lines))
        h = 15 * len(lines) + 6
        y = ay + 14 if below else ay - 14 - h
        d.text(text, x, y, w, h, size=SMALL, align=align)
    a1, a2, a3 = lanes[0][1], lanes[1][1], lanes[2][1]
    # the refund date through the three lanes
    d.box("", ref_x - 0.75, 30, 1.5, a3 + 24 - 30, fill=CORAL, stroke=CORAL, rounded=False)
    d.box(f"<b>Refund date {_date(rf['date'])}</b>", ref_x - 86, 6, 172, 24, fill=CORAL_TINT, stroke=CORAL, stroke_width=1.5, size=SMALL, rounded=False)
    # lane 1: the invoice and its due date
    mark(xt(inv["date"]), a1, "dot", BLUE)
    tag([f"Invoice {inv['num']}", inv["date"]], xt(inv["date"]) - 8, a1, 190)
    mark(xt(inv["due"]), a1, "dot", BLUE)
    tag(["Due date", inv["due"]], xt(inv["due"]) - 70, a1, 140, "center")
    # lane 2: each receipt with the application of it
    for k, (app, rc) in enumerate(zip(apps, rcpt)):
        mark(xt(rc["date"]), a2, "dot", TEAL)
        mark(xt(app[0]), a2, "diamond", TEAL)
        lag = _day(app[0]) - _day(rc["date"])
        lines = [f"Receipt {rc['num']}", f"{rc['date']}, applied {app[0]}", f"{lag} day{'s' if lag != 1 else ''} later, {_money(app[1])}"]
        if k == 0:
            tag(lines, xt(rc["date"]) - 8, a2, 230)
        else:
            tag(lines, ref_x - 12 - 230, a2, 230, "right")
    # lane 3: the credit and the refund
    mark(xt(cm["date"]), a3, "square", AMBER)
    tag([f"Credit memo {cm['num']}", cm["date"]], xt(cm["date"]) - 100, a3, 200, "center")
    mark(ref_x, a3, "square", CORAL)
    tag([f"Refund {rf['num']}", f"{_money(rf['amt'])}"], ref_x - 12 - 190, a3, 190, "right", below=True)
    # the enlarged stretch: the days around the refund, with the three readings
    z0, z1 = _day(rf["date"]) - 1, _day(last_app) + 1
    py = a3 + 78.0
    pbottom = py + 338
    d.box("", 0, py, 860, pbottom - py, fill=WHITE, stroke=GRAY, stroke_width=1.5, rounded=False, dashed=True)
    arrow_x = (ref_x + xt(last_app)) / 2
    a_from, a_to = d.anchor(arrow_x, a3 + 20), d.anchor(arrow_x, py)
    d.edge(a_from, a_to, color=GRAY, width=1.5, end="blockThin", meta=STRAIGHT)
    d.text("<i>enlarged below</i>", arrow_x - 112, py - 22, 100, 18, size=SMALL, color=GRAY, align="right")
    d.text(f"<i>Enlarged, one tick a day: {date.fromordinal(z0).isoformat()} to {date.fromordinal(z1).isoformat()}</i>", 8, py + 4, 480, 20, size=SMALL, color=GRAY)
    px0, px1 = 60.0, 820.0
    zt = lambda iso: px0 + (_day(iso) - z0) * (px1 - px0) / (z1 - z0)                                            # noqa: E731
    zay = py + 78
    d.box("", px0 - 10, zay - 1, px1 - px0 + 20, 3, fill=INK, stroke=INK, rounded=False)
    for k in range(z0, z1 + 1):
        iso = date.fromordinal(k).isoformat()
        d.box("", zt(iso) - 0.5, zay - 5, 1, 11, fill=INK, stroke=INK, rounded=False)
        d.text(_date(iso), zt(iso) - 40, zay + 22, 80, 18, size=SMALL, color=GRAY, align="center")
    x_ref, x_app = zt(rf["date"]), zt(last_app)
    mark(x_ref, zay - 12, "dot", TEAL, 16)
    mark(x_ref, zay + 12, "square", CORAL, 16)
    mark(x_app, zay, "diamond", TEAL, 14)
    d.text(f"<b>Receipt {esc(rcpt[-1]['num'])}</b>", x_ref - 14 - 190, zay - 30, 190, 18, size=SMALL, align="right", valign="middle")
    d.text(f"<b>Refund {esc(rf['num'])}</b>", x_ref - 14 - 190, zay - 1, 190, 18, size=SMALL, align="right", valign="middle")
    d.text(f"<b>Application</b><br>of that receipt", x_app - 70, zay - 50, 140, 36, size=SMALL, align="center")
    # the three readings as brackets
    lane_y = [zay + 80, zay + 134, zay + 188]
    for xx in (x_ref, x_app):
        d.box("", xx - 0.5, zay + 40, 1, lane_y[0] - zay - 40, fill=GRAY, stroke=GRAY, rounded=False)
    specs = [(lane_y[0], x_ref, x_app, CORAL, False, "Last application (the case's definition)",
              f"The last application, {_date(last_app)}, is {gap_days} days after the refund: counted as before the payment"),
             (lane_y[1], x_ref - 22, x_ref + 22, TEAL, False, "Last receipt",
              f"The last receipt, {_date(last_rcpt)}, is the refund's own date: not counted"),
             (lane_y[2], x_ref - 22, x_ref + 22, GRAY, True, "The same day, the order unknown",
              f"A receipt and a refund of {_date(rf['date'])}: the records hold dates and not times, so which came first is unknown")]
    for y, xa, xb, color, dashed, name, text in specs:
        meta = "dashed=1;dashPattern=4 3;" if dashed else ""
        for xx in (xa, xb):
            a, b = d.anchor(xx, y), d.anchor(xx, y + 10)
            d.edge(a, b, color=color, width=2.5, meta="edgeStyle=none;" + meta)
        a, b = d.anchor(xa, y + 10), d.anchor(xb, y + 10)
        d.edge(a, b, color=color, width=2.5, meta="edgeStyle=none;" + meta)
        if xb - xa > 100:
            d.text(f"<b>{esc(name)}</b><br>{text}", xa, y + 14, 860 - xa - 8, 36, size=SMALL)
        else:
            d.text(f"<b>{esc(name)}</b><br>{text}", xb + 12, y - 6, 860 - xb - 20, 36, size=SMALL)
    # a key for the markers
    ky = pbottom + 10
    cx = 0.0
    for kind, color, label in (("dot", TEAL, "A receipt, by its ReceiptDate"), ("diamond", TEAL, "An application of it to the invoice, by its ApplicationDate"),
                               ("square", CORAL, "The refund")):
        if kind == "dot":
            d.vertex("", f"ellipse;html=1;fillColor={color};strokeColor={WHITE};strokeWidth=1.5;", cx, ky + 3, 14, 14)
        elif kind == "diamond":
            d.vertex("", f"rhombus;html=1;fillColor={WHITE};strokeColor={color};strokeWidth=2.5;", cx - 1, ky + 2, 16, 16)
        else:
            d.box("", cx, ky + 3, 14, 14, fill=color, stroke=WHITE, stroke_width=1.5, rounded=False)
        width = text_width(label, SMALL, False) + 10
        d.text(esc(label), cx + 22, ky, width, 20, size=SMALL, valign="middle")
        cx += 22 + width + 16
    _note(d, f"Refund {rf['num']} of {_money(rf['amt'])} is the first of the {parted} refunds that the last-application reading counts and the last-receipt reading does not (the "
             f"{len(t.before)} and {len(t.before_rcpt)} of Requirement 4) whose invoice has two receipts applied a day or two after each. Its credit, {cm['num']}, was approved "
             f"before the second receipt and posted to 2060, so the refund can come before the payment it depends on. No table holds the date a payment cleared the bank, so "
             f"none of the three readings can say when the money was usable.", ky + 28)
    return d


# -- Power BI pieces: the Credits Analytics project -----------------------------------------------------------------------------

def _page_ids() -> list[str]:
    return json.loads((REPORT / "pages" / "pages.json").read_text(encoding="utf-8"))["pageOrder"]


def _page(page: str) -> dict:
    return json.loads((REPORT / "pages" / page / "page.json").read_text(encoding="utf-8"))


def _page_name(page: str) -> str:
    return _page(page)["displayName"]


def _visual(page: str, name: str) -> dict:
    return json.loads((REPORT / "pages" / page / "visuals" / name / "visual.json").read_text(encoding="utf-8"))


def _title(page: str, name: str) -> str:
    """The title the model answer gives a visual (a doubled apostrophe in the file is one in the title)."""
    value = _visual(page, name)["visual"]["visualContainerObjects"]["title"][0]["properties"]["text"]["expr"]["Literal"]["Value"]
    return value.strip("'").replace("''", "'")


def _projections(page: str, name: str, role: str) -> list[str]:
    return [p["nativeQueryRef"] for p in _visual(page, name)["visual"]["query"]["queryState"][role]["projections"]]


def _query_tabs() -> list[str]:
    """The DAX query tabs of the project in their order (Tests, Checks)."""
    return json.loads((MODEL / ".pbi" / "daxQueries.json").read_text(encoding="utf-8"))["tabOrder"]


def _expected(label: str) -> float:
    """The Expected value of one row of the project's Checks tab (DAX query view), found by its label."""
    text = (MODEL / "DAXQueries" / "Checks.dax").read_text(encoding="utf-8")
    m = re.search(r'"Check", "' + re.escape(label) + r'", "Expected", ([-0-9.eE]+),', text)
    assert m, f"the Checks tab has no row {label!r}"
    return float(m.group(1))


def _us_date(iso: str) -> str:
    """A date as the Short Date format shows it in a US locale: 2/17/2024."""
    y, m, d = iso.split("-")
    return f"{int(m)}/{int(d)}/{y}"


def _page_tabs(active: str, d: Diagram, y: float) -> None:
    """The page tabs under the canvas: the pages around the active one, as the report orders them."""
    order = _page_ids()
    i = order.index(active)
    shown = order[max(i - 1, 0):i + 2]
    pbi.page_tabs(d, 4, y, [_page_name(p) for p in shown], _page_name(active))


def _frame(d: Diagram, x: float, y: float, w: float, h: float, title: str | None, selected: bool = False,
           lines: int | None = None) -> tuple[float, float, float, float]:
    """A visual's container.  The title may take two lines, as a long title does in Power BI.  Returns the inner area."""
    d.box("", x, y, w, h, fill=WHITE, stroke=BLUE if selected else RULE, stroke_width=2 if selected else 1, rounded=False)
    if not title:
        return x + 6, y + 6, w - 12, h - 12
    label = f"<b>{esc(title)}</b>"
    lines = lines or kit.lines_for(label, w - 12, SMALL, bold=True)
    th = 16 * lines + 4
    d.text(label, x + 6, y + 4, w - 12, th, size=SMALL)
    return x + 6, y + 4 + th + 2, w - 12, h - th - 12


def _is_num(value: str) -> bool:
    return value[:1].isdigit() or value[:1] in "-−" and value[1:2].isdigit()


def _heads(head: str, width: float, bold: bool = True) -> str:
    """A column heading as HTML: an explicit <br> breaks it, and a heading too wide for its column is cut with an ellipsis."""
    if "<br>" in head:
        return "<br>".join(esc(p) for p in head.split("<br>"))
    return esc(_clip(head, width, margin=8, bold=bold))


def _grid_line(d: Diagram, x: float, y: float, w: float, h: float) -> None:
    """A one-pixel gridline (a box with a stroke would draw two)."""
    d.vertex("", f"rounded=0;html=1;fillColor={RULE};strokeColor=none;", x, y, w, h)


def _scrollbar(d: Diagram, right: float, top: float, height: float, shown: int, total: int) -> None:
    """A vertical scroll bar at a visual's right edge: the thumb is the share of the rows in view."""
    d.box("", right - 10, top, 6, height, fill=GRAY_TINT, stroke=GRAY_TINT, rounded=False)
    d.box("", right - 10, top, 6, max(height * shown / total, 12), fill=GRAY, stroke=GRAY, rounded=False)


def _table_visual(d: Diagram, x: float, y: float, title: str, headers: list[str], widths: list[float], rows: list[list[str]],
                  aligns: list[str] | None = None, title_lines: int | None = None, head_lines: int = 1, row_h: float = 22,
                  total: list[str] | None = None, row_fills: dict[int, str] | None = None) -> dict:
    """A table visual: a title, a header row, the rows (banded; row_fills tints chosen rows) and an optional bold total row.  aligns names
    each column's side.  Returns the cell geometry keyed (row, column) with row -1 for the headers, and the visual's bottom and right edges."""
    w = sum(widths) + 12
    label = f"<b>{esc(title)}</b>"
    th = 16 * (title_lines or kit.lines_for(label, w - 12, SMALL, bold=True)) + 4
    hh = row_h if head_lines == 1 else 15 * head_lines + 8
    n = len(rows) + (1 if total else 0)
    h = th + 12 + hh + n * row_h
    ix, iy, _, _ = _frame(d, x, y, w, h, title, lines=title_lines)
    xs = [ix]
    for wd in widths[:-1]:
        xs.append(xs[-1] + wd)
    aligns = aligns or ["right" if any(_is_num(r[c]) for r in rows) else "left" for c in range(len(headers))]
    geometry = {}
    for c, (head, cx, wd) in enumerate(zip(headers, xs, widths)):
        d.box(f"<b>{_heads(head, wd)}</b>", cx, iy, wd, hh, fill=WHITE, stroke=WHITE, size=SMALL, align=aligns[c], rounded=False, valign="bottom")
        geometry[(-1, c)] = (cx, iy, wd, hh)
    d.box("", ix, iy + hh - 1, sum(widths), 1, fill=INK, stroke=INK, rounded=False)
    for r, row in enumerate(rows + ([total] if total else [])):
        ry = iy + hh + row_h * r
        is_total = total is not None and r == len(rows)
        if is_total:
            d.box("", ix, ry, sum(widths), 1, fill=INK, stroke=INK, rounded=False)
        for c, (value, cx, wd) in enumerate(zip(row, xs, widths)):
            plain = r % 2 == 0 or is_total
            fill = (row_fills or {}).get(r) or (WHITE if plain else GRAY_TINT)
            shown = esc(_clip(value, wd, margin=8 if aligns[c] == "right" else 16, bold=is_total))
            d.box(f"<b>{shown}</b>" if is_total else shown, cx, ry, wd, row_h, fill=fill, stroke=fill, size=SMALL, align=aligns[c], rounded=False)
            geometry[(r, c)] = (cx, ry, wd, row_h)
    return dict(geometry=geometry, bottom=y + h, right=x + w, left=ix, width=sum(widths), top_rows=iy + hh, row_h=row_h)


def _year_matrix(d: Diagram, x: float, y: float, title: str, row_w: float, block: list[float], years: list[str], value_heads: list[str],
                 row_heads: list[str], cells: list[list[list[str]]], total: list[list[str]], row_h: float = 22, head2: float = 38) -> dict:
    """A matrix with one level of rows and the years on columns, each year holding the same values side by side: a title, the year headings
    over the value headings, the rows and a total row.  cells[r][y] are the strings of one row and one year.  Returns the geometry."""
    inner_w = row_w + len(years) * sum(block)
    w = inner_w + 12
    label = f"<b>{esc(title)}</b>"
    th = 16 * kit.lines_for(label, w - 12, SMALL, bold=True) + 4
    head1, tot_h = 22, 24
    h = th + 12 + head1 + head2 + row_h * len(row_heads) + tot_h
    ix, iy, _, _ = _frame(d, x, y, w, h, title)
    xs = [ix + row_w + i * sum(block) for i in range(len(years))]
    cy = iy
    for i, year in enumerate(years):
        d.text(f"<b>{year}</b>", xs[i], cy, sum(block), head1, size=SMALL, align="center", valign="middle")
        if i:
            d.box("", xs[i] - 0.5, cy, 1, head1 + head2 + row_h * len(row_heads) + tot_h, fill=RULE, stroke=RULE, rounded=False)
    cy += head1
    for i in range(len(years)):
        cx = xs[i]
        for c, head in enumerate(value_heads):
            d.text("<br>".join(esc(p) for p in head.split("<br>")), cx, cy, block[c] - 4, head2, size=SMALL, align="right", valign="bottom")
            cx += block[c]
    cy += head2
    d.box("", ix, cy - 1, inner_w, 1, fill=INK, stroke=INK, rounded=False)
    geometry = {}
    for r, head in enumerate(row_heads):
        ry = cy + r * row_h
        d.text("<br>".join(esc(p) for p in head.split("<br>")), ix, ry, row_w, row_h, size=SMALL, align="left", valign="middle")
        for i in range(len(years)):
            cx = xs[i]
            for c in range(len(value_heads)):
                d.text(esc(cells[r][i][c]), cx, ry, block[c] - 4, row_h, size=SMALL, align="right", valign="middle")
                geometry[(r, i, c)] = (cx, ry, block[c], row_h)
                cx += block[c]
        d.box("", ix, ry + row_h - 1, inner_w, 1, fill=RULE, stroke=RULE, rounded=False)
    ty = cy + row_h * len(row_heads)
    d.box("", ix, ty, inner_w, 1, fill=INK, stroke=INK, rounded=False)
    d.text("<b>Total</b>", ix, ty, row_w, tot_h, size=SMALL, align="left", valign="middle")
    for i in range(len(years)):
        cx = xs[i]
        for c in range(len(value_heads)):
            d.text(f"<b>{esc(total[i][c])}</b>", cx, ty, block[c] - 4, tot_h, size=SMALL, align="right", valign="middle")
            geometry[("total", i, c)] = (cx, ty, block[c], tot_h)
            cx += block[c]
    return dict(geometry=geometry, bottom=y + h, right=x + w, left=ix, top_total=ty, width=inner_w, top_rows=cy, head_top=iy)


def _cards(d: Diagram, x: float, y: float, h: float, items: list[tuple[str, str]], widths: list[float]) -> dict:
    """A card visual, one card for each field: the value above the field's name (a name may break with <br>).  Returns the geometry."""
    w = sum(widths) + 6 * (len(items) - 1) + 12
    d.box("", x, y, w, h, fill=WHITE, stroke=RULE, rounded=False)
    geo, cx = [], x + 6
    for (label, value), cw in zip(items, widths):
        d.box("", cx, y + 6, cw, h - 12, fill=WHITE, stroke=RULE, rounded=True)
        d.text(f"<b>{esc(value)}</b>", cx + 2, y + 10, cw - 4, 22, size=14, align="center")
        d.text("<br>".join(esc(p) for p in label.split("<br>")), cx, y + 36, cw, h - 44, size=SMALL, color=GRAY, align="center")
        geo.append((cx, y + 6, cw, h - 12))
        cx += cw + 6
    return dict(bottom=y + h, right=x + w, cards=geo)


def _lines_chart(d: Diagram, x: float, y: float, w: float, h: float, title: str, cats: list[str], series: list[tuple[str, list[float | None], str, str]],
                 ticks: list[float], fmt, every: int = 6) -> dict:
    """A line chart of several series, each a (name, values, color, marker) with the marker a circle or a diamond (so the lines differ
    without their color); a value of None is a blank and leaves a gap.  The first tick is where the value axis starts.  Returns the
    plot geometry and the position of each series' points."""
    ix, iy, iw, ih = _frame(d, x, y, w, h, title)
    lx = ix
    for name, _, color, marker in series:
        d.box("", lx, iy + 8, 22, 3, fill=color, stroke=color, rounded=False)
        _dot(d, lx + 11, iy + 9.5, color, marker, 9)
        nw = text_width(esc(name), SMALL, False) + 10
        d.text(esc(name), lx + 28, iy - 1, nw, 20, size=SMALL)
        lx += 28 + nw + 18
    iy += 22
    ih -= 22
    left, plot_w = ix + 50, iw - 58
    top, plot_h = iy + 8, ih - 34
    low, high = ticks[0], ticks[-1]
    y_of = lambda v: top + (high - v) * plot_h / (high - low)                                                  # noqa: E731
    for t in ticks:
        _grid_line(d, left, y_of(t), plot_w, 1)
        d.text(fmt(t), ix, y_of(t) - 10, 44, 20, size=SMALL, color=GRAY, align="right", valign="middle")
    step = plot_w / len(cats)
    xs = [left + step * (i + 0.5) for i in range(len(cats))]
    for i, (cx, label) in enumerate(zip(xs, cats)):
        if i % every == 0:
            d.text(esc(label), cx - step * every / 2, top + plot_h + 4, step * every, 20, size=SMALL, color=GRAY, align="center")
    points = []
    for name, values, color, marker in series:
        run: list[tuple[float, float]] = []
        pts = []
        for cx, v in zip(xs, values):
            if v is None:
                if len(run) > 1:
                    _polyline(d, run, color, 2.5)
                run = []
                pts.append(None)
            else:
                run.append((cx, y_of(v)))
                pts.append((cx, y_of(v)))
        if len(run) > 1:
            _polyline(d, run, color, 2.5)
        for p in pts:
            if p:
                _dot(d, p[0], p[1], color, marker, 7)
        points.append(pts)
    return dict(left=left, top=top, width=plot_w, height=plot_h, xs=xs, y_of=y_of, points=points)


def _dot(d: Diagram, cx: float, cy: float, color: str, marker: str = "circle", size: float = 8) -> None:
    shape = "rhombus" if marker == "diamond" else "ellipse"
    d.vertex("", f"{shape};html=1;fillColor={color};strokeColor={color};", cx - size / 2, cy - size / 2, size, size)


def _hbars(d: Diagram, x: float, y: float, w: float, h: float, title: str, cats: list[str], values: list[float], ticks: list[float],
           label_w: float = 40) -> dict:
    """A clustered bar chart, the categories down the side in the order given, with a value axis."""
    ix, iy, iw, ih = _frame(d, x, y, w, h, title)
    left, top = ix + label_w, iy + 4
    plot_w = iw - label_w - 14
    plot_h = ih - 30
    x_of = lambda v: left + v * plot_w / ticks[-1]                                                            # noqa: E731
    slot = plot_h / len(cats)
    for t in ticks:
        _grid_line(d, x_of(t), top, 1, plot_h)
        d.text(f"{t:g}", x_of(t) - 22, top + plot_h + 2, 44, 18, size=SMALL, color=GRAY, align="center")
    bars = []
    for i, (name, v) in enumerate(zip(cats, values)):
        by = top + slot * i + slot * 0.16
        d.text(esc(name), ix, by - 2, label_w - 8, slot * 0.68 + 4, size=SMALL, align="right", valign="middle")
        d.box("", left, by, max(x_of(v) - left, 1), slot * 0.68, fill=BLUE, stroke=BLUE, rounded=False)
        bars.append((left, by, x_of(v) - left, slot * 0.68))
    return dict(left=left, top=top, slot=slot, bars=bars)


# -- IF-19-5-A: the Timing page ------------------------------------------------------------------------------------------------

def if_19_5_a() -> Diagram:
    """The Timing page: the line chart of the return rate by the month of the credit and by the month of the shipment, the table of the
    rate by year on both bases, the table of the last fiscal year by month of the shipment and the cards of the days after shipment and
    the lines returned in a later year."""
    data = ctx.data()
    n = _n()
    g5, r5 = ctx.values("ch19.g5"), ctx.values("ch19.r5")
    page = "timing"
    C = data.C
    titles = dict(chart=_title(page, "rateByMonth"), years=_title(page, "rateYears"), months=_title(page, "shipMonths"))
    assert titles["chart"] == g5["t"]["timing"] and titles["years"] == "Return rate by year" and titles["months"] == f"Fiscal {C} by month of the shipment", titles
    assert _projections(page, "rateByMonth", "Category") == ["YearMonth"] and _projections(page, "rateByMonth", "Y") == ["Return Rate", "Return Rate by Ship Date"]
    assert _projections(page, "rateYears", "Values") == ["Year", "Return Rate", "Return Rate by Ship Date"]
    assert _projections(page, "shipMonths", "Values") == ["YearMonth", "Credits by Ship Date", "Revenue by Ship Date", "Return Rate by Ship Date"]
    cards = _projections(page, "lagCards", "Data")
    assert cards == ["Fewest Days after Shipment", "Average Days after Shipment", "Most Days after Shipment", "Lines Returned in a Later Year",
                     "Credits Returned in a Later Year"], cards
    flt = _visual(page, "shipMonths")["filterConfig"]["filters"][0]
    assert flt["field"]["Column"]["Property"] == "Year" and flt["filter"]["Where"][0]["Condition"]["In"]["Values"][0][0]["Literal"]["Value"] == f"{C}L", "the table is filtered to the last year"
    mt = _g().measure_table()
    assert all(m in mt for m in cards + ["Return Rate", "Return Rate by Ship Date", "Credits by Ship Date", "Revenue by Ship Date"]), "the measures the page uses are the model's"
    # the monthly series, by the month of the credit and by the month of the shipment
    q = data.q
    cred = dict(q("SELECT substr(cm.CreditMemoDate, 1, 7), SUM(l.LineTotal) FROM CreditMemoLine l JOIN CreditMemo cm ON cm.CreditMemoID = l.CreditMemoID GROUP BY 1"))
    rev = dict(q("SELECT substr(si.InvoiceDate, 1, 7), SUM(l.LineTotal) FROM SalesInvoiceLine l JOIN SalesInvoice si ON si.SalesInvoiceID = l.SalesInvoiceID GROUP BY 1"))
    cs = dict(q("SELECT substr(s.ShipmentDate, 1, 7), SUM(l.LineTotal) FROM CreditMemoLine l JOIN SalesReturnLine srl ON srl.SalesReturnLineID = l.SalesReturnLineID "
                "JOIN ShipmentLine sl ON sl.ShipmentLineID = srl.ShipmentLineID JOIN Shipment s ON s.ShipmentID = sl.ShipmentID GROUP BY 1"))
    rs = dict(q("SELECT substr(s.ShipmentDate, 1, 7), SUM(sil.LineTotal) FROM SalesInvoiceLine sil JOIN ShipmentLine sl ON sl.ShipmentLineID = sil.ShipmentLineID "
                "JOIN Shipment s ON s.ShipmentID = sl.ShipmentID GROUP BY 1"))
    months = [f"{y}-{m:02d}" for y in data.years for m in range(1, 13)]
    assert all(mo in rev and mo in rs for mo in months) and len(months) == 36
    rate_c = [cred[mo] / rev[mo] if mo in cred else None for mo in months]
    rate_s = [cs[mo] / rs[mo] if mo in cs else None for mo in months]
    assert abs(rate_c[0] - g5["first_rate"]) < 1e-12 and abs(rate_s[-1] - r5["dec"]) < 1e-12, "the first month by credit date and the last by ship date are the notes'"
    assert abs(rate_s[-1] - _expected(f"return rate of the shipments of {C}-12 (censored)")) < 1e-8
    by = lambda dct, y: sum(v for k, v in dct.items() if k.startswith(str(y)))                                # noqa: E731
    year_c = [by(cred, y) / by(rev, y) for y in data.years]
    year_s = [by(cs, y) / by(rs, y) for y in data.years]
    assert all(abs(a - b) < 1e-9 for a, b in zip(year_c, g5["rate_year"])) and all(abs(a - b) < 1e-9 for a, b in zip(year_s, r5["by_ship_year"]))
    for k, y in enumerate(data.years):
        assert abs(year_s[k] - _expected(f"return rate by shipment year {y}")) < 1e-8
    all_c, all_s = sum(cred.values()) / sum(rev.values()), sum(cs.values()) / sum(rs.values())
    assert abs(all_c - g5["rate_all"]) < 1e-9 and abs(all_s - g5["rate_ship"]) < 1e-9 and abs(all_s - _expected("return rate by ship date with no date filter (credits over the shipped lines' revenue)")) < 1e-8
    last = [mo for mo in months if mo.startswith(str(C))]
    for mo in last:
        assert abs(cs[mo] - _expected(f"credits on the shipments of {mo}")) < 0.006 and abs(rs[mo] - _expected(f"revenue of the shipments of {mo}")) < 0.006, mo
    lag = r5["lag"]
    assert (lag["lo"], lag["hi"]) == (_expected("fewest days from shipment to return"), _expected("most days from shipment to return")) and abs(lag["mean"] - _expected("average days from shipment to return")) < 1e-6
    assert r5["cross"] == _expected("return lines in a later year than their shipment") and abs(r5["cross_amt"] - _expected("credits returned in a later year than their shipment")) < 0.006
    pct = lambda v: f"{v * 100:.3f}%"                                                                         # noqa: E731
    card_items = [("Fewest Days<br>after Shipment", f"{lag['lo']:.0f}"), ("Average Days<br>after Shipment", f"{lag['mean']:.1f}"),
                  ("Most Days<br>after Shipment", f"{lag['hi']:.0f}"), ("Lines Returned<br>in a Later Year", f"{r5['cross']:,}"),
                  ("Credits Returned<br>in a Later Year", f"{r5['cross_amt']:,.2f}")]
    years_rows = [[str(y), pct(a), pct(b)] for y, a, b in zip(data.years, year_c, year_s)]
    years_total = ["Total", pct(all_c), pct(all_s)]
    ship_rows = [[mo, f"{cs[mo]:,.2f}", f"{rs[mo]:,.2f}", pct(cs[mo] / rs[mo])] for mo in last]
    ship_total = ["Total", f"{sum(cs[mo] for mo in last):,.2f}", f"{sum(rs[mo] for mo in last):,.2f}", pct(year_s[-1])]
    pad = 16.0
    RW = 340.0
    LW = 860 - 2 * pad - RW - 12
    series = [("Return Rate", [None if v is None else v * 100 for v in rate_c], BLUE, "circle"),
              ("Return Rate by Ship Date", [None if v is None else v * 100 for v in rate_s], AMBER, "diamond")]
    ticks = pbi.nice_ticks(0, max(v for _, vals, _, _ in series for v in vals if v is not None), 4)

    def draw(d: Diagram) -> dict:
        ty = _table_visual(d, pad + LW + 12, pad, titles["years"], ["Year", "Return Rate", "Return Rate<br>by Ship Date"], [66, 110, 152], years_rows,
                           aligns=["right", "right", "right"], title_lines=1, head_lines=2, total=years_total)
        tm = _table_visual(d, pad + LW + 12, ty["bottom"] + 12, titles["months"], ["Year<br>Month", "Credits by<br>Ship Date", "Revenue by<br>Ship Date", "Return Rate<br>by Ship Date"],
                           [64, 78, 96, 90], ship_rows, aligns=["left", "right", "right", "right"], title_lines=1, head_lines=2, total=ship_total)
        right_bottom = tm["bottom"]
        cards_h = 96.0
        chart_h = right_bottom - pad - 12 - cards_h
        ch = _lines_chart(d, pad, pad, LW, chart_h, titles["chart"], [mo for mo in months], series, ticks, lambda t: f"{t:.1f}%", every=6)
        cd = _cards(d, pad, pad + chart_h + 12, cards_h, card_items, [74, 88, 74, 96, 108])
        return dict(chart=ch, cards=cd, years=ty, months=tm, bottom=right_bottom + pad, chart_h=chart_h)

    canvas_h = draw(Diagram("probe"))["bottom"]
    d = Diagram("The Timing Page")
    pbi.canvas(d, 0, 0, 860, canvas_h)
    parts = draw(d)
    ch = parts["chart"]
    # the first and the last points the answer names, labelled as data labels would be
    for si, pi, dx, dy, align in ((0, 0, 8, -26, "left"), (1, len(months) - 1, -64, -28, "right")):
        px, py = ch["points"][si][pi]
        txt = pct((rate_c if si == 0 else rate_s)[pi])
        tw = text_width(f"<b>{txt}</b>", SMALL, True) + 12
        lx = px + dx if align == "left" else px - tw + 6
        _label(d, f"<b>{txt}</b>", lx, py + dy, tw, 18, color=INK, align="center")
    # the months of the last fiscal year outlined on the table, the last month in particular
    tm = parts["months"]
    gx, gy, gw, gh = tm["geometry"][(len(last) - 1, 0)]
    gx2, gy2, gw2, gh2 = tm["geometry"][(len(last) - 1, 3)]
    d.outline(gx - 1, gy - 1, gx2 + gw2 - gx + 2, gh + 2, CORAL, width=2.5, front=True)
    px, py = ch["points"][1][len(months) - 1]
    d.outline(px - 9, py - 9, 18, 18, CORAL, width=2.5, front=True)
    _page_tabs(page, d, canvas_h + 4)
    mlo, mhi = min(rate_s[len(months) - 12:-1]), max(rate_s[len(months) - 12:-1])
    assert abs(mlo - r5["mlo"]) < 1e-12 and abs(mhi - r5["mhi"]) < 1e-12
    _note(d, f"Report view, page {_page_name(page)}, drawn here in a different arrangement (the model answer's page also holds a text box with the written answer, not drawn). "
             f"Blue circles: the rate by the month of the credit, {pct(rate_c[0])} in {months[0]}, the first month of the first year, which had nothing earlier to return. Amber "
             f"diamonds: the rate by the month of the shipment, which falls to {pct(rate_s[-1])} in {months[-1]} (outlined in coral, with its row in the table) against {pct(mlo)} to "
             f"{pct(mhi)} for the other months of {C}, because the returns of the last shipments arrive after the data end. The tables show {pct(all_c)} overall by the credit date against "
             f"{pct(all_s)} by the ship date, whose revenue leaves out the {g5['n_noship']:,} lines that never ship.", canvas_h + 34)
    return d


# -- IF-19-5-B: the Refund Ranking page --------------------------------------------------------------------------------------------

def _flag_names() -> list[str]:
    """The five flag names of the Flags table, in the order of its Order column (typed with Enter data in the project)."""
    import base64
    import zlib
    text = (MODEL / "definition" / "tables" / "Flags.tmdl").read_text(encoding="utf-8")
    blob = re.search(r'Binary\.FromText\("([^"]+)"', text).group(1)
    rows = json.loads(zlib.decompress(base64.b64decode(blob), -15))
    return [r[0] for r in sorted(rows, key=lambda r: int(r[1]))]


def if_19_5_b() -> Diagram:
    """The Refund Ranking page: the bar chart of the refunds at each score, the matrix of the flags by year, the refunds by the approver's
    job title, the two cards for the refunds scoring 4 or 5, and the table of every refund with its five flags and score, sorted from the
    highest score with its first row selected."""
    import csv
    data = ctx.data()
    n = _n()
    g5, r5, r4 = ctx.values("ch19.g5"), ctx.values("ch19.r5"), ctx.values("ch19.r4")
    c, t = n.cycle(data), n.tests(data)
    page = "ranking"
    titles = dict(scores=_title(page, "scores"), flags=_title(page, "flagYears"), approvers=_title(page, "approvers"), table=_title(page, "refunds"))
    assert titles["table"] == g5["t"]["ranking"], titles
    assert _projections(page, "scores", "Category") == ["Score"] and _projections(page, "scores", "Y") == ["Refunds"]
    assert _projections(page, "flagYears", "Rows") == ["Flag"] and _projections(page, "flagYears", "Columns") == ["Year"]
    assert _projections(page, "flagYears", "Values") == ["Flagged Refunds", "Flagged Amount"]
    assert _projections(page, "approvers", "Rows") == ["RefundApproverTitle"] and _projections(page, "approvers", "Values") == ["Refunds"]
    assert _projections(page, "highCards", "Data") == ["High-Score Refunds", "High-Score Amount"]
    cols = _projections(page, "refunds", "Values")
    assert cols == ["RefundNumber", "RefundDate", "CustomerID", "Sum of Amount", "AboveLimit", "CreditAboveLimit", "MethodMismatch", "BeforePayment",
                    "PastDueAtLeastRefund", "Score"], cols
    flags = _flag_names()
    assert flags == ["Refund above authority", "Credit above authority", "Method mismatch", "Before the last application", "Past due at least the refund"], flags
    mt = _g().measure_table()
    assert all(f'"{f}"' in mt["Flagged Refunds"] for f in flags), "the Flags table's names are the measure's"
    # the refunds with their flags: RefundFlags.csv
    rows = list(csv.DictReader(FLAGS_CSV.open(encoding="utf-8-sig")))
    assert len(rows) == len(c.rf) == 222 and all(int(r["Score"]) == sum(int(r[k]) for k in ("AboveLimit", "CreditAboveLimit", "MethodMismatch", "BeforePayment", "PastDueAtLeastRefund"))
                                                for r in rows), "the score is the sum of the five flags"
    key = ["AboveLimit", "CreditAboveLimit", "MethodMismatch", "BeforePayment", "PastDueAtLeastRefund"]
    years = [str(y) for y in data.years]
    year_of = lambda r: r["RefundDate"][:4]                                                                   # noqa: E731
    count = {(k, y): sum(1 for r in rows if int(r[k]) == 1 and year_of(r) == y) for k in key for y in years}
    amount = {(k, y): round(sum(float(r["Amount"]) for r in rows if int(r[k]) == 1 and year_of(r) == y), 2) for k in key for y in years}
    tot_n = {k: sum(count[(k, y)] for y in years) for k in key}
    tot_a = {k: round(sum(amount[(k, y)] for y in years), 2) for k in key}
    assert tot_n["AboveLimit"] == r4["rf_above"]["n"] and [count[("AboveLimit", y)] for y in years] == r4["rf_above"]["years"]
    assert tot_n["BeforePayment"] == r4["before"]["n"] and [count[("BeforePayment", y)] for y in years] == r4["before"]["years"]
    assert tot_n["MethodMismatch"] == r4["mism"]["n"] and tot_n["PastDueAtLeastRefund"] == r4["pd2"]["n"] and tot_n["CreditAboveLimit"] > 0
    assert abs(tot_a["AboveLimit"] - r4["rf_above"]["amt"]) < 0.005 and abs(tot_a["BeforePayment"] - r4["before"]["amt"]) < 0.005
    scores = [sum(1 for r in rows if int(r["Score"]) == s) for s in range(6)]
    assert scores == r5["scores"], (scores, r5["scores"])
    hi = [r for r in rows if int(r["Score"]) >= 4]
    assert len(hi) == r5["hi"] and abs(sum(float(r["Amount"]) for r in hi) - r5["hi_amt"]) < 0.005
    # the approvers' job titles by year
    app = {}
    for title, y, k in data.q("SELECT e.JobTitle, substr(rf.RefundDate, 1, 4), COUNT(*) FROM CustomerRefund rf JOIN Employee e ON e.EmployeeID = rf.ApprovedByEmployeeID GROUP BY 1, 2"):
        app.setdefault(title, {})[y] = k
    titles_sorted = sorted(app)
    assert [[app[t_].get(y, 0) for y in years] for t_ in titles_sorted] == [v for _, v in r5["approvers"]], "the approvers by year are the note's"
    sorted_rows = sorted(rows, key=lambda r: (-int(r["Score"]), -float(r["Amount"]), r["RefundNumber"]))
    big = max(rows, key=lambda r: float(r["Amount"]))
    assert big["RefundNumber"] == r5["big"] and int(big["Score"]) == r5["big_score"] == _expected(f"{r5['big']} score") and int(sorted_rows[0]["Score"]) > r5["big_score"], \
        "the largest refund is not at the top of the ranking"
    shown = sorted_rows[:8]
    assert [int(r["Score"]) for r in shown] == [5] * 8 or max(int(r["Score"]) for r in shown) == 5
    pad = 16.0
    inner = 860 - 2 * pad

    def draw(d: Diagram) -> dict:
        block = [58, 92]
        cells = [[[(f"{count[(k, y)]:,}" if count[(k, y)] else ""), (f"{amount[(k, y)]:,.2f}" if count[(k, y)] else "")] for y in years] +
                 [[f"{tot_n[k]:,}", f"{tot_a[k]:,.2f}"]] for k in key]
        blanks = [["", ""] for _ in range(len(years) + 1)]
        fm = _year_matrix(d, pad, pad, titles["flags"], 828 - 12 - 4 * sum(block), block, years + ["Total"], ["Flagged<br>Refunds", "Flagged<br>Amount"], flags, cells, blanks)
        y2 = fm["bottom"] + 12
        h2 = 196.0
        sc = _hbars(d, pad, y2, 252, h2, titles["scores"], [str(s) for s in range(6)], [float(v) for v in scores], pbi.nice_ticks(0, max(scores), 4), label_w=34)
        am = [[[str(app[t_].get(y, 0))] for y in years] + [[str(sum(app[t_].values()))]] for t_ in titles_sorted]
        a_total = [[str(sum(app[t_].get(y, 0) for t_ in titles_sorted))] for y in years] + [[str(sum(sum(v.values()) for v in app.values()))]]
        ax = pad + 252 + 12
        am_w = 176 + 4 * 58 + 12
        ap = _year_matrix(d, ax, y2, titles["approvers"], 176, [58], years + ["Total"], ["Refunds"], titles_sorted, am, a_total, head2=24)
        cx = ax + am_w + 12
        cw = 860 - pad - cx
        d.box("", cx, y2, cw, h2, fill=WHITE, stroke=RULE, rounded=False)
        for i, (label, val) in enumerate((("High-Score Refunds", f"{len(hi):,}"), ("High-Score Amount", f"{r5['hi_amt']:,.2f}"))):
            cy = y2 + 6 + i * (h2 - 12) / 2
            d.box("", cx + 6, cy + (3 if i else 0), cw - 12, (h2 - 12) / 2 - 3, fill=WHITE, stroke=RULE, rounded=True)
            d.text(f"<b>{esc(val)}</b>", cx + 8, cy + 14 + (3 if i else 0), cw - 16, 22, size=14, align="center")
            d.text(esc(label), cx + 8, cy + 42 + (3 if i else 0), cw - 16, 18, size=SMALL, color=GRAY, align="center")
        y3 = max(fm["bottom"], y2 + h2) + 12
        trows = [[r["RefundNumber"], _us_date(r["RefundDate"]), r["CustomerID"], f"{float(r['Amount']):,.2f}", r["AboveLimit"], r["CreditAboveLimit"], r["MethodMismatch"],
                  r["BeforePayment"], r["PastDueAtLeastRefund"], r["Score"]] for r in shown]
        heads = ["RefundNumber", "RefundDate", "CustomerID", "Sum of<br>Amount", "AboveLimit", "Credit<br>AboveLimit", "Method<br>Mismatch", "Before<br>Payment",
                 "PastDueAt<br>LeastRefund", "Score"]
        widths = [112, 82, 82, 80, 80, 84, 80, 74, 88, 52]
        tb = _table_visual(d, pad, y3, titles["table"], heads, widths, trows, aligns=["left", "right", "right", "right", "right", "right", "right", "right", "right", "right"],
                           title_lines=1, head_lines=2, row_fills={0: BLUE_TINT})
        _scrollbar(d, pad + sum(widths) + 12, tb["top_rows"], tb["row_h"] * len(trows), len(trows), len(rows))
        return dict(matrix=fm, scores=sc, approvers=ap, table=tb, bottom=tb["bottom"] + pad, y3=y3)

    canvas_h = draw(Diagram("probe"))["bottom"]
    d = Diagram("The Refund Ranking Page")
    pbi.canvas(d, 0, 0, 860, canvas_h)
    parts = draw(d)
    tb = parts["table"]
    x0, y0, w0, h0 = tb["geometry"][(0, 0)]
    xl_, yl_, wl_, hl_ = tb["geometry"][(0, 9)]
    d.outline(x0 - 1, y0 - 1, xl_ + wl_ - x0 + 2, h0 + 2, BLUE, width=2.5, front=True)
    # the score column outlined, and the matrix's Total column of the flags
    xs, ys, ws_, hs = tb["geometry"][(-1, 9)]
    xe, ye, we, he = tb["geometry"][(len(shown) - 1, 9)]
    d.outline(xs - 1, ys - 1, ws_ + 2, ye + he - ys + 2, TEAL, width=2.5, front=True)
    mx = parts["matrix"]
    gx, gy, gw, gh = mx["geometry"][(0, len(years), 0)]
    gx2, gy2, gw2, gh2 = mx["geometry"][(len(flags) - 1, len(years), 1)]
    d.outline(gx - 2, mx["head_top"] - 1, gx2 + gw2 - gx + 4, gy2 + gh2 - mx["head_top"] + 2, TEAL, width=2.5, front=True)
    _page_tabs(page, d, canvas_h + 4)
    top_row = shown[0]
    _note(d, f"Report view, page {_page_name(page)}, in a different arrangement from the model answer's (the page also holds a text box with the written answer, not drawn). The bars count "
             f"the refunds at scores 0 to 5: {', '.join(f'{v:,}' for v in scores)}. The matrix totals the flags: its Total row is blank because the flag measures switch on one selected flag, "
             f"and its Total column (outlined in teal, like the Score column of the table) is the sum of the years. The {len(hi)} refunds scoring 4 or 5 ({r5['hi_amt']:,.2f}) are the "
             f"first to review. The table lists all {len(rows)} refunds, {len(shown)} in view, sorted by score and then amount; the first row, {top_row['RefundNumber']}, is selected "
             f"(outlined in blue) before a reviewer drills through to its detail. The largest refund, {r5['big']}, scores {r5['big_score']} and is not at the top. The headings of the "
             f"flags are broken over two lines to fit.", canvas_h + 34)
    return d


# -- IF-19-6-A: the Sample sheet ------------------------------------------------------------------------------------------------

def _as_shown(ws, ref: str) -> str:
    """What a cell shows, as _show does, with a date in the workbook's yyyy-mm-dd format."""
    v = ws[ref].value
    if hasattr(v, "date"):
        return v.date().isoformat()
    return _show(ws, ref)


def if_19_6_a() -> Diagram:
    """Rows 11 to 22 (the named parameters and the two tables of tbl-19-03, with the cells the plan reads), 30 to 38 (the key items with
    the reason each is key) and 42 to 48 (the rule for zero to three deviations) of the Sample sheet."""
    from collections import defaultdict
    data = ctx.data()
    n = _n()
    g, r6 = ctx.values("ch19.g6"), ctx.values("ch19.r6")
    s = r6["s"]
    c, t = n.cycle(data), n.tests(data)
    ws = _book("Sample")
    ok = lambda cond, msg: None                                                                      # noqa: E731
    names = [("SampleConfidence", "B12", n.CONFIDENCE), ("TolerableRate", "B13", n.TOLERABLE), ("ExpectedRate", "B14", 0),
             ("SampleSize", "B15", g["size"]), ("KeyAmount", "B17", n.KEY_AMOUNT), ("KeyCustomerTotal", "B18", n.KEY_CUSTOMER)]
    for name, ref, v in names:
        assert _g().defined(ok, name) == f"Sample!${ref[0]}${ref[1:]}", name
        _agree(ws, ref, v, tol=1e-9)
    assert _g().formula(ok, "Sample", "B15").startswith("=INDEX($E$15:$H$17,MATCH(ExpectedRate"), "the size is read from the table by two lookups"
    _agree(ws, "B16", 0, tol=1e-9)
    for ref, v in (("E15", g["size"]), ("F15", g["size_tol10"]), ("G15", g["size_conf90"]), ("E17", g["size_2pct"]), ("E20", 0)):
        _agree(ws, ref, v, tol=1e-9)
    assert [ws[f"{col}{r}"].value for r in (13, 14) for col in "EFGH"] == [0.95, 0.95, 0.9, 0.9, 0.05, 0.1, 0.05, 0.1], "the table's columns"
    assert [ws[f"D{r}"].value for r in (15, 16, 17)] == [0, 0.01, 0.02] and ws["A15"].value.startswith("Sample size")
    for ref, v in (("B25", g["n_all"]), ("B26", g["nk"]), ("B27", g["n_rest"]), ("C25", n.amount(c.rf.values(), "amt")), ("C26", s["key_total"]),
                   ("C27", s["rest_total"]), ("B28", s["share"])):
        _agree(ws, ref, v, tol=0.006 if v > 1 else 1e-9)
    # the key items: the refunds of 5,000 or more and every refund of a customer whose refunds over the window exceed 10,000
    tot = defaultdict(float)
    for r in c.rf.values():
        tot[r["cust"]] += r["amt"]
    before = {r["id"] for r in t.before}
    key = s["key"]
    assert len(key) == g["nk"] == 7
    keys = []
    for k, r in enumerate(key):
        ref = 32 + k
        reason = ("Amount and customer" if r["amt"] >= n.KEY_AMOUNT and tot[r["cust"]] > n.KEY_CUSTOMER
                  else "Amount" if r["amt"] >= n.KEY_AMOUNT else "Customer total")
        _agree(ws, f"A{ref}", r["num"])
        _agree(ws, f"B{ref}", r["date"])
        _agree(ws, f"C{ref}", r["cust"], tol=0.5)
        _agree(ws, f"D{ref}", r["amt"], tol=0.006)
        _agree(ws, f"E{ref}", round(tot[r["cust"]], 2), tol=0.006)
        _agree(ws, f"F{ref}", reason)
        _agree(ws, f"G{ref}", r["id"] in before)
        keys.append(reason)
    assert ws["A39"].value is None and ws["A41"].value is None and ws["A32"].value == key[0]["num"]
    # the rule: the upper deviation limit for zero to three deviations
    for k in range(4):
        _agree(ws, f"A{44 + k}", k, tol=0.5)
        _agree(ws, f"B{44 + k}", s["udl"][k], tol=1e-9)
        _agree(ws, f"C{44 + k}", "Yes" if s["udl"][k] <= n.TOLERABLE else "No")
    assert ws["A43"].value.startswith("Deviations found") and ws["A42"].value.startswith("Evaluation rule, written before testing")
    assert [ws[f"C{44 + k}"].value for k in range(4)] == ["Yes", "No", "No", "No"] and ws["A48"].value.startswith("Rule: rely on the control only if")
    focus = max(key, key=lambda r: tot[r["cust"]])["cust"]
    assert focus == g["focus"] and sum(1 for r in key if r["cust"] == focus) == g["focus_n"] and round(tot[focus], 2) == g["focus_total"]
    # the sheet
    letters = list("ABCDEFGH")
    widths = [26, 180, 92, 130, 92, 84, 138, 64, 50]
    assert abs(sum(widths) - 856) < 6
    formula = _formula_of("Sample", "B15")
    d = Diagram("The Sample Sheet")
    _formula_bar(d, 0, 0, 860, "SampleSize", formula)
    sy = 40.0
    _ws(d, 0, sy, letters, widths, [], show_letters=True)
    y = _gap(d, 0, sy + 22, sum(widths), "Rows 1 to 10: the three attributes needing outside evidence (A1 to A3) and the data attribute D1, with the evidence and the deviation of each")

    def val(ref: str, **kw) -> dict:
        shown = _as_shown(ws, ref)
        if re.fullmatch(r"\d{4}-\d\d-\d\d", shown):
            return dict(text=shown, html=_nowrap(shown), align="right", clip=False, **kw)
        return _num(shown, **kw)

    def lab(r: int, **kw) -> dict:
        return dict(text=ws[f"A{r}"].value, clip=True, **kw)
    top = [dict(n="11", cells=[dict(text=ws["A11"].value, bold=True, span=8, clip=False)])]
    for r in range(12, 23):
        left = {12: [lab(12), val("B12"), ""], 13: [lab(13), val("B13"), ""], 14: [lab(14), val("B14"), ""], 15: [lab(15), val("B15"), ""],
                16: [lab(16), val("B16"), ""], 17: [lab(17), val("B17"), ""], 18: [lab(18), val("B18"), ""], 19: [lab(19), val("B19"), ""]}.get(r)
        if r == 20:
            left = [lab(20), dict(text=ws["B20"].value, span=2, clip=True)]
        if left is None:
            left = ["", "", ""]
        if r in (12, 19):
            right = [dict(text=ws[f"D{r}"].value, bold=True, span=5, clip=False)]
        elif r in (13, 14):
            right = [dict(text=ws[f"D{r}"].value, clip=False)] + [_num(_as_shown(ws, f"{col}{r}")) for col in "EFGH"]
        else:
            right = [_num(_as_shown(ws, f"D{r}"))] + [_num(_as_shown(ws, f"{col}{r}")) for col in "EFGH"]
        top.append(dict(n=str(r), cells=left + right))
    geo1 = _ws(d, 0, y, letters, widths, top, show_letters=False)
    i1 = {int(row["n"]): i for i, row in enumerate(top)}
    pop = f"Rows 23 to 29: the population, {g['n_all']} refunds ({_money(n.amount(c.rf.values(), 'amt'))}), of which {g['nk']} key items in full and {g['n_rest']} remaining; a sample of {g['size']} is {s['share'] * 100:.1f}% of them"
    y = _gap(d, 0, geo1["bottom"], sum(widths), pop)
    head_cells = [(ws[f"{col}31"].value, w) for col, w in zip("ABCDEFG", widths[1:8])]
    wrapped = [_wrapped(txt, w, bold=True) for txt, w in head_cells]
    key_rows = [dict(n="30", cells=[dict(text=ws["A30"].value, bold=True, span=8, clip=False)]),
                dict(n="31", cells=[wrapped[0], wrapped[1], wrapped[2], wrapped[3], wrapped[4], wrapped[5], wrapped[6]])]
    for k in range(len(key)):
        r = 32 + k
        key_rows.append(dict(n=str(r), cells=[dict(text=ws[f"A{r}"].value, clip=False), val(f"B{r}"), val(f"C{r}"), val(f"D{r}"), val(f"E{r}"),
                                              dict(text=ws[f"F{r}"].value, html=_nowrap(ws[f"F{r}"].value)), dict(text=_as_shown(ws, f"G{r}"), align="center", clip=False)]))
    geo2 = _ws(d, 0, y, letters, widths, key_rows, show_letters=False)
    i2 = {int(row["n"]): i for i, row in enumerate(key_rows)}
    y = _gap(d, 0, geo2["bottom"], sum(widths), "Rows 39 to 41: empty")
    rule_head = [_wrapped(ws["A43"].value, widths[1], bold=True), _wrapped(ws["B43"].value, widths[2], bold=True), _wrapped(ws["C43"].value, widths[3], bold=True)]
    rule_rows = [dict(n="42", cells=[dict(text=ws["A42"].value.strip(), bold=True, span=8, clip=True)]), dict(n="43", cells=rule_head)]
    for k in range(4):
        r = 44 + k
        rule_rows.append(dict(n=str(r), cells=[val(f"A{r}"), val(f"B{r}"), dict(text=ws[f"C{r}"].value, clip=False)]))
    rule_rows.append(dict(n="48", cells=[dict(text=ws["A48"].value, span=8, clip=True)]))
    geo3 = _ws(d, 0, y, letters, widths, rule_rows, show_letters=False)
    i3 = {int(row["n"]): i for i, row in enumerate(rule_rows)}
    # the named cells, each with its name in a tag beside it
    for name, ref, _ in names:
        r = int(ref[1:])
        _outline(d, geo1, (i1[r], 1), (i1[r], 1), TEAL, inset=1.5)
        x0, y0, w0, h0 = geo1[(i1[r], 2)]
        tw = text_width(f"<b>{name}</b>", SMALL, True) + 12
        assert tw < w0 - 8, (name, tw)
        d.box(f"<b>{name}</b>", x0 + 4, y0 + 2, tw, h0 - 4, fill=TEAL, stroke=TEAL, color=WHITE, size=SMALL)
    # the cells the plan reads: the size and the deviations allowed, and the one row of the rule that supports reliance
    _outline(d, geo1, (i1[15], 4), (i1[15], 4), CORAL, inset=1.5)
    _outline(d, geo1, (i1[20], 4), (i1[20], 4), CORAL, inset=1.5)
    _outline(d, geo1, (i1[16], 1), (i1[16], 1), CORAL, inset=1.5)
    _outline(d, geo3, (i3[44], 0), (i3[44], 2), CORAL, inset=1.5)
    # the reason each refund is a key item; the rule's other rows
    _outline(d, geo2, (i2[32], 5), (i2[38], 5), BLUE, inset=1.5)
    _outline(d, geo3, (i3[45], 0), (i3[47], 2), BLUE, inset=1.5)
    bottom = _gap(d, 0, geo3["bottom"], sum(widths), f"Rows 49 to {ws.max_row}: the frozen draw of {g['size']} refunds, D1 applied to the draw and to the population, A3 on its own population, and the model answer")
    bottom = _tabs(d, bottom + 8, "Sample", ["Materiality", "Risks", "Tests", "Sample", "Timeline"])
    by_customer = sum(1 for r in key if r["cust"] == focus)
    _note(d, f"Outlined in teal: the six named parameter cells, each with its name in the tag beside it (the Name Box shows B15's, whose formula reads the size from the table by two "
             f"lookups). In coral: the cells the plan reads, a size of {g['size']} at {n.CONFIDENCE * 100:.0f}% confidence, a {n.TOLERABLE * 100:.0f}% tolerable rate and no expected "
             f"deviation, {0} deviations allowed, and the only row of the rule that supports reliance. In blue: why each of the {len(key)} key items is key ({by_customer} are customer "
             f"{focus}'s, whose refunds total {_money(tot[focus])}) and the rows of the rule whose limit is above the tolerable rate. Long labels are cut and columns narrower than the workbook's.", bottom + 4)
    return d


# -- IF-19-6-B: why almost any draw gives the same answer ------------------------------------------------------------------------

def if_19_6_b() -> Diagram:
    """Bars of the probability of each deviation count in a draw of 59 from the remaining refunds (a hypergeometric distribution), the
    plausible range shaded and bracketed, the model answer's frozen draw marked, and the upper deviation limit labeled under the axis at
    each end of the range, at the frozen draw and at no deviation."""
    from scipy.stats import beta, hypergeom
    data = ctx.data()
    n = _n()
    g, r6 = ctx.values("ch19.g6"), ctx.values("ch19.r6")
    s = r6["s"]
    ws = _book("Sample")
    size, rest, dev = g["size"], g["n_rest"], len(s["dev"])
    hg = hypergeom(rest, dev, size)
    lo, hi, k_drawn = s["lo"], s["hi"], g["k_drawn"]
    udl = lambda k: float(beta.ppf(n.CONFIDENCE, k + 1, size - k))                                   # noqa: E731
    # the sheet's own figures, asserted against the distribution
    for ref, v in (("B119", dev), ("B123", float(hg.mean())), ("B124", lo), ("B125", hi), ("B126", s["cdf_lo"]), ("B127", s["cdf_hi"]),
                   ("B128", udl(lo)), ("B129", udl(round(hg.mean()))), ("B130", udl(hi)), ("B131", float(hg.pmf(0))), ("B115", k_drawn), ("B117", udl(k_drawn))):
        _agree(ws, ref, v, tol=1e-9 if abs(v) < 1 else 1e-6)
    assert lo <= k_drawn <= hi and abs(float(hg.cdf(hi) - hg.cdf(lo - 1)) - s["inside"]) < 1e-12, "the model's draw is inside the plausible range"
    assert udl(lo) > n.TOLERABLE and udl(0) <= n.TOLERABLE < udl(1), "every plausible draw's limit is above the tolerable rate; only no deviation is at or below it"
    kmax = 30
    assert float(hg.sf(kmax)) < 1e-4 and float(hg.pmf(kmax)) < 1e-4, "the bars end where the distribution does"
    pmf = [float(hg.pmf(k)) for k in range(kmax + 1)]
    top = 0.16
    assert max(pmf) < top
    p_low, p_high = float(hg.cdf(lo - 1)), float(hg.sf(hi))
    d = Diagram("Why Almost Any Draw Gives the Same Answer")
    d.text(f"<b>Deviations in a draw of {size} from the {rest} remaining refunds, {dev} of which deviate</b>", 0, 0, 860, 22, size=14)
    PY, PH = 100.0, 330.0
    plot = kit.Plot(d, 0, PY, 860, PH, y_ticks=[0, 0.04, 0.08, 0.12, 0.16], categories=[str(k) for k in range(kmax + 1)],
                    y_fmt=lambda v: f"{v * 100:.0f}%", bottom=44)
    d.text(f"Deviations found in the draw of {size}", plot.px + plot.pw - 300, PY + PH - 22, 300, 18, size=SMALL, color=GRAY, align="right")
    slot = plot.pw / (kmax + 1)
    xs = lambda k: plot.px + k * slot                                                                    # noqa: E731
    # the plausible range behind the bars (translucent, so the gridlines show)
    d.vertex("", f"rounded=0;html=1;fillColor={AMBER_TINT};strokeColor=none;fillOpacity=80;", xs(lo), plot.py, (hi - lo + 1) * slot, plot.ph)
    colors = [CORAL if k == k_drawn else BLUE if lo <= k <= hi else GRAY for k in range(kmax + 1)]
    plot.bars(pmf, colors=colors, ratio=0.72)
    # the bracket over the range
    bx0, bx1, by = xs(lo) + 2, xs(hi + 1) - 2, PY - 12
    for a, b in ((bx0, bx1),):
        _polyline(d, [(a, by + 10), (a, by), (b, by), (b, by + 10)], AMBER, 2.5)
    msg = (f"<b>{s['inside'] * 100:.0f}% of draws find {lo} to {hi} deviations</b>: the chance of fewer is {p_low * 100:.1f}% and of more {p_high * 100:.1f}%. "
           f"A draw expects {hg.mean():.1f}.")
    mw = 520
    d.text(msg, (bx0 + bx1) / 2 - mw / 2, by - 40, mw, 36, size=SMALL, align="center", valign="bottom")
    # the model answer's frozen draw
    cw = 232.0
    cx = xs(kmax + 1) - cw - 8
    cy = plot.py + 14
    drawn = f"<b>The model answer's frozen draw</b> finds {k_drawn} deviations ({k_drawn / size * 100:.1f}% of the draw), an upper limit of {udl(k_drawn) * 100:.1f}%."
    ch = kit.height_for(drawn, cw, SMALL, pad=14)
    d.box(drawn, cx, cy, cw, ch, fill=CORAL_TINT, stroke=CORAL, stroke_width=1.5, size=SMALL, align="left", valign="top", rounded=False)
    bar_x, bar_y = plot.x_of(k_drawn) + 6, plot.y_of(pmf[k_drawn]) - 3
    a, b = d.anchor(cx, cy + ch / 2), d.anchor(bar_x, bar_y)
    d.edge(a, b, color=CORAL, width=2, end="blockThin", meta="edgeStyle=none;")
    # no deviation at all
    zw = 224.0
    zero = (f"<b>No deviation at all</b> has a probability of {g['p_zero']}. Its upper limit, {udl(0) * 100:.2f}%, is the only one at or below the "
            f"{n.TOLERABLE * 100:.0f}% tolerable rate.")
    zh = kit.height_for(zero, zw, SMALL, pad=14)
    d.box(zero, xs(0) + 6, plot.py + 14, zw, zh, fill=TEAL_TINT, stroke=TEAL, stroke_width=1.5, size=SMALL, align="left", valign="top", rounded=False)
    a, b = d.anchor(plot.x_of(0), plot.py + 14 + zh), d.anchor(plot.x_of(0), plot.y_of(0) - 3)
    d.edge(a, b, color=TEAL, width=2, end="blockThin", meta="edgeStyle=none;")
    # the upper deviation limit, labeled under the axis at the marked counts
    ly = PY + PH + 8
    marks = [(0, "none", TEAL), (lo, "lowest", AMBER), (k_drawn, "frozen", CORAL), (hi, "highest", AMBER)]
    for k, role, color in marks:
        lab = f"<b>{udl(k) * 100:.2f}%</b>"
        wd = max(text_width(lab, SMALL, True), text_width(role, SMALL, False)) + 20
        cxk = plot.x_of(k)
        left = min(max(cxk - wd / 2, 0), 860 - wd)
        d.box(f"{lab}<br>{role}", left, ly, wd, 40, fill=WHITE, stroke=color, stroke_width=2.5, size=SMALL, rounded=False)
        a, b = d.anchor(cxk, plot.py + plot.ph + 20), d.anchor(cxk, ly)
        d.edge(a, b, color=color, width=1.5, meta="edgeStyle=none;")
    d.text("<i>Upper deviation limit at that count,<br>95% confidence</i>", 620, ly + 4, 240, 36, size=SMALL, color=GRAY)
    # a key
    ky = ly + 54
    cx0 = 0.0
    for label, fill, stroke in (("Draws inside the range", BLUE, BLUE), ("Draws outside it", GRAY, GRAY), ("The frozen draw", CORAL, CORAL),
                                ("The plausible range", AMBER_TINT, AMBER)):
        d.box("", cx0, ky + 3, 22, 14, fill=fill, stroke=stroke, stroke_width=1.5, rounded=False)
        wd = text_width(esc(label), SMALL, False) + 8
        d.text(esc(label), cx0 + 28, ky, wd, 20, size=SMALL, valign="middle")
        cx0 += 28 + wd + 18
    _note(d, f"The distribution of the deviations in a draw of {size} (the sample size) from the {rest} remaining refunds, {dev} of which are dated before the last application of the customer's "
             f"payments, as the sheet computes it (a hypergeometric distribution; the shaded range is the tightest with each tail below 5%). Whatever the draw, the upper limit is "
             f"{udl(lo) * 100:.0f}% to {udl(hi) * 100:.0f}%, far above the {n.TOLERABLE * 100:.0f}% tolerable rate, so every plausible draw reaches the conclusion the whole remainder reached: the control "
             f"cannot be relied on. The bar at no deviation is a sliver, too thin to see.", ky + 28)
    return d


# -- IF-19-7-A: the largest refund, from the shipment to the settling of the invoice ----------------------------------------------

def if_19_7_a() -> Diagram:
    """The largest refund's timeline (test S6 and the Timeline sheet): one lane for each department (the warehouse, billing, customer
    service, administration) with the shipment, invoice, receipts, return, credit, clawbacks and refund at their dates, the credit and
    refund dates marked through the lanes, and under them the unpaid balance of the invoice with its value on both dates."""
    data = ctx.data()
    n = _n()
    g = ctx.values("ch19.g7")
    c = n.cycle(data)
    big = n.largest_refund(data)
    cm = c.cm[big["cm"]]
    inv = c.inv[cm["inv"]]
    sr = c.sr[cm["ret"]]
    apps = sorted(c.apps[inv["id"]])
    claws = sorted((a for a in c.adj if a["cm"] == cm["id"]), key=lambda a: a["num"])
    ships = data.q("SELECT DISTINCT s.ShipmentNumber, s.ShipmentDate, s.DeliveryDate, s.ShippedBy FROM SalesInvoiceLine sil JOIN ShipmentLine sl ON "
                   "sl.ShipmentLineID = sil.ShipmentLineID JOIN Shipment s ON s.ShipmentID = sl.ShipmentID WHERE sil.SalesInvoiceID = ? "
                   "ORDER BY s.ShipmentDate, s.ShipmentNumber", inv["id"])
    assert len(ships) == g["ships"] == 1 and big["num"] == g["big"] and len(apps) == g["n_apps"] == 3 and len(claws) == 2
    ship_no, ship_day, deliv_day, carrier = ships[0]
    # the same steps as S6 and the Timeline sheet
    s6 = _g().test(data, "S6")
    assert len(s6) == g["n_steps"] == 10
    ws = _book("Timeline")
    t0 = _row_of(ws, "Date") + 1
    sheet = [(ws.cell(r, 1).value.date().isoformat(), ws.cell(r, 2).value, ws.cell(r, 3).value) for r in range(t0, t0 + 11)]
    assert sorted((day, doc) for day, _, doc, _, _ in s6) == sorted((day, doc) for day, step, doc in sheet if step != "Invoice settled")
    settled_day = max(a[0] for a in apps)
    assert sheet[-1] == (settled_day, "Invoice settled", inv["num"]) and abs(sum(a[1] for a in apps) - inv["gt"]) < 0.005, "the invoice was settled on the last application"
    assert ws.cell(t0 + 1, 2).value == "Sales invoice"
    billing = ws.cell(t0 + 1, 5).value                                                  # who the sheet says issued the invoice
    # the departments, from the records
    cc = lambda e: c.emp[e]["cc"]                                                                   # noqa: E731
    assert cc(sr["recv"]) == "Warehouse" and cc(cm["appr"]) == cc(claws[0]["appr"]) == "Customer Service" and cc(big["appr"]) == "Administration"
    assert all(cc(c.rcpt[a[3]]["rec"]) == cc(a[2]) == "Customer Service" for a in apps)
    unpaid = lambda day: round(inv["gt"] - sum(a[1] for a in apps if a[0] <= day), 2)                # noqa: E731
    assert unpaid(cm["date"]) == unpaid(big["date"]) == g["unpaid"] and abs(ws["B25"].value - g["unpaid"]) < 0.006
    by_rcpt = sorted(apps, key=lambda a: c.rcpt[a[3]]["date"])
    a1, a2, a3 = by_rcpt
    r1, r2_, r3 = (c.rcpt[a[3]] for a in by_rcpt)
    assert ship_day < inv["date"] < r1["date"] < sr["date"] < cm["date"] < big["date"] < r2_["date"] < r3["date"] < settled_day
    assert a1[0] < big["date"] < a2[0] == a3[0] == settled_day and all(a["date"] == cm["date"] for a in claws)
    assert a3[2] == cm["appr"] and a2[2] != cm["appr"], "the last receipt was applied by the credit's approver"
    after_due = _day(big["date"]) - _day(inv["due"])
    assert after_due > 0, "the refund is dated after the invoice's due date"
    weekday = date.fromisoformat(big["date"]).strftime("%A")
    limit_cm, limit_rf = c.emp[cm["appr"]]["limit"], c.emp[big["appr"]]["limit"]
    md = n.md

    def who(e: int) -> str:
        title = c.emp[e]["title"]
        return f"representative {e}" if title == n.CSR else f"the {title}" if c.holders[title] == 1 else f"a {title}"
    meth = lambda r: n.method_word(r["method"])                                                        # noqa: E731
    lanes = [("Warehouse", TEAL), (billing, BLUE), ("Customer Service", AMBER), ("Administration", CORAL)]
    ev = [
        dict(lane=0, day=ship_day, kind="dot", side="above", align="left",
             lines=[ship_no, f"shipped {md(ship_day)} by {carrier}", f"delivered {md(deliv_day)}"]),
        dict(lane=1, day=inv["date"], kind="dot", side="above", align="left",
             lines=[inv["num"], f"{_money(inv['gt'])}, due {md(inv['due'])}", f"customer {inv['cust']}"]),
        dict(lane=2, day=r1["date"], kind="dot", side="above", align="right",
             lines=[r1["num"], f"{md(r1['date'])}, {meth(r1)} {_money(a1[1])}", f"applied by {who(a1[2])}"]),
        dict(lane=0, day=sr["date"], kind="dot", side="above", align="right",
             lines=[sr["num"], f"{sr['reason']}, received by", who(sr["recv"])]),
        dict(lane=2, day=cm["date"], kind="dot", side="above", align="right",
             lines=[cm["num"], f"{_money(cm['gt'])}, the largest credit", f"approved by {who(cm['appr'])}", f"(limit {limit_cm:,.0f}), all posted to 2060"]),
        dict(lane=2, day=cm["date"], kind="claw", side="below", align="right",
             lines=[f"Clawbacks {claws[0]['num']}", f"and {int(claws[1]['num'].rsplit('-', 1)[1])}: {_money(claws[0]['amt'])} and {_money(claws[1]['amt'])}"]),
        dict(lane=3, day=big["date"], kind="dot", side="above", align="right",
             lines=[f"{big['num']}, a {weekday}", f"{_money(big['amt'])} by {meth(big)}, cleared {md(big['cleared'])}",
                    f"approved by {who(big['appr'])} (limit {limit_rf:,.0f})"]),
        dict(lane=2, day=r2_["date"], kind="dot", side="above", align="edge",
             lines=[r2_["num"], f"{md(r2_['date'])}, {meth(r2_)} {_money(a2[1])}", f"applied {md(a2[0])}"]),
        dict(lane=2, day=r3["date"], kind="dot", side="below", align="edge",
             lines=[r3["num"], f"{md(r3['date'])}, {meth(r3)} {_money(a3[1])}", f"applied {md(a3[0])} by the", "credit's approver"]),
        dict(lane=1, day=settled_day, kind="diamond", side="below", align="edge", lines=[f"Invoice settled {md(settled_day)}"]),
    ]
    d0, d1 = _day(ship_day), _day(settled_day)
    X0, X1 = 142.0, 818.0
    xt = lambda iso: X0 + (_day(iso) - d0) * (X1 - X0) / (d1 - d0)                                    # noqa: E731
    LH = 15
    hgt = lambda e: LH * len(e["lines"]) + 6                                                            # noqa: E731
    top = 62.0
    geo, y = [], top
    for i in range(len(lanes)):
        ab = max([hgt(e) + 12 for e in ev if e["lane"] == i and e["side"] == "above"] or [14])
        be = max([hgt(e) + 12 for e in ev if e["lane"] == i and e["side"] == "below"] or [14])
        geo.append(dict(top=y, line=y + ab, bottom=y + ab + be))
        y += ab + be
    lanes_bottom = y
    d = Diagram("The Largest Refund, Step by Step")
    for i, ((name, color), gm) in enumerate(zip(lanes, geo)):
        d.box("", 0, gm["top"], 860, gm["bottom"] - gm["top"], fill=GRAY_TINT if i % 2 == 0 else WHITE, stroke=RULE, rounded=False)
        d.box("", 0, gm["top"], 6, gm["bottom"] - gm["top"], fill=color, stroke=color, rounded=False)
        d.text(f"<b>{esc(name)}</b>", 12, gm["line"] - 18, 122, 36, size=SMALL, valign="middle")
        d.box("", X0 - 8, gm["line"] - 1, X1 - X0 + 16, 2, fill=RULE, stroke=RULE, rounded=False)
    # the balance panel: the invoice less the cash applied by each date
    panel_top = lanes_bottom + 12
    PH = 84.0
    panel_bottom = panel_top + PH + 24
    d.box("", 0, panel_top, 860, panel_bottom - panel_top, fill=WHITE, stroke=RULE, rounded=False)
    d.box("", 0, panel_top, 6, panel_bottom - panel_top, fill=GRAY, stroke=GRAY, rounded=False)
    d.text("<b>Unpaid on the invoice</b>", 12, panel_top + 6, 122, 40, size=SMALL)
    top_val = 20000.0
    assert inv["gt"] < top_val
    yv = lambda v: panel_top + 12 + (top_val - v) * PH / top_val                                         # noqa: E731
    d.box("", X0 - 8, yv(0) - 1, X1 - X0 + 16, 2, fill=RULE, stroke=RULE, rounded=False)
    pts, level = [(xt(inv["date"]), yv(inv["gt"]))], inv["gt"]
    for day in sorted({a[0] for a in apps}):
        pts.append((xt(day), yv(level)))
        level = unpaid(day)
        pts.append((xt(day), yv(level)))
    assert abs(level) < 0.005
    _polyline(d, pts, INK, 2.5)
    # the credit date and the refund date, through every lane and the panel
    cx_, rx_ = xt(cm["date"]), xt(big["date"])
    d.box("", cx_ - 1, 24, 2, panel_bottom - 24, fill=AMBER, stroke=AMBER, rounded=False)
    d.box("", rx_ - 1, 52, 2, panel_bottom - 52, fill=CORAL, stroke=CORAL, rounded=False)
    d.box(f"<b>Credit date {_date(cm['date'])}</b>", cx_ - 88, 0, 176, 24, fill=AMBER_TINT, stroke=AMBER, stroke_width=1.5, size=SMALL, rounded=False)
    d.box(f"<b>Refund date {_date(big['date'])}</b>", rx_ - 88, 28, 176, 24, fill=CORAL_TINT, stroke=CORAL, stroke_width=1.5, size=SMALL, rounded=False)
    d.text(f"<i>Dates are in {big['date'][:4]}</i>", 0, 6, 250, 20, size=SMALL, color=GRAY)

    def mark(x: float, ly: float, kind: str, color: str, size: float = 14) -> None:
        if kind == "dot":
            d.vertex("", f"ellipse;html=1;fillColor={color};strokeColor={WHITE};strokeWidth=1.5;", x - size / 2, ly - size / 2, size, size)
        elif kind == "diamond":
            d.vertex("", f"rhombus;html=1;fillColor={WHITE};strokeColor={color};strokeWidth=2.5;", x - size / 2 - 1, ly - size / 2 - 1, size + 2, size + 2)
        else:
            d.box("", x - 5, ly - 5, 10, 10, fill=WHITE, stroke=color, stroke_width=2, rounded=False)
    for e in ev:
        gm, (_, color) = geo[e["lane"]], lanes[e["lane"]]
        x, ly = xt(e["day"]), gm["line"]
        if e["kind"] == "claw":
            for k in range(2):
                mark(x, ly + 15 + 13 * k, "claw", color)
        else:
            mark(x, ly, e["kind"], color)
        widest = max([text_width(f"<b>{esc(e['lines'][0])}</b>", SMALL, True)] + [text_width(esc(t), SMALL, False) for t in e["lines"][1:]])
        bw, h = widest + 16, hgt(e)
        if e["kind"] == "claw":
            bx, by, side = x - 14 - bw, ly + 8, "right"
        else:
            bx = {"left": x + 9, "right": x - 9 - bw, "edge": 858 - bw}[e["align"]]
            by = ly - 10 - h if e["side"] == "above" else ly + 10
            side = "right" if e["align"] == "right" else "left"
        html = "<br>".join((f"<b>{esc(t)}</b>" if i == 0 else esc(t)) for i, t in enumerate(e["lines"]))
        d.vertex(html, f"text;html=1;whiteSpace=wrap;align={side};verticalAlign=middle;fontSize={SMALL};fontColor={INK};{FONT}spacing=2;fillColor={WHITE};strokeColor=none;",
                 bx, by, bw, h)
    # the dots on the balance line at the credit and refund dates, with the value
    lvl = yv(unpaid(cm["date"]))
    for xx, color in ((cx_, AMBER), (rx_, CORAL)):
        d.vertex("", f"ellipse;html=1;fillColor={color};strokeColor={WHITE};strokeWidth=1.5;", xx - 7, lvl - 7, 14, 14)
    up = _money(unpaid(cm["date"]))
    txt = f"<b>{up}</b> unpaid on the credit date and on the refund date"
    wd = text_width(txt, SMALL, False) + 14
    _label(d, txt, cx_ - 12 - wd, lvl + 9, wd, 20, align="right")
    _label(d, f"<b>{_money(inv['gt'])}</b> invoiced", xt(inv["date"]) + 10, yv(inv["gt"]) + 8, 120, 20)
    # the date axis, a tick a week
    ax = panel_bottom + 6
    for wk in range(d0, d1 + 1, 7):
        iso = date.fromordinal(wk).isoformat()
        d.box("", xt(iso) - 0.5, ax - 6, 1, 8, fill=INK, stroke=INK, rounded=False)
        d.text(_date(iso), xt(iso) - 36, ax + 2, 72, 18, size=SMALL, color=GRAY, align="center")
    ky = ax + 30
    cx0 = 0.0
    for kind, label in (("dot", "A document at its date"), ("claw", "A clawback, dated as the credit"), ("diamond", "The invoice settled")):
        if kind == "dot":
            d.vertex("", f"ellipse;html=1;fillColor={GRAY};strokeColor={WHITE};strokeWidth=1.5;", cx0, ky + 3, 14, 14)
        elif kind == "claw":
            d.box("", cx0 + 2, ky + 5, 10, 10, fill=WHITE, stroke=GRAY, stroke_width=2, rounded=False)
        else:
            d.vertex("", f"rhombus;html=1;fillColor={WHITE};strokeColor={GRAY};strokeWidth=2.5;", cx0 - 1, ky + 2, 16, 16)
        wd = text_width(esc(label), SMALL, False) + 10
        d.text(esc(label), cx0 + 22, ky, wd, 20, size=SMALL, valign="middle")
        cx0 += 22 + wd + 18
    _note(d, f"Test S6 for the largest refund, {big['num']} ({_money(big['amt'])}) of the largest credit, {cm['num']}: each document in the lane of the department whose people recorded, received or "
             f"approved it, at its date, as the Timeline sheet lists them (the invoice is billing's, as the sheet puts it; no table records who prepared it). The credit was posted all to 2060, so {up} "
             f"of the invoice was still unpaid on the credit date and on the refund date, {after_due} day{'s' if after_due != 1 else ''} after the invoice was due and before the applications of its last two "
             f"receipts. The records show what happened and in what order, not why: the refund is a follow-up, not a conclusion.", ky + 28)
    return d


# -- IF-19-8-A: the Deficiencies sheet --------------------------------------------------------------------------------------------

def _cut_head(text: str, width: float, lines: int = 2, bold: bool = True) -> dict:
    """A wrapped heading cut at `lines` lines with an ellipsis, as a fixed-height header row shows it."""
    wrapped = _wrap(text, width - 28, bold=bold)
    if len(wrapped) > lines:
        last = wrapped[lines - 1]
        wrapped = wrapped[:lines - 1] + [_clip(last + " " + " ".join(wrapped[lines:]), width - 8, margin=20, bold=bold)]
    html = "<br>".join(esc(t) for t in wrapped)
    return dict(html=f"<b>{html}</b>" if bold else html, lines=len(wrapped))


def _cut_cell(text: str, width: float, lines: int = 2) -> dict:
    """A wrapped body cell cut at `lines` lines with an ellipsis (a worksheet row of a fixed height)."""
    return _cut_head(text, width, lines=lines, bold=False)


def if_19_8_a() -> Diagram:
    """The Deficiencies sheet: the three threshold cells (links to the named cells of Requirement 3), one row for each of D1 to D8 (D5 on
    two) with its control point, design or operation, the exposure of the last year (a link to the Tests sheet), the band against the
    thresholds, the compensating control and the classification, and the combined groups with their exposure and potential magnitude."""
    data = ctx.data()
    n = _n()
    g = ctx.values("ch19.g8")
    ws = _book("Deficiencies")
    ok = lambda cond, msg: None                                                                      # noqa: E731
    gl = _g()
    mat, perf, trivial = g["mat"], g["perf"], g["trivial"]
    for ref, name, v in (("D3", "Materiality", mat), ("D4", "PerformanceMateriality", perf), ("D5", "ClearlyTrivial", trivial)):
        assert gl.formula(ok, "Deficiencies", ref) == f"={name}" and gl.defined(ok, name).startswith("Materiality!"), ref
        _agree(ws, ref, v, tol=0.006)
        assert ws[f"A{ref[1:]}"].value == {"Materiality": "Materiality", "PerformanceMateriality": "Performance materiality", "ClearlyTrivial": "Clearly trivial"}[name]
    rows = g["rows"]
    assert [x["id"] for x in rows] == ["D1", "D2", "D3", "D4", "D5", "D5", "D6", "D7", "D8"] and ws["A7"].value == "ID" and ws["H7"].value == "Exposure, 2026"
    for k, x in enumerate(rows):
        r = 8 + k
        assert ws[f"A{r}"].value == x["id"] and ws[f"B{r}"].value == x["pattern"] and ws[f"C{r}"].value == x["point"] and ws[f"D{r}"].value == x["kind"]
        _agree(ws, f"H{r}", x["e26"], tol=0.006)
        assert ws[f"I{r}"].value == x["band"] and ws[f"M{r}"].value == x["cls"]
    assert gl.formula(ok, "Deficiencies", "H8").startswith("=Tests!") and gl.formula(ok, "Deficiencies", "I8").startswith("=IF(H8>=Materiality")
    groups = g["groups"]
    assert [x["id"] for x in groups] == ["G1", "G2", "G3", "G4", "G5"] and ws["A19"].value == "Group"
    for k, x in enumerate(groups):
        r = 20 + k
        assert ws[f"A{r}"].value == x["id"] and ws[f"B{r}"].value == x["accounts"] and ws[f"C{r}"].value == x["defs"] and ws[f"G{r}"].value == x["band"] and ws[f"J{r}"].value == x["cls"]
        _agree(ws, f"D{r}", x["e26"], tol=0.006)
        _agree(ws, f"E{r}", x["potential"], tol=0.006)
    assert (ws["D26"].value, ws["D27"].value, ws["D28"].value) == (g["n_sig"], g["n_mw"], g["overall"]) and ws["A26"].value == "Significant deficiencies"
    # fragment 1: the thresholds and the matrix, with the columns E to G, J, L and N hidden
    l1 = ["A", "B", "C", "D", "H", "I", "K", "M"]
    w1 = [24, 34, 144, 60, 80, 82, 140, 150, 146]
    assert sum(w1) == 860
    d = Diagram("The Deficiencies Sheet")
    _formula_bar(d, 0, 0, 860, "H8", gl.clean(gl.workbook()["Deficiencies"]["H8"].value))
    sy = 40.0
    _ws(d, 0, sy, l1, w1, [], show_letters=True)
    y = _gap(d, 0, sy + 22, 860, _clip(f"Rows 1 and 2: {ws['A1'].value}. {ws['A2'].value}", 860 - 12, margin=20))

    def txt(ref: str, **kw) -> dict:
        return dict(text=ws[ref].value, **kw)
    f1 = []
    for r in (3, 4, 5):
        f1.append(dict(n=str(r), cells=[dict(text=ws[f"A{r}"].value, span=2, clip=False), "", _num(_show(ws, f"D{r}"), bold=True)]))
    f1.append(dict(n="6", cells=[]))
    head = [txt("A7", bold=True, clip=False), _cut_head(ws["B7"].value, w1[2]), _cut_head("Control point", w1[3]), _cut_head("Design or operation", w1[4]),
            _cut_head("Exposure, 2026", w1[5]), _cut_head(ws["I7"].value, w1[6]), _cut_head(ws["K7"].value, w1[7]), _cut_head(ws["M7"].value, w1[8])]
    f1.append(dict(n="7", cells=head, h=38))
    for k in range(len(rows)):
        r = 8 + k
        f1.append(dict(n=str(r), h=32, cells=[txt(f"A{r}", bold=True, clip=False), _cut_cell(ws[f"B{r}"].value, w1[2]), txt(f"C{r}", clip=False), _cut_cell(ws[f"D{r}"].value, w1[4]),
                                              _num(_show(ws, f"H{r}")), _cut_cell(ws[f"I{r}"].value, w1[6]), _cut_cell(ws[f"K{r}"].value, w1[7]),
                                              _cut_cell(ws[f"M{r}"].value, w1[8])]))
    geo1 = _ws(d, 0, y, l1, w1, f1, show_letters=False)
    i1 = {int(row["n"]): i for i, row in enumerate(f1)}
    for r, name in ((3, "Materiality"), (4, "PerformanceMateriality"), (5, "ClearlyTrivial")):
        _outline(d, geo1, (i1[r], 3), (i1[r], 3), TEAL, inset=1.5)
        x0, y0, w0, h0 = geo1[(i1[r], 4)]
        label = f"={name}"
        tw = text_width(f"<b>{label}</b>", SMALL, True) + 16
        d.box(f"<b>{label}</b>", x0 + 6, y0 + 2, tw, h0 - 4, fill=TEAL, stroke=TEAL, color=WHITE, size=SMALL)
    _outline(d, geo1, (i1[8], 4), (i1[16], 4), BLUE, inset=1.5)
    _outline(d, geo1, (i1[8], 5), (i1[16], 5), CORAL, inset=1.5)
    _outline(d, geo1, (i1[8], 7), (i1[16], 7), AMBER, inset=1.5)
    y = _gap(d, 0, geo1["bottom"], 860, "Row 17 is empty. The combined evaluation below is rows 18 to 28 of the same sheet, with other columns shown (the letters differ)")
    # fragment 2: the combined groups, with the columns F, H, I and K to N hidden
    l2 = ["A", "B", "C", "D", "E", "G", "J"]
    w2 = [24, 52, 184, 84, 84, 110, 160, 162]
    assert sum(w2) == 860
    _ws(d, 0, y, l2, w2, [], show_letters=True)
    f2 = [dict(n="18", cells=[dict(text=ws["A18"].value, bold=True, span=7, clip=True)])]
    head2 = [txt("A19", bold=True, clip=False), _cut_head(ws["B19"].value, w2[2]), _cut_head(ws["C19"].value, w2[3]), _cut_head(ws["D19"].value, w2[4]),
             _cut_head(ws["E19"].value, w2[5], lines=2), _cut_head(ws["G19"].value, w2[6]), _cut_head(ws["J19"].value, w2[7])]
    f2.append(dict(n="19", cells=head2, h=38))
    for k in range(5):
        r = 20 + k
        f2.append(dict(n=str(r), cells=[txt(f"A{r}", bold=True, clip=False), txt(f"B{r}"), txt(f"C{r}", clip=False), _num(_show(ws, f"D{r}")), _num(_show(ws, f"E{r}")),
                                        txt(f"G{r}"), txt(f"J{r}")]))
    f2.append(dict(n="25", cells=[]))
    for r in (26, 27, 28):
        value = _show(ws, f"D{r}")
        f2.append(dict(n=str(r), cells=[dict(text=ws[f"A{r}"].value, span=3, clip=False, bold=r == 28),
                                        _num(value) if r < 28 else dict(text=value, span=4, bold=True, clip=False)]))
    geo2 = _ws(d, 0, y + 22, l2, w2, f2, show_letters=False)
    i2 = {int(row["n"]): i for i, row in enumerate(f2)}
    _outline(d, geo2, (i2[20], 3), (i2[24], 4), BLUE, inset=1.5)
    _outline(d, geo2, (i2[20], 5), (i2[24], 5), CORAL, inset=1.5)
    _outline(d, geo2, (i2[20], 6), (i2[24], 6), AMBER, inset=1.5)
    _outline(d, geo2, (i2[26], 3), (i2[27], 3), TEAL, inset=1.5)
    _outline(d, geo2, (i2[28], 3), (i2[28], 3), TEAL, inset=1.5)
    bottom = _tabs(d, geo2["bottom"] + 8, "Deficiencies", ["Tests", "Sample", "Timeline", "Deficiencies", "Report"])
    u12 = groups[0]
    _note(d, f"Outlined in teal: the three threshold cells, links to the named cells of the Materiality sheet (each formula beside it). In blue: each pattern's exposure for fiscal {data.C}, "
             f"a link to the Tests sheet (the formula bar shows H8's), and the combined groups' exposure and potential magnitude; the groups count each document once ({u12['defs']}: "
             f"{_money(u12['e26'])}, not the {_money(g['sum12'])} the two patterns add up to). In coral: the band, a formula that reads the exposure against the three thresholds. In amber: the "
             f"classification. In teal, the counts that follow it. The sheet goes on to row {ws.max_row} with the clean results and the model answer. Hidden columns differ by block, long texts are cut.", bottom + 4)
    return d


# -- IF-19-9-A: the report's page budget --------------------------------------------------------------------------------------------

def _vhead(d: Diagram, text: str, x: float, y: float, w: float, h: float) -> None:
    """A table heading with its text turned to read upward, as a narrow column's heading is set."""
    style = (f"rounded=0;whiteSpace=wrap;html=1;fillColor={GRAY_TINT};strokeColor={RULE};strokeWidth=1;fontColor={INK};fontSize={SMALL};fontStyle=1;"
             f"align=left;verticalAlign=middle;{FONT}spacingLeft=4;spacingRight=4;spacing=0;horizontal=0;")
    d.vertex(esc(text), style, x, y, w, h)


def _skeleton(d: Diagram, x: float, y: float, w: float, label_w: float, heads: list[str], labels: list[str], head_h: float, row_h: float) -> float:
    """A table outline for a page: turned headings over a column of row labels and empty rows with a line of text in each cell.  Returns the bottom."""
    cw = (w - label_w) / len(heads)
    d.box("", x, y, label_w, head_h, fill=GRAY_TINT, stroke=RULE, rounded=False)
    for i, head in enumerate(heads):
        _vhead(d, head, x + label_w + i * cw, y, cw, head_h)
    for r, lab in enumerate(labels):
        ry = y + head_h + r * row_h
        d.box(f"<b>{esc(lab)}</b>" if lab else "", x, ry, label_w, row_h, fill=WHITE, stroke=RULE, size=SMALL, rounded=False)
        for i in range(len(heads)):
            d.box("", x + label_w + i * cw, ry, cw, row_h, fill=WHITE, stroke=RULE, rounded=False)
            d.box("", x + label_w + i * cw + 4, ry + row_h / 2 - 2, cw - 8, 4, fill=RULE, stroke=RULE, rounded=False)
    return y + head_h + row_h * len(labels)


def if_19_9_a() -> Diagram:
    """The report's page budget: the three pages and the appendix of three, side by side, each page outlined with what the requirement puts
    on it: the executive summary, conclusion and scope; the six findings in a table of six elements; the clean results and the action plans;
    and the appendix's six exhibits."""
    g, g8 = ctx.values("ch19.g9"), ctx.values("ch19.g8")
    text = _case_text()
    assert "no more than three pages, with an appendix of exhibits" in text and "with an appendix of no more than three pages" in text, "the page limits are the case's"
    ws = _book("Report")
    exhibits = [("The map of the cycle", "Cycle Map"), ("The reconciliation by year", "Reconciliation"), ("The test results", "Tests"),
                ("The sample plan", "Sample"), ("The evaluation of deficiencies", "Deficiencies"), ("The action plans", "Report")]
    assert all(sheet in _cached().sheetnames for _, sheet in exhibits)
    assert all(name in ws["A2"].value for name in ("Cycle Map", "Reconciliation", "Tests", "Sample", "Deficiencies")) and "action plans" in ws["A2"].value
    h0 = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Finding")
    heads = [ws.cell(h0, c).value for c in range(2, 8)]
    assert heads == ["Condition", "Criteria", "Cause", "Effect", "Rating", "Recommendation"], heads
    findings = g["findings"]
    assert [f["id"] for f in findings] == [f"F{i}" for i in range(1, 7)]
    a0 = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Action")
    assert [ws.cell(a0, c).value for c in (1, 3, 4, 5)] == ["Action", "Owner (role)", "Date", "Expected response"] and ws.cell(a0, 6).value.startswith("Where we expect disagreement")
    assert len(g["plans"]) == g["n_plans"] == 10 and g8["n_clean"] == 8
    pop = g["popn"]
    PW, GAP = 196.0, 25.33
    PAGE_H = 396.0
    TOP = 64.0
    xs = [i * (PW + GAP) for i in range(4)]
    d = Diagram("The Report's Page Budget")
    # the limits, as brackets over the pages
    for x0, x1, label, color in ((0, 3 * PW + 2 * GAP, "The report: no more than three pages", GRAY), (xs[3], xs[3] + PW, "The appendix:<br>no more than three pages", TEAL)):
        _polyline(d, [(x0 + 1, 46), (x0 + 1, 38), (x1 - 1, 38), (x1 - 1, 46)], color, 2.5)
        d.text("<br>".join(f"<b>{esc(t)}</b>" for t in label.split("<br>")), x0, 0, x1 - x0, 36, size=SMALL, align="center", valign="bottom")

    def page(i: int, label: str, color: str) -> float:
        d.box("", xs[i], TOP, PW, PAGE_H, fill=WHITE, stroke=color, stroke_width=1.5, rounded=False)
        d.box(f"<b>{esc(label)}</b>", xs[i], TOP, PW, 28, fill=color, stroke=color, color=WHITE, size=SMALL, rounded=False)
        return TOP + 28 + 8

    def block(i: int, y: float, title: str, body: str, tint: str, lines: int = 0, sub: str = "") -> float:
        fill, stroke = kit.tint(tint)
        html = f"<b>{esc(title)}</b>" + (f"<br>{esc(body)}" if body else "") + (f"<br><i>{esc(sub)}</i>" if sub else "")
        w = PW - 16
        h = kit.height_for(html, w, SMALL, pad=14) + lines * 9 + (4 if lines else 0)
        d.box(html, xs[i] + 8, y, w, h, fill=fill, stroke=stroke, stroke_width=1.5, size=SMALL, align="left", valign="top", rounded=False)
        for k in range(lines):
            d.box("", xs[i] + 14, y + h - 10 - 9 * (lines - k), w - 12 - (0 if k < lines - 1 else 40), 4, fill=stroke, stroke=stroke, rounded=False)
        return y + h + 8
    # page 1: the executive summary, the overall conclusion and the scope with the reconciliations
    y = page(0, "Page 1", GRAY)
    y = block(0, y, "Executive summary", "What was tested, over which years, and as at what date", "blue", lines=2)
    y = block(0, y, "Overall conclusion", "What the tests covered, then what they could not reach, then the provisional rating", "teal", lines=3)
    y = block(0, y, "Scope and reconciliations", f"{pop['returns']} returns, {pop['credits']} credits, {pop['refunds']} refunds and {pop['clawbacks']} clawbacks; "
                                                  f"materiality {_money(pop['mat'])}; each population reconciled by year", "gray", lines=2)
    # page 2: the findings
    y = page(1, "Page 2", GRAY)
    y = block(1, y, f"{len(findings)} findings", "One row each, with six elements", "blue")
    _skeleton(d, xs[1] + 8, y, PW - 16, 26, heads, [f["id"] for f in findings], 100, 30)
    # page 3: the clean results and the action plans
    y = page(2, "Page 3", GRAY)
    y = block(2, y, "Clean results", f"The {g8['n_clean']} tests that found nothing, each with its test", "teal")
    y = block(2, y, "Action plans", f"{g['n_plans']} plans, {g['n_owners']} owners by role", "amber", sub=f"{g['first_date']} to {g['last_date']}")
    _skeleton(d, xs[2] + 8, y, PW - 16, 8, ["Action", "Owner", "Date", "Response", "Disagreement"], [""] * g["n_plans"], 88, 11)
    # the appendix: three sheets, the front one listing the six exhibits
    for k in (2, 1):
        d.box("", xs[3] + 6 * k, TOP + 6 * k, PW - 12, PAGE_H - 6, fill=WHITE, stroke=TEAL, stroke_width=1.5, rounded=False)
    d.box("", xs[3], TOP, PW - 12, PAGE_H - 6, fill=WHITE, stroke=TEAL, stroke_width=1.5, rounded=False)
    d.box("<b>Appendix</b>", xs[3], TOP, PW - 12, 28, fill=TEAL, stroke=TEAL, color=WHITE, size=SMALL, rounded=False)
    y = TOP + 36
    for title, sheet in exhibits:
        fill, stroke = kit.tint("teal")
        d.box(f"<b>{esc(title)}</b><br><i>sheet {esc(sheet)}</i>", xs[3] + 8, y, PW - 28, 44, fill=fill, stroke=stroke, stroke_width=1.5, size=SMALL, align="left", valign="middle", rounded=False)
        y += 52
    # the table elements named
    ny = TOP + PAGE_H + 12
    _note(d, f"The page limits are the case's: a report of no more than three pages and an appendix of no more than three. Page 2's table has the six elements of a finding (condition, criteria, "
             f"cause, effect, rating, recommendation) for each of the model answer's {len(findings)} findings; page 3's has an action, an owner by role, a date, the response expected and where "
             f"disagreement is expected with the evidence that answers it. The six exhibits of the appendix are the workbook's own sheets; the model answer's text (the conclusion, findings, clean "
             f"results and action plans) runs to about {g['narrative']:,} words. Pages are not to scale.", ny)
    return d


# -- Power BI pieces: the Audit Monitoring project --------------------------------------------------------------------------------

MONITOR = SOLUTIONS / "Chapter19-capstone-pbi-monitoring" / "Audit Monitoring"
MON_MODEL = MONITOR / "Audit Monitoring.SemanticModel"
MON_REPORT = MONITOR / "Audit Monitoring.Report" / "definition"


def _mon_expected(label: str) -> float:
    """The Expected value of one row of Audit Monitoring's Checks tab (DAX query view), found by its label."""
    text = (MON_MODEL / "DAXQueries" / "Checks.dax").read_text(encoding="utf-8")
    m = re.search(r'"Check", "' + re.escape(label) + r'", "Expected", ([-0-9.eE]+),', text)
    assert m, f"the Audit Monitoring Checks tab has no row {label!r}"
    return float(m.group(1))


def _mon_visual(page: str, name: str) -> dict:
    return json.loads((MON_REPORT / "pages" / page / "visuals" / name / "visual.json").read_text(encoding="utf-8"))


def _mon_title(page: str, name: str) -> str:
    value = _mon_visual(page, name)["visual"]["visualContainerObjects"]["title"][0]["properties"]["text"]["expr"]["Literal"]["Value"]
    return value.strip("'").replace("''", "'")


def _mon_projections(page: str, name: str, role: str) -> list[str]:
    return [p["nativeQueryRef"] for p in _mon_visual(page, name)["visual"]["query"]["queryState"][role]["projections"]]


def _mon_page_tabs(active: str, d: Diagram, y: float) -> None:
    """The page tabs of Audit Monitoring around the active page, as the report orders them."""
    order = json.loads((MON_REPORT / "pages" / "pages.json").read_text(encoding="utf-8"))["pageOrder"]
    names = {p: json.loads((MON_REPORT / "pages" / p / "page.json").read_text(encoding="utf-8"))["displayName"] for p in order}
    i = order.index(active)
    shown = order[max(i - 1, 0):i + 2]
    pbi.page_tabs(d, 4, y, [names[p] for p in shown], names[active])


# -- IF-19-10-A: the Credits Cycle page --------------------------------------------------------------------------------------------

def if_19_10_a() -> Diagram:
    """The Credits Cycle page of Audit Monitoring: the matrix of the four new tests by year (exceptions, population, rate per 1,000), the
    cards of the register's exceptions, reviewed, covered by a design finding and open, and the table of the design findings with what each
    covers through and when it is to be fixed."""
    data = ctx.data()
    n = _n()
    r10, g10 = ctx.values("ch19.r10"), ctx.values("ch19.g10")
    page = "creditsCycle"
    years = [str(y) for y in data.years]
    tests = r10["tests"]
    names = [t["name"] for t in tests]
    assert names == ["CM AboveLimit", "CM CashAndCredit", "RF AboveLimit", "RF BeforePayment"]
    titles = dict(matrix=_mon_title(page, "newTests"), table=_mon_title(page, "findings"))
    assert titles == dict(matrix=g10["titles"]["newTests"], table=g10["titles"]["findings"])
    assert _mon_projections(page, "newTests", "Rows") == ["TestID"] and _mon_projections(page, "newTests", "Columns") == ["Year"]
    assert _mon_projections(page, "newTests", "Values") == ["Exceptions", "Population", "Rate per 1,000"]
    assert _mon_projections(page, "registerCards", "Data") == ["Exceptions", "Reviewed Exceptions", "Design Finding Exceptions", "Open Exceptions"]
    assert _mon_projections(page, "findings", "Values") == ["TestID", "Finding", "CoversThrough", "FixBy", "Design Finding Exceptions", "Open Exceptions"]
    pop = {"CM": r10["pop_c"], "RF": r10["pop_r"]}
    assert _mon_expected("CreditMemo rows") == sum(pop["CM"]) and _mon_expected("CustomerRefund rows") == sum(pop["RF"])
    one = lambda v: f"{xl.xround(v, 1):.1f}"                                                           # noqa: E731
    cells, per_year = [], [0] * len(years)
    for t in tests:
        p = pop[t["name"][:2]]
        row = []
        for k, y in enumerate(years):
            assert t["years"][k] == _mon_expected(f"{t['name']} exceptions {y}") and abs(t["rates"][k] - _mon_expected(f"{t['name']} rate per 1,000 {y}")) < 1e-4
            assert abs(1000 * t["years"][k] / p[k] - t["rates"][k]) < 1e-9
            row.append([f"{t['years'][k]:,}", f"{p[k]:,}", one(t["rates"][k])])
            per_year[k] += t["years"][k]
        assert t["n"] == _mon_expected(f"{t['name']} exceptions") and abs(t["rate"] - _mon_expected(f"{t['name']} rate per 1,000, three years")) < 1e-4
        row.append([f"{t['n']:,}", f"{sum(p):,}", one(t["rate"])])
        cells.append(row)
    total = [[f"{v:,}", "", ""] for v in per_year] + [[f"{sum(per_year):,}", "", ""]]
    # the register's counts
    new_rows = sum(t["n"] for t in tests)
    assert new_rows == g10["new_rows"] == r10["new"] - r10["old"] and r10["new"] == _mon_expected("Exceptions rows")
    assert g10["reviewed"] == _mon_expected("Test 3: reviewed exceptions (every disposition found its exception)") and g10["open_old"] == _mon_expected("Test 3: open exceptions")
    assert r10["new"] == g10["reviewed"] + new_rows + g10["open_old"], "reviewed, covered and open add up to the register"
    cards = [("Exceptions", f"{r10['new']:,}"), ("Reviewed<br>Exceptions", f"{g10['reviewed']:,}"), ("Design Finding<br>Exceptions", f"{new_rows:,}"), ("Open<br>Exceptions", f"{g10['open_old']:,}")]
    # the findings, read from the Design Findings table
    findings = {f[0]: f for f in g10["finding_rows"]}
    assert sorted(findings) == sorted(names)
    for t in tests:
        assert t["n"] == _mon_expected(f"{t['name']} exceptions covered by its design finding") and _mon_expected(f"{t['name']} open exceptions (none is dated after {n.year_end(data.C)})") == 0
    trows = [[nm, _clip(findings[nm][1], 328, margin=34), _us_date(findings[nm][2]), _us_date(findings[nm][3]), f"{t['n']:,}", ""] for nm, t in zip(names, tests)]
    pad = 10.0
    block = [68, 66, 44]

    def draw(d: Diagram) -> dict:
        mx = _year_matrix(d, pad, pad, titles["matrix"], 116, block, years + ["Total"], ["Exceptions", "Population", "Rate<br>per<br>1,000"], names, cells, total, head2=50)
        cy = mx["bottom"] + 12
        cd = _cards(d, pad, cy, 76, cards, [202, 203, 202, 203])
        ty = cd["bottom"] + 12
        tb = _table_visual(d, pad, ty, titles["table"], ["TestID", "Finding", "CoversThrough", "FixBy", "Design Finding<br>Exceptions", "Open<br>Exceptions"],
                           [124, 328, 104, 80, 110, 82], trows, aligns=["left", "left", "right", "right", "right", "right"], title_lines=1, head_lines=2)
        return dict(matrix=mx, cards=cd, table=tb, bottom=tb["bottom"] + pad)
    canvas_h = draw(Diagram("probe"))["bottom"]
    d = Diagram("The Credits Cycle Page")
    pbi.canvas(d, 0, 0, 860, canvas_h)
    parts = draw(d)
    mx, cd, tb = parts["matrix"], parts["cards"], parts["table"]
    # the total row, blank by design in its population and rate; the card and the column of the covered exceptions; the table's dates
    gx0, gy0, _, gh0 = mx["geometry"][("total", 0, 0)]
    gx1, _, gw1, _ = mx["geometry"][("total", len(years), 2)]
    d.outline(gx0 - 2, gy0 - 1, gx1 + gw1 - gx0 + 4, gh0 + 2, CORAL, width=2.5, front=True)
    cx, cyy, cw, ch = cd["cards"][2]
    d.outline(cx - 1, cyy - 1, cw + 2, ch + 2, TEAL, width=2.5, front=True)
    xs_, ys_, ws_, hs_ = tb["geometry"][(-1, 4)]
    xe, ye, we, he = tb["geometry"][(len(trows) - 1, 4)]
    d.outline(xs_ - 1, ys_ - 1, ws_ + 2, ye + he - ys_ + 2, TEAL, width=2.5, front=True)
    xs_, ys_, ws_, hs_ = tb["geometry"][(-1, 2)]
    xe, ye, we, he = tb["geometry"][(len(trows) - 1, 3)]
    d.outline(xs_ - 1, ys_ - 1, xe + we - xs_ + 2, ye + he - ys_ + 2, AMBER, width=2.5, front=True)
    _mon_page_tabs(page, d, canvas_h + 4)
    through = findings[names[0]][2]
    fix = findings[names[0]][3]
    _note(d, f"Report view, page Credits Cycle of Audit Monitoring, in a different arrangement from the model answer's (its text box with the written answer is not drawn). The matrix gives, for each "
             f"new test, the exceptions, the population and the rate per 1,000 by year and over the three years; the Total row's population and rate are blank (outlined in coral) because the Population "
             f"measure reads one process at a time. The cards count the register: {r10['new']:,} exceptions, of which {g10['reviewed']} are reviewed, {new_rows:,} covered by a design finding (outlined in "
             f"teal, as in the table) and {g10['open_old']} open, as before the four tests; the table's Open Exceptions is blank because every exception of the new tests is covered. In amber, the date each "
             f"finding covers through ({_us_date(through)}) and the date it is to be fixed ({_us_date(fix)}): exceptions dated later stay open for review.", canvas_h + 34)
    return d


# -- IF-19-10-B: how the flags reach the register ----------------------------------------------------------------------------------

def if_19_10_b() -> Diagram:
    """The flow of the continuing tests in Audit Monitoring, left to right: the applications (two grouped queries), the credit memos and the
    refunds feed the two flag queries (CreditMemo and CustomerRefund); two exception queries unpivot the flags; the register appends them to
    Chapter 16's exceptions and merges the dispositions and the design findings; its TestID ties it to the Tests table, whose Process the
    Population measure reads.  The anomaly columns of the refund table stop at that table."""
    mt = lambda rel: (MON_MODEL / "definition" / rel).read_text(encoding="utf-8")                      # noqa: E731
    g = ctx.values("ch19.g10")
    expr = mt("expressions.tmdl")
    tables = {p.stem for p in (MON_MODEL / "definition" / "tables").glob("*.tmdl")}
    loaded = {"CreditMemo", "CustomerRefund", "Exceptions", "Tests", "Design Findings", "Dispositions"}
    not_loaded = ["Appliers", "Last Payments", "CM Exceptions", "RF Exceptions", "Invoice Methods", "Past Due at Refund", "JE Exceptions", "PO Exceptions", "Payroll Exceptions"]
    assert loaded <= tables and not any(q in tables for q in not_loaded), "the queries drawn dashed do not load; the others are tables of the model"
    assert all(re.search(rf"(?m)^expression '?{re.escape(q)}'? =", expr) for q in not_loaded)

    def column(table: str, name: str) -> bool:
        return re.search(rf"(?m)^\tcolumn {re.escape(name)}\b", mt(f"tables/{table}.tmdl")) is not None
    assert all(column("CreditMemo", c) for c in ("AboveLimit", "CashAndCredit")) and all(column("CustomerRefund", c) for c in ("AboveLimit", "BeforePayment", "MethodMismatch", "OffsetMissed"))
    assert "CashAndCredit" in " ".join(g["m_cm"].split()) and "BeforePayment" in g["m_rf"]
    # the flags, unpivoted into the two exception queries; the anomaly columns are not
    assert 'Table.Unpivot(Source, {"AboveLimit", "CashAndCredit"}' in expr and 'Table.Unpivot(Source, {"AboveLimit", "BeforePayment"}' in expr
    assert not re.search(r'Table.Unpivot\([^)]*(MethodMismatch|OffsetMissed)', expr), "the anomaly columns are not unpivoted"
    m_ex = " ".join(mt("tables/Exceptions.tmdl").split())
    assert "Table.Combine({#\"JE Exceptions\", #\"PO Exceptions\", #\"Payroll Exceptions\", #\"CM Exceptions\", #\"RF Exceptions\"})" in m_ex
    assert 'Dispositions, {"TestID", "DocumentNumber"}' in m_ex and '#"Design Findings", {"TestID"}' in m_ex and '"Covered"' in m_ex
    assert re.search(r"fromColumn: Exceptions\.TestID\s+toColumn: Tests\.TestID", mt("relationships.tmdl"))
    pop = g["d_pop"]
    assert "SELECTEDVALUE ( Tests[Process] )" in pop and '"Credit memos", COUNTROWS ( CreditMemo )' in pop and '"Refunds", COUNTROWS ( CustomerRefund )' in pop
    tests = g["test_rows"]
    assert [t[0] for t in tests] == ["CM AboveLimit", "CM CashAndCredit", "RF AboveLimit", "RF BeforePayment"] and {t[2] for t in tests} == {"Credit memos", "Refunds"}
    n_new = g["new_rows"]
    # the layout: five columns left to right
    xw = [(0.0, 132.0), (172.0, 148.0), (362.0, 140.0), (548.0, 140.0), (744.0, 116.0)]
    (x1, w1), (x2, w2), (x3, w3), (x4, w4), (x5, w5) = xw
    y0 = 108.0
    d = Diagram("How the Flags Reach the Register")
    for (x, w), label in zip(xw, ("Sources", "Flag queries", "Exception queries", "The register", "The model")):
        d.text(f"<i>{label}</i>", x, 0, w, 20, size=SMALL, color=GRAY, align="center")
    STRAIGHT_ = "edgeStyle=none;"

    def node(x: float, y: float, w: float, h: float, title: str, body: str, kind: str, align: str = "center") -> str:
        """A query or table: solid border for one the model loads, dashed for a query with Enable load cleared."""
        fill, stroke = {"table": (BLUE_TINT, BLUE), "query": (WHITE, GRAY), "anomaly": (CORAL_TINT, CORAL), "model": (AMBER_TINT, AMBER), "source": (GRAY_TINT, GRAY)}[kind]
        html = f"<b>{esc(title)}</b><br>{esc(body)}" if body else f"<b>{esc(title)}</b>"
        return d.box(html, x, y, w, h, fill=fill, stroke=stroke, stroke_width=1.5, size=SMALL, rounded=False, dashed=kind == "query", align=align)

    def link(a: str, b: str, ya: float, yb: float, a_box: tuple, b_box: tuple, **kw) -> None:
        """A straight arrow from the right edge of box a at height ya to the left edge of box b at yb (boxes are (y, h))."""
        d.arrow(a, b, exit=(1, (ya - a_box[0]) / a_box[1]), entry=(0, (yb - b_box[0]) / b_box[1]), meta=STRAIGHT_, **kw)
    # column 1: the helper queries on the applications, the tables, and the queries behind the anomaly columns
    ap = (y0, 48.0)
    cmt = (y0 + 56, 48.0)
    rft = (y0 + 170, 48.0)
    lp = (y0 + 226, 48.0)
    mp = (y0 + 284, 48.0)
    n_ap = node(x1, ap[0], w1, ap[1], "Appliers", "applications by invoice and applier", "query")
    n_cm = node(x1, cmt[0], w1, cmt[1], "Credit memos", "the workbook's table", "source")
    n_rt = node(x1, rft[0], w1, rft[1], "Refunds", "the workbook's table", "source")
    n_lp = node(x1, lp[0], w1, lp[1], "Last Payments", "last application of each invoice", "query")
    n_mp = node(x1, mp[0], w1, mp[1] + 20, "Invoice Methods, Past Due at Refund", "", "query")
    mp_box = (mp[0], mp[1] + 20)
    # column 2: the flag queries and the anomaly columns of the refund table
    cmf = (y0 - 2, 108.0)
    rff = (y0 + 168, 108.0)
    anm = (y0 + 276, 68.0)
    n_cmf = node(x2, cmf[0], w2, cmf[1], "CreditMemo", "flags AboveLimit and CashAndCredit", "table")
    n_rff = node(x2, rff[0], w2, rff[1], "CustomerRefund", "flags AboveLimit and BeforePayment", "table")
    n_anm = node(x2, anm[0], w2, anm[1], "MethodMismatch, OffsetMissed", "anomaly columns of the same table", "anomaly")
    link(n_ap, n_cmf, ap[0] + 24, ap[0] + 24, ap, cmf)
    link(n_cm, n_cmf, cmt[0] + 24, cmt[0] + 24, cmt, cmf)
    link(n_rt, n_rff, rft[0] + 24, rft[0] + 24, rft, rff)
    link(n_lp, n_rff, lp[0] + 24, lp[0] + 24, lp, rff)
    link(n_mp, n_anm, mp[0] + 34, mp[0] + 34, mp_box, anm)
    # column 3: the two exception queries, Chapter 16's, and where the anomalies stop
    cme = (y0 + 22, 60.0)
    rfe = (y0 + 192, 60.0)
    c16 = (y0 - 70, 56.0)
    n_cme = node(x3, cme[0], w3, cme[1], "CM Exceptions", "unpivot the two flags, keep the 1s", "query")
    n_rfe = node(x3, rfe[0], w3, rfe[1], "RF Exceptions", "unpivot the two flags, keep the 1s", "query")
    n_c16 = node(x3, c16[0], w3, c16[1], "JE, PO and Payroll Exceptions", "Chapter 16's queries", "query")
    link(n_cmf, n_cme, cmf[0] + cmf[1] / 2, cme[0] + cme[1] / 2, cmf, cme)
    link(n_rff, n_rfe, rff[0] + rff[1] / 2, rfe[0] + rfe[1] / 2, rff, rfe)
    stop_x = x2 + w2 + 22
    a, b = d.anchor(x2 + w2, anm[0] + anm[1] / 2), d.anchor(stop_x, anm[0] + anm[1] / 2)
    d.edge(a, b, color=CORAL, width=2, dashed=True, meta=STRAIGHT_)
    d.box("", stop_x, anm[0] + 6, 7, anm[1] - 12, fill=CORAL, stroke=CORAL, rounded=False)
    d.text("<b>Stops at the refund table</b><br>a reason to look, not a rule failure: no unpivot, no row in the register", stop_x + 12, anm[0] - 2, x4 - stop_x - 24, 74, size=SMALL)
    # column 4: the register, with the dispositions and the design findings merged in from below
    reg = (y0 - 74, 340.0)
    dsp = (y0 + 292, 90.0)
    n_reg = node(x4, reg[0], w4, reg[1], "Exceptions: the register", f"the new tests appended to Chapter 16's, {n_new:,} more rows; one row for each test and document", "table", align="center")
    n_dsp = node(x4, dsp[0], w4, dsp[1], "Dispositions, Design Findings", "merged on TestID and document, then on TestID", "table")
    link(n_c16, n_reg, c16[0] + c16[1] / 2, c16[0] + c16[1] / 2, c16, reg)
    link(n_cme, n_reg, cme[0] + cme[1] / 2, cme[0] + cme[1] / 2, cme, reg)
    link(n_rfe, n_reg, rfe[0] + rfe[1] / 2, rfe[0] + rfe[1] / 2, rfe, reg)
    d.arrow(n_dsp, n_reg, exit=(0.5, 0), entry=(0.5, 1), meta=STRAIGHT_)
    # column 5: the Tests table, tied by TestID, and the Population measure, tied by the process name
    tst = (y0 - 14, 76.0)
    pmx = (y0 + 120, 92.0)
    n_tst = node(x5, tst[0], w5, tst[1], "Tests", "four new rows: TestID, Test, Process", "table")
    n_pop = node(x5, pmx[0], w5, pmx[1], "Population", "a measure: counts the CreditMemo or CustomerRefund rows of the process", "model")
    d.arrow(n_reg, n_tst, exit=(1, (tst[0] + tst[1] / 2 - reg[0]) / reg[1]), entry=(0, 0.5), meta=STRAIGHT_)
    d.text("<b>TestID</b>", x4 + w4 + 2, tst[0] + tst[1] / 2 - 22, x5 - x4 - w4 - 4, 18, size=SMALL, align="center")
    d.arrow(n_tst, n_pop, exit=(0.5, 1), entry=(0.5, 0), meta=STRAIGHT_)
    d.text("<b>Process</b>", x5 + w5 / 2 - 70, tst[0] + tst[1] + 6, 62, 18, size=SMALL, align="right")
    # the key
    ky = 28.0
    for i, (label, kind) in enumerate((("A table the model loads", "table"), ("A query with Enable load cleared", "query"), ("Anomaly columns, outside the register", "anomaly"))):
        fill, stroke = {"table": (BLUE_TINT, BLUE), "query": (WHITE, GRAY), "anomaly": (CORAL_TINT, CORAL)}[kind]
        d.box("", 0, ky + 22 * i + 3, 22, 14, fill=fill, stroke=stroke, stroke_width=1.5, rounded=False, dashed=kind == "query")
        d.text(esc(label), 30, ky + 22 * i, 280, 20, size=SMALL, valign="middle")
    bottom = dsp[0] + dsp[1]
    _note(d, f"The model answer's queries, left to right (the Power Query editor's names). Applications grouped two ways, the credit memos and the refunds feed the two flag queries, which are "
             f"tables of the model (hidden); each flag pair is unpivoted into an exception query that does not load, and the register appends those with Chapter 16's three, then merges the "
             f"dispositions and the design findings. Its TestID is related to the Tests table, whose Process (Credit memos or Refunds) the Population measure reads: a process spelled differently "
             f"leaves the population and the rate blank while the exceptions still count. MethodMismatch and OffsetMissed are columns of the refund table only, read by their own measures "
             f"and page.", bottom + 12)
    return d


FIGURES = {
    "IF-19-0-A": if_19_0_a,
    "IF-19-0-B": if_19_0_b,
    "IF-19-1-A": if_19_1_a,
    "IF-19-2-A": if_19_2_a,
    "IF-19-3-A": if_19_3_a,
    "IF-19-4-A": if_19_4_a,
    "IF-19-5-A": if_19_5_a,
    "IF-19-5-B": if_19_5_b,
    "IF-19-6-A": if_19_6_a,
    "IF-19-6-B": if_19_6_b,
    "IF-19-7-A": if_19_7_a,
    "IF-19-8-A": if_19_8_a,
    "IF-19-9-A": if_19_9_a,
    "IF-19-10-A": if_19_10_a,
    "IF-19-10-B": if_19_10_b,
}
